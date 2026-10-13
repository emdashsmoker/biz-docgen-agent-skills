#!/usr/bin/env python3
"""hwpx_report.py: Markdown 원고를 참조 양식에서 추출한 스켈레톤 위에 조판하는 단일 CLI.

명령
  build     원고.md → 문서.hwpx          (python-hwpx 6.x, 스켈레톤 스타일 참조)
  skeleton  참조.hwpx → 스켈레톤.hwpx     (header.xml + 본문 구역 secPr만 남긴 빈 양식)
  export    문서.hwpx → 원고.md           (export_rich_markdown)
  fill      문서.hwpx + 치환.json → 문서   (플레이스홀더 치환, 표 라벨 경로 채우기)
  render    문서.hwpx → PDF + PNG         (rhwp export-pdf + pdftoppm)
  check     문서.hwpx 스키마 검증 + rhwp verify 기대값

동작 확인 환경: python-hwpx 6.3.0, lxml 6.0.2, pillow 12.1.1, python-docx 1.2.0, openpyxl 3.1.5,
Python 3.12.3, poppler 24.02, rhwp v0.8.6. 항상 최신 버전을 먼저 설치하며 위 기록은 회귀 원인 추적용.

형식 고지: 본 제품은 한컴의 HWP 문서 파일(.hwp) 공개 문서를 참고하여 개발하였습니다.
상표 고지: 한/글, 한컴오피스, HWPX는 한글과컴퓨터의 상표
비제휴 고지: 이 도구는 한글과컴퓨터와 제휴 관계가 없는 독립 도구

Why: 입력 종류마다 조판기를 두면 규범 수정이 N곳에 번지므로 조판기는 이 파일 하나로 두고
     다른 입력(DOCX, XLSX, HWPX)은 Markdown 원고로 먼저 수렴시킨다.
     출력 형식별 차이는 기록 단계에만 둔다. 파서, 인라인 표기, 표 배치 산정, 그림 크기, 목차 순서는 형식 무관 함수로 두고
     docx_report.py(docx 작성기)가 같은 함수를 호출한다.
"""
from __future__ import annotations

import argparse
import json
import os
import pathlib
import re
import shutil
import subprocess
import sys
import tempfile
import warnings
import zipfile
from dataclasses import dataclass, field

warnings.filterwarnings("ignore", category=DeprecationWarning)

# ═══════════════════════════════════════════════════════════════════════
# 설정부: 조판 수치와 역할·스타일 대응은 모두 여기서만 조정한다
# ═══════════════════════════════════════════════════════════════════════
# Why: 한컴 HWP·OWPML 공개 문서의 이용 조건은 결과물의 유저인터페이스, 매뉴얼, 도움말, 소스에 이 문구 기재를 요구함.
# 문구는 한컴 다운로드 센터(https://www.hancom.com/support/downloadCenter/hwpOwpml) 원문 그대로 유지. CLI --help 하단에 출력
HANCOM_FORMAT_NOTICE = "본 제품은 한컴의 HWP 문서 파일(.hwp) 공개 문서를 참고하여 개발하였습니다."
# Why: 상표의 귀속과 비제휴 사실을 밝혀 한글과컴퓨터의 공식 도구로 오인되지 않게 함
HANCOM_TRADEMARK_NOTICE = "한/글, 한컴오피스, HWPX는 한글과컴퓨터의 상표"
HANCOM_AFFILIATION_NOTICE = "이 도구는 한글과컴퓨터와 제휴 관계가 없는 독립 도구"
# CLI --help 하단 출력용. 고지마다 1줄 (RawDescriptionHelpFormatter로 줄바꿈 유지)
CLI_NOTICE = "\n".join((HANCOM_FORMAT_NOTICE, HANCOM_TRADEMARK_NOTICE, HANCOM_AFFILIATION_NOTICE))
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
ASSET_DIR = os.path.join(os.path.dirname(SCRIPT_DIR), "assets")

# 역할(role) → 스켈레톤 스타일명. numbering: manual=번호를 본문에 기입, auto=개요번호가 붙음
PROFILES = {
    "chapter": {  # 장 체계 문서 (제N장 / N.N / N.N.N / 가. / □ / ❍ / - / ․ / ※)
        "numbering": "manual",
        "chapter_format": "제{n}장 {t}",
        "section_format": "{c}.{s} {t}",          # 1.1 제목
        "subsection_format": "{c}.{s}.{ss} {t}",  # 1.1.1 제목
        "caption_prefix": "{c}-",  # <표 1-3>, [그림 1-3]. 장마다 번호 재시작
        "chapter": "장 제목", "section": "절 제목", "subsection": "항 제목", "ganada": "목 제목",
        "square": "논점", "circle": "주장", "dash": "근거", "dot": "상세", "note": "주석",
        "table_pos": "표 위치", "table_cell": "표 본문", "table_cell_left": "표 본문 왼쪽",
        "table_title": "표 캡션", "unit": "단위", "source": "출처",
        "fig_pos": "그림 위치", "fig_title": "그림 캡션", "body": "바탕글",
        "toc_chapter": "목차 장", "toc_section": "목차 절",  # 선택: --toc 정적 목차
        "table_cell_bullet": "표 글머리", "table_cell_dash": "표 하이픈",  # 선택: 셀 글머리 줄
    },
    "flat": {  # 장 구분 없는 문서 (N.N / N.N.N / 가. / □ / ◦ / - / ․ / *)
        "numbering": "manual",           # 단독 문서: 1. 2. 단독 숫자와 가나다 (개요 번호 대신 본문 기입)
        "chapter_format": "{t}",
        "section_format": "{s}. {t}",              # 1. 제목
        "subsection_format": "{s}.{ss} {t}",       # 1.1 제목
        "caption_prefix": "",  # <표 3>, [그림 3]
        "chapter": "제목",  # 장 구분 없는 문서: # 은 문서 제목 (가운데, 20pt 굵게)
        "section": "1.1", "subsection": "1.1.1", "ganada": "가. 제목",
        "square": "□ 내용", "circle": "◦ 내용", "dash": "- 내용", "dot": "․ 내용", "note": "* 내용",
        "table_pos": "표 위치", "table_cell": "표 내용", "table_cell_left": "표 내용",
        "table_title": "표 제목", "unit": "단위", "source": "출처",
        "fig_pos": "그림위치", "fig_title": "그림 제목", "body": "바탕글",
    },
}

# 표 조판 규범 (HWPUNIT). Why: ○ 계위 본문 시작선보다 표 좌측이 돌출하면 계위와 어긋나 보인다
TABLE_INDENT_MARGIN = 1500      # 본문 폭에서 뺄 계위 들여쓰기 (표 폭 상한 = 본문 폭 - 이 값)
TABLE_MIN_RATIO = 0.5           # 표 폭 하한 (상한 대비). 라벨 2-3열 표의 과소 폭 방지
COL_MIN_W = 3000                # 열 최소 폭 (약 10.6mm)
COL_UNIT_W = 500                # 가중 1단위(반각 1자) 폭, 표 글자 10pt 기준 (한글 1자 = 2단위)
COL_PAD_W = 1800                # 열별 좌우 여백 + 글리프 오차 여유
COL_ONE_LINE_UNITS = 20         # 이 단위를 넘는 셀이 있는 열은 장문 열로 보고 왼쪽 정렬 (한글 10자)
COL_SEARCH_STEPS = 40           # 열 폭 기준선 이분 탐색 횟수 (HWPUNIT 1 미만까지 수렴)
CELL_MARGIN = (510, 510, 141, 141)  # 셀 안쪽 여백 좌, 우, 위, 아래 (HWPUNIT)
CELL_UNIT_W = 500               # 행 높이 추정용 가중 1단위 폭 (표 글자 10pt)
AMOUNT_CELL = re.compile(r"\d,\d{3}|\d\.\d")  # 천 단위 쉼표나 소수점이 있는 수치. 이런 열만 오른쪽 정렬
FIG_CAPTION_SIDE = "BOTTOM"    # 그림 캡션 위치 (TOP | BOTTOM)
MERGE_LABEL_COLS = 2            # 선행 라벨 열 중 세로 병합 대상 최대 열 수 (연속 동일값 병합 관행)
LONG_TABLE_ROWS = 25            # 초과 시 글자취급 해제 + 엔진 분할 + 제목행 반복
HEADER_FILL = "#D9D9D9"         # 표 머리글 배경 (정책보고서 관행 값)
BORDER_WIDTH = "0.15 mm"        # 셀 선 굵기 (정책보고서 관행 값)
BORDER_COLOR = "#000000"
# Why: 표 좌우 바깥선을 비우는 것이 정책보고서 관행. 좌측 끝 셀은 왼선, 우측 끝 셀은 오른선 없음
EDGE_OPEN_SIDES = True
FIGURE_IN_TABLE = True          # 그림을 투명 1열 표에 넣고 표의 번호 유형을 그림으로 지정 (손 편집 관행)
ROW_BASE_H = 2000               # 1줄 행 높이 (한/글은 내용에 맞춰 자동 확장하므로 최소값 성격)
ROW_LINE_H = 1400               # 접힘·개행 1줄 추가분
NUMERIC_CELL = re.compile(r"^[-+±(]?[\d,]+(?:\.\d+)?%?\)?$")
BULLET_PREFIXES = ("·", "•", "ㆍ", "∙", "○", "◦", "❍", "□", "■", "※", "- ", "▪")

# 그림 크기 규범 (mm, 본문 폭 약 170mm 기준). Why: 좌우를 가득 채운 그림은 본문 계위와 어긋나 보이고
# 세로형은 한 쪽을 넘긴다. 가로형은 폭, 정사각형은 폭, 세로형은 높이로 상한을 둔다
IMG_LANDSCAPE_W_MM = 140.0      # 비율 1.4 이상
IMG_SQUARE_W_MM = 90.0          # 비율 0.8 초과 1.4 미만
IMG_PORTRAIT_H_MM = 110.0       # 비율 0.8 이하
IMG_DEFAULT_DPI = 96.0
IMG_URL_TIMEOUT = 20
COVER_TOP_LINES = 8             # 표지 제목 위 빈 줄 수
COVER_TITLE_PT = 2400           # 표지 제목 24pt
COVER_SUB_PT = 1600             # 부제 16pt
COVER_DATE_PT = 1400            # 일자 14pt
COVER_ORG_PT = 1600             # 기관 16pt
COVER_GAP_LINES = 14            # 일자와 기관 사이 빈 줄 수
TOC_TAB_POS = 48100             # 목차 스타일의 오른쪽 점선 탭 위치 (스켈레톤 tabPr 값)
TOC_SUB_INDENT = 3000           # 목차 소절 들여쓰기 (절 1500의 2배, 정책보고서 관행 값)
TOC_UNIT_W = (700, 600)         # 목차 탭 폭 추정용 가중 1단위 폭 (장, 절). 한/글은 탭을 다시 계산하므로 근사 렌더러 표시용
TOC_TITLE = "목  차"
TOC_LIST_INDENT = 0             # 표 목차와 그림 목차 항목의 왼쪽 여백 (목차 절 스타일 여백을 덮어씀)
CAPTION_GAP = 850               # 캡션과 표 사이 간격 (python-hwpx 기본값, 실코퍼스 다수 값)
BLOCK_GAP_LINES = 0             # 표·그림 블록 앞뒤 빈 문단 수. 간격은 표위치·그림위치 스타일의 문단 여백이 담당

# rhwp 렌더러
RHWP_RELEASE_PAGE = "https://github.com/edwardkim/rhwp/releases/latest"
RENDER_DPI = 80

HP = "{http://www.hancom.co.kr/hwpml/2011/paragraph}"
MANUAL_SYMBOLS = {"square": "□ ", "circle": "○ ", "dash": "- ", "dot": "ㆍ ", "note": "※ "}


# ═══════════════════════════════════════════════════════════════════════
# 1. Markdown 파서: 원고를 (kind, payload) 요소 목록으로 정규화
# ═══════════════════════════════════════════════════════════════════════
@dataclass
class Element:
    kind: str                # heading square circle dash dot note table image pagebreak
    text: str = ""
    level: int = 0           # heading: 1-4, bullet: 들여쓰기 순위
    data: dict = field(default_factory=dict)


HEADING_RE = re.compile(r"^(#{1,4})\s+(.*)$")
BULLET_RE = re.compile(r"^(\s*)[-*+]\s+(.*)$")
IMAGE_RE = re.compile(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$")
TABLE_SEP_RE = re.compile(r"^\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)*\|?$")
CAPTION_RE = re.compile(r"^[<＜]\s*표\s*[^>＞]*[>＞]\s*(.*)$")
FIG_CAPTION_RE = re.compile(r"^\[\s*그림\s*[^\]]*\]\s*(.*)$")
UNIT_RE = re.compile(r"^\(\s*단위\s*[:：].*\)$")
SOURCE_RE = re.compile(r"^(?:[(（]\s*(?:출처|자료)\s*[)）]|(?:출처|자료|주)\s*[:：])")
FOOTNOTE_DEF_RE = re.compile(r"^\[\^([^\]]+)\]:\s*(.*)$")
NUM_PREFIX_RE = re.compile(r"^(제\s*\d+\s*장|\d+(?:\.\d+)*\.?|[가-힣]\.)\s+")
SQUARE_MARK = ("□", "■", "☐")
CIRCLE_MARK = ("○", "❍", "◦", "ㅇ")
NOTE_MARK = ("※", "*")


def _strip_marker(text: str, marks: tuple[str, ...]) -> str | None:
    for m in marks:
        if text.startswith(m):
            return text[len(m):].lstrip()
    return None


def parse_markdown(path: str) -> tuple[list[Element], dict[str, str]]:
    lines = open(path, encoding="utf-8").read().splitlines()
    footnotes: dict[str, str] = {}
    indents = sorted({len(m.group(1).expandtabs(4)) for ln in lines
                      if (m := BULLET_RE.match(ln))})
    elems: list[Element] = []
    i = 0
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        stripped = line.strip()
        i += 1
        if not stripped:
            continue
        if m := FOOTNOTE_DEF_RE.match(stripped):
            footnotes[m.group(1)] = m.group(2)
            continue
        if stripped == "---":
            elems.append(Element("pagebreak"))
            continue
        if m := HEADING_RE.match(stripped):
            elems.append(Element("heading", m.group(2).strip(), len(m.group(1))))
            continue
        if stripped.startswith("|"):
            rows = []
            j = i - 1
            while j < len(lines) and lines[j].strip().startswith("|"):
                rows.append(lines[j].strip())
                j += 1
            i = j
            tbl = _parse_table(rows, elems)
            while i < len(lines) and not lines[i].strip():
                i += 1
            if i < len(lines) and UNIT_RE.match(lines[i].strip()) and not tbl.data.get("unit"):
                tbl.data["unit"] = lines[i].strip()  # 표 뒤에 쓴 단위 줄도 표 위 단위 행으로
                i += 1
            elems.append(tbl)
            continue
        if m := IMAGE_RE.match(stripped):
            elems.append(Element("image", m.group(1).strip(),
                                 data={"path": m.group(2).strip(), "base": os.path.dirname(path)}))
            continue
        if m := CAPTION_RE.match(stripped):
            elems.append(Element("table_caption", m.group(1).strip()))
            continue
        if m := FIG_CAPTION_RE.match(stripped):
            if elems and elems[-1].kind == "image":
                elems[-1].text = m.group(1).strip()
            continue
        if UNIT_RE.match(stripped):
            elems.append(Element("unit", stripped))
            continue
        if stripped.startswith(">"):  # 인용 줄은 내용으로 역할 판정 (출처 표기면 출처, 아니면 주석)
            body = stripped.lstrip("> ").strip()
            kind = "source" if SOURCE_RE.match(body) else "note"
            elems.append(Element(kind, re.sub(r"^[(（]\s*(출처|자료)\s*[)）]\s*", r"\1: ", body)))
            continue
        if SOURCE_RE.match(stripped):
            elems.append(Element("source", re.sub(r"^[(（]\s*(출처|자료)\s*[)）]\s*", r"\1: ", stripped)))
            continue
        if (t := _strip_marker(stripped, SQUARE_MARK)) is not None:
            elems.append(Element("square", t))
            continue
        if (t := _strip_marker(stripped, NOTE_MARK)) is not None and not stripped.startswith("**"):
            elems.append(Element("note", t))
            continue
        if m := BULLET_RE.match(raw):
            indent = len(m.group(1).expandtabs(4))
            rank = indents.index(indent) if indent in indents else len(indents)
            body = m.group(2).strip()
            if (t := _strip_marker(body, CIRCLE_MARK)) is not None:
                elems.append(Element("circle", t))
            elif (t := _strip_marker(body, SQUARE_MARK)) is not None:
                elems.append(Element("square", t))
            elif (t := _strip_marker(body, NOTE_MARK)) is not None:
                elems.append(Element("note", t))
            else:
                kind = ("circle", "dash", "dot")[min(rank, 2)]
                elems.append(Element(kind, body))
            continue
        if (t := _strip_marker(stripped, CIRCLE_MARK)) is not None:
            elems.append(Element("circle", t))
            continue
        elems.append(Element("circle", stripped))  # 일반 문단은 ○ 서술로 처리
    return elems, footnotes


def _parse_table(rows: list[str], elems: list[Element]) -> Element:
    aligns: list[str | None] = []
    cells: list[list[str]] = []
    for r in rows:
        if TABLE_SEP_RE.match(r):
            for seg in r.strip("|").split("|"):
                s = seg.strip()
                aligns.append("LEFT" if s.startswith(":") and not s.endswith(":")
                              else "RIGHT" if s.endswith(":") and not s.startswith(":")
                              else "CENTER" if s.startswith(":") else None)
            continue
        cells.append([c.strip().replace("<br>", "\n") for c in r.strip().strip("|").split("|")])
    caption, unit = None, None
    for k in range(len(elems) - 1, max(len(elems) - 3, -1), -1):  # 표 앞 2줄 안의 캡션과 (단위) 줄을 순서 무관하게 흡수
        if elems[k].kind == "unit" and unit is None:
            unit = elems.pop(k).text
        elif elems[k].kind == "table_caption" and caption is None:
            caption = elems.pop(k).text
        else:
            break
    return Element("table", caption or "", data={"rows": cells, "aligns": aligns, "unit": unit})


# ═══════════════════════════════════════════════════════════════════════
# 2. 원고 린트: 조판 전에 원고 결함을 알린다 (변환은 계속 진행)
# ═══════════════════════════════════════════════════════════════════════
SHORT_SECTION_ELEMS = 8  # 절(##)당 요소 수가 이 값 미만이면 쪽 나눔 시 여백만 남을 가능성


def lint(elems: list[Element]) -> list[str]:
    issues = []
    prev_level = 0
    sec_sizes, cur = [], None
    for e in elems:
        if e.kind == "heading" and e.level <= 2:
            if cur is not None:
                sec_sizes.append(cur)
            cur = 0
        elif cur is not None:
            cur += 1
    if cur is not None:
        sec_sizes.append(cur)
    if sec_sizes:
        short = sum(1 for n in sec_sizes if n < SHORT_SECTION_ELEMS)
        issues.append(f"[쪽 나눔 HITL] 장과 절 {len(sec_sizes)}개 중 짧은 절 {short}개 (요소 {SHORT_SECTION_ELEMS}개 미만): "
                      f"기본값 chapter 유지. 유저가 쪽 나눔 변경을 요청하면 --page-break section|none")
    for idx, e in enumerate(elems):
        if e.kind == "heading":
            if e.level > prev_level + 1 and prev_level:
                issues.append(f"[헤더 건너뜀] '{e.text[:30]}' (#{prev_level} 다음에 #{e.level})")
            prev_level = e.level
        if e.kind == "square":
            nxt = elems[idx + 1].kind if idx + 1 < len(elems) else None
            if nxt not in ("circle", "dash", "dot", "note"):
                issues.append(f"[□ 단독] '{e.text[:30]}' 아래에 ○ 또는 - 서술이 없음")
        if e.kind == "table" and len(e.data["rows"]) > LONG_TABLE_ROWS:
            issues.append(f"[장표 HITL] {len(e.data['rows'])}행 표: 유저에게 --long-table engine|split 택일을 질의")
        if e.kind == "table" and not e.text:
            issues.append("[표 캡션 없음] 표 앞 줄에 '<표> 제목' 형식 캡션 권장")
    return issues


# ═══════════════════════════════════════════════════════════════════════
# 3. 조판기: python-hwpx로 스켈레톤 위에 요소를 배치
# ═══════════════════════════════════════════════════════════════════════
def _weighted_len(text: str) -> int:
    return sum(2 if ("\uac00" <= ch <= "\ud7a3" or "\u4e00" <= ch <= "\u9fff") else 1 for ch in text)


def _line_units(cell: str) -> int:
    """셀의 가장 긴 줄의 가중 길이 (한 줄에 모두 넣을 때의 필요 단위)."""
    return max((_weighted_len(ln.strip()) for ln in cell.split("\n")), default=0)


def _word_units(cell: str) -> int:
    """셀의 가장 긴 어절의 가중 길이 (어절이 잘리지 않을 최소 단위)."""
    return max((_weighted_len(w) for ln in cell.split("\n") for w in ln.split()), default=0)


def calc_col_widths(rows: list[list[str]], max_w: int) -> list[int]:
    """열 폭 배분. 합계가 표의 실폭이다.
    모든 열이 한 줄에 들어가면 열마다 한 줄 폭을 쓴다(표가 하한보다 좁으면 비례 확대).
    넘치면 표 폭을 상한으로 두고 '열 폭 = min(한 줄 폭, 기준선)'의 합이 상한과 같아지는 기준선을 이분 탐색으로 찾는다.
    Why: 짧은 열은 한 줄 폭을 그대로 받고 긴 열만 기준선에서 접힌다. 어절 폭은 모든 열에 보장하여 낱말 중간 줄바꿈을 막는다."""
    ncol = max(len(r) for r in rows)
    cols = [[r[c] if c < len(r) else "" for r in rows] for c in range(ncol)]
    one_line = [max(COL_MIN_W, max(map(_line_units, col)) * COL_UNIT_W + COL_PAD_W) for col in cols]
    floor = [max(COL_MIN_W, max(map(_word_units, col)) * COL_UNIT_W + COL_PAD_W) for col in cols]
    if sum(one_line) <= max_w:
        target = max(sum(one_line), int(max_w * TABLE_MIN_RATIO))
        widths = [w * target / sum(one_line) for w in one_line]
    elif sum(floor) >= max_w:
        target = max_w
        widths = [w * target / sum(floor) for w in floor]
    else:
        target = max_w
        lo, hi = 0.0, float(max(one_line))
        for _ in range(COL_SEARCH_STEPS):
            cap = (lo + hi) / 2
            if sum(max(f, min(o, cap)) for o, f in zip(one_line, floor)) < target:
                lo = cap
            else:
                hi = cap
        widths = [max(f, min(o, hi)) for o, f in zip(one_line, floor)]
    out = [int(w) for w in widths]
    out[out.index(max(out))] += target - sum(out)  # 반올림 잔여는 가장 넓은 열이 흡수
    return out


def header_rows(rows: list[list[str]]) -> int:
    """머리글 행 수. 첫 행 다음에 '^'로 위 머리글과 병합하는 행이 이어지면 다단 머리글."""
    n = 1
    while n < len(rows) and any(v.strip() == "^" for v in rows[n]) and not any(
            NUMERIC_CELL.match(v.replace(" ", "")) for v in rows[n] if v.strip() and v.strip() not in ("^", "<")):
        n += 1
    return n


def merge_spans(rows: list[list[str]], mode: str, hdr: int = 1) -> list[tuple[int, int, int]]:
    """세로 병합 구간 (col, start_row, end_row). 본문 행 기준(머리글 제외).
    mode auto: 선행 라벨 열(첫 숫자 열 앞, 최대 MERGE_LABEL_COLS)에서 연속 동일값을 병합.
    상위 열 병합 구간 안에서만 하위 열을 병합하여 MECE 유지. 셀 값 '^'는 위 셀과 병합 지시."""
    body = rows[hdr:]
    ncol = len(rows[0])
    spans: list[tuple[int, int, int]] = []
    for c in range(ncol):  # 다단 머리글의 세로 병합 (행 0 기준, 음수 오프셋으로 표기)
        k = 0
        while k + 1 < hdr and rows[k + 1][c].strip() == "^":
            k += 1
        if k:
            spans.append((c, -hdr, -hdr + k))
    if not body:
        return spans
    label_cols = []
    if mode == "auto":
        for c in range(min(ncol, MERGE_LABEL_COLS)):
            vals = [r[c].strip() for r in body if r[c].strip() and r[c].strip() not in ("^", "<")]
            if vals and all(NUMERIC_CELL.match(v.replace(" ", "")) for v in vals):
                break
            label_cols.append(c)
    parent_bounds = [(0, len(body) - 1)]
    for c in range(ncol):
        explicit = any(r[c].strip() == "^" for r in body)
        if c not in label_cols and not explicit:
            continue
        new_bounds = []
        for lo, hi in parent_bounds:
            start = lo
            for i in range(lo + 1, hi + 2):
                same = i <= hi and (body[i][c].strip() == "^" or
                                    (c in label_cols and body[i][c].strip() == body[start][c].strip() and body[start][c].strip()))
                if not same:
                    if i - 1 > start:
                        spans.append((c, start, i - 1))
                    new_bounds.append((start, i - 1))
                    start = i
        parent_bounds = new_bounds
    return spans


def hmerge_spans(rows: list[list[str]]) -> list[tuple[int, int, int]]:
    """가로 병합 구간 (row, start_col, end_col). 셀 값 '<'는 왼쪽 셀과 병합 지시. 머리글 행 포함."""
    spans = []
    for ri, r in enumerate(rows):
        c = 0
        while c < len(r):
            if r[c].strip() == "<":
                c += 1
                continue
            end = c
            while end + 1 < len(r) and r[end + 1].strip() == "<":
                end += 1
            if end > c:
                spans.append((ri, c, end))
            c = end + 1
    return spans


def col_aligns(rows: list[list[str]], explicit: list[str | None]) -> list[str]:
    """정렬 결정 단위는 열. 명시 마커 > 금액과 소수 숫자 열 RIGHT > 그 외 숫자 열(연번, 연도, 건수) CENTER > 글머리와 장문 포함 LEFT > CENTER."""
    ncol = max(len(r) for r in rows)
    out = []
    for c in range(ncol):
        body = [r[c] for r in rows[1:] if c < len(r) and r[c].strip()]
        if c < len(explicit) and explicit[c]:
            out.append(explicit[c])
        elif body and all(NUMERIC_CELL.match(v.replace(" ", "")) for v in body):
            out.append("RIGHT" if any(AMOUNT_CELL.search(v) for v in body) else "CENTER")
        elif any(v.lstrip().startswith(BULLET_PREFIXES) or "\n" in v
                 or _weighted_len(v) > COL_ONE_LINE_UNITS for v in body):
            out.append("LEFT")
        else:
            out.append("CENTER")
    return out


# ── 형식 무관 조판 판단 (hwpx와 docx 작성기가 함께 사용) ─────────────────────
INLINE_RE = re.compile(
    r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`|<sup>.*?</sup>|<sub>.*?</sub>|\^[^^\s]+\^"
    r"|\[\^[^\]]+\])")


def inline_segments(text: str) -> list[tuple[str, str]]:
    """인라인 표기를 (종류, 글자) 목록으로 분해. 종류는 text, bold, italic, code, sup, sub, footnote(각주 키)"""
    out = []
    for seg in INLINE_RE.split(text):
        if not seg:
            continue
        if seg.startswith("**"):
            out.append(("bold", seg[2:-2]))
        elif seg.startswith("*"):
            out.append(("italic", seg[1:-1]))
        elif seg.startswith("`"):
            out.append(("code", seg[1:-1]))
        elif seg.startswith("<sup>") or seg.startswith("^"):
            out.append(("sup", seg[5:-6] if seg.startswith("<") else seg[1:-1]))
        elif seg.startswith("<sub>"):
            out.append(("sub", seg[5:-6]))
        elif seg.startswith("[^"):
            out.append(("footnote", seg[2:-1]))
        else:
            out.append(("text", seg))
    return out


def edge_sides(c0: int, c1: int, ncol: int) -> list[str]:
    """셀 괘선 4변 중 그릴 변. 표 좌우 바깥선은 비움(EDGE_OPEN_SIDES)"""
    sides = ["top", "bottom", "left", "right"]
    if EDGE_OPEN_SIDES and c0 == 0:
        sides.remove("left")
    if EDGE_OPEN_SIDES and c1 == ncol - 1:
        sides.remove("right")
    return sides


def cell_line_count(val: str, width: int) -> int:
    """셀 글자 폭과 셀 폭으로 추정한 접힘 줄 수"""
    inner = max(width - CELL_MARGIN[0] - CELL_MARGIN[1], 1)
    return max(1, val.count("\n") + 1,
               sum(-(-_weighted_len(ln) * CELL_UNIT_W // inner) for ln in val.split("\n")))


def lines_to_height(lines: int) -> int:
    return ROW_BASE_H + ROW_LINE_H * (lines - 1)


def row_heights(rows: list[list[str]], widths: list[int], hdr: int, vspans: list, hspans: list) -> list[int]:
    """병합을 반영한 행별 최소 높이. 셀 글자 폭과 열 폭으로 접힘 줄 수를 추정 (한/글과 Word는 내용에 맞춰 확장).
    행 인덱스는 머리글 포함 원고 행 기준.
    Why: 자동 세로 병합 구간의 라벨을 행마다 세면 병합된 모든 행이 라벨 줄 수만큼 커짐.
    세로 병합 라벨은 행 높이 산정에서 빼고, 구간 행 높이 합이 라벨 높이보다 작을 때만 부족분을 구간 행에 고르게 더함."""
    anchors_v = {(hdr + r0, c): hdr + r1 for c, r0, r1 in vspans}
    covered = {(r, c) for (a, c), b in anchors_v.items() for r in range(a + 1, b + 1)}
    span_w = {}
    for ri, c0, c1 in hspans:
        span_w[(ri, c0)] = sum(widths[c0:c1 + 1])
        covered |= {(ri, c) for c in range(c0 + 1, c1 + 1)}
    heights = []
    for ri, row in enumerate(rows):
        lines = [cell_line_count(v, span_w.get((ri, ci), widths[ci])) for ci, v in enumerate(row)
                 if (ri, ci) not in covered and (ri, ci) not in anchors_v]
        heights.append(lines_to_height(max(lines, default=1)))
    for (a, c), b in anchors_v.items():
        need = lines_to_height(cell_line_count(rows[a][c], widths[c]))
        short = need - sum(heights[a:b + 1])
        if short > 0:
            n = b - a + 1
            for k in range(n):
                heights[a + k] += short // n + (1 if k < short % n else 0)
    return heights


def cell_lines(val: str, role_fn) -> tuple[str, list[str | None]]:
    """셀 값의 줄별 글머리 역할과 글머리 기호를 뗀 값. '^', '<' 병합 지시 셀은 빈 값"""
    lines_txt = val.split("\n")
    roles = [role_fn(ln) for ln in lines_txt]
    if any(roles):
        val = "\n".join(re.sub(r"^\s*[◦○ㅇ·•ㆍ-]\s*", "", ln) if br else ln for ln, br in zip(lines_txt, roles))
    return ("" if val.strip() in ("^", "<") else val), roles


def cell_role(align: str, is_head: bool) -> tuple[str, str | None]:
    """셀 문단의 역할과 정렬 덮어쓰기. 장문 열은 왼쪽 스타일, 금액 열은 오른쪽 정렬, 머리글은 가운데"""
    role = "table_cell_left" if align == "LEFT" and not is_head else "table_cell"
    return role, ("RIGHT" if not is_head and align == "RIGHT" else None)


def table_spans(rows: list[list[str]], hdr: int, merge_mode: str) -> tuple[list, list]:
    """세로 병합 (col, r0, r1)과 가로 병합 (row, c0, c1). 세로 병합과 겹치는 가로 병합은 생략하고 알림"""
    if merge_mode == "none":
        return [], []
    vspans = merge_spans(rows, merge_mode, hdr)
    vcells = {(r, c) for c, r0, r1 in vspans for r in range(r0, r1 + 1)}
    hspans = []
    for ri, c0, c1 in hmerge_spans(rows):
        if any((ri - hdr, c) in vcells for c in range(c0, c1 + 1)):
            print(f"WARN: 표 {ri}행 {c0}-{c1}열 가로 병합이 세로 병합과 겹쳐 생략")
            continue
        hspans.append((ri, c0, c1))
    return vspans, hspans


def split_table(e: Element, rows: list[list[str]]) -> list[Element]:
    """장표 물리 분할. 청크마다 제목행을 다시 붙이고 캡션에 (순번)을 덧붙임"""
    chunks = [rows[i:i + LONG_TABLE_ROWS - 1] for i in range(1, len(rows), LONG_TABLE_ROWS - 1)]
    return [Element("table", f"{e.text} ({k})" if e.text else "",
                    data={**e.data, "rows": [rows[0]] + chunk, "split_done": True})
            for k, chunk in enumerate(chunks, 1)]


def front_matter_items(cover: dict | None, toc: bool, headings: list[tuple[int, str]], toc_depth: int,
                       lists: list[tuple[str, list[str]]] | None = None) -> list[dict]:
    """표지와 목차(표 목차, 그림 목차 포함)의 문단 순서. 작성기는 이 목록을 순서대로 기록만 함
    kind: blank, cover(text, pt, bold), toc_title(text, page_break), toc_entry(level, text), list_entry(text)"""
    items: list[dict] = []
    if cover:
        items += [{"kind": "blank"}] * COVER_TOP_LINES
        items.append({"kind": "cover", "text": cover.get("title", ""), "pt": COVER_TITLE_PT, "bold": True})
        if cover.get("subtitle"):
            items += [{"kind": "blank"}, {"kind": "cover", "text": cover["subtitle"], "pt": COVER_SUB_PT, "bold": False}]
        items += [{"kind": "blank"}] * 2
        if cover.get("date"):
            items.append({"kind": "cover", "text": cover["date"], "pt": COVER_DATE_PT, "bold": False})
        items += [{"kind": "blank"}] * COVER_GAP_LINES
        if cover.get("org"):
            items.append({"kind": "cover", "text": cover["org"], "pt": COVER_ORG_PT, "bold": True})
    if toc:
        items.append({"kind": "toc_title", "text": TOC_TITLE, "page_break": bool(cover)})
        items += [{"kind": "toc_entry", "level": lv, "text": text} for lv, text in headings if lv <= toc_depth]
        for title, entries in lists or []:
            if entries:
                items.append({"kind": "toc_title", "text": title, "page_break": True})
                items += [{"kind": "list_entry", "text": t} for t in entries]
    return items


def prepare_image(e: Element) -> tuple[bytes, str, float, float] | None:
    """그림 요소의 (바이트, 형식, 폭 mm, 높이 mm). 파일이 없으면 알리고 None. 한/글 미지원 webp는 jpg로 변환"""
    path = e.data["path"]
    if path.startswith(("http://", "https://")):
        import urllib.request
        req = urllib.request.Request(path, headers={"User-Agent": "Mozilla/5.0"})
        with urllib.request.urlopen(req, timeout=IMG_URL_TIMEOUT) as resp:
            data = resp.read()
            ctype = resp.headers.get("Content-Type", "")
        fmt = {"image/jpeg": "jpg", "image/png": "png", "image/gif": "gif", "image/bmp": "bmp",
               "image/webp": "webp"}.get(ctype.split(";")[0].strip(), os.path.splitext(path)[1].lstrip(".").lower() or "png")
    else:
        if not os.path.isabs(path):
            path = os.path.join(e.data["base"], path)
        if not os.path.exists(path):
            print(f"WARN: 이미지 없음 {path}")
            return None
        data = open(path, "rb").read()
        fmt = os.path.splitext(path)[1].lstrip(".").lower() or "png"
    if fmt == "jpeg":
        fmt = "jpg"
    w_mm, h_mm = image_mm(data)
    if fmt == "webp":
        from PIL import Image
        import io
        buf = io.BytesIO()
        Image.open(io.BytesIO(data)).convert("RGB").save(buf, "JPEG", quality=90)
        data, fmt = buf.getvalue(), "jpg"
    return data, fmt, w_mm, h_mm


def image_mm(data: bytes) -> tuple[float, float]:
    """그림 크기 규범 적용 (가로형 폭, 정사각형 근방 폭, 세로형 높이 상한)"""
    try:
        from PIL import Image
        import io
        im = Image.open(io.BytesIO(data))
        px_w, px_h = im.size
        dpi = im.info.get("dpi", (IMG_DEFAULT_DPI, IMG_DEFAULT_DPI))[0] or IMG_DEFAULT_DPI
    except Exception:
        px_w, px_h, dpi = 800, 600, IMG_DEFAULT_DPI
    w_mm, h_mm = px_w / dpi * 25.4, px_h / dpi * 25.4
    ratio = w_mm / h_mm
    if ratio >= 1.4:
        scale = IMG_LANDSCAPE_W_MM / w_mm
    elif ratio > 0.8:
        scale = IMG_SQUARE_W_MM / w_mm
    else:
        scale = IMG_PORTRAIT_H_MM / h_mm
    return round(w_mm * scale, 1), round(h_mm * scale, 1)


def cell_line_role(line: str, profile: dict) -> str | None:
    """셀 안의 ◦, - 줄에 셀 글머리 역할을 대응 (프로파일에 역할이 있을 때만)."""
    st = line.lstrip()
    if st.startswith(("◦", "○", "ㅇ", "·", "•", "ㆍ")) and profile.get("table_cell_bullet"):
        return "table_cell_bullet"
    if st.startswith("- ") and profile.get("table_cell_dash"):
        return "table_cell_dash"
    return None


class Builder:
    def __init__(self, skeleton: str, profile: dict, page_break: str = "chapter"):
        try:
            from hwpx import HwpxDocument
        except ImportError:
            sys.exit("ERROR: python-hwpx 미설치. `pip install -U python-hwpx --break-system-packages` 후 재실행")
        try:
            with zipfile.ZipFile(skeleton) as z:
                bad = z.testzip()
            if bad:
                sys.exit(f"ERROR: 스켈레톤 zip 손상 항목 {bad}. 스킬 패키지를 다시 설치")
            self.doc = HwpxDocument.open(os.path.abspath(skeleton))
        except SystemExit:
            raise
        except Exception as ex:
            sys.exit(f"ERROR: 스켈레톤 로드 실패 ({type(ex).__name__}: {ex}). "
                     "스켈레톤 파일은 정상이므로 python-hwpx 버전 문제 가능성. `pip show python-hwpx` 확인 후 최신 버전 재설치. "
                     "python-hwpx API를 직접 호출하여 문서를 새로 조립하지 말고 이 스크립트를 최신 API에 맞게 수정할 것")
        self.p = profile
        self.page_break = page_break  # chapter | section | none
        self.styles = {s.name: (sid, s) for sid, s in self.doc.styles.items()}
        missing = [v for k, v in profile.items()
                   if k not in ("numbering", "chapter_format", "section_format", "subsection_format", "caption_prefix") and v is not None and v not in self.styles]
        self.headings: list[tuple[int, str]] = []
        self.toc_numbers: dict[str, int] = {}
        if missing:
            sys.exit(f"ERROR: 스켈레톤에 없는 스타일 {missing}: profile 수정 또는 `skeleton` 명령으로 목록 확인")
        self.n = {"chapter": 0, "section": 0, "subsection": 0, "ganada": 0, "table": 0, "figure": 0,
                  "total_section": 0}
        self.header = self.doc.headers[0]
        self.body_w = self._body_width()
        self.table_max_w = self.body_w - TABLE_INDENT_MARGIN
        self._align_pp: dict[tuple[str, str], str] = {}
        self._bold_cp: dict[str, str] = {}
        self.pending_break = False
        self.manual_symbol = self._manual_symbols()
        self._warned: set[str] = set()
        self.numbered_square = False
        self.toc = False
        self.long_table_mode = "engine"
        self.after_block = False

    def _manual_symbols(self) -> dict[str, str]:
        """스타일의 heading이 BULLET이면 한/글이 기호를 그리고, NONE이면 본문에 기호를 넣어야 한다."""
        out = {}
        for role, sym in MANUAL_SYMBOLS.items():
            if self.p.get(role) is None:
                continue
            pp_id = str(self.styles[self.p[role]][1].para_pr_id_ref)
            pp = self.header.paragraph_property(pp_id)
            htype = getattr(getattr(pp, "heading", None), "type", None) or getattr(pp, "heading_type", None)
            if htype is None:  # 객체 속성이 없으면 XML에서 직접 읽는다
                el = self.header.element.find(f".//{{http://www.hancom.co.kr/hwpml/2011/head}}paraPr[@id='{pp_id}']")
                hd = el.find("{http://www.hancom.co.kr/hwpml/2011/head}heading") if el is not None else None
                htype = hd.get("type") if hd is not None else "NONE"
            if str(htype).upper() != "BULLET":
                out[role] = sym
        return out

    # ── 기본 요소 ───────────────────────────────────────────────────────
    def _body_width(self) -> int:
        sp = self.doc.sections[0].properties
        return sp.page_size.width - sp.page_margins.left - sp.page_margins.right

    def refs(self, role: str) -> dict:
        if self.p.get(role) is None and role not in self._warned:
            self._warned.add(role)
            print(f"WARN: 프로파일에 '{role}' 스타일 없음 → 바탕글로 출력. 파생, 스켈레톤 추가, 기존 계위 흡수 중 유저 택일 필요")
        sid, s = self.styles[self.p[role] or self.p["body"]]
        return {"style_id_ref": sid, "para_pr_id_ref": s.para_pr_id_ref, "char_pr_id_ref": s.char_pr_id_ref}

    INLINE_RE = INLINE_RE

    def _run_style(self, base_cp: str, **fmt) -> str:
        """기준 charPr을 복제하고 굵게·기울임·첨자 요소만 더한다.
        Why: 라이브러리 ensure_run은 굵게 여부만 맞는 기존 charPr(크기 상이)을 반환하여 머리글 글씨가 커진다."""
        key = (base_cp,) + tuple(sorted(fmt.items()))
        if key in self._bold_cp:
            return self._bold_cp[key]
        HH = "{http://www.hancom.co.kr/hwpml/2011/head}"
        base = self.header.element.find(f".//{HH}charPr[@id='{base_cp}']")
        tags = {"bold": "bold", "italic": "italic", "sup": "supscript", "sub": "subscript"}
        wanted = [tags[k] for k, v in (("bold", fmt.get("bold")), ("italic", fmt.get("italic")),
                                       (fmt.get("script"), True)) if v and k in tags]

        def same(e):
            if e.get("height") != base.get("height") or e.get("textColor") != base.get("textColor"):
                return False
            fr, br = e.find(f"{HH}fontRef"), base.find(f"{HH}fontRef")
            if fr is None or br is None or fr.get("hangul") != br.get("hangul"):
                return False
            has = {c.tag.split("}")[1] for c in e}
            return all(t in has for t in wanted) and not ({"bold", "italic", "supscript", "subscript"} & has) - set(wanted)

        def add(e):
            import lxml.etree as ET
            anchor = e.find(f"{HH}offset")
            idx = list(e).index(anchor) + 1 if anchor is not None else len(e)
            for t in wanted:
                e.insert(idx, ET.Element(f"{HH}{t}"))
                idx += 1

        el = self.header.ensure_char_property(predicate=same, modifier=add, base_char_pr_id=base_cp)
        self._bold_cp[key] = el.get("id")
        return self._bold_cp[key]

    def para(self, role: str, text: str, footnotes: dict | None = None, no_heading: bool = False, **attrs):
        """스타일 3종 참조를 함께 지정해야 서식이 붙는다.
        인라인: **굵게** *기울임* ^윗첨자^ <sup></sup> <sub></sub> `그대로` [^n]각주(해당 위치)."""
        r = self.refs(role)
        if role == "square" and self.numbered_square and re.match(r"^\d+\.\s", text):
            r["para_pr_id_ref"] = self._no_bullet_pp(str(r["para_pr_id_ref"]))  # 1. 2. 3. 숫자 헤더
        elif no_heading:  # 번호를 본문에 기입하므로 스타일의 자동 개요 번호를 끈다
            r["para_pr_id_ref"] = self._no_bullet_pp(str(r["para_pr_id_ref"]))
        else:
            text = self.manual_symbol.get(role, "") + text

        if self.pending_break:
            attrs["pageBreak"] = "1"
            self.pending_break = False
        p = self.doc.add_paragraph("", **r, **attrs)
        cp = str(r["char_pr_id_ref"])
        for kind, seg in inline_segments(text):
            if kind == "bold":
                p.add_run(seg, char_pr_id_ref=self._run_style(cp, bold=True))
            elif kind == "italic":
                p.add_run(seg, char_pr_id_ref=self._run_style(cp, italic=True))
            elif kind in ("sup", "sub"):
                p.add_run(seg, char_pr_id_ref=self._run_style(cp, script=kind))
            elif kind == "footnote":
                if footnotes is not None:
                    p.add_footnote(footnotes.get(seg, ""))
            else:  # text, code
                p.add_run(seg, char_pr_id_ref=cp)
        return p

    def _centered_refs(self, role: str) -> dict:
        """표·그림 문단은 스타일 정렬과 무관하게 가운데. 커스텀 양식의 표위치 스타일이 좌측이어도 표가 좌측에 붙지 않음."""
        r = self.refs(role)
        HH = "{http://www.hancom.co.kr/hwpml/2011/head}"
        el = self.header.element.find(f".//{HH}paraPr[@id='{r['para_pr_id_ref']}']")
        al = el.find(f"{HH}align") if el is not None else None
        if al is None or al.get("horizontal") != "CENTER":
            r["para_pr_id_ref"] = self._aligned_pp(str(r["para_pr_id_ref"]), "CENTER")
        return r

    def _aligned_pp(self, base_pp: str, align: str) -> str:
        key = (str(base_pp), align)
        if key not in self._align_pp:
            self._align_pp[key] = self.header.ensure_paragraph_format(base_para_pr_id=base_pp, alignment=align)
        return self._align_pp[key]

    def _no_bullet_pp(self, base_pp: str) -> str:
        key = (base_pp, "NOBULLET")
        if key not in self._align_pp:
            self._align_pp[key] = self.header.ensure_paragraph_format(
                base_para_pr_id=base_pp, heading={"type": "NONE", "idRef": "0", "level": "0"})
        return self._align_pp[key]


    def gap(self):
        for _ in range(BLOCK_GAP_LINES):
            self.doc.add_paragraph("", **self.refs("body"))


    def _cover_char(self, height: int, bold: bool) -> str:
        base = str(self.refs("chapter")["char_pr_id_ref"])
        HH = "{http://www.hancom.co.kr/hwpml/2011/head}"

        def mod(e):
            import lxml.etree as ET
            e.set("height", str(height))
            b = e.find(f"{HH}bold")
            if bold and b is None:
                anchor = e.find(f"{HH}offset")
                e.insert(list(e).index(anchor) + 1, ET.Element(f"{HH}bold"))
            if not bold and b is not None:
                e.remove(b)
        font = self.header.element.find(f".//{HH}charPr[@id='{base}']").find(f"{HH}fontRef").get("hangul")
        return self.header.ensure_char_property(
            predicate=lambda e: e.get("height") == str(height) and e.find(f"{HH}fontRef").get("hangul") == font
            and (e.find(f"{HH}bold") is not None) == bold, modifier=mod, base_char_pr_id=base).get("id")

    def _front_matter(self, cover: dict | None):
        """표지와 목차를 앞부분 구역(section0)에 두고 본문을 새 구역으로 옮긴다. 본문 쪽번호는 1부터 시작."""
        import copy
        import lxml.etree as ET
        s0 = self.doc.sections[0]
        s1 = self.doc.add_section(after=0)
        old = s1.paragraphs[0].element
        old.getparent().replace(old, copy.deepcopy(s0.paragraphs[0].element))
        for p in list(s0.element)[1:]:
            s0.element.remove(p)
            s1.element.append(p)
        self.body_section = s1
        sn = s1.paragraphs[0].element.find(f".//{HP}startNum")
        if sn is not None:
            sn.set("page", "1")
        body_ref = self.refs("body")
        center = self._aligned_pp(str(body_ref["para_pr_id_ref"]), "CENTER")

        def front(text="", cp=None, pp=center, **attrs):
            p = self.doc.add_paragraph("", **body_ref, **attrs)
            p.element.set("paraPrIDRef", pp)
            if text:
                p.add_run(text, char_pr_id_ref=cp or str(body_ref["char_pr_id_ref"]))
            s0.element.append(p.element)  # add_paragraph는 마지막 구역에 붙이므로 앞부분 구역으로 이동
            return p
        if self.toc and not (self.p.get("toc_chapter") and self.p.get("toc_section")):
            print("WARN: 프로파일에 목차 스타일(toc_chapter, toc_section)이 없어 --toc 생략")
            self.toc = False
        sub_pp = list_pp = None
        for it in front_matter_items(cover, self.toc, self.headings, getattr(self, "toc_depth", 2), self.toc_lists()):
            kind = it["kind"]
            if kind == "blank":
                front()
            elif kind == "cover":
                front(it["text"], self._cover_char(it["pt"], it["bold"]))
            elif kind == "toc_title":
                ch = self.refs("chapter")
                front(it["text"], str(ch["char_pr_id_ref"]), self._aligned_pp(str(ch["para_pr_id_ref"]), "CENTER"),
                      pageBreak="1" if it["page_break"] else "0")
            else:  # toc_entry, list_entry
                lv = it.get("level", 2)
                text = it["text"]
                role = "toc_chapter" if lv == 1 else "toc_section"
                r = self.refs(role)
                pp = str(r["para_pr_id_ref"])
                left = TOC_SUB_INDENT if lv >= 3 else (0 if lv == 1 else 1500)
                if kind == "list_entry":
                    if list_pp is None:
                        list_pp = self.header.ensure_paragraph_format(base_para_pr_id=pp, margins={"left": TOC_LIST_INDENT})
                    pp, left = list_pp, TOC_LIST_INDENT
                elif lv >= 3:
                    if sub_pp is None:
                        sub_pp = self.header.ensure_paragraph_format(base_para_pr_id=pp, margins={"left": TOC_SUB_INDENT})
                    pp = sub_pp
                p = front(text, str(r["char_pr_id_ref"]), pp)
                p.element.set("styleIDRef", str(r["style_id_ref"]))
                num = self.toc_numbers.get(text)
                if num:
                    tnode = [t for t in p.element.iter(f"{HP}t") if t.text][-1]
                    unit = TOC_UNIT_W[0] if lv == 1 else TOC_UNIT_W[1]
                    width = max(TOC_TAB_POS - left - _weighted_len(text) * unit - len(str(num)) * unit, 1000)
                    ET.SubElement(tnode, f"{HP}tab", width=str(width), leader="3", type="2").tail = str(num)
        for sec in (s0, s1):
            sec.mark_dirty() if hasattr(sec, "mark_dirty") else None

    def toc_lists(self) -> list[tuple[str, list[str]]]:
        """표 목차와 그림 목차 (제목, 항목 목록). 기본 조판기는 목록 없음"""
        return []

    def set_header(self, spec: str | None):
        """머리말 (기본 OFF). "텍스트;align=LEFT|CENTER|RIGHT". 본문 구역에만 적용.
        python-hwpx가 header를 secPr 안과 ctrl 두 곳에 쓰므로 secPr 안 사본을 제거하고 ctrl 사본의 문단 정렬을 지정."""
        if not spec:
            return
        parts = spec.split(";")
        text = parts[0].strip()
        align = next((kv.split("=", 1)[1].strip().upper() for kv in parts[1:] if kv.strip().lower().startswith("align=")), "RIGHT")
        sec = self.doc.sections[-1]
        hf = self.doc.set_header_text(text, section=sec)
        body_pp = str(self.refs("body")["para_pr_id_ref"])
        pp = self._aligned_pp(body_pp, align)
        HPH = f"{HP}header"
        first = sec.paragraphs[0].element
        sec_pr = first.find(f".//{HP}secPr")
        if sec_pr is not None:
            for h in list(sec_pr.iter(HPH)):
                h.getparent().remove(h)
        for h in first.iter(HPH):
            for pe in h.iter(f"{HP}p"):
                pe.set("paraPrIDRef", pp)

    def set_page_number(self, on: bool):
        """꼬리말 쪽번호. 기본 OFF (취합 시 번호 충돌 방지). ON이면 가운데 아래 '- N -'."""
        for sec in self.doc.sections:  # 앞부분 구역은 항상 제거, 본문 구역만 on에 따라 삽입
            fp = sec.paragraphs[0].element
            for ctrl in list(fp.iter(f"{HP}ctrl")):
                if ctrl.find(f"{HP}pageNum") is not None:
                    ctrl.getparent().remove(ctrl)
        first = self.doc.sections[-1].paragraphs[0].element
        for ctrl in list(first.iter(f"{HP}ctrl")):
            if ctrl.find(f"{HP}pageNum") is not None:
                ctrl.getparent().remove(ctrl)
        if on:
            import lxml.etree as ET
            run = first.find(f"{HP}run")
            ctrl = ET.Element(f"{HP}ctrl")
            ET.SubElement(ctrl, f"{HP}pageNum", pos="BOTTOM_CENTER", formatType="DIGIT", sideChar="-")
            # 한/글 저작 순서(구역 정의, 단 정의, 쪽번호)를 따름. 쪽번호를 단 정의 앞에 두면 한/글이
            # 첫 구역의 여백을 기본값(좌우 30mm)으로 읽어 본문 폭이 150mm로 줄고 표가 인쇄 영역을 넘음
            anchor = next((c for c in run.findall(f"{HP}ctrl") if c.find(f"{HP}colPr") is not None),
                          run.find(f"{HP}secPr"))
            run.insert(list(run).index(anchor) + 1 if anchor is not None else 0, ctrl)

    def _cell_line_role(self, line: str) -> str | None:
        return cell_line_role(line, self.p)

    def _border_fill(self, sides: list[str], fill: str | None = None) -> str:
        """테두리 정의를 만들고 4변 상태를 요소에서 직접 확인·강제한다.
        Why: 라이브러리 버전에 따라 active_borders 해석이 달라 투명 행에 선이 남는 사례가 있었다."""
        HH = "{http://www.hancom.co.kr/hwpml/2011/head}"
        bid = self.header.ensure_border_fill(border_color=BORDER_COLOR, border_width=BORDER_WIDTH,
                                             active_borders=sides, fill_color=fill)
        el = self.header.element.find(f".//{HH}borderFill[@id='{bid}']")
        if el is not None:
            for side in ("left", "right", "top", "bottom"):
                b = el.find(f"{HH}{side}Border")
                if b is not None:
                    b.set("type", "SOLID" if side in sides else "NONE")
        return str(bid)

    def _transparent_row(self, t, ri: int, ncol: int, text: str, role: str, align: str):
        """표 안의 투명 병합 행 (단위, 자료). 손으로 편집할 때 표와 함께 이동하도록 표 내부에 둔다."""
        no_border = self._border_fill([])
        u = self.refs(role)
        t.set_cell_text(ri, 0, text)
        for pe in t.cell(ri, 0).element.iter(f"{HP}p"):
            pe.set("paraPrIDRef", self._aligned_pp(str(u["para_pr_id_ref"]), align))
            pe.set("styleIDRef", str(u["style_id_ref"]))
            for run in pe.iter(f"{HP}run"):
                run.set("charPrIDRef", str(u["char_pr_id_ref"]))
        for ci in range(ncol):
            t.cell(ri, ci).element.set("borderFillIDRef", str(no_border))
            t.cell(ri, ci).set_size(height=ROW_BASE_H)
        if ncol > 1:
            t.merge_cells(ri, 0, ri, ncol - 1)

    def attach_caption(self, holder, kind: str, text: str, side: str, width: int):
        """개체 부착 캡션(한/글 캡션 기능). 번호는 autoNum 필드이며 장 시작 시 재시작."""
        role = "table_title" if kind == "TABLE" else "fig_title"
        r = self.refs(role)
        cap = holder.set_caption("", side=side, width=width, gap=CAPTION_GAP, char_pr_id_ref=r["char_pr_id_ref"])
        cp = cap.paragraphs[0]
        cp.element.set("paraPrIDRef", str(r["para_pr_id_ref"]))
        cp.element.set("styleIDRef", str(r["style_id_ref"]))
        for run in [c for c in cp.element if c.tag == f"{HP}run"]:
            cp.element.remove(run)
        open_b, close_b = ("<표 ", "> ") if kind == "TABLE" else ("[그림 ", "] ")
        prefix = self.p.get("caption_prefix", "").format(c=max(self.n["chapter"], 1))
        run = cp.add_run(open_b + prefix, char_pr_id_ref=r["char_pr_id_ref"])
        import lxml.etree as ET
        an = ET.SubElement(ET.SubElement(run.element, f"{HP}ctrl"), f"{HP}autoNum", num="1", numType=kind)
        ET.SubElement(an, f"{HP}autoNumFormat", type="DIGIT", userChar="", prefixChar="", suffixChar="", supscript="0")
        cp.add_run(close_b + text, char_pr_id_ref=r["char_pr_id_ref"])

    # ── 헤더 계위 ───────────────────────────────────────────────────────
    def heading(self, level: int, text: str):
        text = NUM_PREFIX_RE.sub("", text).strip()
        role = ("chapter", "section", "subsection", "ganada")[level - 1]
        manual = self.p["numbering"] == "manual"
        if level == 1:
            self.n.update(chapter=self.n["chapter"] + 1, section=0, subsection=0, ganada=0)
            if self.page_break in ("chapter", "section") and self.n["chapter"] > 1:
                self.pending_break = True
            label = self.p["chapter_format"].format(n=self.n["chapter"], t=text)
        elif level == 2:
            self.n.update(section=self.n["section"] + 1, subsection=0, ganada=0)
            label = (self.p.get("section_format", "{c}.{s} {t}").format(c=max(self.n["chapter"], 1), s=self.n["section"], t=text)
                     if manual else text)
            # 절 단위 나눔: 장 직후 첫 절과 문서 첫 절은 이미 새 쪽이므로 제외
            if self.page_break == "section" and not (self.n["section"] == 1 and self.n["chapter"] > 0) \
                    and self.n["total_section"] > 0:
                self.pending_break = True
            self.n["total_section"] += 1
        elif level == 3:
            self.n.update(subsection=self.n["subsection"] + 1, ganada=0)
            label = (self.p.get("subsection_format", "{c}.{s}.{ss} {t}").format(
                c=max(self.n["chapter"], 1), s=max(self.n["section"], 1), ss=self.n["subsection"], t=text) if manual else text)
        else:
            self.n["ganada"] += 1
            label = f"{'가나다라마바사아자차카타파하'[(self.n['ganada'] - 1) % 14]}. {text}"
        p = self.para(role, label, no_heading=manual and level in (2, 3))
        if level <= 3:
            self.headings.append((level, label))
        if level == 1 and self.p.get("caption_prefix"):  # 장별 <표 N-1>부터 재시작
            p.add_new_num(kind="TABLE", number=1)
            p.add_new_num(kind="PICTURE", number=1)

    # ── 표 ──────────────────────────────────────────────────────────────
    def table(self, e: Element):
        rows = [r + [""] * (max(len(x) for x in e.data["rows"]) - len(r)) for r in e.data["rows"]]
        nrow, ncol = len(rows), len(rows[0])
        widths = calc_col_widths(rows, self.table_max_w)
        aligns = col_aligns(rows, e.data["aligns"])
        hdr = header_rows(rows)
        table_w = sum(widths)
        if nrow > LONG_TABLE_ROWS and self.long_table_mode == "split":  # 물리 분할: 청크마다 제목행 재부착
            for sub in split_table(e, rows):
                self.table(sub)
            return
        self.gap()
        r = self._centered_refs("table_pos")
        unit = e.data.get("unit")
        sources = e.data.get("sources", [])
        head = 1 if unit else 0  # 단위 행: 머리글 위 투명 병합 셀 (손 편집 관행, 표 우측 끝 정렬)
        t = self.doc.add_table(nrow + head + len(sources), ncol, width=table_w, **r)
        if e.text:
            self.attach_caption(t, "TABLE", e.text, "TOP", table_w)
        t.set_column_widths(widths)
        bf = {}
        if unit:
            self._transparent_row(t, 0, ncol, unit, "unit", "RIGHT")
        for k, src in enumerate(sources):  # 자료 행: 표 아래 투명 병합 셀, 좌측 정렬
            self._transparent_row(t, head + nrow + k, ncol, src, "source", "LEFT")
        for ci in range(ncol):
            sides = edge_sides(ci, ci, ncol)
            bf[(ci, True)] = self._border_fill(sides, HEADER_FILL)
            bf[(ci, False)] = self._border_fill(sides)
        cell_cp = str(self.refs("table_cell")["char_pr_id_ref"])
        vspans, hspans = table_spans(rows, hdr, self.merge_mode)
        heights = row_heights(rows, widths, hdr, vspans, hspans)
        total_h = sum(heights)
        for ri, row in enumerate(rows):
            row_h = heights[ri]
            for ci, val in enumerate(row):
                val, bullet_roles = cell_lines(val, self._cell_line_role)
                t.set_cell_text(ri + head, ci, val, split_paragraphs=True)
                cell = t.cell(ri + head, ci)
                cell.set_size(height=row_h)
                is_head = ri < hdr
                role, align = cell_role(aligns[ci], is_head)
                ref = self.refs(role)
                pp = str(ref["para_pr_id_ref"])
                if align:
                    pp = self._aligned_pp(pp, align)
                cp = self._run_style(cell_cp, bold=True) if is_head else str(ref["char_pr_id_ref"])
                for li, pe in enumerate(cell.element.iter(f"{HP}p")):
                    br = bullet_roles[li] if li < len(bullet_roles) and not is_head else None
                    bref = self.refs(br) if br else None
                    pe.set("paraPrIDRef", str(bref["para_pr_id_ref"]) if bref else pp)
                    pe.set("styleIDRef", str(bref["style_id_ref"] if bref else ref["style_id_ref"]))
                    for run in pe.iter(f"{HP}run"):
                        run.set("charPrIDRef", str(bref["char_pr_id_ref"]) if bref else cp)
                cell.element.set("borderFillIDRef", str(bf[(ci, is_head)]))
                cell.element.set("hasMargin", "1")  # cellMargin 값이 적용되려면 필요
                cm = cell.element.find(f"{HP}cellMargin")
                if cm is not None:
                    for side, mv in zip(("left", "right", "top", "bottom"), CELL_MARGIN):
                        cm.set(side, str(mv))
                if is_head:
                    cell.element.set("header", "1")
        sz = t.element.find(f"{HP}sz")
        if sz is not None:  # 표 높이가 행 높이 합보다 크면 렌더러가 마지막 행을 늘여 그린다
            sz.set("height", str(total_h + ROW_BASE_H * (head + len(sources))))
        pos = t.element.find(f"{HP}pos")
        if pos is not None:  # 글자취급 표는 문단 정렬을 따르므로 pos 정렬도 CENTER로 맞춰 해석 충돌 제거
            pos.set("horzAlign", "CENTER")
        for c, r0, r1 in vspans:  # 연속 동일 라벨 세로 병합 (관행). 음수는 머리글 행
            a, b_ = head + hdr + r0, head + hdr + r1
            for rr in range(a + 1, b_ + 1):
                t.set_cell_text(rr, c, "")
            t.merge_cells(a, c, b_, c)
        for ri, c0, c1 in hspans:  # '<' 지시 가로 병합
            for cc in range(c0 + 1, c1 + 1):
                t.set_cell_text(head + ri, cc, "")
            t.merge_cells(head + ri, c0, head + ri, c1)
        # 병합 뒤 실제 스팬 기준으로 좌우 바깥선 재지정 (병합 셀이 표 끝까지 닿으면 그 쪽 선을 비운다)
        for tc in t.element.iter(f"{HP}tc"):
            addr = tc.find(f"{HP}cellAddr")
            span = tc.find(f"{HP}cellSpan")
            if addr is None or span is None:
                continue
            c0 = int(addr.get("colAddr")); r0 = int(addr.get("rowAddr"))
            c1 = c0 + int(span.get("colSpan")) - 1
            if r0 < head or r0 >= head + nrow:
                continue  # 단위 행과 출처 행은 투명 (nrow는 머리글 포함 행 수)
            is_head = r0 < head + hdr
            tc.set("borderFillIDRef", self._border_fill(edge_sides(c0, c1, ncol), HEADER_FILL if is_head else None))
        if nrow > LONG_TABLE_ROWS and not e.data.get("split_done"):  # 장표: 글자취급 해제, 한/글이 쪽 경계에서 행 단위 분할
            t.element.set("repeatHeader", "1")
            t.element.set("pageBreak", "CELL")
            if pos is not None:
                pos.set("treatAsChar", "0")
                pos.set("flowWithText", "1")
        t.mark_dirty()
        self.after_block = True

    # ── 그림 ────────────────────────────────────────────────────────────
    def image(self, e: Element):
        prepared = prepare_image(e)
        if prepared is None:
            return
        data, fmt, w_mm, h_mm = prepared
        self.gap()
        r = self._centered_refs("fig_pos")
        w_hu, h_hu = round(w_mm * 7200 / 25.4), round(h_mm * 7200 / 25.4)
        sources = e.data.get("sources", [])
        if not FIGURE_IN_TABLE:
            pic = self.doc.add_picture(data, fmt, width_mm=w_mm, height_mm=h_mm, **r)
            if e.text:
                self.attach_caption(pic, "PICTURE", e.text, "BOTTOM", w_hu)
            for src in sources:
                self.para("source", src)
            self.after_block = True
            return
        # 투명 1열 표에 그림을 넣고 표를 그림 번호 유형으로 지정. 내부 그림은 번호 집계에서 제외
        t = self.doc.add_table(1 + len(sources), 1, width=w_hu, **r)
        item = self.doc.add_image(data, fmt)
        item_id = getattr(item, "item_id", item)  # 라이브러리 버전에 따라 객체 또는 문자열
        cell = t.cell(0, 0)
        pic = cell.paragraphs[0].add_picture(str(item_id), width=w_hu, height=h_hu, char_pr_id_ref=r["char_pr_id_ref"])
        img = next((el for el in pic.element.iter() if el.tag.endswith("}img")), None)
        if img is None or not img.get("binaryItemIDRef"):
            sys.exit(f"ERROR: 그림 바이너리 참조 누락 ({e.data['path']}). python-hwpx add_image 반환값 확인 필요")
        pic.element.set("numberingType", "NONE")
        for pe in cell.element.iter(f"{HP}p"):
            pe.set("paraPrIDRef", str(r["para_pr_id_ref"]))
            pe.set("styleIDRef", str(r["style_id_ref"]))
        cell.element.set("borderFillIDRef", self._border_fill([]))
        cell.set_size(height=h_hu + 2 * ROW_LINE_H)
        for k, src in enumerate(sources):
            self._transparent_row(t, 1 + k, 1, src, "source", "LEFT")
        t.element.set("numberingType", "PICTURE")
        pos = t.element.find(f"{HP}pos")
        if pos is not None:
            pos.set("horzAlign", "CENTER")
        sz = t.element.find(f"{HP}sz")
        if sz is not None:
            sz.set("height", str(h_hu + 2 * ROW_LINE_H + ROW_BASE_H * len(sources)))
        if e.text:
            self.attach_caption(t, "PICTURE", e.text, FIG_CAPTION_SIDE, w_hu)
        t.mark_dirty()
        self.after_block = True

    @staticmethod
    def _image_mm(data: bytes, fmt: str) -> tuple[float, float]:
        return image_mm(data)

    # ── 실행 ────────────────────────────────────────────────────────────
    def build(self, elems: list[Element], footnotes: dict, out: str):
        skip = 0
        for idx, e in enumerate(elems):
            if skip:
                skip -= 1
                continue
            if e.kind in ("table", "image"):  # 직후의 ※, *, 자료: 줄은 블록 하단 투명 행으로 부착
                srcs, consumed = [], 0
                for nx in elems[idx + 1:]:
                    if nx.kind in ("note", "source"):
                        srcs.append(nx.text)
                    elif nx.kind == "unit" and e.kind == "table" and not e.data.get("unit"):
                        e.data["unit"] = nx.text  # 표 뒤에 적힌 단위 줄도 표 위 단위 행으로 끌어올림
                    else:
                        break
                    consumed += 1
                e.data["sources"] = srcs
                skip = consumed
            if e.kind == "heading":
                self.heading(e.level, e.text)
            elif e.kind in ("square", "circle", "dash", "dot", "note"):
                self.last_para = self.para(e.kind, e.text, footnotes)
            elif e.kind == "unit":
                self.para("unit", e.text)
            elif e.kind == "source":  # 표·그림에 붙지 않은 출처 줄은 직전 문단의 각주로 이동
                if self.last_para is not None:
                    self.last_para.add_footnote(e.text)
                else:
                    self.para("note", e.text, footnotes)
            elif e.kind == "table":
                self.table(e)
            elif e.kind == "image":
                self.image(e)
            elif e.kind == "pagebreak":
                self.pending_break = True
        if self.toc or self.cover:
            self._front_matter(self.cover)
        self.set_page_number(self.page_number)
        self.set_header(self.header_spec)
        report = self.doc.validate()
        issues = getattr(report, "issues", [])
        if issues:
            print(f"WARN: 스키마 검증 이슈 {len(issues)}건")
            for it in list(issues)[:5]:
                print("   ", it)
        self.doc.save_to_path(os.path.abspath(out))
        self.doc.close()


def repair_manifest(path: str) -> int:
    """content.hpf의 opf:item을 header.xml의 binItem 및 BinData 파일과 대조하고 누락분을 등록한다.
    Why: 라이브러리 버전에 따라 다중 이미지 삽입 시 첫 항목만 매니페스트에 등록되는 사례가 있고,
    이 경우 한/글이 나머지 그림을 열지 못한다. 정합하면 파일을 건드리지 않는다."""
    MEDIA = {"jpg": "image/jpeg", "jpeg": "image/jpeg", "png": "image/png", "gif": "image/gif",
             "bmp": "image/bmp", "tif": "image/tiff", "tiff": "image/tiff", "svg": "image/svg+xml"}
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        data = {n: z.read(n) for n in names}
    hpf = data["Contents/content.hpf"].decode("utf-8")
    header = data.get("Contents/header.xml", b"").decode("utf-8")
    have = set(re.findall(r'<opf:item[^>]*href="BinData/([^"]+)"', hpf))
    wanted: list[str] = []
    for f in re.findall(r'<hh:binItem[^>]*BinData="([^"]+)"', header):  # 1순위: header.xml binItem
        if f not in wanted:
            wanted.append(f)
    for n in names:  # 2순위: 실제 BinData 파일 (binItem에 없는 파일까지 포괄)
        if n.startswith("BinData/") and n.split("/", 1)[1] not in wanted:
            wanted.append(n.split("/", 1)[1])
    missing = [f for f in wanted if f not in have and f"BinData/{f}" in names]
    if not missing:
        return 0
    items = "".join(
        f'<opf:item id="{os.path.splitext(f)[0]}" href="BinData/{f}" '
        f'media-type="{MEDIA.get(os.path.splitext(f)[1].lstrip(".").lower(), "application/octet-stream")}" '
        f'isEmbeded="1"/>' for f in missing)
    hpf = hpf.replace("</opf:manifest>", items + "</opf:manifest>", 1)
    data["Contents/content.hpf"] = hpf.encode("utf-8")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr(zipfile.ZipInfo("mimetype"), data["mimetype"], compress_type=zipfile.ZIP_STORED)
        for n in names:
            if n != "mimetype":
                z.writestr(n, data[n])
    return len(missing)


def cmd_build(a):
    profile = PROFILES.get(a.profile) if a.profile in PROFILES else json.load(open(a.profile, encoding="utf-8"))
    skeleton = a.skeleton or os.path.join(ASSET_DIR, f"{a.profile}.hwpx")
    if not os.path.exists(skeleton):
        sys.exit(f"ERROR: 스켈레톤 없음 {skeleton}")
    elems, footnotes = parse_markdown(a.input)
    for msg in lint(elems):
        print("LINT:", msg)
    out = a.output or os.path.splitext(a.input)[0] + ".hwpx"
    def make(numbers=None):
        b = Builder(skeleton, profile, page_break=a.page_break)
        b.numbered_square = a.numbered_square
        b.toc = a.toc
        b.page_number = a.page_number
        b.toc_depth = a.toc_depth
        b.cover = dict(kv.split("=", 1) for kv in a.cover.split(";") if "=" in kv) if a.cover else None
        b.header_spec = a.header
        b.merge_mode = a.merge
        b.long_table_mode = a.long_table
        b.toc_numbers = numbers or {}
        b.build(elems, footnotes, out)
        return b
    b = make()
    fixed = repair_manifest(out)
    if fixed:
        print(f"WARN: 매니페스트 미등록 이미지 {fixed}건 자동 보정 (python-hwpx 버전 확인 권장)")
    if a.toc and a.page_number:  # 2패스: 렌더링으로 제목의 쪽 위치를 찾아 목차에 기입 (근사값, 한/글에서 확정)
        numbers = toc_page_numbers(out, [t for _, t in b.headings], a.rhwp)
        if numbers:
            make(numbers)
            repair_manifest(out)
            listed = [t for lv, t in b.headings if lv <= b.toc_depth]  # Why: 제목 수집은 항까지, 목차는 toc_depth까지
            print(f"목차 쪽번호 {sum(t in numbers for t in listed)}/{len(listed)}건 기입 (근사 렌더러 기준)")
        else:
            print("WARN: rhwp 또는 pdftotext 부재로 목차 쪽번호 미기입")
    print(f"OK: {out} (요소 {len(elems)}, 각주 {len(footnotes)})")


def toc_page_numbers(hwpx_path: str, headings: list[str], rhwp_opt: str | None) -> dict[str, int]:
    if not shutil.which("pdftotext") or not (rhwp_opt or find_rhwp() or find_soffice()):
        return {}
    pdf = os.path.splitext(hwpx_path)[0] + ".toc.pdf"
    engine = resolve_engine("auto", rhwp_opt)
    work = render_copy(hwpx_path, strip_newnum=(engine == "soffice"))
    try:
        to_pdf(work, pdf, engine, rhwp_opt, text_as_paths=False)
    except SystemExit:
        os.remove(work)
        return {}
    os.remove(work)
    txt = subprocess.run(["pdftotext", "-layout", pdf, "-"], capture_output=True, text=True).stdout
    os.remove(pdf)
    return match_toc_pages(txt.split("\f"), headings)


def _toc_norm(t: str) -> str:
    return re.sub(r"[\s.]+", "", t)  # 밴드는 번호와 제목을 다른 칸에 두므로 마침표도 무시


TOC_TITLES = {"목차", "표목차", "그림목차"}   # 목차 쪽의 제목 줄(공백 제거 후 전체 일치)
CHAPTER_TOKEN_RE = re.compile(r"제\d+장|[ⅠⅡⅢⅣⅤⅥⅦⅧⅨⅩ]")   # 간지와 밴드에서 제목 사이에 끼어드는 장 번호
TOC_TOP_LINES = 6   # 장 제목을 우선 찾는 쪽 머리 줄 수. 장은 새 쪽에서 시작하므로 쪽 머리에 제목이 옴
TOC_TITLE_WRAP_LINES = 3   # 장 제목 1개가 걸치는 최대 줄 수. 간지의 긴 제목은 두 줄로 접히고 그 사이에 장 번호 줄이 끼어듦


def match_toc_pages(pages: list[str], headings: list[str], extras: list[tuple[str, str]] | None = None,
                    top_titles: set[str] | None = None) -> dict[str, int]:
    """렌더링 PDF의 쪽별 텍스트에서 목차 항목의 본문 쪽번호(본문 첫 쪽 = 1)를 찾음
    headings: 제목 문구. extras: (목차 항목 문구, 찾을 문구) 목록. 캡션은 번호를 빼고 닫는 괄호와 제목으로 찾음
    top_titles: 쪽 머리에서 먼저 찾을 제목(장 제목). 본문 표에 같은 문구가 앞서 나와도 장 시작 쪽을 고름
    Why: 근사 렌더러는 캡션 자동번호를 비워서 그림. 번호 없이도 찾도록 닫는 괄호 뒤 제목을 찾을 문구로 씀.
    간지의 긴 장 제목은 두 줄로 접히고 pdftotext가 장 번호를 두 줄 사이에 끼워 넣으므로 장 번호를 지우고 찾음"""
    norm = _toc_norm

    def clean(t: str) -> str:  # 렌더러가 "제 1장"처럼 글자 사이에 빈칸을 넣으므로 빈칸 제거 후 장 번호를 지움
        return CHAPTER_TOKEN_RE.sub("", norm(t))

    def is_title_page(page: str) -> bool:  # 본문 문장 속 "목차" 낱말은 제목 줄이 아니므로 제외
        return any(norm(ln) in TOC_TITLES for ln in page.splitlines())
    keys = [norm(t) for t in headings + [e for e, _ in extras or []]]

    def entry_index(line: str) -> int | None:
        return next((i for i, k in enumerate(keys) if k and line.startswith(k)), None)

    def page_indexes(page: str) -> list[int | None]:
        return [entry_index(ln) for ln in (norm(x) for x in page.splitlines()) if ln]
    # 목차 구간: 첫 목차 제목 쪽부터 제목 쪽과 이어짐 쪽이 연속되는 곳까지
    # 이어짐 쪽은 모든 줄이 목차 항목이고 첫 줄이 앞 쪽 마지막 항목보다 뒤의 항목인 쪽
    # Why 순서 조건: 간지 쪽(장 제목 1줄)도 목차 항목 줄만 있는 쪽이라 순서로 구분
    first = next((i for i, page in enumerate(pages) if is_title_page(page)), None)
    body_start = 0
    if first is not None:
        i, last = first, -1
        while i < len(pages):
            idx = page_indexes(pages[i])
            if is_title_page(pages[i]):
                last = max([x for x in idx if x is not None] + [last])
            elif idx and None not in idx and idx[0] > last:
                last = max(idx)
            else:
                break
            i += 1
        body_start = i
    top = top_titles or set()

    def top_lines(page: str) -> list[str]:
        return [x for x in page.splitlines() if x.strip()][:TOC_TOP_LINES]

    def title_alone_at_top(page: str, key: str) -> bool:
        # Why: 밴드와 간지의 장 제목은 줄을 단독으로 차지함. 표 셀 줄은 같은 행의 다른 열 글자와 한 줄에 추출됨.
        # 부분 일치만 쓰면 쪽 머리 표 셀의 장 제목 문구를 장 시작으로 오판함(6차 시범 R 트랙 재현)
        lines = top_lines(page)
        return any(clean("".join(lines[i:j])) == key
                   for i in range(len(lines)) for j in range(i + 1, min(i + TOC_TITLE_WRAP_LINES, len(lines)) + 1))
    found: dict[str, int] = {}
    for group in ([(h, h) for h in headings], extras or []):
        pos = body_start
        for entry, needle in group:
            key = clean(needle)
            if not key:
                continue
            hit = None
            if entry in top:
                hit = next((i for i in range(pos, len(pages)) if title_alone_at_top(pages[i], key)), None)
            if entry in top and hit is None:  # 제목 줄이 다른 글자와 섞여 추출된 쪽은 종전의 부분 일치로 찾음
                hit = next((i for i in range(pos, len(pages)) if key in clean("\n".join(top_lines(pages[i])))), None)
            if hit is None:
                hit = next((i for i in range(pos, len(pages)) if key in clean(pages[i])), None)
            if hit is not None:
                found[entry] = hit - body_start + 1
                pos = hit
    return found


# ═══════════════════════════════════════════════════════════════════════
# 4. skeleton: 참조 hwpx에서 스타일 시트와 본문 구역 설정만 남긴 빈 양식 추출
# ═══════════════════════════════════════════════════════════════════════
def cmd_skeleton(a):
    src = zipfile.ZipFile(a.input)
    secs = sorted(n for n in src.namelist() if re.match(r"Contents/section\d+\.xml", n))
    sec_name = f"Contents/section{a.section}.xml" if a.section is not None else secs[-1]
    s = src.read(sec_name).decode("utf-8")
    first = s[: s.find("</hp:p>") + 7]
    first = re.sub(r"<hp:t\b[^>]*>.*?</hp:t>", "<hp:t/>", first, flags=re.S)
    first = re.sub(r"<hp:t\b[^>]*/>", "<hp:t/>", first)
    first = re.sub(r'(<hp:p\b[^>]*?)paraPrIDRef="\d+"', r'\1paraPrIDRef="0"', first, count=1)
    first = re.sub(r'(<hp:p\b[^>]*?)pageBreak="1"', r'\1pageBreak="0"', first, count=1)
    hpf = src.read("Contents/content.hpf").decode("utf-8")
    hpf = re.sub(r'<opf:item id="section[1-9]\d*"[^>]*/>', "", hpf)
    hpf = re.sub(r'<opf:itemref idref="section[1-9]\d*"[^>]*/>', "", hpf)
    hpf = re.sub(r'<opf:item [^>]*href="BinData/[^"]*"[^>]*/>', "", hpf)
    hpf = re.sub(r'(<opf:meta name="(creator|lastsaveby|date|CreatedDate|ModifiedDate)"[^>]*>)[^<]*(</opf:meta>)',
                 r"\1\3", hpf)  # 작성자와 일시 메타데이터 제거
    with zipfile.ZipFile(a.output, "w", zipfile.ZIP_DEFLATED) as z:
        for it in src.infolist():
            n = it.filename
            if n.startswith("BinData/") or re.match(r"Contents/section[1-9]", n) or n == "Preview/PrvImage.png":
                continue
            data = src.read(n)
            if n == "Contents/section0.xml":
                data = (first + "</hs:sec>").encode("utf-8")
            elif n == "Preview/PrvText.txt":
                data = b""  # 원본 미리보기 텍스트 제거
            elif n == "Contents/content.hpf":
                data = hpf.encode("utf-8")
            elif n == "Contents/header.xml":
                data = re.sub(rb'secCnt="\d+"', b'secCnt="1"', data)
            if n == "mimetype":
                z.writestr(zipfile.ZipInfo("mimetype"), data, compress_type=zipfile.ZIP_STORED)
            else:
                z.writestr(it, data)
    from hwpx import HwpxDocument
    doc = HwpxDocument.open(os.path.abspath(a.output))
    names = [s.name for s in doc.styles.values()]
    bullets = {bid: b.char for bid, b in doc.styles.bullets.items()} if hasattr(doc.styles, "bullets") else {}
    doc.close()
    profile = {"numbering": "manual", "chapter_format": "제{n}장 {t}"}
    guess = {"chapter": ("장 제목", "장"), "section": ("절 제목", "1.1", "절"), "subsection": ("1.1.1",), "ganada": ("가.", "가. 제목"),
             "square": ("□", "□ 내용"), "circle": ("❍ 내용1", "◦ 내용", "○"), "dash": ("- 내용2", "- 내용"),
             "dot": ("․ 내용3", "․ 내용"), "note": ("※ 내용4", "* 내용"), "table_pos": ("표위치", "표 위치"),
             "table_cell": ("표내용", "표 내용"), "table_cell_left": ("표내용(양쪽정렬)", "표내용", "표 내용"),
             "table_title": ("표 제목",), "unit": ("표,그림의 단위", "표 제목"), "source": ("표.그림의 자료및주", "출처", "* 내용"),
             "fig_pos": ("그림위치",), "fig_title": ("그림 제목",), "body": ("바탕글",)}
    for role, cands in guess.items():
        profile[role] = next((c for c in cands if c in names), None)
    prof_path = os.path.splitext(a.output)[0] + ".profile.json"
    json.dump(profile, open(prof_path, "w", encoding="utf-8"), ensure_ascii=False, indent=2)
    print(f"OK: {a.output} ({os.path.getsize(a.output)} bytes, 구역 {sec_name})")
    print("스타일:", names)
    print("글머리:", bullets)
    print(f"프로파일 초안: {prof_path}: None 항목은 스타일명을 직접 지정한 뒤 build --profile 로 전달")


# ═══════════════════════════════════════════════════════════════════════
# 5. export / fill / render / check
# ═══════════════════════════════════════════════════════════════════════
ROLE_MD = {"chapter": "# ", "section": "## ", "subsection": "### ", "ganada": "#### ", "square": "□ ",
           "circle": "- ", "dash": "    + ", "dot": "        + ", "note": "※ ", "source": "> ",
           "table_title": "<표> ", "fig_title": "[그림] "}  # source 역할은 "> " 접두
LEADING_MARK_RE = re.compile(r"^(제\s*\d+\s*장|\d+(?:\.\d+)*\.?|[가-힣]\.|[□■○❍◦ㅇ\-ㆍ·※*]|<표[^>]*>|\[그림[^\]]*\])\s*")


def guess_role(text: str) -> str | None:
    """스타일명이 없는 문서(수동 기호 타이핑)용 역할 추정."""
    for pat, role in ((r"^제\s*\d+\s*장\b", "chapter"), (r"^\d+\.\d+\.\d+\s", "subsection"), (r"^\d+\.\d+\s", "section"),
                      (r"^[가-힣]\.\s", "ganada"), (r"^[□■]", "square"), (r"^[○❍◦ㅇ]", "circle"),
                      (r"^-\s", "dash"), (r"^[ㆍ·•]", "dot"), (r"^[※*]", "note"), (r"^(자료|출처)\s*[:：]", "source"),
                      (r"^<표", "table_title"), (r"^\[그림", "fig_title")):
        if re.match(pat, text):
            return role
    return None


SYMBOL_ROLE = {"□": "square", "■": "square", "○": "circle", "❍": "circle", "◦": "circle", "ㅇ": "circle",
               "-": "dash", "\u00ad": "dash", "ㆍ": "dot", "·": "dot", "․": "dot", "※": "note", "*": "note"}
TEXT_ROLE_RE = [(re.compile(r"^제\s*\d+\s*장\s"), "chapter"), (re.compile(r"^\d+\.\d+\.\d+\s"), "subsection"),
                (re.compile(r"^\d+\.\d+\s"), "section"), (re.compile(r"^[가-힣]\.\s"), "ganada"),
                (re.compile(r"^[<＜]\s*표"), "table_title"), (re.compile(r"^\[\s*그림"), "fig_title"),
                (re.compile(r"^(자료|출처)\s*[:：]"), "source")]


def _guess_role(text: str, bullet_char: str | None) -> str | None:
    """스타일명이 규범과 다른 문서(참조 보고서류)용: 글머리 문자 → 선행 기호 → 번호 패턴 순으로 계위 추정."""
    if bullet_char and bullet_char in SYMBOL_ROLE:
        return SYMBOL_ROLE[bullet_char]
    if text and text[0] in SYMBOL_ROLE and (len(text) == 1 or text[1] == " "):
        return SYMBOL_ROLE[text[0]]
    for rx, role in TEXT_ROLE_RE:
        if rx.match(text):
            return role
    return None


def export_hierarchy(doc, profile: dict, image_dir: str | None) -> str:
    """스타일명 역대응 후 실패분은 기호·번호 패턴으로 추정하여 계위 규범 Markdown으로 복원 (재편집 후 build 왕복용)."""
    rev = {v: k for k, v in profile.items() if isinstance(v, str) and k in ROLE_MD}
    names = {sid: st.name for sid, st in doc.styles.items()}
    HH = "{http://www.hancom.co.kr/hwpml/2011/head}"
    hx = doc.headers[0].element
    bullet_char = {b.get("id"): b.get("char") for b in hx.iter(f"{HH}bullet")}
    pp_bullet = {}
    for pp in hx.iter(f"{HH}paraPr"):
        hd = pp.find(f"{HH}heading")
        if hd is not None and hd.get("type") == "BULLET":
            pp_bullet[pp.get("id")] = bullet_char.get(hd.get("idRef"))
    images = {it.item_id: it for it in doc.list_images()} if hasattr(doc, "list_images") else {}
    out, img_n = [], 0
    if image_dir:
        os.makedirs(image_dir, exist_ok=True)
    for sec in doc.sections:
        for p in sec.paragraphs:
            pics = list(p.element.iter(f"{HP}pic"))
            if pics and image_dir:
                for pic in pics:
                    ref = next((x for x in pic.iter() if x.tag.endswith("}img")), None)
                    item = images.get(ref.get("binaryItemIDRef")) if ref is not None else None
                    if item is None:
                        continue
                    img_n += 1
                    fn = f"img{img_n:03d}.{item.format}"
                    open(os.path.join(image_dir, fn), "wb").write(doc.package.read(item.href))
                    out.append(f"![]({os.path.basename(image_dir)}/{fn})")
            if p.tables:
                for t in p.tables:
                    grid, seen = [], set()
                    for row in t.get_cell_map():  # 병합 셀은 첫 위치에만 내용, 나머지는 빈 칸
                        line = []
                        for pos in row:
                            key = id(pos.cell.element)
                            line.append("" if key in seen else pos.cell.text.replace("\n", "<br>"))
                            seen.add(key)
                        grid.append(line)
                    cap = t.caption.text.strip() if getattr(t, "caption", None) else ""
                    if cap:
                        out.append("<표> " + LEADING_MARK_RE.sub("", cap))
                    out.append("| " + " | ".join(grid[0]) + " |")
                    out.append("|" + "---|" * len(grid[0]))
                    out.extend("| " + " | ".join(r) + " |" for r in grid[1:])
                    out.append("")
                continue
            text = p.text.strip()
            if not text:
                continue
            role = rev.get(names.get(str(p.style_id_ref), ""), None) or guess_role(text)
            if role:
                out.append(ROLE_MD[role] + LEADING_MARK_RE.sub("", text))
            else:
                out.append(text)
    return "\n".join(out) + "\n"


def cmd_export(a):
    from hwpx import HwpxDocument
    doc = HwpxDocument.open(os.path.abspath(a.input))
    if a.profile:
        profile = PROFILES.get(a.profile) if a.profile in PROFILES else json.load(open(a.profile, encoding="utf-8"))
        md = export_hierarchy(doc, profile, a.image_dir)
        out = a.output or os.path.splitext(a.input)[0] + ".md"
        open(out, "w", encoding="utf-8").write(md)
        print(f"OK: {out} ({len(md)}자, 계위 복원)")
        return
    kw = {"image_dir": a.image_dir, "image_ref_prefix": os.path.basename(a.image_dir) + "/"} if a.image_dir else {}
    md = doc.export_rich_markdown(detect_headings=True, **kw)
    out = a.output or os.path.splitext(a.input)[0] + ".md"
    open(out, "w", encoding="utf-8").write(md)
    print(f"OK: {out} ({len(md)}자)")


def cmd_fill(a):
    from hwpx import HwpxDocument
    spec = json.load(open(a.mapping, encoding="utf-8"))
    doc = HwpxDocument.open(os.path.abspath(a.input))
    total = 0
    for old, new in spec.get("replace", {}).items():
        total += doc.replace_text_in_runs(old, new) or 0
    if spec.get("cells"):  # "라벨 > right" 경로로 표 셀 채우기
        res = doc.fill_by_path(spec["cells"])
        print(f"표 셀: 적용 {res['applied_count']} 실패 {res['failed_count']}")
    doc.save_to_path(os.path.abspath(a.output or a.input))
    print(f"OK: 텍스트 치환 {total}건 → {a.output or a.input}")


def find_rhwp() -> str | None:
    for cand in (os.path.join(SCRIPT_DIR, "rhwp", "rhwp"), os.path.join(SCRIPT_DIR, "rhwp", "rhwp.exe"),
                 shutil.which("rhwp"), os.path.expanduser("~/.local/bin/rhwp")):
        if cand and os.path.isfile(cand) and os.access(cand, os.X_OK):
            return cand
    return None


def rhwp_supports(rhwp: str, command: str, flag: str | None = None) -> bool:
    """명령·플래그 존재를 help 출력으로 확인. 미성숙 프로젝트라 버전마다 CLI가 달라질 수 있다."""
    res = subprocess.run([rhwp, command, "--help"], capture_output=True, text=True)
    text = res.stdout + res.stderr
    if res.returncode != 0 and "Usage" not in text and "사용법" not in text:
        return False
    return flag is None or flag in text


def writable_dir(path: str) -> str:
    """산출물을 둘 디렉터리. 원본 옆이 읽기 전용(/mnt/skills, uploads)이면 현재 작업 디렉터리를 쓴다."""
    d = os.path.dirname(os.path.abspath(path))
    return d if os.access(d, os.W_OK) else os.getcwd()


def render_copy(src: str, strip_newnum: bool = False) -> str:
    """검수 전용 사본. 근사 렌더러 한계 회피용 치환만 넣고 납품물은 무변경.
    - 소프트 하이픈 글머리 → 하이픈 (근사 렌더러가 그리지 않음)
    - strip_newnum: 본문 중간 newNum 컨트롤 제거 (H2Orestart 전용. 첫 문단 밖의 newNum에서 변환 실패)"""
    work = os.path.join(writable_dir(src), "_render_copy.hwpx")  # ASCII 파일명 (H2Orestart 로케일 문제 회피)
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(work, "w", zipfile.ZIP_DEFLATED) as zout:
        for it in zin.infolist():
            d = zin.read(it.filename)
            if it.filename.startswith("Contents/section") and it.filename.endswith(".xml"):
                t = d.decode("utf-8").replace("\u00ad", "-")
                if strip_newnum:
                    t = re.sub(r'<hp:run charPrIDRef="\d+"><hp:ctrl><hp:newNum[^>]*/></hp:ctrl></hp:run>', "", t)
                d = t.encode("utf-8")
            if it.filename == "mimetype":
                zout.writestr(zipfile.ZipInfo("mimetype"), d, compress_type=zipfile.ZIP_STORED)
            else:
                zout.writestr(it.filename, d)
    return work


def find_soffice() -> str | None:
    return shutil.which("soffice") or shutil.which("libreoffice")


def resolve_engine(engine: str, rhwp_opt: str | None) -> str:
    rhwp = rhwp_opt or find_rhwp()
    if engine == "auto":
        return "rhwp" if rhwp else ("soffice" if find_soffice() else "")
    return engine


def to_pdf(src: str, pdf: str, engine: str, rhwp_opt: str | None, text_as_paths: bool = True) -> str:
    """hwpx → PDF. engine: rhwp | soffice (auto는 resolve_engine으로 사전 확정). 사용한 엔진명을 반환.
    text_as_paths=False: 글자를 path로 굳히지 않아 PDF에 텍스트 레이어가 남는다.
    Why: 목차 쪽번호는 pdftotext로 제목의 쪽 위치를 읽어 산정한다. path 변환본은 추출 결과가 0자."""
    rhwp = rhwp_opt or find_rhwp()
    soffice = find_soffice()
    engine = resolve_engine(engine, rhwp_opt)
    if engine == "rhwp":
        if not rhwp:
            sys.exit("ERROR: rhwp 미확보: SKILL.md 준비 절의 절차로 확보 후 --rhwp 경로 지정 또는 scripts/rhwp/ 에 배치")
        if not rhwp_supports(rhwp, "export-pdf"):
            sys.exit("ERROR: 이 rhwp 버전에 export-pdf 명령이 없음: `rhwp --help`로 PDF 출력 명령을 확인하여 수동 실행")
        cmd = [rhwp, "export-pdf", src, "-o", pdf]
        if text_as_paths and rhwp_supports(rhwp, "export-pdf", "--text-as-paths"):
            cmd.append("--text-as-paths")
        res = subprocess.run(cmd, capture_output=True, text=True)
        if res.returncode != 0 or not os.path.exists(pdf):
            sys.exit(f"ERROR: rhwp 변환 실패 (문서 결함 신호일 수 있음): {(res.stderr or res.stdout)[-400:]}")
        return "rhwp"
    if engine == "soffice":
        if not soffice:
            sys.exit("ERROR: LibreOffice 미확보: SKILL.md 준비 절의 폴백 설치 절차 참조")
        outdir = os.path.join(writable_dir(src), "_lo_out")
        shutil.rmtree(outdir, ignore_errors=True)
        env = {**os.environ, "LANG": os.environ.get("LANG") or "C.UTF-8"}
        # 호출마다 별도 프로필 사용: 고정 프로필은 동시 실행되는 다른 조판 세션과 잠금이 충돌
        profile = tempfile.mkdtemp(prefix="lo-profile-")
        cmd = [soffice, f"-env:UserInstallation={pathlib.Path(profile).as_uri()}", "--headless",
               "--convert-to", "pdf:writer_pdf_Export", "--outdir", outdir, src]
        try:
            res = subprocess.run(cmd, capture_output=True, text=True, env=env, timeout=600)
        finally:
            shutil.rmtree(profile, ignore_errors=True)
        made = os.path.join(outdir, os.path.splitext(os.path.basename(src))[0] + ".pdf")
        is_hwpx = src.lower().endswith(".hwpx")
        if not os.path.exists(made):
            sys.exit(("ERROR: LibreOffice(H2Orestart) 변환 실패: H2Orestart 확장 미설치이거나 문서 결함(요소 순서, 참조 ID) 신호: "
                      if is_hwpx else "ERROR: LibreOffice 변환 실패 (문서 결함 신호일 수 있음): ") + (res.stdout + res.stderr)[-300:])
        shutil.move(made, pdf)
        shutil.rmtree(outdir, ignore_errors=True)
        return "soffice(H2Orestart)" if is_hwpx else "soffice"
    sys.exit("ERROR: 렌더러 없음: rhwp 또는 LibreOffice+H2Orestart 확보 후 재실행 (SKILL.md 준비 절)")


PAGE_PNG_RE = re.compile(r"p-\d+\.png")   # pdftoppm 쪽 PNG 이름(접두 p). render가 지우고 세는 대상


def cmd_render(a):
    src = a.input
    is_docx = src.lower().endswith(".docx")
    engine = "soffice" if is_docx else resolve_engine(a.engine, a.rhwp)  # docx는 LibreOffice만 사용
    work = src if a.keep_soft_hyphen or is_docx else render_copy(src, strip_newnum=(engine == "soffice"))
    stem = os.path.join(writable_dir(src), os.path.splitext(os.path.basename(src))[0])
    outdir = a.outdir or stem + "_pages"
    os.makedirs(outdir, exist_ok=True)
    # Why: 이전 렌더의 쪽 PNG만 지움. 폴더 전체를 비우면 --outdir로 원본이나 다른 파일이 있는 폴더를 줄 때 그 파일까지 삭제됨
    for f in os.listdir(outdir):
        if PAGE_PNG_RE.fullmatch(f):
            os.remove(os.path.join(outdir, f))
    # Why: --outdir 지정 시 PDF도 그 폴더에 둠. 엔진 교차 확인 렌더가 기존 PDF를 덮어쓰지 않게 함
    pdf = os.path.join(outdir, os.path.basename(stem) + ".pdf") if a.outdir else stem + ".pdf"
    used = to_pdf(work, pdf, engine, a.rhwp)
    if work != src:
        os.remove(work)
    if shutil.which("pdftoppm"):
        subprocess.run(["pdftoppm", "-png", "-r", str(a.dpi), pdf, os.path.join(outdir, "p")], check=True)
        n_png = len([f for f in os.listdir(outdir) if PAGE_PNG_RE.fullmatch(f)])
        print(f"OK[{used}]: {pdf} → {outdir}/ ({n_png}쪽)")
    else:
        print(f"OK[{used}]: {pdf} (pdftoppm 부재로 PNG 생략)")


def cmd_check(a):
    from hwpx import HwpxDocument
    doc = HwpxDocument.open(os.path.abspath(a.input))
    report = doc.validate()
    issues = list(getattr(report, "issues", []))
    with zipfile.ZipFile(a.input) as z:
        hpf = z.read("Contents/content.hpf").decode("utf-8")
        have = set(re.findall(r'<opf:item[^>]*href="BinData/([^"]+)"', hpf))
        binitems = set(re.findall(r'<hh:binItem[^>]*BinData="([^"]+)"', z.read("Contents/header.xml").decode("utf-8")))
        files = {n.split("/", 1)[1] for n in z.namelist() if n.startswith("BinData/")}
        sec = "".join(z.read(n).decode("utf-8") for n in z.namelist() if n.startswith("Contents/section"))
    refs = set(re.findall(r'binaryItemIDRef="([^"]+)"', sec))
    stems = {os.path.splitext(f)[0] for f in files}
    if files - have:
        issues.append(f"매니페스트 미등록 이미지: {sorted(files - have)}")
    if binitems - files:
        issues.append(f"binItem에 대응하는 BinData 파일 없음: {sorted(binitems - files)}")
    if refs - stems:
        issues.append(f"본문 참조에 대응하는 BinData 없음: {sorted(refs - stems)}")
    print(f"이미지: BinData {len(files)}, binItem {len(binitems)}, 매니페스트 {len(have)}, 본문 참조 {len(refs)}")
    print(f"스키마 검증: 이슈 {len(issues)}건")
    for it in issues[:10]:
        print("   ", it)
    doc.close()
    rhwp = find_rhwp()
    if rhwp and a.expect and not rhwp_supports(rhwp, "verify"):
        print("rhwp verify 명령이 이 버전에 없어 생략")
    elif rhwp and a.expect:
        import shlex
        cmd = [rhwp, "verify", a.input, "--json"] + shlex.split(a.expect)
        res = subprocess.run(cmd, capture_output=True, text=True)
        try:
            d = json.loads(res.stdout)
            print("rhwp verify:", d.get("verdict"), [(e["kind"], e.get("actual"), e["pass"]) for e in d["expectations"]])
        except json.JSONDecodeError:
            print("rhwp verify 출력 해석 실패:", res.stdout[:200], res.stderr[:200])
    sys.exit(1 if issues else 0)


# ═══════════════════════════════════════════════════════════════════════
# 6. 입력 어댑터: DOCX, XLSX를 계위 규범 Markdown 원고로 (본문 재타이핑 없이)
# ═══════════════════════════════════════════════════════════════════════
def _docx_para_key(p) -> str:
    """문단의 계위 판별 키: 스타일명 | 왼쪽 들여쓰기 | 자동번호 여부 | 선두 기호."""
    pf = p.paragraph_format
    indent = int((pf.left_indent or 0) / 12700 * 10) if pf.left_indent is not None else 0
    numbered = p._p.pPr is not None and p._p.pPr.numPr is not None
    lead = p.text.lstrip()[:1] if p.text.strip() else ""
    lead = lead if lead in "□■○◦❍ㅇ-ㆍ·※*" else ""
    return f"{p.style.name}|{indent}|{'num' if numbered else ''}|{lead}"


def cmd_docx2md(a):
    import docx
    d = docx.Document(a.input)
    if not a.map:  # 1단계: 프로파일 출력 → 유저와 대응 확정
        keys: dict[str, list[str]] = {}
        for p in d.paragraphs:
            if p.text.strip():
                keys.setdefault(_docx_para_key(p), []).append(p.text.strip()[:40])
        print("키 = 스타일명|들여쓰기(pt*10)|자동번호|선두기호  →  건수, 표본. 대응 JSON 예: {\"본문|0||□\": \"square\"}")
        for k, v in sorted(keys.items(), key=lambda kv: -len(kv[1])):
            print(f"  {k!r}: {len(v)}건  예) {v[0]}")
        print("역할: chapter section subsection ganada square circle dash dot note source skip")
        return
    mp = json.load(open(a.map, encoding="utf-8"))
    md_of = {"chapter": "# ", "section": "## ", "subsection": "### ", "ganada": "#### ", "square": "□ ",
             "circle": "- ", "dash": "    + ", "dot": "        + ", "note": "※ ", "source": "> "}
    out = []
    body = d.element.body
    pi, ti = 0, 0
    for child in body.iterchildren():
        tag = child.tag.split("}")[1]
        if tag == "p":
            p = d.paragraphs[pi]
            pi += 1
            if not p.text.strip():
                continue
            role = mp.get(_docx_para_key(p), "circle")
            if role == "skip":
                continue
            out.append(md_of[role] + LEADING_MARK_RE.sub("", p.text.strip()))
        elif tag == "tbl":
            t = d.tables[ti]
            ti += 1
            rows = [[c.text.strip().replace("\n", "<br>") for c in r.cells] for r in t.rows]
            out.append("| " + " | ".join(rows[0]) + " |")
            out.append("|" + "---|" * len(rows[0]))
            out.extend("| " + " | ".join(r) + " |" for r in rows[1:])
            out.append("")
    o = a.output or os.path.splitext(a.input)[0] + ".md"
    open(o, "w", encoding="utf-8").write("\n".join(out) + "\n")
    print(f"OK: {o} (문단 {pi}, 표 {ti}): 표 캡션 <표> 줄과 그림은 원고에서 보완")


def cmd_xlsx2md(a):
    import openpyxl
    wb = openpyxl.load_workbook(a.input, data_only=True)
    out = []
    for ws in wb.worksheets:
        rows = [[("" if v is None else (f"{v:,}" if isinstance(v, (int, float)) and not isinstance(v, bool) else str(v)).strip())
                 for v in r] for r in ws.iter_rows(values_only=True)]
        rows = [r for r in rows if any(r)]
        if not rows:
            continue
        ncol = max(len(r) for r in rows)
        if ncol > 6 and not a.allow_wide:
            print(f"WARN: 시트 '{ws.title}' 열 {ncol}개: 행열 전환 또는 표 분할을 유저에게 질의 (--allow-wide 로 그대로 진행)")
        out.append(f"<표> {ws.title}")
        out.append("| " + " | ".join(rows[0]) + " |")
        out.append("|" + "---|" * len(rows[0]))
        out.extend("| " + " | ".join(r + [""] * (ncol - len(r))) + " |" for r in rows[1:])
        out.append("")
    o = a.output or os.path.splitext(a.input)[0] + ".md"
    open(o, "w", encoding="utf-8").write("\n".join(out) + "\n")
    print(f"OK: {o} (시트 {len(wb.worksheets)}): 캡션 제목, (단위: ) 줄, 자료: 줄은 원고에서 보완")


# ═══════════════════════════════════════════════════════════════════════
def main():
    ap = argparse.ArgumentParser(description="hwpx_report CLI", epilog=CLI_NOTICE,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)

    b = sub.add_parser("build", help="원고.md → 문서.hwpx")
    b.add_argument("input")
    b.add_argument("output", nargs="?")
    b.add_argument("--profile", default="chapter", help="chapter | flat | 프로파일 JSON 경로")
    b.add_argument("--skeleton", help="스켈레톤 hwpx (기본 assets/<profile>.hwpx)")
    b.add_argument("--page-break", default="chapter", choices=["chapter", "section", "none"],
                   help="새 쪽 시작 단위: chapter=# 장마다(기본), section=# 장과 ## 절마다, none=붙여쓰기 (유저 택일)")
    b.add_argument("--long-table", default="engine", choices=["engine", "split"],
                   help="25행 초과 표: engine=한/글 엔진 분할+제목행 반복, split=청크별 표로 물리 분할 (유저 택일)")
    b.add_argument("--toc", action="store_true", help="문서 앞에 장과 절의 정적 목차 삽입 (기본 OFF, 유저 요청 시)")
    b.add_argument("--page-number", action="store_true", help="꼬리말 가운데 쪽번호 '- N -' (기본 OFF, 취합 시 충돌 방지)")
    b.add_argument("--toc-depth", type=int, default=2, choices=[1, 2, 3], help="목차 계위 깊이: 1=장, 2=절(기본), 3=소절 (유저 택일)")
    b.add_argument("--cover", help='표지 (기본 OFF): "title=제목;subtitle=부제;date=YYYY. M.;org=기관명"')
    b.add_argument("--header", help='머리말 (기본 OFF): "텍스트;align=LEFT|CENTER|RIGHT"')
    b.add_argument("--merge", default="auto", choices=["auto", "explicit", "none"],
                   help="셀 병합: auto=선행 라벨 열 연속 동일값 + '^'(세로)·'<'(가로) 지시(기본), explicit=지시만, none=병합 없음")
    b.add_argument("--rhwp", help="목차 쪽번호 산정에 쓸 rhwp 경로")
    b.add_argument("--numbered-square", action="store_true",
                   help="'□ 1. 제목' 처럼 숫자로 시작하는 □ 줄을 글머리 없는 숫자 헤더로 출력 (장 없는 문서용)")
    b.set_defaults(func=cmd_build)

    s = sub.add_parser("skeleton", help="참조.hwpx → 빈 양식 + 프로파일 초안")
    s.add_argument("input")
    s.add_argument("output")
    s.add_argument("--section", type=int, help="본문 구역 번호 (기본 마지막 구역)")
    s.set_defaults(func=cmd_skeleton)

    e = sub.add_parser("export", help="문서.hwpx → 원고.md")
    e.add_argument("input")
    e.add_argument("output", nargs="?")
    e.add_argument("--image-dir")
    e.add_argument("--profile", help="지정 시 스타일명을 역대응하여 계위 규범 Markdown으로 복원 (chapter | flat | JSON)")
    e.set_defaults(func=cmd_export)

    f = sub.add_parser("fill", help='치환 JSON {"replace": {...}, "cells": {"라벨 > right": "값"}}')
    f.add_argument("input")
    f.add_argument("mapping")
    f.add_argument("-o", "--output")
    f.set_defaults(func=cmd_fill)

    r = sub.add_parser("render", help="rhwp export-pdf + pdftoppm")
    r.add_argument("input")
    r.add_argument("--outdir")
    r.add_argument("--dpi", type=int, default=RENDER_DPI)
    r.add_argument("--rhwp")
    r.add_argument("--keep-soft-hyphen", action="store_true", help="검수 사본 치환 없이 원본 그대로 렌더")
    r.add_argument("--engine", default="auto", choices=["auto", "rhwp", "soffice"],
                   help="auto=rhwp 우선, 없으면 LibreOffice+H2Orestart 폴백")
    r.set_defaults(func=cmd_render)

    c = sub.add_parser("check", help="스키마 검증 + rhwp verify")
    c.add_argument("input")
    c.add_argument("--expect", default="", help='rhwp verify 인자 문자열 예: "--expect-min-tables 2 --expect-contains 목차"')
    c.set_defaults(func=cmd_check)

    d = sub.add_parser("docx2md", help="DOCX → 원고.md (1단계: 프로파일 출력, 2단계: --map JSON으로 변환)")
    d.add_argument("input")
    d.add_argument("output", nargs="?")
    d.add_argument("--map", help="계위 대응 JSON 경로")
    d.set_defaults(func=cmd_docx2md)

    x = sub.add_parser("xlsx2md", help="XLSX → 원고.md (시트마다 파이프 표)")
    x.add_argument("input")
    x.add_argument("output", nargs="?")
    x.add_argument("--allow-wide", action="store_true")
    x.set_defaults(func=cmd_xlsx2md)


    a = ap.parse_args()
    a.func(a)


if __name__ == "__main__":
    main()
