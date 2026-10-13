#!/usr/bin/env python3
"""juja-builder (scripts/proposal.py): 연구용역 제안서 작성 파이프라인의 단일 CLI.

명령
  rfp       RFP(hwpx|pdf|docx|md) → 01_RFP원문.md(변환 텍스트), 01_요구항목표.md와 01_배점표.md 빈 표 (모델이 원문을 읽고 채움)
  deck      제안 발표자료(pptx|pdf) → 발표자료/(기준본 PDF, 원문.md, 쪽 PNG, 축소판 모음, 내장 이미지)
  figure    발표자료의 쪽 전체 또는 일부를 삽입용 도식 PNG로 캡처 (--grid는 자르기 값을 읽는 눈금 PNG)
  init      작업폴더 생성 (--kind proposal|report|plan. state.json, _양식/proposal.hwpx 기본 스켈레톤. 제안서는 00_제안전략서.md, 회사자료.md, 원고/41 61 62 공통 파트 템플릿 추가. 사전진단보고서는 00_진단서.md, 기획보고서는 00_기획설계.md 추가)
  ledger    01_배점표.md의 결정 섹션을 집계하여 02_섹션대장.md의 통제배점과 상태 갱신
  check     rfp | density | wording | tone | numbers | fill | volume | all (all은 산출 유형별 점검 목록)
  build     원고/*.md 합본 → 납품/{제안서|사전진단보고서|기획보고서}.hwpx, .docx (--format hwpx|docx|both. 스켈레톤은 실행마다 재생성)
  render    hwpx 또는 docx → PDF + 페이지 PNG (hwpx는 rhwp 우선, 부재 시 LibreOffice와 H2Orestart. docx는 LibreOffice)
  selftest  표본 원고 조판과 렌더링으로 렌더 경로 판정 (착수 관문. 결과는 state.json checks.render)
  skeleton  발주처 양식.hwpx → 스켈레톤.hwpx + 프로파일 초안

구성: hwpx_report.py는 형식 무관 조판 계층(원고 파서, 인라인 표기, 표 배치 산정, 그림 크기, 목차 순서)과 hwpx 작성기,
docx_report.py는 docx 작성기. 이 파일은 산출 유형별 판단(제목 번호, 밴드, 캡션 번호, 표 폭 지시)과 판정 규칙을 담당한다.
Why 분리: 내용 원천(MD 원고)과 양식 원천(설정부)을 하나로 두고 판단은 1회만 계산하여 hwpx와 docx가 같은 결과를 기록하게 한다.
동작 확인 환경: Python 3.11-3.12, python-hwpx 6.x, python-docx 1.2.0, lxml, pillow, poppler 24.02, rhwp v0.8.6, LibreOffice(docx 렌더)

형식 고지: 본 제품은 한컴의 HWP 문서 파일(.hwp) 공개 문서를 참고하여 개발하였습니다.
상표 고지: 한/글, 한컴오피스, HWPX는 한글과컴퓨터의 상표
비제휴 고지: 이 도구는 한글과컴퓨터와 제휴 관계가 없는 독립 도구
"""
from __future__ import annotations

import argparse
import glob
import json
import os
import re
import shutil
import subprocess
import sys
import copy
import warnings
import zipfile
from datetime import datetime

warnings.filterwarnings("ignore", category=DeprecationWarning)
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
ASSET_DIR = os.path.join(os.path.dirname(SCRIPT_DIR), "assets")
sys.path.insert(0, SCRIPT_DIR)
import hwpx_report as lib  # noqa: E402  파서, 표 배치 산정 같은 형식 무관 계층과 hwpx 작성기
import docx_report as docx_lib  # noqa: E402  docx 작성기
from lxml import etree  # noqa: E402  python-hwpx 의존성으로 함께 설치됨

HP = lib.HP
HH = "{http://www.hancom.co.kr/hwpml/2011/head}"
ROMAN = ["Ⅰ", "Ⅱ", "Ⅲ", "Ⅳ", "Ⅴ", "Ⅵ", "Ⅶ", "Ⅷ", "Ⅸ", "Ⅹ", "Ⅺ", "Ⅻ"]
GANADA = "가나다라마바사아자차카타파하"

# ═══════════════════════════════════════════════════════════════════════
# 설정부 1: 기본 스켈레톤과 제안서 프로파일
# ═══════════════════════════════════════════════════════════════════════
# 기본 스켈레톤은 init과 build가 python-hwpx의 빈 새 문서 위에 아래 정의를 얹어 작업폴더/_양식/에 생성한다.
# Why 코드 생성: 양식 파일을 스킬에 동봉하지 않고 모든 스타일 값을 이 설정부에서 재현 가능하게 둔다
SKELETON_FONT_BODY = "나눔바른고딕"   # fontface id 0: 본문, 표, 주석
# fontface id 1: 제목 계열. 굵은 글꼴이 같은 이름으로 묶인 글꼴을 써야 한/글과 근사 렌더러 모두에서 굵게가 표시됨
# (나눔스퀘어는 굵기별로 이름이 갈려 굵게 지정이 렌더러마다 다르게 표시됨)
SKELETON_FONT_HEAD = "나눔바른고딕"
SKELETON_TEXT_COLOR = "#000000"
SKELETON_BULLET_OFFSET_TYPE = "HWPUNIT"  # 기호와 본문 사이 간격 단위. PERCENT는 기호 폭 비례라 좁은 기호에서 간격이 사라짐
SKELETON_BULLET_OFFSET = "600"           # 기호와 본문 사이 간격 (6pt)
SKELETON_TOC_TAB_POS = 48100             # 목차 쪽번호 오른쪽 탭 위치 (본문 폭 48190에서 여유 90)
# 쪽 설정 (HWPUNIT, 1mm = 283.46)
SKELETON_PAGE_MARGIN = {"left": 5669, "right": 5669, "top": 4252, "bottom": 4252, "header": 2835, "footer": 2835}
# 계위 기호 (한/글 글머리 기능). 「행정업무의 운영 및 혁신에 관한 규정 시행규칙」 제2조의 특수기호 □, ○, -, · 체계
# 4계위(·)와 주석(※)은 기호를 본문 글자로 넣음(MANUAL_SYMBOLS). 좁은 기호도 간격이 렌더러마다 같게 보임
SKELETON_GLYPHS = {"square": "□", "circle": "○", "dash": "-", "cell_bullet": "·", "cell_dash": "-"}
SKELETON_HANG_DOT = -800     # 4계위 "· " 내어쓰기 (12pt에서 기호와 빈칸 폭)
SKELETON_HANG_NOTE = -1100   # 주석 "※ " 내어쓰기 (10pt에서 기호와 빈칸 폭)
# 스타일 정의. 이 표가 스켈레톤 스타일과 PROFILES의 역할 → 스타일명 대응의 단일 원천
# font: body|head, pt: 글자 크기, align, left/indent/prev/next: HWPUNIT, ls: 줄 간격 %,
# bullet: 글머리 키(SKELETON_GLYPHS), keep: 다음 문단과 함께, widow: 외톨이줄 보호
SKELETON_STYLES = [
    # 이름           역할 키          font    pt    bold   align      left  indent prev  next  ls   bullet         keep widow
    ("장 제목",      "chapter",       "head", 20.0, True,  "LEFT",    0,    0,     0,    0,    140, None,          1, 1),
    ("절 제목",      "section",       "head", 15.0, True,  "LEFT",    0,    0,     2500, 300,  140, None,          1, 1),
    ("항 제목",      "subsection",    "head", 14.0, True,  "LEFT",    0,    0,     2000, 200,  140, None,          1, 1),
    ("목 제목",      "ganada",        "head", 13.0, True,  "LEFT",    300,  0,     1500, 100,  140, None,          1, 1),
    ("논점",         "square",        "head", 13.0, True,  "JUSTIFY", 500,  0,     1500, 0,    140, "square",      1, 1),
    ("주장",         "circle",        "body", 12.0, False, "JUSTIFY", 1000, 0,     1000, 0,    140, "circle",      0, 1),
    ("근거",         "dash",          "body", 12.0, False, "JUSTIFY", 2410, 0,     700,  0,    140, "dash",        0, 1),
    ("상세",         "dot",           "body", 12.0, False, "JUSTIFY", 3410, SKELETON_HANG_DOT, 700, 0, 140, None, 0, 1),
    ("주석",         "note",          "body", 10.0, False, "JUSTIFY", 3000, SKELETON_HANG_NOTE, 400, 0, 140, None, 0, 1),
    ("표 위치",      "table_pos",     "body", 10.0, False, "CENTER",  0,    0,     800,  1000, 160, None,          0, 0),
    ("표 본문",      "table_cell",    "body", 10.0, False, "CENTER",  0,    0,     0,    0,    130, None,          0, 0),
    ("표 본문 왼쪽", "table_cell_left", "body", 10.0, False, "LEFT",  0,    0,     0,    0,    130, None,          0, 0),
    ("표 글머리",    "table_cell_bullet", "body", 10.0, False, "LEFT", 0,   0,     0,    150,  130, "cell_bullet", 0, 0),
    ("표 하이픈",    "table_cell_dash", "body", 10.0, False, "LEFT",  1000, 0,     0,    150,  130, "cell_dash",   0, 0),
    ("표 캡션",      "table_title",   "head", 10.5, True,  "CENTER",  0,    0,     0,    0,    150, None,          1, 0),
    ("그림 캡션",    "fig_title",     "head", 10.5, True,  "CENTER",  0,    0,     0,    0,    150, None,          1, 0),
    ("그림 위치",    "fig_pos",       "body", 10.0, False, "CENTER",  0,    0,     1000, 1000, 160, None,          0, 0),
    ("단위",         "unit",          "body", 9.0,  False, "RIGHT",   0,    0,     300,  0,    150, None,          0, 0),
    ("출처",         "source",        "body", 9.0,  False, "LEFT",    0,    0,     300,  0,    140, None,          0, 0),
    ("밴드 위치",    "band_pos",      "body", 10.0, False, "CENTER",  0,    0,     1200, 600,  160, None,          1, 0),   # 밴드를 담는 문단 (앞뒤 여백, 다음 문단과 함께)
    ("밴드 칸",      "band_cell",     "head", 12.0, True,  "LEFT",    0,    0,     0,    0,    130, None,          0, 0),   # 밴드 번호 칸과 제목 칸의 글자 스타일
    ("목차 장",      "toc_chapter",   "head", 14.0, True,  "LEFT",    0,    0,     1000, 0,    155, None,          0, 0),
    ("목차 절",      "toc_section",   "body", 12.0, False, "LEFT",    1500, 0,     500,  0,    155, None,          0, 0),
    ("문서 제목",    "doc_title",     "head", 20.0, True,  "CENTER",  0,    0,     0,    1500, 140, None,          1, 1),   # 기획보고서 첫 줄 (build --title)
]
# 기본 스타일: python-hwpx 빈 새 문서에 이미 있는 스타일. 스켈레톤 생성 시 이 값으로 덮어써서 hwpx와 docx가 같은 값을 읽음
SKELETON_BASE_STYLES = [
    ("바탕글",       "body",          "body", 10.0, False, "JUSTIFY", 0,    0,     0,    0,    160, None,          0, 0),
    ("쪽 번호",      "page_number",   "body", 10.0, False, "JUSTIFY", 0,    0,     0,    0,    160, None,          0, 0),
    ("머리말",       "header_text",   "body", 9.0,  False, "JUSTIFY", 0,    0,     0,    0,    150, None,          0, 0),
    ("각주",         "footnote",      "body", 9.0,  False, "LEFT",    0,    -1310, 0,    0,    130, None,          0, 0),   # 출처 URL과 서지 표기의 자간 벌어짐 방지
]
SKELETON_PAGE_SIZE = (59528, 84186)   # A4 폭, 높이 (HWPUNIT)
# 나눔바른고딕 글리프 폭(em). docx 글머리의 기호 폭 산정용. 한/글은 글꼴에서 직접 측정하므로 hwpx에는 쓰지 않음
SKELETON_GLYPH_EM = {"□": 0.892, "○": 0.892, "-": 0.583, "·": 0.446, "※": 0.892, " ": 0.224}
# 빈칸 폭(em). 스켈레톤 글자 모양은 '글꼴에 어울리는 빈칸'을 끈 상태(useFontSpace=0)라 한/글이 빈칸을 글자 크기의 절반으로 고정
# docx는 빈칸에 자간을 더해 같은 폭을 맞춤
SKELETON_SPACE_EM = 0.5
# lib.Builder는 프로파일 값 전부를 스타일명으로 검증하므로 스타일명이 아닌 키는 조판기에 넘길 때 뺌 (_pt로 끝나는 키 포함)
PROFILE_EXTRA_KEYS = ("item_format", "chapter_label", "chapter_band", "toc_lists")
PROFILES = {
    "proposal": {
        **lib.PROFILES["chapter"],
        # 역할 → 스타일명은 SKELETON_STYLES에서 가져옴. 계위 기호, 글꼴, 문단 여백은 스켈레톤의 스타일이 정함
        **{role: name for name, role, *_ in SKELETON_STYLES},
        "body": "바탕글",                    # python-hwpx 빈 새 문서의 기본 스타일
        "chapter_format": "{r}. {t}",       # Ⅰ. 제안 개요
        "chapter_label": "{r}",             # 장 밴드 번호 칸
        "section_format": "{s}. {t}",       # 1. 제안 배경 및 필요성
        "subsection_format": "{g}. {t}",    # 가. 항
        "item_format": "{i}) {t}",          # 1) 목
        "caption_prefix": None,             # None=무번호 캡션 <제목> (제안서 관행). "{c}-"를 주면 <표 1-1> 자동번호
        "chapter_pt": 2000, "section_pt": 1600, "subsection_pt": 1400,  # 위계 확보용 크기. None이면 스켈레톤 크기
    },
}
# 기획보고서: 제안서와 같은 계위(Ⅰ., 1., 가., 1))와 무번호 캡션. 문서 제목은 build --title
PROFILES["plan"] = dict(PROFILES["proposal"])
# 사전진단보고서: 최종보고서 관행(제1장, 1.1, 1.1.1, 가.)과 장별 캡션 번호(<표 1-1>). 목차에 표 목차와 그림 목차 포함
PROFILES["report"] = {
    **PROFILES["proposal"],
    "chapter_format": "제{n}장 {t}", "chapter_label": "제{n}장",
    "section_format": "{c}.{s} {t}", "subsection_format": "{c}.{s}.{ss} {t}", "item_format": "{g}. {t}",
    "caption_prefix": "{c}-", "toc_lists": True,
    "chapter_band": {"label_w": 7600, "label_pt": 1800},   # "제10장"까지 한 줄. 폭은 Module 밴드와 같아 제목 시작선이 맞음
}
# 산출 유형: init --kind로 state.json에 기록. build의 기본 프로파일과 납품 파일명을 정함
KINDS = {"proposal": {"profile": "proposal", "out": "제안서"},
         "report": {"profile": "report", "out": "사전진단보고서"},
         "plan": {"profile": "plan", "out": "기획보고서"}}
TOC_LIST_TITLES = ("표 목차", "그림 목차")
# 색 테마: 밴드에만 적용. build --theme으로 선택. base는 장 번호 칸 채움, tint는 Module 라벨 칸 채움,
# dark는 라벨 글자와 Task 라벨 테두리
THEMES = {
    "navy": {"base": "#6B8EC2", "tint": "#E3E9F3", "dark": "#35578F"},
    "mono": {"base": "#262626", "tint": "#D9D9D9", "dark": "#1A1A1A"},
}
DEFAULT_THEME = "navy"
# 표 서식은 테마와 무관하게 조판기 기본값(검정 0.15mm 괘선, 회색 #D9D9D9 머리글)을 유지. Why: 편집자가 한/글에서 표를 고칠 때 관행 서식과 같아야 함
lib.CELL_MARGIN = (510, 510, 280, 280)   # 셀 위아래 여백 1mm. 글머리 셀 첫 줄이 괘선에 붙지 않게 함
lib.FIG_CAPTION_SIDE = "BOTTOM"          # 표 캡션은 위, 그림 캡션은 아래 (관행). 캡션은 투명 표에 붙으므로 출처 행 다음에 옴
lib.TOC_UNIT_W = (630, 540)              # 나눔바른고딕 14pt, 12pt의 반각 1자 폭 근사값
# 밴드: 번호 칸, 빈 칸, 제목 칸의 1행 구성. 무게는 채움(장), 옅은 채움(Module), 테두리(Task) 순으로 낮춤
# 색 값은 테마 키(base, tint, dark) 또는 "white". label_w: 번호 칸 폭(HWPUNIT). Module과 Task는 같은 폭이어서 제목 시작선이 맞음
# rows: 1줄 행 높이의 배수
BAND_STYLE = {
    "chapter": {"label_w": 4200, "fill": "base", "border": None, "label_color": "white", "label_pt": 2000,
                "title_pt": 2000, "rows": 2.2},
    "module":  {"label_w": 7600, "fill": "tint", "border": None, "label_color": "dark", "label_pt": 1150,
                "title_pt": 1350, "rows": 1.5},
    "task":    {"label_w": 7600, "fill": None, "border": "dark", "label_color": "dark", "label_pt": 1100,
                "title_pt": 1200, "rows": 1.3},
    # 간지: 장 밴드와 같은 구성을 키워 쪽 가운데 부근에 배치. "제10장"까지 한 줄
    "divider": {"label_w": 11000, "fill": "base", "border": None, "label_color": "white", "label_pt": 2600,
                "title_pt": 2600, "rows": 3.5},
}
DIVIDER_TOP_LINES = 18   # 간지 밴드 위 빈 바탕글 줄 수 (본문 높이 약 700pt 중 288pt). 밴드가 쪽 가운데보다 약간 위에 옴
BAND_GAP_W = 1000        # 번호 칸과 제목 칸 사이의 빈 칸 폭 (약 3.5mm). 셀 여백 대신 빈 칸을 두어 렌더러마다 간격이 같게 보임
BAND_BORDER_W = "0.12 mm"
# 행 높이 추정의 1줄 추가분. 표 글자 10pt × 줄 간격 130% = 1300에 셀 글머리 문단 뒤 여백 150을 더한 값
lib.ROW_LINE_H = 1450
BAND_POS_ROLE = "band_pos"
NOTE_PREFIX = "※ "       # 표와 그림 아래 부기 중 출처 줄이 아닌 줄의 표지
# 4계위(상세)와 주석은 스켈레톤에 글머리 정의가 없어 조판기가 기호를 본문 글자로 넣음. 내어쓰기는 스켈레톤 문단 모양이 담당
lib.MANUAL_SYMBOLS = {**lib.MANUAL_SYMBOLS, "dot": "· ", "note": "※ "}
TASK_BAND_RE = re.compile(r"^(Module|Task|Step)\s*([\d\-.]+)[.\s]+(.*)$", re.I)
LONG_SPEC_RE = re.compile(r"^\(\s*장표\s*[:：]\s*(engine|split)(?:\s+(\d+))?\s*\)$")   # 장표 지시 줄: (장표: engine) 글자취급 해제와 엔진 분할, (장표: split 10) 10행 단위 물리 분할
TABLE_PAGE_LINES = 30            # 표 추정 줄 수가 이 값 이상이면 한 쪽을 넘을 가능성으로 [장표 HITL] 경고 (본문 폭 표 기준의 근사값)
TABLE_LINE_UNITS = 80            # 표 한 줄의 가중 글자 폭 (한글 40자 근사). 열 폭 배분의 기준
WIDTH_SPEC_RE = re.compile(r"^\(\s*(?:폭|열폭)\s*[:：]\s*(.+?)\s*\)$")   # 표 폭 지시 줄: (폭: 실적) 그룹명 또는 (폭: 3,2,2) 비율, (폭: 3,2,2, 80%) 비율과 표 폭

# ═══════════════════════════════════════════════════════════════════════
# 설정부 2: 작업폴더 파일명과 판정 기준
# ═══════════════════════════════════════════════════════════════════════
F_STATE = "state.json"
F_STRATEGY = "00_제안전략서.md"
F_DIAGNOSIS = "00_진단서.md"   # 사전진단보고서의 설계 문서 (제안전략서에 대응)
F_PLAN_DESIGN = "00_기획설계.md"   # 기획보고서의 설계 문서
# 산출 유형별 설계 문서 양식: 유형 → (assets 양식, 작업폴더 파일명). 제안서는 cmd_init의 전용 분기
DESIGN_TEMPLATES = {"report": ("diagnosis.md", F_DIAGNOSIS), "plan": ("plan-design.md", F_PLAN_DESIGN)}
F_COMPANY = "회사자료.md"
F_RFP_TEXT = "01_RFP원문.md"
F_REQ = "01_요구항목표.md"
F_SCORE = "01_배점표.md"
F_LEDGER = "02_섹션대장.md"
F_STORYLINE = "02_스토리라인.md"
F_REPORT = "05_점검기록.md"
D_DECK = "발표자료"
D_MANUSCRIPT = "원고"
D_DELIVER = "납품"
D_SELFTEST = "_렌더시험"   # selftest 산출 폴더 (표본 원고, 표본 hwpx, PDF, PNG)
D_TEMPLATE = "_양식"       # 기본 스켈레톤 폴더. init과 build가 생성
F_SKELETON = "proposal.hwpx"
SELFTEST_MIN_PAGES = 1     # 렌더 시험 PASS 조건: PNG 1쪽 이상
# Why 표본에 장, 절, 밴드, □, ❍, 근거, 표를 모두 둠: 렌더러가 제안서 고유 요소를 그리지 못하는 결함을 착수 시점에 드러내기 위함
SELFTEST_SAMPLE = """# 렌더 시험

## 표본 절

Module 1. 표본 모듈

Task 1-1. 표본 과업

□ 표본 논점

- (표본) 계위와 글머리 표시 확인
    + 근거 줄 표시 확인
        + 상세 줄 표시 확인

<표> 표본 표

| 항목 | 내용 |
|---|---|
| 셀 | 한글 표시 확인 |
"""
COMMON_TEMPLATES = {"41_제안사일반현황.md": "overview.md", "61_수행조직.md": "organization.md", "62_관리부문.md": "management.md"}
PREFIX_DENSITY = "5"          # check density 대상 (수행부문)
PREFIX_FILL = ("4", "6")      # check fill 대상 (공통 파트)
PREFIX_VOLUME = ("3", "4", "5", "6")    # check volume 대상. □ 단독 판정은 전 원고, 구성 단위 하한은 오리지널 파트(3, 5)
PREFIX_ORIGINAL = ("3", "5")
# 순번 접두(선택). RFP 지정 목차의 장 순서가 기본 순서와 다를 때 `10-51_추진체계.md`처럼 붙인다.
# Why: 파일명 첫 숫자가 합본 순서와 판정 유형을 겸하면 지정 목차의 장 순서를 따를 수 없다. 순서는 순번 접두가, 유형은 그 뒤 첫 숫자가 맡는다
ORDER_PREFIX_RE = re.compile(r"^\d+-")

CONCRETE_MARK_RE = re.compile(r"\([^)]{0,6}(대상|자료|방법|기준|절차|범위|규모|구성|설계|산정|선정|편성|기록)[^)]{0,6}\)")
QUANTITY_MANY = 6   # Task 하나에 수치 표기가 이보다 많으면 나열 과다 경고
QUANTITY_RE = re.compile(r"\d[\d,]*\s*(건|명|회|개소|개|종|차|주|일|개월|기관|과제|표본|응답)")

# 분량 구성 단위 (정성제안서의 일반적 구성: □ 1개당 ❍ 2-3개, 근거 하이픈 또는 ※ 2개 이상, 시각자료 1개)
VOLUME_MIN_CIRCLE = 2       # □ 아래 ❍ 하한
VOLUME_MIN_EVIDENCE = 2     # □ 아래 하이픈+※ 합계 하한
VOLUME_MIN_BLOCK_PER_TASK = 1   # Task 1개당 표 또는 그림 하한
VOLUME_TASK_MIN_CHARS = 900     # Task 1개 본문 글자 수 하한 (약 1.5쪽). 통상 Task는 2-5쪽
VOLUME_SQUARE_MIN_CHARS = 350   # □ 1개 글자 수 하한 (약 0.5쪽)
# 계위 개수 상한 (SKILL.md 분량 구성 단위 절). 초과는 check volume 경고이며 판정에 반영하지 않음
SQUARE_PER_SECTION_MAX = 4       # 절(##)이나 항(###) 안의 □. 5개 이상이면 항으로 묶음
CIRCLE_PER_SQUARE_MAX = 4        # □ 아래 ❍
DASH_PER_CIRCLE_MAX = 3          # ❍ 아래 하이픈. 4개 이상이면 표로 전환
SQUARE_TITLE_MAX = 40           # □ 제목 상한 (13pt 굵게, 본문 폭 기준 1줄. rhwp v0.8.6 렌더링 기준)
LINE_MAX = {"circle": 86, "dash": 84, "dot": 82, "note": 50}   # ❍ 하이픈 ․는 2줄 상한, ※는 1줄 상한 (공백 포함 글자 수. rhwp v0.8.6 렌더링 기준)
SCHEDULE_MARK = "■"              # 추진일정표 기간 셀 표기. 조판 시 음영 셀로 전환
SCHEDULE_FILL = "#D9D9D9"      # 일정표 기간 셀 채움 (표 머리글과 같은 회색)

# 판정 기준 (check wording/tone)
FORBIDDEN_DEFAULT = ["가용가능", "할 수 있다", "할 수 있음", "고려하고 있", "검토할 예정", "노력하겠", "가능할 것으로", "필요시 검토"]
SENTENCE_WARN = [(r"—", "em dash"), (r"가 아닌 ", "대비 구문 'X가 아닌 Y'"), (r"이 아니라 ", "대비 구문 'X가 아니라 Y'"),
                 (r"(으며|하며|이며)\s", "연결어미 복문"), (r"(?<=[가-힣])(하고|되고)\s", "연결어미 복문 (-하고)"), (r"(그리고|그러나|그래서|하지만)\s", "접속부사"),
                 (r"(본 제안서|본 절에서|앞서 언급|다음과 같이|살펴보면|정리하면|요약하면)", "메타 서술 표지"),
                 (r"(MECE|피라미드 원칙|제1원칙|맥킨지식)", "사고 프레임워크명 노출")]
FRAMEWORK_FAIL = [(r"\b(SCQA|SCR|SCQ|PREP)\b", "배경 골격 명칭 노출"),
                  (r"\((S|C|Q|A|R)\s*(상황|전개|질문|답|해결)\)", "골격 단 이름 노출"),
                  (r"(상황|전개|질문|해결)\s*단계로\s*(구성|서술)", "골격 단 이름 노출")]
TONE_FAIL = [(r"(그간|그동안)[^.\n]{0,25}(소홀|미흡|부족|치중|그쳐|그친)", "발주처 정책 미비 시사"),
             (r"(형식적|관행적|소극적)[^.\n]{0,15}(운영|협의|추진|대응)", "발주처 관행 비판"),
             (r"(제안요청서|RFP)(에|의|가|는)[^.\n]{0,15}(더하여|보완하여|미비|흠결|불명확)", "RFP 교정"),
             # 역산(?!업): "지역산업연관표", "지역산업진흥원" 같은 합성어의 부분 일치 제외
             (r"(역산(?!업)|채점|평가위원|통제 ?배점|배점 귀속|완결조건|섹션 ?대장|적대적|채움본|미배치)", "내부 로직 어휘 노출 (역산은 일정 문맥에서도 금지. 체계적 일정 수립, 효율적 일정 편성으로 표현)"),
             (r"(?<![가-힣])(이면 ?니즈|표면 ?니즈|표면 ?목적|판단문|문형|구성 ?단위|변별력|결정 ?장소|지시 ?어휘|조판|치환|오리지널 ?파트|공통 ?파트)", "지시 어휘 유출 (스킬 지침의 용어. 그 자리에 들어갈 실제 내용을 서술. references/lexicon.md 대체표)"),
             (r"(작업분해도|WBS|마일스톤|임계경로|워크 ?패키지|이터레이션|페인포인트|딜리버리|인사이트|온보딩)", "관리 외래어"),
             (r"(?<![가-힣])(포인트|세팅|케이스|니즈|픽스|컨펌|어필|메리트)(?=[\s,.)을를은는이가의로와과에별도]|$)", "준구어 어휘 (references/lexicon.md 대체표의 격식어로 교체)"),
             (r"(입력값|출력값|인계|게이트|관문|입력(이|으로|되| ?자료)|출력(이|으로|물)|연결을 확보|매핑|인터페이스|파이프라인|연계 구조|산출물 연계)", "소프트웨어 설계 어투 (아키텍처 설명이 아닌 수행계획에서는 '활용 자료', '~에 활용', '대응표'로 전환)"),
             (r"(과업의 답|답이 되|답을 제시|전제 ?조건|누락 ?없이|빠짐 ?없이|대응 관계|과업의 핵심 ?답|과업 ?방향|수행 ?여건|일정 ?정합|근거 ?누적|Task 배치|도달 ?상태|도달 ?시점|분석형|조사형|도출형|지원형|수행내용 구성|열 번호|행 번호|(?<![가-힣])열은 |주차 단위|N주 단위)", "골격 언어와 내부 판단의 노출 (스킬의 설계 어휘와 작성자의 점검 언어. 결과 내용을 직접 서술)"),
             (r"(Task|Module|모듈)\s*\d+(-\d+)?(\s*(의|에서|을|를|과|와|은|는|에|로|으로)\s|\s+[가-힣])|(?<![가-힣])(앞|위|이전|전술|후술|다음|아래|해당|상기|하기)\s*(절|장|항)(?=[\s.,)에의을를은는이가과와로도만까부]|$)|(앞|이전|전술|후술|해당|상기)\s*(모듈|Task|Module)|(을|를)\s*참조|참조하", "교차 참조 (다른 절과 모듈과 Task를 번호로 가리키지 않음. 필요한 내용을 그 자리에 다시 서술. 번호 참조는 총괄표·조견표·일정표·업무분장표 셀에만 허용)"),
             (r"\d+\s*개\s*(Module|Task|모듈|영역|축|전략|방안|과제|단계|원칙|방향|관점|부문|파트)", "숫자 선언 (자기 구조의 개수 선언. 항목 자체를 제시)"),
             (r"\d+\s*종(?!류|합|료|목|사)", "N종 선언 (자료명과 대상명을 실명으로 나열. 목록표가 같은 자리에 있을 때만 허용)"),
             (r"(쓰임새|쓰임|짜임새|얼개|뼈대|밑그림|출발점|손질|채비|매만지|다듬|목표 ?값|(일정|보고|회의|과업)[을를의]? ?동기화|일정[을를]? ?정렬)", "비관행 어휘 (references/lexicon.md 대체어 사용)"),
             (r"\d+\s*(대|가지)\s*(전략|방안|과제|핵심|원칙|방향)", "숫자 선언 (개수 선언 없이 항목 자체를 제시. 외부 계획의 원문 인용이면 판정 오판 사유를 기록)"),
             (r"본 (과업|연구|제안|사업|용역|제안서)", "본 과업 선언 (자기 지시 없이 내용으로 진입. 지칭이 필요하면 '과업' 또는 사업명 실명)"),
             (r"(발주처|주관기관|주관부처)[^.\n]{0,20}(협조|제공|지원)[^.\n]{0,12}(전제|조건)|(을|를)\s*전제로|(확보|제공)(되어야|돼야)\s*(수행|추진|착수)", "전제 서술 (발주처에 조건을 제시하는 문장. 자료 확보 방법을 제안사의 절차로 서술)"),
             (r"(범위|과업|연구)(에|에서)\s*(포함하지|포함되지|다루지)\s*않|제외\s*범위", "제외 선언 (발주처가 묻지 않은 제외 범위를 선언하지 않음. RFP 명시 제외만 범위 표에 반영)"),
             (r"(RFP|제안요청서|과업지시서)(가|는|이|에서)\s*(제시|요구|명시|규정|정)", "RFP 주어 서술 (RFP를 문장 주어로 두지 않음. 제시된 내용 자체를 서술)"),
             (r"(아래|위|다음|상기|하기|해당)\s*(표|그림|도표)(는|은|에서|와 같이|과 같이)", "표 지시 메타 서술 (표가 말하는 바를 ❍로 서술하고 표를 배치. '아래 표는'을 삭제)"),
             (r"(첫째|둘째|셋째|넷째|다섯째|여섯째|첫\s*번째|두\s*번째|세\s*번째|네\s*번째|다섯\s*번째|하나는|다른 하나는|마지막으로)", "순서 열거 서술 (첫째 둘째로 잇지 않음. 항목마다 ❍를 두거나 표로 전환)"),
             (r"(제안요청서|RFP|과업지시서|요구사항|요구항목|평가항목|평가기준)(의|에|에서)?\s*(요구|내용|취지|목적|기준)?(에|을|를|과|와)?\s*(대응|부합|충실|준수|맞추|맞게|따라|반영하여|고려하여)", "RFP 대응 선언 (요구항목 원문을 Task 명칭과 서술에 직접 사용)"),
             (r"것(임|이다|입니다|이었음)?[.)]?\s*$", "구어체 '것' 종결 ('남은 것', '해야 할 것', '과업의 포인트인 것'. 명사형 종결로 전환: '~ 잔존', '~ 필요', '~이 과업의 핵심')"),
             (r"(는데|은데|한데|인데|(?<!가)운데|더라|거든|잖아)[.)]?\s*$", "구어체 종결 ('-는데', '-인데'는 설명체 어미. 문장을 끊고 명사형으로 종결)"),
             (r"(는|은|을|를|할|될|된|인|한)\s*(게|거|걸|건)(?!\s*\d)(\s|$)", "구어 축약 ('~하는 게' → '~하는 것이' 또는 명사형, '~할 거' → '~ 예정', '~하는 건' → '~은')"),
             (r"\s(듯함|듯하다|듯|셈|편|만함|터)[.)]?\s*$|(로|고|다고)\s*(봄|보임|보여짐|여겨짐)[.)]?\s*$|(?<![가-힣])(식으로|쪽으로|쪽에서)\s", "전문어로 오인되는 구어체 ('~듯', '~인 셈', '~하는 편', '~로 봄', '~식으로', '~쪽'. 근거를 들어 판정하거나 '~로 판단', '~로 확인'으로 전환)"),
             (r"(^|\s)(일단|약간|좀|꽤|딱|그냥|어쨌든|아무튼|나름)\s", "구어 부사 (삭제하거나 수치·기준으로 대체)"),
             (r"^```", "코드블록 펜스 (산식은 2열 표)")]
# 계위 줄의 마침표. 계위 1개에 문장 1개. 마침표로 두 문장을 잇거나 줄 끝에 마침표를 붙이는 형태를 FAIL로 판정
# 판정 대상은 □, ❍(-), 근거(+), ※ 줄. 표 셀, 각주, 출처 줄, 캡션은 날짜와 서지 표기에 마침표가 쓰이므로 제외
HIER_LINE_RE = re.compile(r"^\s*(□|[-+*]\s|※)")
PERIOD_EXEMPT_RE = re.compile(r"^\s*(\[\^|※\s*출처|출처\s*[:：]|<표>|!\[)")
PERIOD_MID_RE = re.compile(r"[가-힣)\]\"”'’%]\.\s+[가-힣(A-Za-z\d]")   # "확정. 이후" 형태. 소수점(3.5)과 URL은 제외됨
PERIOD_END_RE = re.compile(r"[가-힣)\]\"”'’%]\.\s*$")
# 구어 의존명사(듯, 셈, 편, 터, 만함)는 앞 공백을 요구. "개편", "재편", "데이터", "센터" 같은 명사형 종결과의 충돌 방지
TONE_WARN = [(r"(최대|최초|유일|최고|최상)의?\s", "최상급 단정. 비교 근거 확인"),
             (r"\bRFP\b|제안요청서|과업지시서", "RFP 언급. 요구항목 조견표의 머리글 외에는 과업 내용으로 서술. 기관 지칭이 필요하면 기관 실명"),
             (r"(획기적|혁신적|선도적|압도적|진정한)", "과장 수식어")]
# 기관 호칭 판정. Why: '발주처'는 제안사 쪽에서 상대를 부르는 작업 용어. 원고에 옮겨지면 독자를 제3자로 세우는 공급자 시각의 문장이 되므로 실명으로 유도
CLIENT_TERM = r"발주\s?(처|기관|부서)"
CLIENT_SUBJECT_RE = re.compile(CLIENT_TERM + r"(는|은|가|이|에서는|께서는)(?![가-힣])")
CLIENT_TERM_RE = re.compile(CLIENT_TERM)
# 수행실적 표의 열 이름은 과거 과업의 계약 상대를 가리키는 관행 표기이므로 판정에서 제외
CLIENT_HEADER_CELL_RE = re.compile(r"\|\s*" + CLIENT_TERM + r"\s*(?=\|)")
CLIENT_SLOT_RE = re.compile(r"\[기관명\]")
CLIENT_SUBJECT_LABEL = "발주처 주어 서술 (주어를 지운 뒤 서술부에 대상과 상태를 기입. 지칭이 필요하면 00_제안전략서.md 기관 호칭 절의 실명)"
CLIENT_TERM_LABEL = "'발주처' 호칭 잔여. 기관 실명으로 기입 (00_제안전략서.md 기관 호칭 절)"
CLIENT_SLOT_LABEL = "기관 호칭 자리 잔여 (00_제안전략서.md 기관 호칭 절의 표기로 바꿔 씀)"
SUM_LABEL_RE = re.compile(r"^\s*(합\s*계|계|총\s*계|소\s*계|합계)\s*$")
NUM_CELL_RE = re.compile(r"^[\d,]+(\.\d+)?%?$")
STAGES = ["strategy", "rfp", "ledger", "core", "common", "check", "build"]
# 유형별 check all 대상. 밀도(Task 단위)와 공통 파트 치환은 제안서 전용. 워딩은 제안전략서의 핵심 워딩 사전이 원천
CHECKS_BY_KIND = {"proposal": ("rfp", "density", "volume", "wording", "tone", "numbers", "fill"),
                  "report": ("rfp", "volume", "tone", "numbers", "fill"),
                  "plan": ("volume", "tone", "numbers", "fill")}   # 보고서와 기획서의 fill은 [확인] 잔여 경고만
# check tone 화이트리스트 원천. 핵심 워딩 사전(00_제안전략서.md)과 모듈총괄표(02_모듈총괄표.md)의 Module·Task 명칭.
# RFP 원문 용어와 유저가 확정한 골격 명칭은 어조 판정 대상이 아님. 최초 check tone 시 state.json에 고정하여 이후 임의 추가를 차단
# 모듈총괄표는 02_섹션대장, 02_스토리라인과 같은 2단계 설계 산출물이므로 작업폴더 루트에 둔다.
# Why: 원고/ 안의 번호 파일로 두면 표의 위치가 파일명 정렬에 묶여 지정 목차에서 앞 장(제안사 일반현황 등) 아래로 들어간다.
# 표의 위치는 원고의 삽입 지시 줄 `(삽입: 모듈총괄표)`가 정하고 build가 합본 시 표로 바꾼다. 표의 원본은 1곳
F_MODULE_TABLE = "02_모듈총괄표.md"
MODULE_TABLE_SUFFIX = "모듈총괄표.md"   # 구 방식(원고/50_모듈총괄표.md) 작업폴더 호환용 식별 접미
INSERT_MODULE_RE = re.compile(r"^\(\s*삽입\s*[:：]\s*모듈총괄표\s*\)$")
WHITELIST_MASK = "◇"
WHITELIST_MIN_LEN = 4   # 2-3자 일반명사("조사", "매핑")를 마스킹하면 "조사형" 같은 다른 판정을 가리므로 제외
TABLE_SEP_RE = lib.TABLE_SEP_RE
Element = lib.Element


# ═══════════════════════════════════════════════════════════════════════
# 1. 원고 파싱: 조판기 파서 위에 Module/Task 밴드만 추가
# ═══════════════════════════════════════════════════════════════════════
def parse_markdown(path: str) -> tuple[list[Element], dict[str, str]]:
    """lib.parse_markdown 위에 제안서 표기를 얹는다.
    - `(폭: ...)` 표 폭 지시 줄은 lib 파서에 넘기기 전에 걷어내고 다음 표의 순번에 기록 (캡션과 표 사이에 두어도 캡션 부착이 깨지지 않도록)
    - `Module 1. 제목` 줄과 `Task 1-1. 제목` 줄은 band 요소로 바꾼다"""
    raw = open(path, encoding="utf-8").read().split("\n")
    kept, specs, longs, table_idx, in_table = [], {}, {}, 0, False
    for ln in raw:
        st = ln.strip()
        if (m := WIDTH_SPEC_RE.match(st)):
            specs[table_idx] = m.group(1)
            continue
        if (m := LONG_SPEC_RE.match(st)):
            longs[table_idx] = (m.group(1), int(m.group(2)) if m.group(2) else None)
            continue
        is_row = st.startswith("|")
        if is_row and not in_table:
            table_idx += 1
        in_table = is_row
        kept.append(ln)
    src = path
    if specs or longs:  # 이미지 상대 경로가 원고 폴더 기준이므로 같은 폴더에 임시 사본을 둔다
        src = os.path.join(os.path.dirname(os.path.abspath(path)), f".{os.path.basename(path)}.tmp")
        open(src, "w", encoding="utf-8").write("\n".join(kept))
    try:
        elems, footnotes = lib.parse_markdown(src)
    finally:
        if src != path and os.path.exists(src):
            os.remove(src)
    ti = 0
    for e in elems:
        if e.kind == "table":
            if ti in specs:
                e.data["width_spec"] = specs[ti]
            if ti in longs:
                e.data["long_spec"] = longs[ti]
            ti += 1
        elif e.kind in ("circle", "square") and (m := TASK_BAND_RE.match(e.text)):
            e.kind = "band"
            e.data["kind"] = "module" if m.group(1).lower() == "module" else "task"
    return elems, footnotes


def estimate_table_lines(rows: list[list[str]]) -> int:
    """표 높이를 줄 수로 근사. 열 폭은 열별 최대 가중 길이에 비례 배분하고 셀 줄 수는 <br> 분절마다 접힘 줄 수를 합산."""
    ncol = max(len(r) for r in rows)
    need = [0] * ncol
    for r in rows:
        for c, v in enumerate(r):
            need[c] = max(need[c], max((lib._weighted_len(s.strip()) for s in v.split("<br>")), default=0), 4)
    total = sum(need) or 1
    caps = [max(6, TABLE_LINE_UNITS * n // total) for n in need]
    lines = 0
    for r in rows:
        row_lines = 1
        for c, v in enumerate(r):
            segs = [s.strip() for s in v.split("<br>")] or [""]
            row_lines = max(row_lines, sum(max(1, -(-lib._weighted_len(s) // caps[c])) for s in segs))
        lines += row_lines
    return lines


def lint(elems: list[Element]) -> list[str]:
    out = lib.lint(elems) + lib.lint_counts(elems) if hasattr(lib, "lint_counts") else lib.lint(elems)
    for e in elems:
        if e.kind != "table" or e.data.get("long_spec"):
            continue
        rows = e.data.get("rows") or []
        est = estimate_table_lines(rows)
        if est >= TABLE_PAGE_LINES and len(rows) <= lib.LONG_TABLE_ROWS:
            out.append(f"[장표 HITL] '{e.text[:24]}' {len(rows)}행 추정 {est}줄: 한 쪽을 넘을 가능성. 유저에게 (장표: engine) 글자취급 해제 또는 (장표: split N) 물리 분할 택일을 질의")
    return out


# ═══════════════════════════════════════════════════════════════════════
# 2. 조판: 형식 무관 판단(ProposalLayout)과 형식별 작성기(ProposalBuilder는 hwpx, ProposalDocxBuilder는 docx)
# ═══════════════════════════════════════════════════════════════════════
class ProposalLayout:
    """제안서, 사전진단보고서, 기획보고서의 조판 판단. hwpx와 docx 작성기가 함께 상속하여 같은 결과를 기록
    판단 대상: 제목 번호와 쪽 나눔, 밴드 구성, 캡션 문구와 번호, 표 폭 지시와 장표 지시, 표 아래 부기, 요소 순회
    작성기가 구현할 기록 함수: para, heading_para, band_emit, chapter_started, add_footnote, after_table, after_image,
    _front_matter, set_page_number, set_header, save"""

    def _cell_line_role(self, line: str) -> str | None:
        """기호 하나만 있는 셀 줄(비교 매트릭스의 ○ 등)은 글머리로 보지 않는다.
        Why: lib는 ○로 시작하는 줄을 셀 글머리로 판정해 기호를 지우고 글머리 스타일을 입힌다. ●, ◐와 크기와 정렬이 어긋났다."""
        if len(line.strip()) <= 1:
            return None
        return super()._cell_line_role(line)

    def heading(self, level: int, text: str):
        """로마자 장(Ⅰ.), 아라비아 절(1.), 가나다 항(가.), 목(1)). 번호는 본문에 기입하고 개요 번호는 끈다.
        보고서 프로파일은 제1장, 1.1, 1.1.1, 가. 체계. 서식 문자열은 프로파일이 정함"""
        text = lib.NUM_PREFIX_RE.sub("", text).strip()
        role = ("chapter", "section", "subsection", "ganada")[level - 1]
        pt = None
        if level == 1:
            self.n.update(chapter=self.n["chapter"] + 1, section=0, subsection=0, ganada=0)
            # 문서의 첫 장은 쪽 나눔 제외. 장 번호가 1이 아닌 부분 원고(--chapter-start)도 같음
            if self.page_break in ("chapter", "section") and self.n.get("chapters_done", 0) > 0:
                self.pending_break = True
            self.n["chapters_done"] = self.n.get("chapters_done", 0) + 1
            c = self.n["chapter"]
            label = self.p["chapter_format"].format(n=c, r=ROMAN[(c - 1) % len(ROMAN)], t=text)
            pt = self.p.get("chapter_pt")
        elif level == 2:
            self.n.update(section=self.n["section"] + 1, subsection=0, ganada=0)
            label = self.p["section_format"].format(c=max(self.n["chapter"], 1), s=self.n["section"], t=text)
            if self.page_break == "section" and not (self.n["section"] == 1 and self.n["chapter"] > 0) and self.n["total_section"] > 0:
                self.pending_break = True
            self.n["total_section"] += 1
            pt = self.p.get("section_pt")
        elif level == 3:
            self.n.update(subsection=self.n["subsection"] + 1, ganada=0)
            ss = self.n["subsection"]
            label = self.p["subsection_format"].format(c=max(self.n["chapter"], 1), s=max(self.n["section"], 1), ss=ss, g=GANADA[(ss - 1) % 14], t=text)
            pt = self.p.get("subsection_pt")
        else:
            self.n["ganada"] += 1
            i = self.n["ganada"]
            label = self.p.get("item_format", "{g}. {t}").format(i=i, g=GANADA[(i - 1) % 14], t=text)
        if level == 1:  # 장 제목은 Module, Task와 같은 밴드. 목차와 쪽번호 대조용 표기는 label 그대로 기록
            self.cap_n = {"TABLE": 0, "PICTURE": 0}
            band_label = self.p.get("chapter_label", "{r}").format(n=c, r=ROMAN[(c - 1) % len(ROMAN)])
            if getattr(self, "divider", False):  # 간지: 장 제목만 담은 쪽. 본문은 다음 쪽의 첫 절부터 시작
                self.pending_break = False
                p = self.divider_emit(self.band_spec(lib.Element("band", label, data={"kind": "divider", "label": band_label, "title": text})),
                                      new_page=self._body_started)
                self.after_block = True
                self.pending_break = True
            else:
                p = self.band(lib.Element("band", label, data={"kind": "chapter", "label": band_label, "title": text}))
        else:
            p = self.heading_para(role, label, pt)
        if level <= 3:
            self.headings.append((level, label))
        if level == 1:
            self.chapter_started(p)

    def band_spec(self, e: Element) -> dict:
        """밴드 1행 3칸(번호 칸, 빈 칸, 제목 칸)의 폭, 높이, 칸별 (글자, 크기, 글자색, 정렬, 채움, 테두리)"""
        kind = e.data["kind"]
        label, title = e.data.get("label"), e.data.get("title")
        if label is None:  # Module, Task 원고 줄에서 번호와 제목을 분리
            m = TASK_BAND_RE.match(e.text)
            label, title = (f"{m.group(1)} {m.group(2).rstrip('.')}", m.group(3).strip()) if m else ("", e.text)
        st = {**BAND_STYLE[kind], **(self.p.get("chapter_band", {}) if kind == "chapter" else {})}
        color = {**THEMES[getattr(self, "theme", DEFAULT_THEME)], "white": "#FFFFFF"}.get
        return {"widths": [st["label_w"], BAND_GAP_W, self.table_max_w - st["label_w"] - BAND_GAP_W],
                "height": int(lib.ROW_BASE_H * st["rows"]),
                "cells": [(label, st["label_pt"], color(st["label_color"]), "CENTER",
                           color(st["fill"]) if st["fill"] else None, color(st["border"]) if st["border"] else None),
                          ("", st["title_pt"], "#000000", "LEFT", None, None),
                          (title, st["title_pt"], "#000000", "LEFT", None, None)]}

    def band(self, e: Element):
        """장, Module, Task 제목을 1행 3칸 표로 배치. 밴드가 담긴 문단(또는 첫 문단)을 돌려주어 장의 캡션 번호 재시작에 쓴다."""
        p = self.band_emit(self.band_spec(e))
        self.after_block = True
        return p

    def caption_label(self, kind: str, text: str) -> dict:
        """캡션 문구. 번호 없는 프로파일은 표 <제목>, 그림 [제목]. 번호 있는 프로파일은 <표 1-3> 제목, [그림 1-3] 제목
        반환: open(번호 앞), num(번호, 없으면 None), close(번호 뒤), text(전체). 번호 캡션은 표 목차와 그림 목차에 등록"""
        prefix = self.p.get("caption_prefix")
        if prefix is None:
            label = f"<{text}>" if kind == "TABLE" else f"[{text}]"
            return {"open": label, "num": None, "close": "", "text": label}
        self.cap_n[kind] = self.cap_n.get(kind, 0) + 1
        num = self.cap_n[kind]
        open_b, close_b = ("<표 ", "> ") if kind == "TABLE" else ("[그림 ", "] ")
        head = open_b + prefix.format(c=max(self.n["chapter"], 1))
        full = f"{head}{num}{close_b}{text}"
        self.captions.append((kind, full, close_b.strip() + text))
        return {"open": head, "num": num, "close": close_b + text, "text": full}

    def toc_lists(self) -> list[tuple[str, list[str]]]:
        if not self.p.get("toc_lists"):
            return []
        return [(TOC_LIST_TITLES[0], [t for k, t, _ in self.captions if k == "TABLE"]),
                (TOC_LIST_TITLES[1], [t for k, t, _ in self.captions if k == "PICTURE"])]

    def unify_table_widths(self, elems: list[Element]):
        """표 폭 지시를 열 폭으로 바꾼다. 지시는 원고의 `(폭: ...)` 줄이며 모델이 양식 반복 여부를 보고 정성적으로 지정한다.
        (폭: 그룹명) 같은 그룹명의 표는 그룹 전 행을 합쳐 1회 산정한 폭을 공유 (최대 필요폭 기준이므로 좁은 폭은 전파되지 않음)
        (폭: 3,2,2) 비율로 표 폭 상한을 배분. (폭: 3,2,2, 80%) 상한의 80%를 비율로 배분
        지시가 없는 표는 lib 규칙대로 표마다 산정. `--unify auto`일 때만 머리글 또는 첫 열 라벨이 같은 표를 자동 그룹으로 묶음
        Why 수동 우선: 어떤 표가 같은 양식인지는 문서 설계의 판단이며 기계 판정은 오인을 만든다."""
        tables = [e for e in elems if e.kind == "table" and e.data.get("rows")]
        groups: dict[tuple, list[Element]] = {}
        for e in tables:
            spec = e.data.get("width_spec")
            if not spec:
                continue
            parts = [x.strip() for x in spec.split(",")]
            if all(re.fullmatch(r"\d+(\.\d+)?%?", x) for x in parts):
                ratio = [float(x) for x in parts if not x.endswith("%")]
                pct = next((float(x[:-1]) / 100 for x in parts if x.endswith("%")), 1.0)
                ncol = len(e.data["rows"][0])
                if len(ratio) != ncol:
                    print(f"WARN: (폭: {spec}) 비율 개수 {len(ratio)}가 열 수 {ncol}와 다름. 표마다 산정으로 대체")
                    continue
                total = int(self.table_max_w * pct)
                widths = [int(total * r / sum(ratio)) for r in ratio]
                widths[-1] += total - sum(widths)
                e.data["col_widths"] = widths
            else:
                groups.setdefault(("G", spec), []).append(e)
        if getattr(self, "unify_mode", "manual") == "auto":
            def header_key(e): return ("H", tuple(c.strip() for c in e.data["rows"][0]))
            def label_key(e): return ("L", len(e.data["rows"][0]), tuple(r[0].strip() for r in e.data["rows"] if r))
            from collections import Counter
            free = [e for e in tables if not e.data.get("width_spec")]
            hc, lc = Counter(header_key(e) for e in free), Counter(label_key(e) for e in free)
            for e in free:
                if hc[header_key(e)] >= 2:
                    groups.setdefault(header_key(e), []).append(e)
                elif lc[label_key(e)] >= 2:
                    groups.setdefault(label_key(e), []).append(e)
        for key, group in groups.items():
            # 열 수가 다른 표(Task 수가 다른 모듈 개요표 등)는 열 수가 같은 표끼리 묶어 통일
            by_ncol: dict[int, list] = {}
            for t in group:
                by_ncol.setdefault(len(t.data["rows"][0]), []).append(t)
            for sub in by_ncol.values():
                widths = lib.calc_col_widths([r for t in sub for r in t.data["rows"]], self.table_max_w)
                for t in sub:
                    t.data["col_widths"] = widths

    def table(self, e: Element):
        """표 폭 지시와 장표 지시를 적용한 뒤 작성기의 표 기록을 호출"""
        forced = e.data.get("col_widths")
        original = lib.calc_col_widths
        if forced:
            lib.calc_col_widths = lambda rows, max_w: list(forced)  # 그룹 공통 폭을 lib의 산정 대신 사용
        long_spec = e.data.get("long_spec")  # (장표: engine|split N) 지시: 행 수와 무관하게 이 표에만 장표 처리 적용
        saved_rows, saved_mode = lib.LONG_TABLE_ROWS, self.long_table_mode
        if long_spec:
            mode, chunk = long_spec
            self.long_table_mode = mode
            lib.LONG_TABLE_ROWS = 0 if mode == "engine" else (chunk or 12) + 1
        try:
            r = super().table(e)
        finally:
            lib.calc_col_widths = original
            lib.LONG_TABLE_ROWS, self.long_table_mode = saved_rows, saved_mode
        self.after_table(e)
        return r

    def image(self, e: Element):
        r = super().image(e)
        self.after_image(e)
        return r

    @staticmethod
    def special_cell(val: str) -> str | None:
        """셀 값이 ■ 하나뿐이면 추진일정표의 기간 셀. 글자 대신 채움으로 기간을 보이는 관행. 반환은 채움색"""
        return SCHEDULE_FILL if val.strip() == SCHEDULE_MARK else None

    def build(self, elems: list[Element], footnotes: dict, out: str):
        """요소 순회. 표와 그림 직후의 부기와 출처 줄은 블록 하단 행으로 붙이고 본문 뒤 출처 줄은 직전 문단의 각주로 옮김"""
        self.unify_table_widths(elems)
        self.last_para = None
        self._body_started = False   # 본문 첫 요소 기록 여부. 간지의 쪽 나눔 판단에 사용
        if getattr(self, "doc_title", None):
            self.para("doc_title", self.doc_title)
            self._body_started = True
        skip = 0
        for idx, e in enumerate(elems):
            if skip:
                skip -= 1
                continue
            if e.kind in ("table", "image"):
                srcs, consumed = [], 0
                for nx in elems[idx + 1:]:
                    if nx.kind in ("note", "source"):  # 출처 줄은 그대로, 그 밖의 부기는 ※ 표지를 유지
                        text = inline_note_footnotes(nx.text, footnotes)
                        srcs.append(text if nx.kind == "source" or lib.SOURCE_RE.match(text) else NOTE_PREFIX + text)
                    elif nx.kind == "unit" and e.kind == "table" and not e.data.get("unit"):
                        e.data["unit"] = nx.text
                    else:
                        break
                    consumed += 1
                e.data["sources"] = srcs
                skip = consumed
            if e.kind == "heading":
                self.heading(e.level, e.text)
            elif e.kind == "band":
                self.band(e)
            elif e.kind in ("square", "circle", "dash", "dot", "note"):
                self.last_para = self.para(e.kind, e.text, footnotes)
            elif e.kind == "unit":
                self.para("unit", e.text)
            elif e.kind == "source":
                if self.last_para is not None:
                    self.add_footnote(self.last_para, e.text)
                else:
                    self.para("note", e.text, footnotes)
            elif e.kind == "table":
                self.table(e)
            elif e.kind == "image":
                self.image(e)
            elif e.kind == "pagebreak":
                self.pending_break = True
            if e.kind != "pagebreak":
                self._body_started = True
        if self.toc or self.cover:
            self._front_matter(self.cover)
        self.set_page_number(self.page_number)
        self.set_header(self.header_spec)
        self.save(out)


class ProposalBuilder(ProposalLayout, lib.Builder):
    """hwpx 작성기. 판단은 ProposalLayout, 기록은 python-hwpx"""

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.cap_n = {"TABLE": 0, "PICTURE": 0}
        self.captions: list[tuple[str, str, str]] = []

    def _sized_char(self, height: int, bold: bool) -> str:
        """장 스타일 글꼴을 기준으로 크기와 굵기만 바꾼 charPr (헤더 위계, 밴드 글자)."""
        return self._cover_char(height, bold)

    def _border_fill(self, sides: list[str], fill: str | None = None) -> str:
        """lib의 테두리 정의에 hatchColor를 6자리로 정규화한다.
        Why: python-hwpx는 hatchColor를 8자리(#FF000000)로 쓰는데 한/글 저작 파일은 6자리만 쓴다. 채움색 오표시 예방."""
        bid = super()._border_fill(sides, fill)
        el = self.header.element.find(f".//{HH}borderFill[@id='{bid}']")
        if el is not None:
            HC = "{http://www.hancom.co.kr/hwpml/2011/core}"
            for wb in el.iter(f"{HC}winBrush"):
                if len(wb.get("hatchColor", "")) == 9:
                    wb.set("hatchColor", "#" + wb.get("hatchColor")[3:])
        return bid

    def attach_caption(self, holder, kind: str, text: str, side: str, width: int):
        """caption_prefix가 None이면 번호 없는 캡션 <제목>을 개체에 부착한다 (제안서 관행). 값이 있으면 lib의 자동번호 캡션."""
        label = self.caption_label(kind, text)
        if label["num"] is not None:
            return super().attach_caption(holder, kind, text, side, width)
        role = "table_title" if kind == "TABLE" else "fig_title"
        r = self.refs(role)
        cap = holder.set_caption("", side=side, width=width, gap=lib.CAPTION_GAP, char_pr_id_ref=r["char_pr_id_ref"])
        cp = cap.paragraphs[0]
        cp.element.set("paraPrIDRef", str(r["para_pr_id_ref"]))
        cp.element.set("styleIDRef", str(r["style_id_ref"]))
        for run in [c for c in cp.element if c.tag == f"{HP}run"]:
            cp.element.remove(run)
        cp.add_run(label["text"], char_pr_id_ref=r["char_pr_id_ref"])

    def _band_char(self, height: int, color: str) -> str:
        """장 스타일 글꼴 기준으로 크기와 굵기와 글자색을 바꾼 charPr (밴드 전용. 흰 글자는 장 밴드)."""
        base = str(self.refs("chapter")["char_pr_id_ref"])

        def mod(e):
            e.set("height", str(height))
            e.set("textColor", color)
            if e.find(f"{HH}bold") is None:
                anchor = e.find(f"{HH}offset")
                e.insert(list(e).index(anchor) + 1, etree.Element(f"{HH}bold"))
        font = self.header.element.find(f".//{HH}charPr[@id='{base}']").find(f"{HH}fontRef").get("hangul")
        return self.header.ensure_char_property(
            predicate=lambda e: e.get("height") == str(height) and e.get("textColor") == color
            and e.find(f"{HH}fontRef").get("hangul") == font and e.find(f"{HH}bold") is not None,
            modifier=mod, base_char_pr_id=base).get("id")

    def heading_para(self, role: str, label: str, pt: int | None):
        p = self.para(role, label, no_heading=True)
        if pt:
            cp = self._sized_char(pt, True)
            for run in p.element.iter(f"{HP}run"):
                run.set("charPrIDRef", cp)
        return p

    def chapter_started(self, p):
        if self.p.get("caption_prefix"):  # 장별 <표 N-1>부터 재시작
            p.add_new_num(kind="TABLE", number=1)
            p.add_new_num(kind="PICTURE", number=1)

    def band_emit(self, spec: dict):
        """1행 3칸 표를 밴드 위치 문단에 담음. Why 표: 문단 음영보다 표 셀 채움이 렌더러 간 표시가 일관됨"""
        self.gap()
        r = self._centered_refs(BAND_POS_ROLE)
        attrs = {}
        if self.pending_break:
            attrs["pageBreak"] = "1"
            self.pending_break = False
        p = self.doc.add_paragraph("", **r, **attrs)
        t = p.add_table(1, len(spec["cells"]), width=self.table_max_w)
        t.element.set("numberingType", "NONE")  # 밴드는 표 번호 집계에서 제외 (캡션 자동번호 오염 방지)
        t.set_column_widths(spec["widths"])
        ref = self.refs("band_cell")
        height = spec["height"]
        empty = self._band_fill(None, None)
        fills = [self._band_fill(fill, border) if (fill or border) or ci == 0 else empty
                 for ci, (_, _, _, _, fill, border) in enumerate(spec["cells"])]
        for ci, (text, pt, color, align, _, _) in enumerate(spec["cells"]):
            t.set_cell_text(0, ci, text)
            cell = t.cell(0, ci)
            cp = self._band_char(pt, color)
            for pe in cell.element.iter(f"{HP}p"):
                pe.set("paraPrIDRef", self._aligned_pp(str(ref["para_pr_id_ref"]), align))
                pe.set("styleIDRef", str(ref["style_id_ref"]))
                for run in pe.iter(f"{HP}run"):
                    run.set("charPrIDRef", cp)
            cell.element.set("borderFillIDRef", fills[ci])
            cell.set_size(height=height)
            sub = cell.element.find(f"{HP}subList")
            if sub is not None:
                sub.set("vertAlign", "CENTER")
        sz = t.element.find(f"{HP}sz")
        if sz is not None:
            sz.set("height", str(height))
        pos = t.element.find(f"{HP}pos")
        if pos is not None:
            pos.set("horzAlign", "CENTER")
        t.mark_dirty()
        return p

    def divider_emit(self, spec: dict, new_page: bool):
        """간지 쪽: 빈 바탕글 줄 뒤에 간지 밴드. 이 쪽의 머리말, 꼬리말, 쪽번호는 감춤(한/글 감추기)"""
        body = self.refs("body")
        first = None
        for i in range(DIVIDER_TOP_LINES):
            q = self.doc.add_paragraph("", **body, **({"pageBreak": "1"} if i == 0 and new_page else {}))
            first = first or q
        first.add_page_hiding(header=True, footer=True, page_num=True)
        return self.band_emit(spec)

    def _band_fill(self, fill: str | None, border: str | None) -> str:
        """밴드 칸의 채움과 테두리. 테두리 색은 표 괘선과 별도로 지정한다."""
        sides = ["top", "bottom", "left", "right"] if border else []
        bid = str(self.header.ensure_border_fill(border_color=border or lib.BORDER_COLOR, border_width=BAND_BORDER_W,
                                                 active_borders=sides, fill_color=fill))
        el = self.header.element.find(f".//{HH}borderFill[@id='{bid}']")
        if el is not None:
            for side in ("left", "right", "top", "bottom"):
                b = el.find(f"{HH}{side}Border")
                if b is not None:
                    b.set("type", "SOLID" if side in sides else "NONE")
            HC = "{http://www.hancom.co.kr/hwpml/2011/core}"
            for wb in el.iter(f"{HC}winBrush"):
                if len(wb.get("hatchColor", "")) == 9:
                    wb.set("hatchColor", "#" + wb.get("hatchColor")[3:])
        return bid

    def _last_tbl(self):
        tbls = list(self.doc.sections[-1].element.iter(f"{HP}tbl"))
        return tbls[-1] if tbls else None

    def after_table(self, e: Element):
        """■ 셀을 음영 셀로 바꾸고, 번호 캡션 프로파일에서 캡션 없는 표는 번호 집계에서 뺌 (표 목차 번호와 한/글 자동번호 일치)"""
        tbl = self._last_tbl()
        if tbl is None:
            return
        if self.p.get("caption_prefix") and not e.text and tbl.get("numberingType") == "TABLE":
            tbl.set("numberingType", "NONE")
        rows = e.data.get("rows") or []
        if not any(self.special_cell(c) for row in rows for c in row):
            return
        ncol = int(tbl.get("colCnt"))
        for tc in tbl.iter(f"{HP}tc"):
            texts = [t for t in tc.iter(f"{HP}t")]
            fill = self.special_cell("".join((t.text or "") for t in texts))
            if not fill:
                continue
            for t in texts:
                t.text = ""
            # Why: 음영 셀도 표 좌우 바깥선 규칙을 따라야 첫 열과 마지막 열에 세로선이 생기지 않음
            c0 = int(tc.find(f"{HP}cellAddr").get("colAddr"))
            c1 = c0 + int(tc.find(f"{HP}cellSpan").get("colSpan")) - 1
            tc.set("borderFillIDRef", self._border_fill(lib.edge_sides(c0, c1, ncol), fill))

    def after_image(self, e: Element):
        tbl = self._last_tbl()
        if tbl is not None and self.p.get("caption_prefix") and not e.text and tbl.get("numberingType") == "PICTURE":
            tbl.set("numberingType", "NONE")

    def add_footnote(self, p, text: str):
        p.add_footnote(text)

    def save(self, out: str):
        report = self.doc.validate()
        issues = getattr(report, "issues", [])
        if issues:
            print(f"WARN: 스키마 검증 이슈 {len(issues)}건")
            for it in list(issues)[:5]:
                print("   ", it)
        self.doc.save_to_path(os.path.abspath(out))
        self.doc.close()


class ProposalDocxBuilder(ProposalLayout, docx_lib.DocxBuilder):
    """docx 작성기. 판단은 ProposalLayout, 기록은 docx_report.DocxBuilder"""

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.cap_n = {"TABLE": 0, "PICTURE": 0}
        self.captions: list[tuple[str, str, str]] = []


def load_state(work: str) -> dict:
    path = os.path.join(work, F_STATE)
    if os.path.exists(path):
        return json.load(open(path, encoding="utf-8"))
    return {"created": datetime.now().isoformat(timespec="seconds"), "stages": {}, "checks": {}, "files": {}}


def save_state(work: str, st: dict):
    st["updated"] = datetime.now().isoformat(timespec="seconds")
    json.dump(st, open(os.path.join(work, F_STATE), "w", encoding="utf-8"), ensure_ascii=False, indent=2)


def record(work: str, key: str, passed: bool, detail: dict):
    st = load_state(work)
    st["checks"][key] = {"pass": passed, "at": datetime.now().isoformat(timespec="seconds"), **detail}
    save_state(work, st)


def manuscript_kind(path: str) -> str:
    """순번 접두를 뗀 파일명. 첫 글자가 판정 유형(3, 5 오리지널. 4, 6 공통)"""
    return ORDER_PREFIX_RE.sub("", os.path.basename(path))


def manuscript_files(work: str, prefix: str = "") -> list[str]:
    """합본 순서는 전체 파일명 순. prefix는 순번 접두를 뗀 이름에 적용"""
    files = sorted(glob.glob(os.path.join(work, D_MANUSCRIPT, "*.md")))
    return [f for f in files if manuscript_kind(f).startswith(prefix)]


def legacy_module_table(work: str) -> str | None:
    found = [f for f in manuscript_files(work) if manuscript_kind(f).endswith(MODULE_TABLE_SUFFIX)]
    return found[0] if found else None


def module_table_path(work: str) -> str:
    root = os.path.join(work, F_MODULE_TABLE)
    return root if os.path.exists(root) else (legacy_module_table(work) or root)


def module_table_block(work: str) -> str:
    """02_모듈총괄표.md에서 캡션(<표>)부터 표의 마지막 행까지를 꺼낸다. 캡션과 표 사이의 (장표: ), (폭: ) 지시 줄을 포함한다."""
    lines = open(os.path.join(work, F_MODULE_TABLE), encoding="utf-8").read().split("\n")
    rows = [i for i, ln in enumerate(lines) if ln.strip().startswith("|")]
    if not rows:
        sys.exit(f"ERROR: {F_MODULE_TABLE}에 표가 없음")
    caps = [i for i, ln in enumerate(lines[:rows[0]]) if ln.strip().startswith("<표>")]
    return "\n".join(lines[(caps[-1] if caps else rows[0]):rows[-1] + 1])


def expand_module_table(work: str, text: str) -> tuple[str, int]:
    """원고의 삽입 지시 줄을 표로 바꾼다. 반환은 (바꾼 본문, 바꾼 횟수)"""
    out, count = [], 0
    for ln in text.split("\n"):
        if INSERT_MODULE_RE.match(ln.strip()):
            out += ["", module_table_block(work), ""]
            count += 1
        else:
            out.append(ln)
    return "\n".join(out), count


# ═══════════════════════════════════════════════════════════════════════
# 기본 스켈레톤 생성 (설정부 1의 SKELETON_* 정의를 python-hwpx 빈 새 문서에 적용)
# ═══════════════════════════════════════════════════════════════════════
SK_HC = "{http://www.hancom.co.kr/hwpml/2011/core}"


class SkeletonBaseError(Exception):
    """python-hwpx의 빈 새 문서가 없거나 구조가 예상과 다를 때 발생"""


def _skeleton_base_path() -> str:
    import hwpx
    return os.path.join(os.path.dirname(hwpx.__file__), "data", "Skeleton.hwpx")


def _sk_container(head, tag):
    el = head.find(f".//{HH}{tag}")
    if el is None:
        raise SkeletonBaseError(f"header.xml에 {tag} 없음")
    return el


def _sk_renumber(container):
    container.set("itemCnt", str(len(container)))


def _sk_set_fonts(head):
    """모든 언어의 글꼴 목록을 본문(0)과 제목(1) 두 글꼴로 교체. 기본 스타일(바탕글, 각주 등)은 본문 글꼴"""
    for ff in head.iter(f"{HH}fontface"):
        for f, face in zip(ff.findall(f"{HH}font"), (SKELETON_FONT_BODY, SKELETON_FONT_HEAD)):
            f.set("face", face)
            ti = f.find(f"{HH}typeInfo")
            if ti is not None:
                ti.set("familyType", "FCAT_GOTHIC")
    for cp in head.iter(f"{HH}charPr"):
        ref = cp.find(f"{HH}fontRef")
        for k in list(ref.attrib):
            ref.set(k, "0")


def _sk_add_char(head, template, font_id, pt, bold) -> str:
    cps = _sk_container(head, "charProperties")
    el = copy.deepcopy(template)
    el.set("id", str(len(cps)))
    el.set("height", str(int(round(pt * 100))))
    el.set("textColor", SKELETON_TEXT_COLOR)
    ref = el.find(f"{HH}fontRef")
    for k in list(ref.attrib):
        ref.set(k, str(font_id))
    for b in el.findall(f"{HH}bold"):
        el.remove(b)
    if bold:
        el.insert(list(el).index(el.find(f"{HH}offset")) + 1, etree.Element(f"{HH}bold"))
    cps.append(el)
    _sk_renumber(cps)
    return el.get("id")


def _sk_add_bullet(head, char) -> str:
    bullets = head.find(f".//{HH}bullets")
    if bullets is None:  # 빈 새 문서에는 글머리 목록이 없으므로 번호 목록 뒤에 만든다
        bullets = etree.Element(f"{HH}bullets", itemCnt="0")
        _sk_container(head, "numberings").addnext(bullets)
    el = etree.SubElement(bullets, f"{HH}bullet", id=str(len(bullets) + 1), char=char, useImage="0")
    etree.SubElement(el, f"{HH}paraHead", level="0", align="LEFT", useInstWidth="0", autoIndent="1",
                     widthAdjust="0", textOffsetType=SKELETON_BULLET_OFFSET_TYPE, textOffset=SKELETON_BULLET_OFFSET,
                     numFormat="DIGIT", charPrIDRef="4294967295", checkable="0")
    _sk_renumber(bullets)
    return el.get("id")


def _sk_add_toc_tab(head) -> str:
    """목차 문단의 오른쪽 정렬 점선 탭. 한/글이 쪽번호를 오른쪽 끝에 맞추는 기준"""
    tabs = _sk_container(head, "tabProperties")
    tp = etree.SubElement(tabs, f"{HH}tabPr", id=str(len(tabs)), autoTabLeft="0", autoTabRight="0")
    sw = etree.SubElement(tp, f"{HP}switch")
    case = etree.SubElement(sw, f"{HP}case")
    case.set(f"{HP}required-namespace", "http://www.hancom.co.kr/hwpml/2016/HwpUnitChar")
    etree.SubElement(case, f"{HH}tabItem", pos=str(SKELETON_TOC_TAB_POS), type="RIGHT", leader="DOT", unit="HWPUNIT")
    default = etree.SubElement(sw, f"{HP}default")
    etree.SubElement(default, f"{HH}tabItem", pos=str(SKELETON_TOC_TAB_POS * 2), type="RIGHT", leader="DOT")
    _sk_renumber(tabs)
    return tp.get("id")


def _sk_add_para(head, template, align, left, indent, prev, nxt, ls, bullet_id, keep, widow, tab_id=None) -> str:
    pps = _sk_container(head, "paraProperties")
    el = copy.deepcopy(template)
    el.set("id", str(len(pps)))
    if tab_id is not None:
        el.set("tabPrIDRef", tab_id)
    el.find(f"{HH}align").set("horizontal", align)
    hd = el.find(f"{HH}heading")
    hd.set("type", "BULLET" if bullet_id else "NONE")
    hd.set("idRef", bullet_id or "0")
    hd.set("level", "0")
    bs = el.find(f"{HH}breakSetting")
    bs.set("breakNonLatinWord", "KEEP_WORD")  # 어절 단위 줄바꿈
    bs.set("keepWithNext", str(keep))
    bs.set("widowOrphan", str(widow))
    for m in el.iter(f"{HH}margin"):
        for tag, val in (("intent", indent), ("left", left), ("prev", prev), ("next", nxt)):
            m.find(f"{SK_HC}{tag}").set("value", str(val))
    for s in el.iter(f"{HH}lineSpacing"):
        s.set("value", str(ls))
    pps.append(el)
    _sk_renumber(pps)
    return el.get("id")


def _sk_add_style(head, name, pp, cp):
    styles = _sk_container(head, "styles")
    sid = str(len(styles))
    etree.SubElement(styles, f"{HH}style", id=sid, type="PARA", name=name, engName="", paraPrIDRef=pp,
                     charPrIDRef=cp, nextStyleIDRef=sid, langID="1042", lockForm="0")
    _sk_renumber(styles)


def _sk_build_header(xml: bytes) -> bytes:
    head = etree.fromstring(xml)
    _sk_set_fonts(head)
    char_tpl = head.find(f".//{HH}charPr[@id='0']")
    para_tpl = head.find(f".//{HH}paraPr[@id='0']")
    if char_tpl is None or para_tpl is None:
        raise SkeletonBaseError("header.xml에 id 0 글자 모양 또는 문단 모양 없음")
    char_tpl, para_tpl = copy.deepcopy(char_tpl), copy.deepcopy(para_tpl)
    toc_tab = _sk_add_toc_tab(head)
    for name, role, font, pt, bold, align, left, indent, prev, nxt, ls, bkey, keep, widow in SKELETON_STYLES:
        cp = _sk_add_char(head, char_tpl, 1 if font == "head" else 0, pt, bold)
        bullet_id = _sk_add_bullet(head, SKELETON_GLYPHS[bkey]) if bkey else None
        pp = _sk_add_para(head, para_tpl, align, left, indent, prev, nxt, ls, bullet_id, keep, widow,
                          toc_tab if role.startswith("toc_") else None)
        _sk_add_style(head, name, pp, cp)
    _sk_apply_base_styles(head)
    return etree.tostring(head, xml_declaration=True, encoding="UTF-8", standalone=True)


def _sk_apply_base_styles(head):
    """빈 새 문서의 기본 스타일(바탕글, 쪽 번호, 머리말, 각주)을 설정부 값으로 맞춤. 값이 같으면 XML은 그대로"""
    for name, _role, _font, pt, bold, align, left, indent, prev, nxt, ls, _b, keep, widow in SKELETON_BASE_STYLES:
        st = head.find(f".//{HH}style[@name='{name}']")
        if st is None:
            raise SkeletonBaseError(f"header.xml에 기본 스타일 {name} 없음")
        cp = head.find(f".//{HH}charPr[@id='{st.get('charPrIDRef')}']")
        pp = head.find(f".//{HH}paraPr[@id='{st.get('paraPrIDRef')}']")
        cp.set("height", str(int(round(pt * 100))))
        cp.set("textColor", SKELETON_TEXT_COLOR)
        for b in cp.findall(f"{HH}bold"):
            cp.remove(b)
        if bold:
            cp.insert(list(cp).index(cp.find(f"{HH}offset")) + 1, etree.Element(f"{HH}bold"))
        pp.find(f"{HH}align").set("horizontal", align)
        bs = pp.find(f"{HH}breakSetting")
        bs.set("keepWithNext", str(keep))
        bs.set("widowOrphan", str(widow))
        for m in pp.iter(f"{HH}margin"):
            # hp:switch의 default 분기는 구형 단위(HWPUNIT의 2배)로 기록. 한/글 저작 파일과 같은 방식
            scale = 2 if m.getparent().tag == f"{HP}default" else 1
            for tag, val in (("intent", indent), ("left", left), ("prev", prev), ("next", nxt)):
                m.find(f"{SK_HC}{tag}").set("value", str(val * scale))
        for sp in pp.iter(f"{HH}lineSpacing"):
            sp.set("value", str(ls))


def _sk_build_section(xml: bytes) -> bytes:
    s = xml.decode("utf-8")
    m = re.search(r"<hp:margin [^>]*/>", s)
    if m is None:
        raise SkeletonBaseError("section0.xml에 쪽 여백 요소 없음")
    new = ('<hp:margin header="{header}" footer="{footer}" gutter="0" left="{left}" right="{right}" '
           'top="{top}" bottom="{bottom}"/>').format(**SKELETON_PAGE_MARGIN)
    s = s.replace(m.group(0), new, 1)
    pm = re.search(r'<hp:pagePr [^>]*>', s)
    if pm is None:
        raise SkeletonBaseError("section0.xml에 쪽 크기 요소 없음")
    sized = re.sub(r'width="\d+"', f'width="{SKELETON_PAGE_SIZE[0]}"', pm.group(0))
    sized = re.sub(r'height="\d+"', f'height="{SKELETON_PAGE_SIZE[1]}"', sized)
    s = s.replace(pm.group(0), sized, 1)
    # 첫 문단의 줄 배치 캐시(horzsize)는 빈 새 문서의 기본 여백(좌우 30mm) 기준 값으로 남아 있음
    # 한/글 저작 파일처럼 캐시를 실제 본문 폭과 일치시켜 여백 설정과 캐시가 어긋나지 않게 함
    text_width = SKELETON_PAGE_SIZE[0] - SKELETON_PAGE_MARGIN["left"] - SKELETON_PAGE_MARGIN["right"]
    s = re.sub(r'(<hp:lineseg [^>]*?horzsize=")\d+(")', rf'\g<1>{text_width}\g<2>', s)
    return s.encode("utf-8")


def default_skeleton_path(base_dir: str) -> str:
    return os.path.join(base_dir, D_TEMPLATE, F_SKELETON)


def write_default_skeleton(base_dir: str) -> str:
    """기본 스켈레톤을 base_dir/_양식/에 생성하고 경로를 반환. 호출마다 같은 내용으로 다시 씀
    Why 매번 재생성: 스킬을 갱신해도 작업폴더의 양식이 설정부와 어긋나지 않게 함. 양식 변경은 설정부 수정이나 build --skeleton"""
    base = _skeleton_base_path()
    if not os.path.exists(base):
        sys.exit(f"ERROR: python-hwpx 빈 새 문서 없음 {base}. python-hwpx 설치 상태와 버전 확인 (준비 절의 API 호출부 수정)")
    out = default_skeleton_path(base_dir)
    os.makedirs(os.path.dirname(out), exist_ok=True)
    tmp = out + ".tmp"
    try:
        with zipfile.ZipFile(base) as zin, zipfile.ZipFile(tmp, "w") as zout:
            for n in zin.namelist():
                d = zin.read(n)
                if n == "Contents/header.xml":
                    d = _sk_build_header(d)
                elif n == "Contents/section0.xml":
                    d = _sk_build_section(d)
                elif n == "Preview/PrvText.txt":
                    d = b""
                method = zipfile.ZIP_STORED if n == "mimetype" else zipfile.ZIP_DEFLATED
                zout.writestr(n, d, compress_type=method)
    # 구조 변경의 유형: 필수 요소 부재(SkeletonBaseError, AttributeError), 부품 누락(KeyError), 순서 불일치(ValueError), XML 손상
    except (SkeletonBaseError, AttributeError, KeyError, ValueError, etree.XMLSyntaxError) as e:
        if os.path.exists(tmp):
            os.remove(tmp)
        sys.exit(f"ERROR: 기본 스켈레톤 생성 실패 ({e}). python-hwpx 빈 새 문서의 구조 변경 추정 (준비 절의 API 호출부 수정)")
    os.replace(tmp, out)
    return out


def cmd_init(a):
    work = a.work
    kind = a.kind
    st = load_state(work)
    if st.get("kind") and st["kind"] != kind:
        sys.exit(f"ERROR: 이 작업폴더의 산출 유형은 {st['kind']}. 유형을 바꾸려면 새 작업폴더에서 init")
    os.makedirs(os.path.join(work, D_MANUSCRIPT), exist_ok=True)
    os.makedirs(os.path.join(work, D_DELIVER), exist_ok=True)
    made = []
    if kind == "proposal":  # 제안전략서, 회사자료, 공통 파트 템플릿은 제안서 전용
        for src, dst in (("strategy.md", F_STRATEGY), ("company.md", F_COMPANY)):
            target = os.path.join(work, dst)
            if not os.path.exists(target):
                shutil.copy(os.path.join(ASSET_DIR, src), target)
                made.append(dst)
        for dst, src in COMMON_TEMPLATES.items():
            target = os.path.join(work, D_MANUSCRIPT, dst)
            if not os.path.exists(target):
                shutil.copy(os.path.join(ASSET_DIR, "common", src), target)
                made.append(f"{D_MANUSCRIPT}/{dst}")
    elif kind in DESIGN_TEMPLATES:  # 보고서의 진단서, 기획서의 기획 설계
        src, dst = DESIGN_TEMPLATES[kind]
        target = os.path.join(work, dst)
        if not os.path.exists(target):
            shutil.copy(os.path.join(ASSET_DIR, src), target)
            made.append(dst)
    made.append(os.path.relpath(write_default_skeleton(work), work))
    st = load_state(work)
    st["kind"] = kind
    st["stages"]["init"] = datetime.now().isoformat(timespec="seconds")
    save_state(work, st)
    print(f"OK: {work} ({kind}) 생성 파일 {made}")
    if kind == "proposal":
        print("다음: `selftest`로 렌더 경로 PASS 확인 후 `rfp` 실행과 착수 질의. 00_제안전략서.md와 회사자료.md의 [확인] 항목은 유저 답변으로 채움")
    else:
        guide = "references/report.md" if kind == "report" else "references/planning.md"
        print(f"다음: `selftest`로 렌더 경로 PASS 확인 후 {guide}의 절차로 진행")


# ═══════════════════════════════════════════════════════════════════════
# 5. RFP 변환: 원문 텍스트와 빈 표 템플릿 (추출은 모델이 수행)
# ═══════════════════════════════════════════════════════════════════════
MERGE_UP = "^"      # 원고 표기와 같은 병합 표지: 위 셀과 세로 병합
MERGE_LEFT = "<"    # 왼쪽 셀과 가로 병합


def _hwpx_paragraph_lines(p) -> list[str]:
    """문단 1개를 글자 줄과 표 줄로 변환. 표는 문단 안의 run에 들어 있으므로 글자 순서대로 끊어서 배출"""
    out, buf = [], []
    for run in p.findall(f"{HP}run"):
        for child in run:
            tag = etree.QName(child).localname
            if tag == "t":
                buf.append("".join(child.itertext()))
            elif tag == "tbl":
                if "".join(buf).strip():
                    out.append("".join(buf).strip())
                buf = []
                out += [""] + _hwpx_table_lines(child) + [""]
            else:  # 글상자 등 하위 문단을 가진 개체는 그 문단을 이어서 배출
                for sub in child.iter(f"{HP}subList"):
                    for sp in sub.findall(f"{HP}p"):
                        out += _hwpx_paragraph_lines(sp)
    if "".join(buf).strip():
        out.append("".join(buf).strip())
    return out


def _hwpx_cell_text(tc) -> str:
    """셀의 문단을 <br>로 이음. 셀 안의 표는 글자만 이어 붙임(표 안 표는 마크다운 표로 표현 불가)"""
    parts = []
    for sub in tc.findall(f"{HP}subList"):
        for sp in sub.findall(f"{HP}p"):
            text = " ".join(line for line in _hwpx_paragraph_lines(sp) if line and not line.startswith("|"))
            nested = [" ".join(t.itertext()).strip() for t in sp.iter(f"{HP}tbl")]
            parts.append(" ".join([text] + [n for n in nested if n]).strip())
    return "<br>".join(x for x in parts if x).replace("|", "/")


def _hwpx_table_lines(tbl) -> list[str]:
    """셀 주소와 병합 범위로 격자를 채움. python-hwpx export_markdown은 세로 병합 뒤의 셀을 앞으로 당겨 배점이 어긋남"""
    rows, cols = int(tbl.get("rowCnt", "0")), int(tbl.get("colCnt", "0"))
    if not rows or not cols:
        return []
    grid = [["" for _ in range(cols)] for _ in range(rows)]
    for tr in tbl.findall(f"{HP}tr"):
        for tc in tr.findall(f"{HP}tc"):
            addr, span = tc.find(f"{HP}cellAddr"), tc.find(f"{HP}cellSpan")
            r, c = int(addr.get("rowAddr")), int(addr.get("colAddr"))
            rs = int(span.get("rowSpan", "1")) if span is not None else 1
            cs = int(span.get("colSpan", "1")) if span is not None else 1
            if r >= rows or c >= cols:
                continue
            grid[r][c] = _hwpx_cell_text(tc)
            for rr in range(r, min(r + rs, rows)):
                for cc in range(c, min(c + cs, cols)):
                    if (rr, cc) != (r, c):
                        grid[rr][cc] = MERGE_LEFT if rr == r else MERGE_UP
    if all(cell in ("", MERGE_UP, MERGE_LEFT) for row in grid for cell in row):
        return []  # 쪽 배치용 빈 표는 원문 정보가 없어 생략
    lines = ["| " + " | ".join(grid[0]) + " |", "|" + "---|" * cols]
    return lines + ["| " + " | ".join(row) + " |" for row in grid[1:]]


def hwpx_to_markdown(path: str) -> str:
    with zipfile.ZipFile(path) as z:
        names = sorted((n for n in z.namelist() if re.match(r"Contents/section\d+\.xml$", n)),
                       key=lambda n: int(re.search(r"\d+", os.path.basename(n)).group()))
        lines = []
        for n in names:
            root = etree.fromstring(z.read(n))
            for p in root.findall(f"{HP}p"):
                lines += _hwpx_paragraph_lines(p)
    return "\n".join(lines)


def read_document(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    if ext == ".hwpx":
        return hwpx_to_markdown(path)
    if ext == ".pdf":
        if not shutil.which("pdftotext"):
            sys.exit("ERROR: pdftotext 부재. `apt-get install -y poppler-utils` 후 재실행")
        return subprocess.run(["pdftotext", "-layout", path, "-"], capture_output=True, text=True).stdout
    if ext == ".docx":
        import docx
        d = docx.Document(path)
        out = [p.text for p in d.paragraphs]
        for t in d.tables:
            rows = [[c.text.strip().replace("\n", " ") for c in r.cells] for r in t.rows]
            if rows:
                out.append("| " + " | ".join(rows[0]) + " |")
                out.append("|" + "---|" * len(rows[0]))
                out.extend("| " + " | ".join(r) + " |" for r in rows[1:])
        return "\n".join(out)
    if ext in (".md", ".txt"):
        return open(path, encoding="utf-8").read()
    sys.exit(f"ERROR: 지원하지 않는 입력 형식 {ext} (hwpx, pdf, docx, md)")


REQ_TABLE_HEADER = ["번호", "상위 항목", "요구항목 원문", "배치 섹션", "비고"]
SCORE_TABLE_HEADER = ["평가항목", "세부평가항목", "평가기준 원문", "배점", "결정 섹션"]


def _md_table(header: list[str], rows: list[list[str]]) -> list[str]:
    out = ["| " + " | ".join(header) + " |", "|" + "---|" * len(header)]
    out += ["| " + " | ".join(r) + " |" for r in rows]
    return out


def cmd_rfp(a):
    """RFP를 텍스트로 변환하고 빈 표 템플릿을 배치한다.
    Why: RFP의 절 구성, 글머리, 표 사용 여부가 발주처마다 달라 정규식 추출은 표형 RFP에서 0건, PDF에서 줄 분절을 낸다.
    요구항목과 배점은 모델이 01_RFP원문.md를 읽고 표에 원문 그대로 옮겨 적는다. 후속 명령(ledger, check rfp)은 표만 읽는다."""
    work = a.work
    if not os.path.isfile(a.input):
        sys.exit(f"ERROR: RFP 파일 없음: {a.input}")
    os.makedirs(work, exist_ok=True)
    text = read_document(a.input)

    head = ["# RFP 원문", f"- 입력 파일 {os.path.basename(a.input)}", f"- 변환 문자 수 {len(text)}",
            f"- 표의 `{MERGE_UP}`는 위 셀, `{MERGE_LEFT}`는 왼쪽 셀과 병합된 칸", "", "## 원문 전체", "", text]
    open(os.path.join(work, F_RFP_TEXT), "w", encoding="utf-8").write("\n".join(head))

    kind = load_state(work).get("kind") or "proposal"
    place = ("02_섹션대장.md의 섹션 ID" if kind == "proposal" else "02_스토리라인.md 목차 설계 표의 절 번호(예시로 2.1)")
    req = ["# 요구항목 전수 추출표",
           "- 01_RFP원문.md의 과업 내용(세부 수행내용, 과업 범위) 절을 읽고 최하위 항목을 행 단위로 원문 그대로 이관. 요약과 재서술 금지",
           "- 상위 항목 열에는 그 항목이 속한 상위 제목(예: 1. 현황 분석 > 가. 정책 동향)을 기입. 배치 섹션 판단의 근거",
           f"- 배치 섹션 열에 {place}를 기입 (복수는 쉼표). 전 행 기입 후 `check rfp`",
           "- 원문의 표 안 항목도 행으로 이관. 셀 안의 `|`는 `/`로 치환. 원문 표의 `^`는 위 셀, `<`는 왼쪽 셀과 병합된 칸", "",
           *_md_table(REQ_TABLE_HEADER, [["R1", "", "", "", ""]])]
    open(os.path.join(work, F_REQ), "w", encoding="utf-8").write("\n".join(req) + "\n")

    sc = ["# 배점표",
          "- 01_RFP원문.md의 제안서 평가 절(평가항목, 배점, 평가기준)을 읽고 세부평가항목 단위로 전개표를 채움. 배점 열은 숫자만",
          "- 결정 섹션은 그 점수를 읽고 매기는 섹션 ID (복수는 쉼표, 균등 분할). 정량 항목(신용등급 등 별지 서식)은 `정량`",
          "- 사업 기본조건(사업명, 기간, 예산, 참가자격, 제출기한)과 작성 규정(모호 표현 금지, 계량화, 분량 상한, 참조표, 약어)은 원문에서 확인하여 00_제안전략서.md에 기입", "",
          "## 전개표", "", *_md_table(SCORE_TABLE_HEADER, [["", "", "", "", ""]])]
    if kind == "proposal":  # 배점표와 ledger는 제안서 전용. 보고서의 배치 섹션은 절 번호
        open(os.path.join(work, F_SCORE), "w", encoding="utf-8").write("\n".join(sc) + "\n")

    st = load_state(work)
    st["stages"]["rfp"] = datetime.now().isoformat(timespec="seconds")
    st["files"]["rfp"] = os.path.abspath(a.input)
    save_state(work, st)
    if kind == "proposal":
        print(f"OK: {F_RFP_TEXT} ({len(text)}자), {F_REQ}와 {F_SCORE} 빈 표 배치")
        print("다음: 01_RFP원문.md를 읽고 01_요구항목표.md와 01_배점표.md 전개표를 채운 뒤 `ledger`")
    else:
        print(f"OK: {F_RFP_TEXT} ({len(text)}자), {F_REQ} 빈 표 배치")
        print("다음: 01_RFP원문.md를 읽고 01_요구항목표.md를 채움. 배치 섹션 열은 2단계 목차의 절 번호")


# ═══════════════════════════════════════════════════════════════════════
# 5-1. 발표자료 입력: pptx와 pdf를 PDF 기준본으로 통일하여 원문을 변환 (본체는 deck_source.py)
# ═══════════════════════════════════════════════════════════════════════
def cmd_deck(a):
    """발표자료를 RFP와 같은 방식으로 다룸. 본문은 파일로 변환. 추출과 판단은 원문 전문을 읽은 모델이 수행"""
    import deck_source
    work = a.work
    os.makedirs(work, exist_ok=True)
    deck_source.parse(a.input, os.path.join(work, D_DECK), allow_garbled=a.allow_garbled)
    st = load_state(work)
    st["stages"]["deck"] = datetime.now().isoformat(timespec="seconds")
    st["files"]["deck"] = os.path.abspath(a.input)
    save_state(work, st)
    deck_source.say(f"다음: {D_DECK}/원문.md 전문 읽기 (01_RFP원문.md와 같은 방식). 도식은 {D_DECK}/sheets의 축소판 모음에서 확인")


def cmd_figure(a):
    import deck_source
    out_dir = os.path.join(a.work, D_DECK)
    if a.grid:
        deck_source.grid(out_dir, a.page)
    else:
        deck_source.figure(out_dir, a.page, a.crop, a.name, not a.no_trim, a.pdf)


# ═══════════════════════════════════════════════════════════════════════
# 6. 섹션대장
# ═══════════════════════════════════════════════════════════════════════
def parse_table_rows(md: str, header_key: str) -> list[list[str]]:
    """header_key를 포함한 파이프 표의 데이터 행을 반환 (구분선 제외)."""
    lines = md.splitlines()
    for i, ln in enumerate(lines):
        if ln.strip().startswith("|") and header_key in ln:
            rows = []
            for ln2 in lines[i + 1:]:
                if not ln2.strip().startswith("|"):
                    break
                if TABLE_SEP_RE.match(ln2.strip()):
                    continue
                rows.append([c.strip() for c in ln2.strip().strip("|").split("|")])
            return rows
    return []


def score_by_section(work: str) -> tuple[dict[str, float], int]:
    rows = parse_table_rows(open(os.path.join(work, F_SCORE), encoding="utf-8").read(), "결정 섹션")
    totals: dict[str, float] = {}
    n = 0
    for r in rows:
        if len(r) < 5 or not r[3].strip():
            continue
        try:
            pts = float(re.sub(r"[^\d.]", "", r[3]))
        except ValueError:
            continue
        ids = [x.strip() for x in re.split(r"[,，/]", r[4]) if x.strip()]
        if not ids:
            continue
        n += 1
        for sid in ids:
            totals[sid] = totals.get(sid, 0) + pts / len(ids)
    return totals, n


def section_sort_key(sid: str) -> tuple:
    """Ⅱ-1, Ⅲ-2 같은 섹션 ID를 장 번호와 절 번호 순으로 정렬."""
    m = re.match(r"([ⅠⅡⅢⅣⅤⅥⅦⅧⅨⅩ]|[IVX]+|\d+)[-.]?(\d+)?", sid)
    if not m:
        return (99, 0, sid)
    ch = m.group(1)
    ci = ROMAN.index(ch) + 1 if ch in ROMAN else (["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X"].index(ch) + 1
                                                  if ch in ("I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X") else int(ch))
    return (ci, int(m.group(2) or 0), sid)


def cmd_ledger(a):
    work = a.work
    totals, n = score_by_section(work)
    path = os.path.join(work, F_LEDGER)
    st = load_state(work)
    existing = parse_table_rows(open(path, encoding="utf-8").read(), "섹션ID") if os.path.exists(path) else []
    by_id = {r[0]: r for r in existing if r and r[0]}
    ids = sorted(dict.fromkeys(list(by_id) + [k for k in totals if k != "정량"]), key=section_sort_key)
    out = ["# 섹션대장", f"- 배점표 전개 행 {n}건 집계. 통제배점은 결정 섹션 열에서 자동 산정, 유형은 오리지널 또는 공통",
           "- 원고파일은 원고/ 아래 파일명. 상태는 check 결과에서 갱신", "",
           "| 섹션ID | 섹션명 | 유형 | 통제배점 | 원고파일 | 상태 |", "|---|---|---|---|---|---|"]
    for sid in ids:
        r = by_id.get(sid, [sid, "", "", "", "", ""]) + [""] * 6
        pts = totals.get(sid)
        pts_s = f"{pts:g}" if pts is not None else r[3]
        status = []
        for key, val in st.get("checks", {}).items():
            if key.startswith("density:") and r[4] and r[4] in key:
                status.append("밀도 " + ("PASS" if val["pass"] else "FAIL"))
        out.append(f"| {sid} | {r[1]} | {r[2]} | {pts_s} | {r[4]} | {' '.join(status) or r[5]} |")
    unassigned = "정량 " + f"{totals['정량']:g}점" if "정량" in totals else ""
    open(path, "w", encoding="utf-8").write("\n".join(out) + "\n")
    st["stages"]["ledger"] = datetime.now().isoformat(timespec="seconds")
    save_state(work, st)
    print(f"OK: {F_LEDGER} 섹션 {len(ids)}개, 배점 합계 {sum(totals.values()):g}점 {unassigned}")
    if not existing:
        print("다음: 섹션명, 유형, 원고파일 열을 채운 뒤 다시 `ledger` (통제배점은 자동 유지)")


# ═══════════════════════════════════════════════════════════════════════
# 7. 점검 (check)
# ═══════════════════════════════════════════════════════════════════════
def _print_result(name: str, passed: bool, fails: list[str], warns: list[str] | None = None):
    print(f"[{name}] {'PASS' if passed else 'FAIL'}")
    for f in fails[:40]:
        print("   FAIL", f)
    if len(fails) > 40:
        print(f"   ... 외 {len(fails) - 40}건")
    for w in (warns or [])[:20]:
        print("   WARN", w)


OUT_OF_SCOPE = "범위 외"   # 사전진단보고서의 산출 범위 밖 요구항목 표기 (report.md 절차 절)
SECTION_NO_RE = re.compile(r"\d+\.\d+")


def rfp_place_ids(work: str) -> tuple[set[str], str, str]:
    """배치 섹션 열의 대조 원천. 제안서는 섹션대장의 섹션 ID, 사전진단보고서는 스토리라인 목차 설계 표의 절 번호
    반환: (유효 값 집합, 원천 이름, 원천 부재 시 안내 문구)"""
    kind = load_state(work).get("kind") or "proposal"
    if kind == "report":
        path = os.path.join(work, F_STORYLINE)
        rows = parse_table_rows(open(path, encoding="utf-8").read(), "작성 방식") if os.path.exists(path) else []
        ids = {m.group() for r in rows if len(r) > 1 for m in SECTION_NO_RE.finditer(r[1])}
        return ids, "스토리라인 목차 설계 표", f"{F_STORYLINE} 목차 설계 표의 절 번호 부재."
    path = os.path.join(work, F_LEDGER)
    ledger = parse_table_rows(open(path, encoding="utf-8").read(), "섹션ID") if os.path.exists(path) else []
    return {r[0] for r in ledger if r and r[0]}, "섹션대장", f"{F_LEDGER} 부재(2단계 ledger 전)."


def check_rfp(work: str) -> tuple[bool, list[str], list[str]]:
    if not os.path.exists(os.path.join(work, F_REQ)):
        return False, [f"{F_REQ} 부재. `rfp` 명령 미실행"], []
    rows = parse_table_rows(open(os.path.join(work, F_REQ), encoding="utf-8").read(), "요구항목 원문")
    ids, source, missing = rfp_place_ids(work)
    fails, warns = [], []
    if not ids:
        warns.append(f"{missing} 배치 섹션 값의 대조 생략. 해당 파일 작성 후 재실행")
    for r in rows:
        if len(r) < 4 or not r[0]:
            continue
        if not r[3].strip():
            fails.append(f"{r[0]} 배치 섹션 미기입: {r[2][:40]}")
            continue
        for sid in re.split(r"[,，/]", r[3]):
            sid = sid.strip()
            if sid and ids and sid not in ids and sid != OUT_OF_SCOPE:
                fails.append(f"{r[0]} 배치 섹션 '{sid}'이 {source}에 없음")
    if not rows:
        fails.append("요구항목표에 행이 없음")
    return not fails, fails, warns


def _segments(lines: list[str]) -> list[tuple[str, list[str]]]:
    """Task 또는 Step 밴드 단위로 원고를 분할. Task 구간은 다음 Task, Module 밴드, 장절 제목(#) 앞에서 끝남.
    Why: 구간을 다음 Task에서만 끊으면 마지막 Task 뒤의 항이나 다음 Module 도입부가 그 Task에 합산되어 오판이 생김"""
    segs: list[tuple[str, list[str]]] = []
    cur_name, cur = None, []
    for ln in lines:
        s = ln.strip()
        m = TASK_BAND_RE.match(s)
        if m or re.match(r"#{1,6}\s", s):
            if cur_name is not None:
                segs.append((cur_name, cur))
            is_task = m is not None and m.group(1).lower() != "module"
            cur_name, cur = (s[:40] if is_task else None), []
        else:
            cur.append(ln)
    if cur_name is not None:
        segs.append((cur_name, cur))
    return segs


def check_density(path: str) -> tuple[bool, list[str], list[str]]:
    lines = open(path, encoding="utf-8").read().splitlines()
    fails, warns = [], []
    elems, _ = parse_markdown(path)
    for msg in lint(elems):
        if "□ 단독" in msg:
            fails.append(msg)
        elif "HITL" not in msg:
            warns.append(msg)
    segs = _segments(lines)
    for name, body in segs:
        text = "\n".join(body)
        miss = []
        if "(목적)" not in text:
            miss.append("(목적) 서술")
        # Why: 원본에서 채택한 그림도 Task의 근거 블록으로 인정 (check volume과 같은 기준)
        if not any(ln.strip().startswith(("|", "![")) for ln in body):
            miss.append("표 또는 그림 (분석 중심 Task는 분석 방법에 맞는 채움본, 도출 중심 Task는 기준표나 시안)")
        if "산출물" not in text:
            miss.append("(산출물) 서술")
        if not QUANTITY_RE.search(text) and not CONCRETE_MARK_RE.search(text):  # 수치 또는 대상·자료·방법·기준의 특정
            miss.append("구체성 근거 (규모 수치 또는 대상·자료·방법·기준·절차 서술)")
        n_num = len(QUANTITY_RE.findall(text))
        if n_num > QUANTITY_MANY:
            warns.append(f"{name}: 수치 표기 {n_num}개. 셀 수 없는 업무에 수치를 붙였는지 확인")
        if miss:
            fails.append(f"{name}: {', '.join(miss)} 누락")
    if not segs and not any(ln.strip().startswith("|") for ln in lines):
        warns.append("표가 없는 원고. 근거 표 1개 이상 권장")
    return not fails, fails, warns


def _strategy_wording(work: str) -> tuple[list[list[str]], list[str]]:
    md = open(os.path.join(work, F_STRATEGY), encoding="utf-8").read()
    terms = parse_table_rows(md, "RFP 용어")
    forbidden = []
    m = re.search(r"^#+\s*금지 표현.*?$(.*?)(?=^#+\s|\Z)", md, re.S | re.M)
    if m:
        forbidden = [re.sub(r"^\s*[-+*]\s*", "", ln).strip() for ln in m.group(1).splitlines() if re.match(r"^\s*[-+*]\s+\S", ln)]
    return terms, forbidden


def check_wording(work: str) -> tuple[bool, list[str], list[str]]:
    terms, forbidden = _strategy_wording(work)
    # 대응 섹션 열의 섹션 ID는 섹션대장의 원고파일로 바꿔 범위를 정함. 파일명 조각을 적은 경우는 그대로 대조
    ledger = os.path.join(work, F_LEDGER)
    sec_files = {r[0]: os.path.basename(r[4]) for r in parse_table_rows(open(ledger, encoding="utf-8").read(), "섹션ID")
                 if len(r) > 4 and r[0] and r[4]} if os.path.exists(ledger) else {}
    files = manuscript_files(work)
    texts = {os.path.basename(f): open(f, encoding="utf-8").read() for f in files}
    fails, warns = [], []
    if not files:
        return False, ["원고/ 에 파일이 없음"], []
    for r in terms:
        if len(r) < 2 or not r[0] or r[0].startswith("["):
            continue
        rfp_term, use_term = r[0], (r[1] or r[0])
        sec = r[2] if len(r) > 2 else ""
        wanted = {sec_files.get(x.strip(), x.strip()) for x in re.split(r"[,，/]", sec) if x.strip()}
        scope = {k: v for k, v in texts.items() if not wanted or any(w in k for w in wanted)} or texts
        hits = sum(v.count(use_term) for v in scope.values())
        if hits == 0:
            fails.append(f"워딩 '{use_term}' (RFP '{rfp_term}') 출현 0회 (범위 {sec or '전체'})")
        variants = {k for k, v in texts.items() if rfp_term != use_term and rfp_term in v}
        if variants:
            warns.append(f"RFP 원표기 '{rfp_term}'와 제안서 표기 '{use_term}' 혼용: {sorted(variants)}")
    for term in dict.fromkeys(FORBIDDEN_DEFAULT + forbidden):
        if not term:
            continue
        for k, v in texts.items():
            for i, ln in enumerate(v.splitlines(), 1):
                if term in ln:
                    fails.append(f"금지 표현 '{term}' {k}:{i}")
    for pat, label in SENTENCE_WARN:
        n = 0
        for k, v in texts.items():
            for i, ln in enumerate(v.splitlines(), 1):
                if ln.strip().startswith("|") or ln.strip().startswith(">"):
                    continue
                if re.search(pat, ln):
                    n += 1
                    if n <= 3:
                        warns.append(f"{label} {k}:{i} {ln.strip()[:50]}")
        if n > 3:
            warns.append(f"{label} 외 {n - 3}건")
    return not fails, fails, warns


def check_residual(work: str) -> tuple[bool, list[str], list[str]]:
    """사전진단보고서와 기획보고서의 [확인] 잔여 위치. 판정에는 반영하지 않고 5단계 유저 질의 목록으로 씀"""
    warns = []
    for f in manuscript_files(work):
        for i, ln in enumerate(open(f, encoding="utf-8").read().splitlines(), 1):
            if "[확인" in ln:
                warns.append(f"[확인] 잔여 {os.path.basename(f)}:{i} {ln.strip()[:60]} (납품 전 유저 질의로 해소)")
    return True, [], warns


MANUAL_HEADING = "## 수동 기록"   # 05_점검기록.md에서 check all 재실행 시 보존하는 절


def write_check_report(work: str, auto_lines: list[str]):
    """자동 판정 결과를 다시 쓰되 수동 기록 절 이하는 보존. 자체 검토, 렌더링 검수, 재량 조정 기록이 재실행으로 지워지지 않게 함"""
    path = os.path.join(work, F_REPORT)
    old = open(path, encoding="utf-8").read() if os.path.exists(path) else ""
    i = old.find(MANUAL_HEADING)
    manual = old[i:] if i >= 0 else f"{MANUAL_HEADING}\n\n- 자체 검토, 렌더링 검수, 재량 조정 기록. check all을 다시 실행해도 이 절 이하는 보존\n"
    open(path, "w", encoding="utf-8").write("\n".join(auto_lines).rstrip() + "\n\n" + manual)


def manuscript_hash(work: str) -> str:
    """check all이 판정한 원고의 지문. build가 판정 이후의 원고 수정을 감지하는 기준. 모듈총괄표는 build가 원고에 삽입하므로 포함"""
    import hashlib
    h = hashlib.sha256()
    for f in manuscript_files(work) + [p for p in [module_table_path(work)] if os.path.exists(p)]:
        h.update(os.path.basename(f).encode("utf-8"))
        h.update(open(f, "rb").read())
    return h.hexdigest()


def check_fill(work: str) -> tuple[bool, list[str], list[str]]:
    fails, warns = [], []
    for f in [x for pre in PREFIX_FILL for x in manuscript_files(work, pre)]:
        name = os.path.basename(f)
        for i, ln in enumerate(open(f, encoding="utf-8").read().splitlines(), 1):
            if "[확인" in ln:
                fails.append(f"치환 잔여 {name}:{i} {ln.strip()[:60]}")
        rows = parse_table_rows(open(f, encoding="utf-8").read(), "참여율")
        if rows:
            hdr_line = next(ln for ln in open(f, encoding="utf-8") if "참여율" in ln and ln.strip().startswith("|"))
            cols = [c.strip() for c in hdr_line.strip().strip("|").split("|")]
            ci = next(i for i, c in enumerate(cols) if "참여율" in c)
            total = 0.0
            for r in rows:
                if len(r) > ci and re.search(r"\d", r[ci]):
                    v = float(re.sub(r"[^\d.]", "", r[ci]) or 0)
                    total += v
                    if v > 100:
                        fails.append(f"참여율 100% 초과 {name}: {r[0]} {v:g}%")
            warns.append(f"{name} 참여율 합계 {total:g}% (인월 산정과 대조)")
    # 오리지널 파트의 [확인]은 작성 중 임시 표지. 판정은 바꾸지 않고 잔여 위치만 알려 납품 전 유저 질의에 사용
    for f in [x for pre in PREFIX_ORIGINAL for x in manuscript_files(work, pre)]:
        name = os.path.basename(f)
        for i, ln in enumerate(open(f, encoding="utf-8").read().splitlines(), 1):
            if "[확인" in ln:
                warns.append(f"오리지널 파트 [확인] 잔여 {name}:{i} {ln.strip()[:60]} (납품 전 유저 질의로 해소)")
    return not fails, fails, warns


ACRONYM_RE = re.compile(r"[A-Z][A-Z0-9&]+")   # 영문 대문자 약어(RFP, ISP)는 짧아도 화이트리스트에 넣음. 한글 짧은 명사는 다른 판정을 가리므로 제외


def _whitelist_terms(work: str) -> list[str]:
    """핵심 워딩 사전(제안서)이나 진단서 RFP 용어 표(사전진단보고서)의 용어, 모듈총괄표의 Module·Task 명칭을 수집한다."""
    terms: list[str] = []
    if os.path.exists(os.path.join(work, F_STRATEGY)):
        for r in _strategy_wording(work)[0]:
            terms += [c.strip() for c in r[:2] if c.strip() and "[확인]" not in c]
    diag = os.path.join(work, F_DIAGNOSIS)
    if os.path.exists(diag):  # 사전진단보고서: 진단서 RFP 용어 표의 원문 용어와 원고 표기 (R&D 평가 과업의 '평가위원' 같은 도메인 용어)
        for r in parse_table_rows(open(diag, encoding="utf-8").read(), "RFP 용어"):
            terms += [c.strip() for c in r[:2] if c.strip() and "[확인]" not in c]
    mod = module_table_path(work)
    if os.path.exists(mod):
        for r in parse_table_rows(open(mod, encoding="utf-8").read(), "Module"):
            for c in r[:2]:
                name = TASK_BAND_RE.sub(lambda m: m.group(3), c.strip())
                if name:
                    terms.append(name)
    terms = [t for t in dict.fromkeys(terms)
             if (len(t) >= WHITELIST_MIN_LEN or ACRONYM_RE.fullmatch(t)) and "[확인]" not in t]
    return sorted(terms, key=len, reverse=True)   # 긴 용어부터 마스킹하여 부분 일치 방지


def _whitelist_ready(work: str) -> bool:
    """고정 시점은 2단계 완료 이후. 모듈총괄표가 있고 핵심 워딩 사전에 [확인]이 남지 않은 상태.
    Why: 골격 확정 전에 고정하면 정상적인 사전 채움과 모듈명 확정이 원천 변경 FAIL로 오판된다."""
    if not os.path.exists(module_table_path(work)):
        return False
    rows = _strategy_wording(work)[0] if os.path.exists(os.path.join(work, F_STRATEGY)) else []
    return bool(rows) and not any("[확인]" in c for r in rows for c in r[:2])


def _lock_whitelist(work: str, terms: list[str], reset: bool) -> tuple[str | None, str | None]:
    """고정 조건 충족 시 화이트리스트를 state.json에 고정. 이후 원천이 바뀌면 (FAIL 사유, None), 미고정이면 (None, WARN 사유)를 반환한다.
    Why: FAIL 회피 목적으로 사전이나 모듈명에 용어를 덧붙이는 행위를 차단. 해제는 유저 지시에 따른 --reset-whitelist만 허용."""
    st = load_state(work)
    locked = st.get("whitelist")
    if locked is None and not reset and not _whitelist_ready(work):
        return None, "화이트리스트 미고정 (모듈총괄표 부재 또는 핵심 워딩 사전에 [확인] 잔여). 2단계 완료 후 첫 check tone에서 고정"
    if locked is None or reset:
        st["whitelist"] = {"terms": terms, "locked_at": datetime.now().isoformat(timespec="seconds")}
        save_state(work, st)
        return None, None
    added = sorted(set(terms) - set(locked["terms"]))
    removed = sorted(set(locked["terms"]) - set(terms))
    if added or removed:
        return f"화이트리스트 원천 변경 (고정 {locked['locked_at']}). 추가 {added[:5]} 삭제 {removed[:5]}. 유저 지시가 있으면 `check tone --reset-whitelist`", None
    return None, None


def _mask_whitelist(ln: str, terms: list[str]) -> str:
    for t in terms:
        if t in ln:
            ln = ln.replace(t, WHITELIST_MASK * len(t))
    return ln


def _client_term_findings(ln: str) -> tuple[list[str], list[str]]:
    """기관 호칭 판정. 주어 자리의 '발주처'와 채우지 않은 호칭 자리는 FAIL, 그 밖의 '발주처'는 WARN.
    Why: '발주처 협의'처럼 관행상 남는 용례가 있어 전면 FAIL은 오탐이 많음. 주어 반복만 확실한 결함."""
    body = CLIENT_HEADER_CELL_RE.sub("|", ln)
    fails, warns = [], []
    if CLIENT_SLOT_RE.search(body):
        fails.append(CLIENT_SLOT_LABEL)
    if CLIENT_SUBJECT_RE.search(body):
        fails.append(CLIENT_SUBJECT_LABEL)
    elif CLIENT_TERM_RE.search(body):
        warns.append(CLIENT_TERM_LABEL)
    return fails, warns


def check_tone(work: str, reset_whitelist: bool = False, kind: str = "proposal") -> tuple[bool, list[str], list[str]]:
    fails, warns = [], []
    terms = _whitelist_terms(work)
    # 화이트리스트 고정은 제안서 전용 (핵심 워딩 사전이 제안전략서에 있음). 보고서와 기획서는 모듈총괄표 명칭만 가림
    lock_fail, lock_warn = _lock_whitelist(work, terms, reset_whitelist) if kind == "proposal" else (None, None)
    if lock_fail:
        fails.append(lock_fail)
    if lock_warn:
        warns.append(lock_warn)
    for f in manuscript_files(work):
        name = os.path.basename(f)
        for i, raw in enumerate(open(f, encoding="utf-8").read().splitlines(), 1):
            ln = _mask_whitelist(raw, terms)
            st = ln.strip()
            in_cell = st.startswith("|")
            for pat, label in TONE_FAIL + FRAMEWORK_FAIL:
                if in_cell and "교차 참조" in label:  # 표 셀의 Task·Module 번호는 총괄표·조견표·일정표의 정상 용법
                    continue
                if (TASK_BAND_RE.match(st) or st.startswith("<표>") or WIDTH_SPEC_RE.match(st)) and "교차 참조" in label:
                    continue
                if re.search(pat, ln):
                    fails.append(f"{label} {name}:{i} {raw.strip()[:50]}")
            for pat, label in TONE_WARN:
                if re.search(pat, ln):
                    warns.append(f"{label} {name}:{i} {raw.strip()[:50]}")
            client_fail, client_warn = _client_term_findings(ln)
            fails += [f"{label} {name}:{i} {raw.strip()[:50]}" for label in client_fail]
            warns += [f"{label} {name}:{i} {raw.strip()[:50]}" for label in client_warn]
            if HIER_LINE_RE.match(raw) and not PERIOD_EXEMPT_RE.match(raw):
                if PERIOD_MID_RE.search(raw):
                    fails.append(f"계위 1개에 문장 2개 (마침표로 이어진 문장. 뒤 문장을 별도 계위로 분리) {name}:{i} {raw.strip()[:50]}")
                elif PERIOD_END_RE.search(raw):
                    fails.append(f"줄 끝 마침표 (명사형 종결에 마침표를 붙이지 않음) {name}:{i} {raw.strip()[:50]}")
    return not fails, fails, warns


def _table_blocks(lines: list[str]) -> list[tuple[int, list[list[str]]]]:
    blocks, i = [], 0
    while i < len(lines):
        if lines[i].strip().startswith("|"):
            j = i
            rows = []
            while j < len(lines) and lines[j].strip().startswith("|"):
                if not TABLE_SEP_RE.match(lines[j].strip()):
                    rows.append([c.strip() for c in lines[j].strip().strip("|").split("|")])
                j += 1
            blocks.append((i + 1, rows))
            i = j
            continue
        i += 1
    return blocks


def check_numbers(work: str) -> tuple[bool, list[str], list[str]]:
    """표의 합계 행과 구성비 열을 재계산한다. 모델의 암산을 대체하는 자리."""
    fails, warns = [], []
    for f in manuscript_files(work):
        name = os.path.basename(f)
        lines = open(f, encoding="utf-8").read().splitlines()
        for start, rows in _table_blocks(lines):
            if len(rows) < 3:
                continue
            body = [r for r in rows[1:] if not SUM_LABEL_RE.match(r[0])]
            total_rows = [r for r in rows[1:] if SUM_LABEL_RE.match(r[0])]
            for tr in total_rows:
                for c in range(1, min(len(tr), len(rows[0]))):
                    if not NUM_CELL_RE.match(tr[c].replace(" ", "")):
                        continue
                    vals = []
                    for r in body:
                        if c < len(r) and NUM_CELL_RE.match(r[c].replace(" ", "")):
                            vals.append(float(r[c].replace(",", "").rstrip("%")))
                    if len(vals) < 2:
                        continue
                    shown = float(tr[c].replace(",", "").replace(" ", "").rstrip("%"))
                    calc = sum(vals)
                    if abs(shown - calc) > max(0.05, abs(calc) * 0.005):
                        fails.append(f"합계 불일치 {name}:{start} 열 '{rows[0][c]}' 표기 {shown:g} 계산 {calc:g}")
            for c in range(1, len(rows[0])):
                if not re.search(r"(비율|비중|구성비|점유율|%)", rows[0][c]):
                    continue
                vals = [float(r[c].replace(",", "").rstrip("%")) for r in body
                        if c < len(r) and NUM_CELL_RE.match(r[c].replace(" ", ""))]
                if len(vals) >= 3 and 90 < sum(vals) < 110 and abs(sum(vals) - 100) > 0.5:
                    warns.append(f"구성비 합계 {sum(vals):.1f}% {name}:{start} 열 '{rows[0][c]}'")
    return not fails, fails, warns




# ═══════════════════════════════════════════════════════════════════════
# 7-2. 분량 구성 단위 판정 (check volume)
# ═══════════════════════════════════════════════════════════════════════
def count_warnings(elems: list[Element]) -> list[str]:
    """형제 항목 수가 계위별 상한을 넘는 지점을 경고로 반환. 그룹마다 상한을 처음 넘는 시점에 1회만 알림.
    Why: 형제 항목이 상한을 넘으면 논점이 흐려지므로 항 분할이나 표 전환의 신호로 사용"""
    out: list[str] = []
    squares = circles = dashes = 0
    sec_name = sq_name = ci_name = ""
    for e in elems:
        if e.kind in ("heading", "band"):
            squares = circles = dashes = 0
            sec_name, sq_name, ci_name = e.text[:24], "", ""
        elif e.kind == "square":
            squares += 1
            if squares == SQUARE_PER_SECTION_MAX + 1:
                out.append(f"[계위 개수] '{sec_name}' 안의 □ {squares}개 이상 (상한 {SQUARE_PER_SECTION_MAX}). 항(###)으로 묶음")
            circles = dashes = 0
            sq_name, ci_name = e.text[:24], ""
        elif e.kind == "circle":
            circles += 1
            if sq_name and circles == CIRCLE_PER_SQUARE_MAX + 1:
                out.append(f"[계위 개수] □ '{sq_name}' 아래 ❍ {circles}개 이상 (상한 {CIRCLE_PER_SQUARE_MAX}). □를 나눔")
            dashes, ci_name = 0, e.text[:24]
        elif e.kind == "dash":
            dashes += 1
            if ci_name and dashes == DASH_PER_CIRCLE_MAX + 1:
                out.append(f"[계위 개수] ❍ '{ci_name}' 아래 하이픈 {dashes}개 이상 (상한 {DASH_PER_CIRCLE_MAX}). 표로 전환")
    return out


def check_volume(path: str, original: bool | None = None, style: bool | None = None) -> tuple[bool, list[str], list[str]]:
    """□ 1개당 ❍ 2개 이상과 근거(하이픈, ※) 2개 이상, Task 1개당 표 또는 그림 1개 이상과 본문 글자 수 하한.
    Why: 잘 쓴 제안서는 문장을 늘리지 않고 구성 단위를 채워 분량을 만든다. 단문 규율과 분량을 동시에 지키는 유일한 경로."""
    elems, _ = parse_markdown(path)
    fails, warns = [], []
    squares: list[dict] = []
    tasks: list[dict] = []
    cur_sq = cur_task = None
    if original is None:
        original = manuscript_kind(path)[:1] in PREFIX_ORIGINAL
    if style is None:  # □ 제목 줄 상한과 계위 개수 경고. 기획보고서는 하한 없이 이 경고만 적용
        style = original
    for e in elems:
        if e.kind == "band" and e.data.get("kind") == "task":
            cur_task = {"name": e.text[:30], "blocks": 0, "chars": 0}
            tasks.append(cur_task)
            cur_sq = None
        elif e.kind == "band" or e.kind == "heading":
            cur_sq = None
        elif e.kind == "square":
            cur_sq = {"name": e.text[:30], "title": e.text, "circle": 0, "evidence": 0, "chars": len(e.text), "block_first": False}
            squares.append(cur_sq)
        elif e.kind in ("circle", "dash", "dot", "note"):
            if len(e.text) > LINE_MAX[e.kind]:
                warns.append(f"{e.kind} 항목 {len(e.text)}자 ({'1줄' if e.kind == 'note' else '2줄'} 상한 {LINE_MAX[e.kind]}자). 문장을 나누어 계위를 더하거나 표로 전환: {e.text[:24]}")
            if cur_sq is not None:
                cur_sq["chars"] += len(e.text)
                if e.kind == "circle":
                    cur_sq["circle"] += 1
                else:
                    cur_sq["evidence"] += 1
        if cur_sq is not None and e.kind in ("table", "image") and cur_sq["circle"] == 0:
            cur_sq["block_first"] = True   # □ 바로 아래에 ❍ 없이 표나 그림이 옴
        if cur_task is not None:
            if e.kind in ("table", "image"):
                cur_task["blocks"] += 1
                cur_task["chars"] += sum(len(c) for row in e.data.get("rows", []) for c in row)
            elif e.kind in ("square", "circle", "dash", "dot", "note"):
                cur_task["chars"] += len(e.text)
    for s in squares:
        if s["circle"] == 0:
            fails.append(f"□ '{s['name']}' 단독. □는 논점이며 아래에 ❍ 서술 1개 이상 필요. 표만 이어지면 표가 말하는 바를 ❍로 먼저 서술")
        elif s["block_first"]:
            fails.append(f"□ '{s['name']}' 바로 아래에 ❍ 없이 표나 그림. 표 앞에 표의 내용을 말하는 ❍를 배치")
        elif original and s["circle"] < VOLUME_MIN_CIRCLE:
            fails.append(f"□ '{s['name']}' 아래 ❍ {s['circle']}개 (하한 {VOLUME_MIN_CIRCLE}). 주장을 나누거나 □를 합침")
        if style and len(s["title"]) > SQUARE_TITLE_MAX:
            warns.append(f"□ '{s['name']}' 제목 {len(s['title'])}자 (1줄 상한 {SQUARE_TITLE_MAX}자). 판단의 핵심어만 남기고 부연은 ❍로 이동")
        if not original:
            continue
        if s["evidence"] < VOLUME_MIN_EVIDENCE:
            warns.append(f"□ '{s['name']}' 근거 줄(하이픈, ※) {s['evidence']}개 (권장 {VOLUME_MIN_EVIDENCE} 이상). 수치, 자료명, 사례를 근거 계위에 추가")
        if s["chars"] < VOLUME_SQUARE_MIN_CHARS:
            warns.append(f"□ '{s['name']}' 글자 수 {s['chars']} (권장 {VOLUME_SQUARE_MIN_CHARS} 이상)")
    if style:
        warns += count_warnings(elems)
    for t in tasks:
        if t["blocks"] < VOLUME_MIN_BLOCK_PER_TASK:
            fails.append(f"{t['name']} 표 또는 그림 {t['blocks']}개. 분석 방법에 맞는 채움본 표나 기준표를 배치")
        if t["chars"] < VOLUME_TASK_MIN_CHARS:
            warns.append(f"{t['name']} 본문 글자 수 {t['chars']} (권장 {VOLUME_TASK_MIN_CHARS} 이상). 분석 대상 실명 목록, 사전조사 채움본, 진단 항목표, 비교 매트릭스로 보강")
    return not fails, fails, warns


def resolve_manuscript(work: str, name: str) -> str:
    """--file 경로 해석. 현재 폴더 기준, 작업폴더 기준, 원고 폴더 기준 순으로 찾음"""
    for cand in (name, os.path.join(work, name), os.path.join(work, D_MANUSCRIPT, os.path.basename(name))):
        if os.path.exists(cand):
            return cand
    sys.exit(f"ERROR: 원고 파일 없음 {name} (현재 폴더, 작업폴더, {D_MANUSCRIPT}/ 기준으로 찾지 못함)")


def cmd_check(a):
    work = a.work
    kind = work_kind(work)
    results = {}

    def run(name: str) -> bool:
        """개별 지정은 유형과 무관하게 실행. all은 유형별 점검 목록(CHECKS_BY_KIND)만 실행"""
        return a.what == name or (a.what == "all" and name in CHECKS_BY_KIND[kind])
    if run("rfp"):
        ok, fails, warns = check_rfp(work)
        _print_result("요구항목 배치", ok, fails, warns)
        record(work, "rfp", ok, {"fails": fails})
        results["요구항목 배치"] = (ok, fails, warns)
    if run("density"):
        targets = [resolve_manuscript(work, a.file)] if a.file else manuscript_files(work, PREFIX_DENSITY)
        for f in targets:
            ok, fails, warns = check_density(f)
            _print_result(f"밀도 {os.path.basename(f)}", ok, fails, warns)
            record(work, f"density:{os.path.basename(f)}", ok, {"fails": fails})
            results[f"밀도 {os.path.basename(f)}"] = (ok, fails, warns)
    if run("volume"):
        if kind == "proposal":
            targets = [resolve_manuscript(work, a.file)] if a.file else [f for p in PREFIX_VOLUME for f in manuscript_files(work, p)]
        else:
            targets = [resolve_manuscript(work, a.file)] if a.file else manuscript_files(work)
        original = None if kind == "proposal" else kind == "report"   # 보고서는 전 원고에 구성 단위 하한 적용
        for f in targets:
            ok, fails, warns = check_volume(f, original, None if kind == "proposal" else True)
            _print_result(f"분량 구성 {os.path.basename(f)}", ok, fails, warns)
            record(work, f"volume:{os.path.basename(f)}", ok, {"fails": fails})
            results[f"분량 구성 {os.path.basename(f)}"] = (ok, fails, warns)
    if run("wording"):
        ok, fails, warns = check_wording(work)
        _print_result("워딩", ok, fails, warns)
        record(work, "wording", ok, {"fails": fails})
        results["워딩"] = (ok, fails, warns)
    if run("tone"):
        ok, fails, warns = check_tone(work, a.reset_whitelist, kind)
        _print_result("어조와 어휘", ok, fails, warns)
        record(work, "tone", ok, {"fails": fails})
        results["어조와 어휘"] = (ok, fails, warns)
    if run("numbers"):
        ok, fails, warns = check_numbers(work)
        _print_result("수치 검산", ok, fails, warns)
        record(work, "numbers", ok, {"fails": fails})
        results["수치 검산"] = (ok, fails, warns)
    if run("fill"):
        name = "공통 파트 치환" if kind == "proposal" else "확인 표지 잔여"
        ok, fails, warns = check_fill(work) if kind == "proposal" else check_residual(work)
        _print_result(name, ok, fails, warns)
        record(work, "fill", ok, {"fails": fails})
        results[name] = (ok, fails, warns)
    if a.what == "all":
        out = ["# 점검기록", f"- 실행 {datetime.now().isoformat(timespec='minutes')}", ""]
        for name, (ok, fails, warns) in results.items():
            out.append(f"## {name} {'PASS' if ok else 'FAIL'}")
            out += [f"- (결함) {f}" for f in fails] + [f"- (경고) {w}" for w in warns] + [""]
        write_check_report(work, out)
        allok = all(v[0] for v in results.values())
        st = load_state(work)
        st["stages"]["check"] = datetime.now().isoformat(timespec="seconds")
        st["checks"]["all"] = {"pass": allok, "hash": manuscript_hash(work)}
        save_state(work, st)
        print(f"{'OK' if allok else 'FAIL'}: {F_REPORT} 기록. {'build 진행 가능' if allok else '결함 수정 후 재실행'}")
        sys.exit(0 if allok else 1)


# ═══════════════════════════════════════════════════════════════════════
# 8. 조판·렌더·스켈레톤 명령 (조판기 본체는 hwpx_report.py)
# ═══════════════════════════════════════════════════════════════════════
def _profile_for_builder(profile: dict) -> dict:
    """lib.Builder는 프로파일 값 전부를 스타일명으로 검증하므로 스타일명이 아닌 키(크기, 서식 문자열, 밴드 설정)를 뺀 사본을 넘긴다."""
    return {k: v for k, v in profile.items() if not k.endswith("_pt") and k not in PROFILE_EXTRA_KEYS}


def docx_spec() -> dict:
    """docx 작성기에 넘기는 양식 정의. hwpx 스켈레톤과 같은 설정부 값(SKELETON_*)만 사용"""
    return {"fonts": {"body": SKELETON_FONT_BODY, "head": SKELETON_FONT_HEAD}, "color": SKELETON_TEXT_COLOR,
            "styles": SKELETON_BASE_STYLES + SKELETON_STYLES, "glyphs": SKELETON_GLYPHS, "glyph_em": SKELETON_GLYPH_EM,
            "bullet_offset": int(SKELETON_BULLET_OFFSET), "toc_tab": SKELETON_TOC_TAB_POS, "band_border": BAND_BORDER_W,
            "space_em": SKELETON_SPACE_EM, "divider_top_lines": DIVIDER_TOP_LINES,
            "page": {"width": SKELETON_PAGE_SIZE[0], "height": SKELETON_PAGE_SIZE[1], **SKELETON_PAGE_MARGIN}}


def work_kind(work: str | None) -> str:
    """작업폴더의 산출 유형 (init --kind). 기록이 없으면 제안서"""
    return (load_state(work).get("kind") if work else None) or "proposal"


def module_table_parent(files: list[str]) -> str | None:
    """모듈총괄표 표가 어느 파일의 제목(# 또는 ##) 아래에 놓이는지 반환. 총괄표 파일 안에서 표보다 앞에 제목이 있으면 그 파일 자신.
    Why: 총괄표 파일에 장 제목이나 절 제목이 없으면 표가 앞 파일의 마지막 절(기본 구성에서는 제안사 일반현황) 아래에 조판된다."""
    parent = None
    for f in files:
        is_table_file = manuscript_kind(f).endswith(MODULE_TABLE_SUFFIX)
        for ln in open(f, encoding="utf-8").read().split("\n"):
            st = ln.strip()
            if is_table_file and (st.startswith("|") or st.startswith("<표>")):
                return parent
            if re.match(r"^#{1,2} ", st):
                parent = f
    return None


FOOTNOTE_REF_RE = re.compile(r"\[\^([^\]]+)\]")


def inline_note_footnotes(text: str, footnotes: dict) -> str:
    """표와 그림 하단 부기 행은 각주를 달 수 없는 투명 표 행이라 각주 표기가 글자 그대로 찍힘.
    각주 내용을 괄호로 풀어 적고 원고 수정을 경고"""
    def repl(m):
        body = footnotes.get(m.group(1), "")
        label = re.sub(r"^f\d+-", "", m.group(1))   # 합본 시 붙인 파일 범위 접두는 원고에 없으므로 빼고 안내
        print(f"WARN: 표 하단 부기 줄의 각주 [^{label}]를 괄호 출처로 조판. 원고에서 `※ 출처:` 줄로 옮김")
        return f" ({body})" if body else ""
    return FOOTNOTE_REF_RE.sub(repl, text)
FOOTNOTE_DEF_LINE_RE = re.compile(r"^\[\^([^\]]+)\]:", re.M)


def scope_footnotes(text: str, scope: str) -> tuple[str, list[str]]:
    """각주 라벨에 파일 범위 접두를 붙임. 합본 파서는 각주 정의를 라벨 하나의 사전으로 모으므로
    파일마다 [^1]을 쓰면 뒤 파일의 정의가 앞 파일 각주를 덮어씀. 반환: (변환 원고, 경고 목록)"""
    defs = FOOTNOTE_DEF_LINE_RE.findall(text)
    warns = [f"각주 정의 중복 [^{x}]" for x in sorted({d for d in defs if defs.count(d) > 1})]
    undefined = sorted(set(FOOTNOTE_REF_RE.findall(text)) - set(defs))
    warns += [f"정의 없는 각주 [^{x}] (빈 각주로 조판됨)" for x in undefined]
    return FOOTNOTE_REF_RE.sub(lambda m: f"[^{scope}-{m.group(1)}]", text), warns


def cmd_build(a):
    work = a.input if os.path.isdir(a.input) else None
    kind = work_kind(work)
    profile_key = a.profile or KINDS[kind]["profile"]
    profile = PROFILES.get(profile_key) if profile_key in PROFILES else json.load(open(profile_key, encoding="utf-8"))
    formats = ("hwpx", "docx") if a.format == "both" else (a.format,)
    if "docx" in formats and (a.skeleton or profile_key not in PROFILES):
        sys.exit("ERROR: docx는 기본 양식(설정부 정의)만 지원. 발주처 양식(--skeleton, 프로파일 JSON) 조판은 --format hwpx")
    skeleton = None
    if "hwpx" in formats:
        if a.skeleton:
            skeleton = a.skeleton
        elif profile_key in PROFILES:  # 기본 프로파일: 작업폴더(단일 원고면 원고가 있는 폴더)의 _양식/에 생성
            skeleton = write_default_skeleton(work or os.path.dirname(os.path.abspath(a.input)))
        else:
            sys.exit("ERROR: 프로파일 JSON을 쓰면 --skeleton으로 스켈레톤 hwpx를 함께 지정")
        if not os.path.exists(skeleton):
            sys.exit(f"ERROR: 스켈레톤 없음 {skeleton}")
    if work:  # 작업폴더: 원고/*.md를 파일명 순으로 합본
        files = manuscript_files(work)
        if not files:
            sys.exit(f"ERROR: {work}/{D_MANUSCRIPT}/ 에 원고 없음")
        merged = os.path.join(work, D_DELIVER, "_합본.md")
        os.makedirs(os.path.dirname(merged), exist_ok=True)
        has_design = os.path.exists(os.path.join(work, F_MODULE_TABLE))
        inserted = 0
        with open(merged, "w", encoding="utf-8") as fo:
            for i, f in enumerate(files, 1):
                text, fn_warns = scope_footnotes(open(f, encoding="utf-8").read(), f"f{i}")
                for w in fn_warns:
                    print(f"WARN: {os.path.basename(f)} {w}")
                if has_design:
                    text, n = expand_module_table(work, text)
                    inserted += n
                fo.write(text.rstrip() + "\n\n")
        # 제안서는 수행 방안 첫머리에 모듈총괄표 1회 필수. 보고서와 기획서는 선택(0회 또는 1회)
        need_one = kind == "proposal"
        if has_design and (inserted != 1 if need_one else inserted > 1) and not legacy_module_table(work):
            msg = (f"모듈총괄표 삽입 지시 줄 {inserted}개. 수행 방안이 시작되는 자리의 ❍ 서술 뒤에 "
                   "`(삽입: 모듈총괄표)` 줄을 1회 기입한 뒤 build")
            if not a.force:
                sys.exit("ERROR: " + msg)
            print("WARN: " + msg)
        src = merged
        out_base = os.path.splitext(a.output)[0] if a.output else os.path.join(work, D_DELIVER, KINDS[kind]["out"])
        parent = module_table_parent(files)
        if kind == "proposal" and parent and manuscript_kind(parent)[:1] in PREFIX_FILL:  # --force는 시험 조판용이므로 중단 대신 경고
            msg = (f"모듈총괄표가 공통 파트 제목({os.path.basename(parent)}) 아래에 놓임. "
                   f"수행 방안이 시작되는 장 제목(#)이나 절 제목(##)을 모듈총괄표 파일 첫 줄로 옮긴 뒤 build")
            if not a.force:
                sys.exit("ERROR: " + msg)
            print("WARN: " + msg)
        st = load_state(work)
        rec = st.get("checks", {}).get("all", {})
        if not rec.get("pass") and not a.force:
            sys.exit("ERROR: `check all` PASS 기록 없음. 결함을 수정하고 check all을 통과한 뒤 build. 시험 조판만 --force")
        if rec.get("hash") and rec["hash"] != manuscript_hash(work) and not a.force:
            sys.exit("ERROR: check all 통과 이후 원고나 모듈총괄표가 바뀜. check all을 다시 실행한 뒤 build")
        print(f"합본 {len(files)}개 파일 → {merged}")
    else:
        src = a.input
        out_base = os.path.splitext(a.output or a.input)[0]
    elems, footnotes = parse_markdown(src)
    msgs = lint(elems)
    for msg in msgs:
        print("LINT:", msg)
    if a.long_table is None and any("[장표 HITL]" in m for m in msgs):
        sys.exit("ERROR: 장표 대상 표가 있으나 --long-table 미지정. 착수 질의 답변(engine 또는 split)을 --long-table로 지정한 뒤 build")

    def make(fmt: str, out: str, numbers: dict[str, int] | None = None) -> ProposalLayout:
        """조판 1회. numbers를 주면 목차에 쪽번호를 기입한 조판본을 다시 쓴다. 원고 요소는 조판 중에 변하므로 사본을 넘김"""
        if fmt == "hwpx":
            b = ProposalBuilder(skeleton, _profile_for_builder(profile), page_break=a.page_break)
            b.p = profile
        else:
            b = ProposalDocxBuilder(docx_spec(), profile, page_break=a.page_break)
        b.numbered_square = False
        b.toc = a.toc
        b.page_number = a.page_number
        b.toc_depth = a.toc_depth
        b.cover = dict(kv.split("=", 1) for kv in a.cover.split(";") if "=" in kv) if a.cover else None
        b.header_spec = a.header
        b.merge_mode = a.merge
        b.long_table_mode = a.long_table or "split"
        b.toc_numbers = numbers or {}
        b.unify_mode = a.unify
        b.theme = a.theme
        b.doc_title = a.title
        b.divider = (a.divider or ("on" if "report" in (kind, profile_key) else "off")) == "on"   # 간지 기본값: 보고서 ON, 그 밖 OFF
        b.n["chapter"] = a.chapter_start - 1   # 부분 원고의 장 번호를 전체 목차에 맞춤
        b.build(copy.deepcopy(elems), footnotes, out)
        return b

    made = {}
    for fmt in formats:
        out = f"{out_base}.{fmt}"
        b = make(fmt, out)
        if fmt == "hwpx":
            fixed = lib.repair_manifest(out)
            if fixed:
                print(f"WARN: 매니페스트 미등록 이미지 {fixed}건 자동 보정")
        if a.toc and a.page_number:  # 2패스: 렌더링으로 제목과 캡션의 쪽 위치를 찾아 목차에 기입 (근사값, 한/글과 Word에서 확정)
            numbers = toc_page_numbers(out, fmt, [t for _, t in b.headings], [(t, needle) for _, t, needle in b.captions],
                                       {t for lv, t in b.headings if lv == 1})
            if numbers:
                make(fmt, out, numbers)
                if fmt == "hwpx":
                    lib.repair_manifest(out)
                # 집계는 목차에 실리는 항목만. Why: 제목은 항(###)까지, 캡션은 전부 탐지 입력에 넣지만 목차는 toc_depth와 표 목차 설정까지만 실음
                listed = [t for lv, t in b.headings if lv <= b.toc_depth] + [t for _, items in b.toc_lists() for t in items]
                print(f"목차 쪽번호 {sum(t in numbers for t in listed)}/{len(listed)}건 기입 ({fmt}, 근사 렌더러 기준)")
            else:
                print(f"WARN: 렌더러 또는 pdftotext 부재로 {fmt} 목차 쪽번호 미기입")
        made[fmt] = out
        print(f"OK: {out} (요소 {len(elems)}, 각주 {len(footnotes)})")
    if len(made) == 2:  # 두 형식 동시 산출: 문단 글자 순서 대조
        diffs = check_parity(made["hwpx"], made["docx"])
        if diffs:
            print(f"WARN: hwpx와 docx 문단 글자 불일치 {len(diffs)}건 (조판 코드 결함 신호)")
            for d in diffs[:5]:
                print("   ", d)
        else:
            print(f"OK: hwpx와 docx 문단 글자 일치 ({len(hwpx_paragraph_texts(made['hwpx'])[0])}문단, 목차 쪽번호 제외)")
    if work:
        st = load_state(work)
        st["stages"]["build"] = datetime.now().isoformat(timespec="seconds")
        for fmt, out in made.items():
            st["files"][fmt] = os.path.abspath(out)
        save_state(work, st)


W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
TOC_NUM_RE = re.compile(r"\t\d+$")   # 목차 쪽번호는 형식별 렌더링 기준이므로 대조에서 제외


def hwpx_paragraph_texts(path: str) -> tuple[list[str], list[str]]:
    """hwpx 본문 문단 글자(표 셀과 캡션 포함, 머리말 제외)와 각주 글자. 문서 순서"""
    body, notes = [], []
    with zipfile.ZipFile(path) as z:
        for n in sorted(x for x in z.namelist() if re.match(r"Contents/section\d+\.xml$", x)):
            for p in etree.fromstring(z.read(n)).iter(f"{HP}p"):
                anc = {a.tag for a in p.iterancestors()}
                if f"{HP}header" in anc or f"{HP}footer" in anc:
                    continue
                buf = []
                for t in p.iter(f"{HP}t"):
                    owner = next(a for a in t.iterancestors() if a.tag == f"{HP}p")
                    if owner is not p:
                        continue
                    buf.append(t.text or "")
                    buf += ["\t" + (c.tail or "") for c in t if c.tag == f"{HP}tab"]
                text = "".join(buf)
                if text.strip():
                    (notes if f"{HP}footNote" in anc else body).append(text)
    return body, notes


def docx_paragraph_texts(path: str) -> tuple[list[str], list[str]]:
    """docx 본문 문단 글자와 각주 글자. 캡션 번호 필드와 각주 번호 뒤 괄호는 hwpx 자동번호와 맞추려고 뺌"""
    with zipfile.ZipFile(path) as z:
        root = etree.fromstring(z.read("word/document.xml"))
        fn = etree.fromstring(z.read("word/footnotes.xml")) if "word/footnotes.xml" in z.namelist() else None
    body = []
    for p in root.iter(f"{W_NS}p"):
        buf, after_ref = [], False
        for el in p.iter(f"{W_NS}t", f"{W_NS}tab", f"{W_NS}footnoteReference"):
            if el.tag == f"{W_NS}footnoteReference":
                after_ref = True
            elif any(a.tag == f"{W_NS}fldSimple" for a in el.iterancestors()):
                continue
            elif el.tag == f"{W_NS}tab":
                buf.append("\t")
            elif after_ref and el.text == docx_lib.FOOTNOTE_SUFFIX:
                after_ref = False
            else:
                buf.append(el.text or "")
        if "".join(buf).strip():
            body.append("".join(buf))
    notes = []
    for note in (fn.findall(f"{W_NS}footnote") if fn is not None else []):
        if not note.get(f"{W_NS}type"):
            text = "".join(t.text or "" for t in note.iter(f"{W_NS}t"))
            notes.append(text[len(docx_lib.FOOTNOTE_SUFFIX):].strip())
    return body, notes


def check_parity(hwpx_path: str, docx_path: str) -> list[str]:
    """같은 원고에서 나온 hwpx와 docx의 문단 글자 순서 대조. 차이 목록을 반환 (없으면 일치)"""
    hb, hn = hwpx_paragraph_texts(hwpx_path)
    db, dn = docx_paragraph_texts(docx_path)
    hb, db = [TOC_NUM_RE.sub("", t) for t in hb], [TOC_NUM_RE.sub("", t) for t in db]
    diffs = [f"{i + 1}번째 문단: hwpx '{a[:30]}' docx '{b[:30]}'" for i, (a, b) in enumerate(zip(hb, db)) if a != b]
    if len(hb) != len(db):
        diffs.append(f"문단 수 hwpx {len(hb)} docx {len(db)}")
    if [t.strip() for t in hn] != dn:
        diffs.append(f"각주 hwpx {len(hn)}건 docx {len(dn)}건 불일치")
    return diffs


def toc_page_numbers(path: str, fmt: str, headings: list[str], captions: list[tuple[str, str]],
                     chapters: set[str] | None = None) -> dict[str, int]:
    """조판본을 텍스트 레이어가 남는 PDF로 렌더링하여 목차 항목의 본문 쪽번호를 찾음. hwpx는 rhwp 우선, docx는 LibreOffice"""
    if not shutil.which("pdftotext"):
        return {}
    pdf = os.path.splitext(path)[0] + ".toc.pdf"
    if fmt == "hwpx":
        if not (lib.find_rhwp() or lib.find_soffice()):
            return {}
        engine = lib.resolve_engine("auto", None)
        src = lib.render_copy(path, strip_newnum=(engine == "soffice"))
    else:
        if not lib.find_soffice():
            return {}
        engine, src = "soffice", path
    try:
        lib.to_pdf(src, pdf, engine, None, text_as_paths=False)
    except SystemExit:
        return {}
    finally:
        if src != path and os.path.exists(src):
            os.remove(src)
    pages = subprocess.run(["pdftotext", "-layout", pdf, "-"], capture_output=True, text=True).stdout.split("\f")
    os.remove(pdf)
    return lib.match_toc_pages(pages, headings, captions, chapters)


_ORIGINAL_COL_ALIGNS = lib.col_aligns


def _col_aligns_symbol_neutral(rows: list[list[str]], explicit: list[str | None]) -> list[str]:
    """기호 하나만 있는 셀(비교 매트릭스의 ○, 일정표의 ■)은 열 정렬 판정에서 글머리로 세지 않는다.
    Why: lib는 ○로 시작하는 셀이 있으면 열을 LEFT로 정한다. 기호 열이 ●, ◐와 함께 좌측으로 쏠렸다. 명시 마커는 종전대로 우선."""
    masked = [rows[0]] + [["" if v.strip() in lib.BULLET_PREFIXES else v for v in r] for r in rows[1:]]
    return _ORIGINAL_COL_ALIGNS(masked, explicit)


lib.col_aligns = _col_aligns_symbol_neutral

_ORIGINAL_RENDER_COPY = lib.render_copy


def _cell_bullet_paras(header_xml: str) -> dict[str, str]:
    """표 셀 글머리 스타일(스타일명에 '표' 포함, heading BULLET)의 paraPr id → 글머리 문자."""
    bullets = dict(re.findall(r'<hh:bullet id="(\d+)"[^>]*char="([^"]*)"', header_xml))
    out = {}
    for m in re.finditer(r'<hh:paraPr id="(\d+)".*?</hh:paraPr>', header_xml, re.S):
        pp_id, body = m.group(1), m.group(0)
        h = re.search(r'<hh:heading type="BULLET" idRef="(\d+)"', body)
        if h and re.search(rf'<hh:style [^>]*name="[^"]*표[^"]*"[^>]*paraPrIDRef="{pp_id}"', header_xml):
            out[pp_id] = bullets.get(h.group(1), "◦").replace("\u00ad", "-")
    return out


def _render_copy_with_bullets(src: str, strip_newnum: bool = False) -> str:
    """검수 사본에서 글머리 정의(header.xml)의 소프트 하이픈을 하이픈으로 바꾸고 표 셀 글머리를 본문 글자로 넣는다.
    Why: 근사 렌더러는 소프트 하이픈 글머리와 표 셀 안의 글머리를 그리지 않아 계위와 셀 항목이 빈 자리로 보인다. 원본 hwpx는 무변경."""
    work = _ORIGINAL_RENDER_COPY(src, strip_newnum)
    tmp = work + ".tmp"
    with zipfile.ZipFile(work) as zin:
        cell_bullets = _cell_bullet_paras(zin.read("Contents/header.xml").decode("utf-8"))
    with zipfile.ZipFile(work) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for it in zin.infolist():
            d = zin.read(it.filename)
            if it.filename == "Contents/header.xml":
                d = d.decode("utf-8").replace('char="\u00ad"', 'char="-"').encode("utf-8")
            # rhwp(v0.8.6 기준)는 셀 글머리를 직접 그린다. 글자 삽입은 soffice 경로(strip_newnum)에만 적용하여 이중 표시를 막는다
            if it.filename.startswith("Contents/section") and it.filename.endswith(".xml") and cell_bullets and strip_newnum:
                t = d.decode("utf-8")
                for pp_id, ch in cell_bullets.items():
                    t = re.sub(rf'(<hp:p [^>]*paraPrIDRef="{pp_id}"[^>]*>.*?<hp:t>)', rf'\g<1>{ch} ', t, flags=re.S)
                d = t.encode("utf-8")
            if it.filename == "mimetype":
                zout.writestr(zipfile.ZipInfo("mimetype"), d, compress_type=zipfile.ZIP_STORED)
            else:
                zout.writestr(it.filename, d)
    os.replace(tmp, work)
    return work


def cmd_render(a):
    lib.render_copy = _render_copy_with_bullets
    return lib.cmd_render(a)


def _run_self(args: list[str]) -> subprocess.CompletedProcess:
    """이 CLI를 하위 프로세스로 실행. Why: build와 render의 sys.exit 오류를 잡아 판정 결과로 기록하기 위함"""
    return subprocess.run([sys.executable, os.path.abspath(__file__), *args], capture_output=True, text=True)


def cmd_selftest(a):
    """표본 원고를 조판하고 렌더링하여 렌더 경로의 동작을 판정. 결과는 state.json checks.render에 기록.
    Why: 렌더러 결함을 컨텍스트 여유가 큰 착수 시점에 발견하여 6단계의 확보 실패를 막음"""
    out_dir = os.path.abspath(os.path.join(a.work, D_SELFTEST))  # Why: state.json 기록 경로를 실행 위치와 무관하게 둠
    os.makedirs(out_dir, exist_ok=True)
    sample_md = os.path.join(out_dir, "표본.md")
    sample_hwpx = os.path.join(out_dir, "표본.hwpx")
    open(sample_md, "w", encoding="utf-8").write(SELFTEST_SAMPLE)

    built = _run_self(["build", sample_md, sample_hwpx, "--skeleton", write_default_skeleton(a.work)])
    if built.returncode != 0 or not os.path.exists(sample_hwpx):
        record(a.work, "render", False, {"stage": "build", "error": (built.stderr or built.stdout)[-400:]})
        sys.exit(f"FAIL[build]: 표본 조판 실패. 의존성 설치 상태 확인 (준비 절)\n{(built.stderr or built.stdout)[-400:]}")

    render_args = ["render", sample_hwpx, "--engine", a.engine] + (["--rhwp", a.rhwp] if a.rhwp else [])
    rendered = _run_self(render_args)
    pages_dir = os.path.splitext(sample_hwpx)[0] + "_pages"
    pages = len(glob.glob(os.path.join(pages_dir, "*.png")))
    engine = m.group(1) if (m := re.search(r"OK\[([^\]]+)\]", rendered.stdout)) else ""
    passed = rendered.returncode == 0 and pages >= SELFTEST_MIN_PAGES
    detail = {"stage": "render", "engine": engine, "pages": pages, "png_dir": pages_dir}
    if not passed:
        detail["error"] = (rendered.stderr or rendered.stdout)[-400:]
    record(a.work, "render", passed, detail)
    if not passed:
        sys.exit(f"FAIL[render]: 엔진 {engine or '없음'}, PNG {pages}쪽. 준비 절의 다음 순위 렌더러로 이동\n{detail['error']}")
    print(f"PASS[{engine}]: 표본 {pages}쪽 렌더링 → {pages_dir}/ (PNG를 열어 글머리, 밴드, 표 표시 확인)")
    # docx 경로: 같은 표본을 docx로 조판하고 LibreOffice로 렌더링. 결과는 checks.render_docx (docx 납품 시에만 관문)
    sample_docx = os.path.join(out_dir, "표본_docx.docx")   # PNG 폴더가 hwpx 표본과 겹치지 않게 이름을 나눔
    built = _run_self(["build", sample_md, sample_docx, "--format", "docx"])
    if built.returncode != 0 or not os.path.exists(sample_docx):
        record(a.work, "render_docx", False, {"stage": "build", "error": (built.stderr or built.stdout)[-400:]})
        print(f"WARN[docx]: 표본 docx 조판 실패. python-docx 설치 확인 (준비 절)\n{(built.stderr or built.stdout)[-400:]}")
        return
    if not lib.find_soffice():
        record(a.work, "render_docx", False, {"stage": "render", "error": "LibreOffice 없음"})
        print("WARN[docx]: LibreOffice 부재로 docx 렌더 검수 불가. docx 납품 시 준비 절의 LibreOffice 설치 수행")
        return
    rendered = _run_self(["render", sample_docx])
    docx_pages_dir = os.path.splitext(sample_docx)[0] + "_pages"
    n_docx = len(glob.glob(os.path.join(docx_pages_dir, "*.png")))
    ok = rendered.returncode == 0 and n_docx >= SELFTEST_MIN_PAGES
    record(a.work, "render_docx", ok, {"stage": "render", "pages": n_docx, "png_dir": docx_pages_dir})
    print(f"{'PASS' if ok else 'WARN'}[docx]: 표본 {n_docx}쪽 렌더링 → {docx_pages_dir}/")


def main():
    ap = argparse.ArgumentParser(description="juja-builder CLI", epilog=lib.CLI_NOTICE,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)

    p = sub.add_parser("init", help="작업폴더 생성")
    p.add_argument("work")
    p.add_argument("--kind", default="proposal", choices=sorted(KINDS), help="산출 유형: proposal(제안서, 기본), report(사전진단보고서), plan(기획보고서)")
    p.set_defaults(func=cmd_init)

    p = sub.add_parser("rfp", help="RFP(hwpx|pdf|docx|md) → 01_RFP원문.md와 요구항목표, 배점표 빈 표")
    p.add_argument("input")
    p.add_argument("work")
    p.set_defaults(func=cmd_rfp)

    p = sub.add_parser("ledger", help="배점표 집계로 섹션대장 생성·갱신")
    p.add_argument("work")
    p.set_defaults(func=cmd_ledger)

    p = sub.add_parser("deck", help="제안 발표자료(pptx|pdf) → 발표자료/ 원문.md와 도식 캡처 재료")
    p.add_argument("input")
    p.add_argument("work")
    p.add_argument("--allow-garbled", action="store_true", help="텍스트 레이어 불완전 PDF를 쪽 PNG 판독 전제로 진행 (유저 결정 시에만)")
    p.set_defaults(func=cmd_deck)

    p = sub.add_parser("figure", help="발표자료 쪽 → 삽입용 도식 PNG (발표자료/figures)")
    p.add_argument("work")
    p.add_argument("page", type=int)
    p.add_argument("--grid", action="store_true", help="자르기 값을 읽기 위한 10% 눈금 PNG만 생성")
    p.add_argument("--crop", help="좌,상,우,하 백분율 (예: 4,28,97,96)")
    p.add_argument("--name", help="파일명 접미")
    p.add_argument("--no-trim", action="store_true", help="바깥 여백 자동 제거를 끔")
    p.add_argument("--pdf", help="캡처 원천 PDF의 교체. pptx 변환본에서 글자가 넘친 도식을 PowerPoint 저장 PDF에서 캡처")
    p.set_defaults(func=cmd_figure)

    p = sub.add_parser("check", help="rfp | density | volume | wording | tone | numbers | fill | all")
    p.add_argument("what", choices=["rfp", "density", "volume", "wording", "tone", "numbers", "fill", "all"])
    p.add_argument("work")
    p.add_argument("--file", help="density/volume 대상 원고 1개")
    p.add_argument("--reset-whitelist", action="store_true", help="tone 화이트리스트 재고정 (유저 지시 시에만)")
    p.set_defaults(func=cmd_check)

    p = sub.add_parser("build", help="작업폴더(원고/*.md 합본) 또는 원고.md → hwpx, docx")
    p.add_argument("input")
    p.add_argument("output", nargs="?", help="산출 경로 (확장자는 --format이 정함)")
    p.add_argument("--profile", default=None, help="proposal | report | plan | 프로파일 JSON 경로 (기본은 작업폴더 산출 유형, 단일 원고는 proposal)")
    p.add_argument("--format", default="hwpx", choices=["hwpx", "docx", "both"], help="산출 형식 (착수 질의 답변). 두 형식은 같은 원고와 같은 양식 정의로 조판")
    p.add_argument("--title", help="본문 첫 줄 문서 제목 (기획보고서)")
    p.add_argument("--chapter-start", type=int, default=1, help="첫 장 번호 (전체 목차 중 일부 장만 산출할 때 전체 기준 번호)")
    p.add_argument("--divider", choices=["on", "off"], default=None, help="장 간지 (기본: 사전진단보고서 on, 제안서와 기획보고서 off)")
    p.add_argument("--skeleton", help="스켈레톤 hwpx (기본은 작업폴더/_양식/proposal.hwpx를 실행마다 생성)")
    p.add_argument("--page-break", default="chapter", choices=["chapter", "section", "none"])
    p.add_argument("--long-table", default=None, choices=["engine", "split"], help="긴 표의 기본 처리(착수 질의 답변). split=청크별 표로 분할하고 제목행 반복, engine=한/글 엔진 분할. 장표 대상 표가 있는데 미지정이면 build 중단")
    p.add_argument("--merge", default="auto", choices=["auto", "explicit", "none"])
    p.add_argument("--toc", action="store_true", help="정적 목차 삽입 (기본 OFF. 납품 목차는 한/글 차례 기능)")
    p.add_argument("--toc-depth", type=int, default=2, choices=[1, 2, 3])
    p.add_argument("--page-number", action="store_true", help="꼬리말 쪽번호 (기본 OFF)")
    p.add_argument("--cover", help='표지: "title=제목;subtitle=부제;date=YYYY. M.;org=기관"')
    p.add_argument("--header", help='머리말: "텍스트;align=RIGHT"')
    p.add_argument("--unify", default="manual", choices=["manual", "auto"], help="표 폭 통일: manual=원고의 (폭: ...) 지시만 적용(기본), auto=같은 양식 표 자동 그룹 추가")
    p.add_argument("--theme", default=DEFAULT_THEME, choices=sorted(THEMES), help="밴드 색 테마: navy(기본), mono(검정과 회색). 표 서식은 테마와 무관")
    p.add_argument("--force", action="store_true")
    p.set_defaults(func=cmd_build)

    p = sub.add_parser("render", help="hwpx 또는 docx → PDF + PNG (hwpx_report.py render 위임. docx는 LibreOffice)")
    p.add_argument("input")
    p.add_argument("--outdir")
    p.add_argument("--dpi", type=int, default=lib.RENDER_DPI)
    p.add_argument("--rhwp")
    p.add_argument("--keep-soft-hyphen", action="store_true")
    p.add_argument("--engine", default="auto", choices=["auto", "rhwp", "soffice"])
    p.set_defaults(func=cmd_render)

    p = sub.add_parser("selftest", help="표본 원고 조판과 렌더링으로 렌더 경로 판정 (착수 관문, state.json checks.render 기록)")
    p.add_argument("work")
    p.add_argument("--rhwp")
    p.add_argument("--engine", default="auto", choices=["auto", "rhwp", "soffice"])
    p.set_defaults(func=cmd_selftest)

    p = sub.add_parser("skeleton", help="발주처 양식.hwpx → 스켈레톤 + 프로파일 초안 (hwpx_report.py skeleton 위임)")
    p.add_argument("input")
    p.add_argument("output")
    p.add_argument("--section", type=int)
    p.set_defaults(func=lib.cmd_skeleton)

    a = ap.parse_args()
    a.func(a)


if __name__ == "__main__":
    main()
