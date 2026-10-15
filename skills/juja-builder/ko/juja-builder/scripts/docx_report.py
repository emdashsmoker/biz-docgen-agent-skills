#!/usr/bin/env python3
"""docx_report.py: 공통 조판 계층의 결과를 docx로 기록하는 작성기.

구성: 원고 파싱, 인라인 표기, 표 배치 산정(열 폭, 행 높이, 병합, 정렬, 괘선 변), 그림 크기, 표지와 목차 순서는
hwpx_report.py의 형식 무관 함수를 그대로 쓴다. 이 파일은 그 결과를 Word 문서 요소로 옮기는 일만 담당한다.
양식 수치(글꼴, 크기, 정렬, 여백, 줄 간격, 계위 기호, 쪽 설정)는 호출부가 넘기는 spec에서만 읽는다.
Why python-docx: 패키지 부품, 관계, 그림, 머리말과 꼬리말은 라이브러리에 맡기고 API가 없는 서식
(글머리 정의, 셀 괘선과 채움, 셀 여백, 필드, 각주 부품)만 OOXML 요소로 보강하여 Word 호환 결함을 줄인다.
동작 확인 환경: python-docx 1.2.0, lxml 6.x, LibreOffice 24.x(렌더 검수)
"""
from __future__ import annotations

import copy
import io
import os
import sys
from datetime import datetime
from xml.sax.saxutils import escape

import hwpx_report as lib

try:
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
    from docx.opc.constants import CONTENT_TYPE as CT, RELATIONSHIP_TYPE as RT
    from docx.opc.packuri import PackURI
    from docx.opc.part import Part
    from docx.oxml import OxmlElement, parse_xml
    from docx.oxml.ns import nsdecls, qn
    from docx.shared import Emu, Pt, RGBColor, Twips
except ImportError:  # 조판 시점에만 필요. hwpx 전용 경로는 python-docx 없이 동작
    Document = None

# ═══════════════════════════════════════════════════════════════════════
# 설정부: 단위 환산과 Word 전용 보정값
# ═══════════════════════════════════════════════════════════════════════
HU_PER_TWIP = 5            # HWPUNIT(1/7200 inch) → twip(1/1440 inch)
EMU_PER_HU = 127           # HWPUNIT → EMU(1/914400 inch)
HU_PER_PT = 100
TWIP_PER_PT = 20
EIGHTHS_PER_MM = 8 * 72 / 25.4   # 괘선 굵기 w:sz 단위(1/8 pt)
SPACER_MIN_TWIP = 20       # 간격 문단 최소 높이(1pt). Word는 두 표 사이에 문단이 없으면 표를 합침
JC = {"LEFT": WD_ALIGN_PARAGRAPH.LEFT, "CENTER": WD_ALIGN_PARAGRAPH.CENTER, "RIGHT": WD_ALIGN_PARAGRAPH.RIGHT,
      "JUSTIFY": WD_ALIGN_PARAGRAPH.JUSTIFY, "DISTRIBUTE": WD_ALIGN_PARAGRAPH.DISTRIBUTE} if Document else {}
FOOTNOTE_SUFFIX = ")"      # 각주 번호 뒤 글자. hwpx 각주 번호 형식(숫자와 닫는 괄호)과 같게 둠
TABLE_IN_MARGIN = (510, 510, 141, 141)          # 표 기본 셀 여백 좌, 우, 위, 아래 (python-hwpx 표 기본값과 같음)
# 서식 요소의 스키마 순서. Word는 순서가 어긋난 요소를 손상으로 판정함
TCPR_ORDER = ("cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar",
              "textDirection", "tcFitText", "vAlign", "hideMark")
TBLPR_ORDER = ("tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize", "tblW",
               "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook")
RPR_ORDER = ("rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike", "outline", "shadow",
             "emboss", "imprint", "noProof", "snapToGrid", "vanish", "webHidden", "color", "spacing", "w", "kern",
             "position", "sz", "szCs", "highlight", "u", "effect", "bdr", "shd", "fitText", "vertAlign", "rtl", "cs",
             "em", "lang", "eastAsianLayout", "specVanish", "oMath")
PPR_ORDER = ("pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
             "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct",
             "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind",
             "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection", "textAlignment",
             "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange")
SPACE_SPLIT_RE = __import__("re").compile(r"( +)")
SECTPR_ORDER = ("headerReference", "footerReference", "footnotePr", "endnotePr", "type", "pgSz", "pgMar", "paperSrc",
                "pgBorders", "lnNumType", "pgNumType", "cols", "formProt", "vAlign", "noEndnote", "titlePg",
                "textDirection", "bidi", "rtlGutter", "docGrid", "printerSettings")


def seq_name(kind: str) -> str:
    """캡션 번호 필드(SEQ) 이름. caption 라벨의 seq, 없으면 label. Word 필드 이름에 빈칸을 쓸 수 없어 밑줄로 바꿈"""
    cap = lib.label("caption.table" if kind == "TABLE" else "caption.figure")
    return "_".join((cap["seq"] or cap["label"]).split())


def tw(hu: float) -> int:
    return int(round(hu / HU_PER_TWIP))


def _insert_ordered(parent, child, order: tuple[str, ...]):
    """스키마 순서를 지켜 자식 요소를 넣는다. 같은 이름의 기존 요소는 교체"""
    name = child.tag.split("}")[1]
    old = parent.find(qn(f"w:{name}"))
    if old is not None:
        parent.replace(old, child)
        return child
    later = order[order.index(name) + 1:]
    for i, c in enumerate(parent):
        if c.tag.split("}")[-1] in later:
            parent.insert(i, child)
            return child
    parent.append(child)
    return child


def _w(tag: str, **attrs) -> "OxmlElement":
    el = OxmlElement(f"w:{tag}")
    for k, v in attrs.items():
        el.set(qn(f"w:{k}"), str(v))
    return el


class DocxBuilder:
    """lib.Builder(hwpx)와 같은 기록 함수를 docx로 구현. 판단 함수(표 배치, 인라인, 목차 순서)는 lib를 호출"""

    def __init__(self, spec: dict, profile: dict, page_break: str = "chapter"):
        if Document is None:
            sys.exit("ERROR: python-docx 미설치. `pip install -U python-docx --break-system-packages` 후 재실행")
        self.doc = Document()
        self.spec, self.p, self.page_break = spec, profile, page_break
        self.n = {"chapter": 0, "section": 0, "subsection": 0, "ganada": 0, "table": 0, "figure": 0, "total_section": 0}
        self.headings: list[tuple[int, str]] = []
        self.toc_numbers: dict[str, int] = {}
        self.pending_break = False
        self.after_block = False
        self.toc, self.cover, self.page_number, self.header_spec, self.toc_depth = False, None, False, None, 2
        self.merge_mode, self.long_table_mode = "auto", "engine"
        self.last_para = None
        self._warned: set[str] = set()
        self._footnotes: list[tuple[int, str]] = []
        page = spec["page"]
        self.body_w = page["width"] - page["left"] - page["right"]
        self.table_max_w = self.body_w - lib.TABLE_INDENT_MARGIN
        self.rows = {r[0]: r for r in spec["styles"]}          # 스타일명 → 정의 행
        self.base_names = {r[1]: r[0] for r in spec["styles"]}  # 역할 → 스타일명 (머리말, 쪽 번호, 각주)
        self._clean_template()
        self._setup_page()
        self._setup_defaults()
        self._num_ids = self._setup_numbering()
        self._setup_styles()
        self.manual_symbol = {role: sym for role, sym in lib.MANUAL_SYMBOLS.items()
                              if self.p.get(role) and self.p[role] in self.rows and not self.rows[self.p[role]][11]}
        # hwpx 구역의 첫 문단(구역 설정을 담는 빈 바탕글 줄)과 같은 빈 줄을 본문 첫머리에 둠
        self._new_para("body")

    # ── 양식: 쪽, 기본값, 글머리, 스타일 (spec 한 곳에서 읽음) ────────────────
    def _clean_template(self):
        """python-docx 기본 템플릿의 미리보기 그림, 사용자 정의 XML, 작성자 정보를 지움"""
        pkg = self.doc.part.package
        for rid, rel in list(pkg.rels.items()):
            if rel.reltype == RT.THUMBNAIL:
                del pkg.rels[rid]
        for rid, rel in list(self.doc.part.rels.items()):
            if rel.reltype == RT.CUSTOM_XML:
                self.doc.part.drop_rel(rid)
        settings = self.doc.settings.element
        zoom = settings.find(qn("w:zoom"))
        if zoom is not None:  # 기본 템플릿의 zoom에 스키마 필수 속성(percent)이 빠져 있음
            zoom.set(qn("w:percent"), "100")
        for tag in ("savePreviewPicture", "rsids", "attachedTemplate"):
            for el in settings.findall(qn(f"w:{tag}")):
                settings.remove(el)
        cp = self.doc.core_properties
        cp.author = cp.last_modified_by = cp.comments = cp.title = cp.subject = cp.keywords = ""
        cp.revision = 1
        cp.created = cp.modified = datetime.now().replace(microsecond=0)

    def _setup_page(self):
        pg = self.spec["page"]
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Twips(tw(pg["width"])), Twips(tw(pg["height"]))
        sec.left_margin, sec.right_margin = Twips(tw(pg["left"])), Twips(tw(pg["right"]))
        # 한/글의 위쪽 여백은 쪽 끝에서 머리말 영역까지, 머리말 여백은 그 영역 높이. Word의 위쪽 여백은 쪽 끝에서 본문까지
        sec.top_margin, sec.header_distance = Twips(tw(pg["top"] + pg["header"])), Twips(tw(pg["top"]))
        sec.bottom_margin, sec.footer_distance = Twips(tw(pg["bottom"] + pg["footer"])), Twips(tw(pg["bottom"]))
        grid = sec._sectPr.find(qn("w:docGrid"))
        if grid is not None:  # 문자 격자 해제 (줄 간격이 격자에 맞춰 늘어나지 않게 함)
            sec._sectPr.remove(grid)

    def _font(self, key: str) -> str:
        return self.spec["fonts"][key]

    def _setup_defaults(self):
        body = self.rows[self.p["body"]]
        f = self._font(body[2])
        styles = self.doc.styles.element
        old = styles.find(qn("w:docDefaults"))
        new = parse_xml(
            f'<w:docDefaults {nsdecls("w")}><w:rPrDefault><w:rPr>'
            f'<w:rFonts w:ascii="{f}" w:eastAsia="{f}" w:hAnsi="{f}" w:cs="{f}"/>'
            f'<w:sz w:val="{int(body[3] * 2)}"/><w:szCs w:val="{int(body[3] * 2)}"/>'
            f'<w:lang w:val="{lib.label("language.latin")}" w:eastAsia="{lib.label("language.east_asia")}" w:bidi="ar-SA"/>'
            f'</w:rPr></w:rPrDefault>'
            # 한/글은 한글과 영문 사이 간격을 자동으로 넣지 않음. Word의 자동 간격을 끔
            f'<w:pPrDefault><w:pPr><w:autoSpaceDE w:val="0"/><w:autoSpaceDN w:val="0"/>'
            f'<w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr></w:pPrDefault></w:docDefaults>')
        styles.replace(old, new) if old is not None else styles.insert(0, new)

    def _bullet_geometry(self, row) -> tuple[int, int]:
        """글머리 문단의 (본문 시작 위치, 내어쓰기) HWPUNIT. 한/글은 기호를 왼쪽 여백에 두고 기호 폭과 간격 뒤에서 본문을 시작"""
        glyph = self.spec["glyphs"][row[11]]
        gw = self.spec["glyph_em"].get(glyph, 1.0) * row[3] * HU_PER_PT
        hang = gw + self.spec["bullet_offset"]
        return row[6] + hang, hang

    def _setup_numbering(self) -> dict[str, int]:
        """계위 기호 정의(글머리)를 스타일마다 1개씩 만든다. 기호는 본문 글자가 아닌 Word 글머리로 표시"""
        numbering = self.doc.part.numbering_part.element
        ids: dict[str, int] = {}
        first_num = numbering.find(qn("w:num"))
        # 기본 템플릿의 목록 스타일(List Bullet 등)이 쓰는 번호와 겹치지 않게 기존 최대 번호 뒤부터 부여
        base_a = max([int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))] or [0])
        base_n = max([int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))] or [0])
        for i, row in enumerate(r for r in self.spec["styles"] if r[11]):
            aid, nid = base_a + i + 1, base_n + i + 1
            glyph = self.spec["glyphs"][row[11]].replace("­", "-")
            font = self._font(row[2])
            text_at, hang = self._bullet_geometry(row)
            absn = parse_xml(
                f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{aid}"><w:multiLevelType w:val="singleLevel"/>'
                f'<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="bullet"/><w:lvlText w:val="{escape(glyph)}"/>'
                f'<w:lvlJc w:val="left"/><w:pPr><w:ind w:left="{tw(text_at)}" w:hanging="{tw(hang)}"/></w:pPr>'
                f'<w:rPr><w:rFonts w:ascii="{font}" w:eastAsia="{font}" w:hAnsi="{font}" w:hint="eastAsia"/></w:rPr>'
                f'</w:lvl></w:abstractNum>')
            num = parse_xml(f'<w:num {nsdecls("w")} w:numId="{nid}"><w:abstractNumId w:val="{aid}"/></w:num>')
            if first_num is not None:
                first_num.addprevious(absn)
            else:
                numbering.append(absn)
            numbering.append(num)
            ids[row[0]] = nid
        return ids

    def _style_line(self, row, pt: float | None = None) -> int:
        """한/글 퍼센트 줄 간격을 고정 줄 간격으로 옮김 (글자 크기 × 퍼센트). 글꼴 높이에 따른 렌더러 차이를 없앰"""
        return int(round((pt or row[3]) * row[10] / 100 * TWIP_PER_PT))

    def _setup_styles(self):
        self.style_names: dict[str, str] = {}
        normal = self.doc.styles["Normal"]
        for row in self.spec["styles"]:
            name = row[0]
            if name == self.p["body"]:
                st = normal  # 바탕글은 Word 기본 스타일(Normal)로 둠. 이름을 바꾸면 Word가 기본 스타일로 인식하지 못함
            else:
                st = self.doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
                st.base_style = normal
                st.quick_style = True
            self.style_names[name] = st.name
            self._format_style(st, row)
            if st is normal:  # 한/글처럼 한글과 영문, 숫자 사이에 자동 간격을 넣지 않음 (LibreOffice는 기본값보다 스타일 값을 읽음)
                ppr = st.element.get_or_add_pPr()
                for tag in ("autoSpaceDE", "autoSpaceDN"):
                    _insert_ordered(ppr, _w(tag, val=0), PPR_ORDER)
        for name in ("Heading 1", "Heading 2", "Heading 3", "Title"):  # 기본 템플릿 제목 스타일의 테마 색과 글꼴이 섞이지 않게 숨김
            try:
                self.doc.styles[name].hidden = True
            except KeyError:
                pass

    def _format_style(self, st, row):
        name, _role, fkey, pt, bold, align, left, indent, prev, nxt, ls, bullet, keep, widow = row
        font = self._font(fkey)
        st.font.name = font
        rfonts = st.element.get_or_add_rPr().get_or_add_rFonts()
        rfonts.set(qn("w:eastAsia"), font)
        rfonts.set(qn("w:cs"), font)
        st.font.size = Pt(pt)
        st.font.bold = bool(bold)
        st.font.color.rgb = RGBColor.from_string(self.spec["color"].lstrip("#"))
        pf = st.paragraph_format
        pf.alignment = JC[align]
        if bullet:
            text_at, hang = self._bullet_geometry(row)
            pf.left_indent, pf.first_line_indent = Twips(tw(text_at)), Twips(-tw(hang))
            numpr = st.element.get_or_add_pPr().get_or_add_numPr()
            numpr.get_or_add_ilvl().val = 0
            numpr.get_or_add_numId().val = self._num_ids[name]
        else:
            self._indent(pf, left, indent)
        pf.space_before, pf.space_after = Twips(tw(prev)), Twips(tw(nxt))
        pf.line_spacing = Twips(self._style_line(row))
        pf.keep_with_next = bool(keep)
        pf.widow_control = bool(widow)
        if name in (self.p.get("toc_chapter"), self.p.get("toc_section")):
            pf.tab_stops.add_tab_stop(Twips(tw(self.spec["toc_tab"])), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)

    @staticmethod
    def _indent(pf, left: int, indent: int):
        """한/글 여백을 Word 들여쓰기로 옮김. 한/글 내어쓰기(음수)는 첫 줄을 여백에 두고 둘째 줄부터 안으로 들임"""
        if indent < 0:
            pf.left_indent, pf.first_line_indent = Twips(tw(left - indent)), Twips(tw(indent))
        else:
            pf.left_indent, pf.first_line_indent = Twips(tw(left)), Twips(tw(indent))

    # ── 기본 요소 ───────────────────────────────────────────────────────
    def _row(self, role: str):
        name = self.p.get(role)
        if name is None and role not in self._warned:
            self._warned.add(role)
            print(f"WARN: 프로파일에 '{role}' 스타일 없음 → 바탕글로 출력")
        return self.rows[name or self.p["body"]]

    def _new_para(self, role: str, align: str | None = None):
        row = self._row(role)
        p = self.doc.add_paragraph(style=self.style_names[row[0]])
        if align:
            p.paragraph_format.alignment = JC[align]
        if self.pending_break:
            p.paragraph_format.page_break_before = True
            self.pending_break = False
        return p

    def _run(self, p, text: str, role: str, pt: float | None = None, bold: bool | None = None,
             color: str | None = None, font_key: str | None = None):
        """크기를 바꾼 글자(제목 위계, 밴드, 표지). 문단 줄 간격도 바뀐 크기 기준으로 다시 계산"""
        if pt is not None:
            p.paragraph_format.line_spacing = Twips(self._style_line(self._row(role), pt))
        return self._text(p, text, pt or self._row(role)[3], bold=bold, color=color, font_key=font_key, size=pt)

    def _space_twips(self, pt: float) -> int:
        """한/글 고정 빈칸 폭(글자 크기의 절반)과 글꼴 빈칸 폭의 차이. docx 빈칸에 자간으로 더함"""
        return int(round((self.spec["space_em"] - self.spec["glyph_em"][" "]) * pt * TWIP_PER_PT))

    def _text(self, p, text: str, pt: float, **fmt):
        """글자를 빈칸 단위로 나누어 기록. 빈칸 조각에는 한/글 빈칸 폭을 맞추는 자간을 붙임
        fmt: bold, italic, script(sup|sub), color, font_key"""
        runs = []
        for part in SPACE_SPLIT_RE.split(text):
            if not part:
                continue
            r = p.add_run(part)
            if fmt.get("bold") is not None:
                r.font.bold = fmt["bold"]
            if fmt.get("italic"):
                r.font.italic = True
            if fmt.get("script") == "sup":
                r.font.superscript = True
            elif fmt.get("script") == "sub":
                r.font.subscript = True
            if fmt.get("color"):
                r.font.color.rgb = RGBColor.from_string(fmt["color"].lstrip("#"))
            if fmt.get("size"):
                r.font.size = Pt(fmt["size"])
            if fmt.get("font_key"):
                f = self._font(fmt["font_key"])
                r.font.name = f
                r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), f)
            if part[0] == " " and self._space_twips(pt):
                _insert_ordered(r._element.get_or_add_rPr(), _w("spacing", val=self._space_twips(pt)), RPR_ORDER)
            runs.append(r)
        return runs

    def para(self, role: str, text: str, footnotes: dict | None = None, no_heading: bool = False, **_):
        p = self._new_para(role)
        if no_heading:  # 번호를 본문에 기입하므로 스타일의 글머리를 끈다
            numpr = p._p.get_or_add_pPr().get_or_add_numPr()
            numpr.get_or_add_numId().val = 0
        else:
            text = self.manual_symbol.get(role, "") + text
        pt = self._row(role)[3]
        for kind, seg in lib.inline_segments(text):
            if kind == "footnote":
                if footnotes is not None:
                    self.add_footnote(p, footnotes.get(seg, ""))
                continue
            self._text(p, seg, pt, bold=True if kind == "bold" else None, italic=kind == "italic",
                       script=kind if kind in ("sup", "sub") else None)
        return p

    def add_footnote(self, p, text: str):
        """각주 참조. 번호는 Word 자동 번호, 뒤의 닫는 괄호는 글자로 붙여 hwpx 각주 번호 형식(1))과 맞춤"""
        fid = len(self._footnotes) + 1
        r = p.add_run()
        r.font.superscript = True
        r._r.append(_w("footnoteReference", id=fid))
        p.add_run(FOOTNOTE_SUFFIX).font.superscript = True
        self._footnotes.append((fid, text))

    def gap(self):
        for _ in range(lib.BLOCK_GAP_LINES):
            self._new_para("body")

    def _spacer(self, role: str, which: str):
        """표와 그림을 담는 문단의 앞뒤 여백을 빈 문단 높이로 옮김. 한/글은 표를 문단 안 글자처럼 두지만 Word 표는 문단 밖에 놓임"""
        row = self._row(role)
        height = row[8] if which == "before" else row[9]
        if which == "before" and self.pending_break:  # 쪽 나눔 직후 쪽 첫머리의 위 여백은 한/글처럼 두지 않음
            height = 0
        p = self._new_para(role)
        pf = p.paragraph_format
        pf.space_before = pf.space_after = Twips(0)
        pf.line_spacing = Twips(max(tw(height), SPACER_MIN_TWIP))
        return p

    # ── 표 ──────────────────────────────────────────────────────────────
    def _cell_line_role(self, line: str) -> str | None:
        return lib.cell_line_role(line, self.p)

    def special_cell(self, val: str) -> str | None:
        return None

    def caption_label(self, kind: str, text: str) -> dict:
        cap = lib.label("caption.table" if kind == "TABLE" else "caption.figure")
        label = f"{cap['open']}{text}{cap['close']}"   # 표는 <제목>, 그림은 [제목] (caption 라벨의 괄호)
        return {"open": label, "num": None, "close": "", "text": label}

    def _caption(self, kind: str, text: str, side: str):
        role = "table_title" if kind == "TABLE" else "fig_title"
        label = self.caption_label(kind, text)
        p = self._new_para(role)
        gap = Twips(tw(lib.CAPTION_GAP))
        if side == "TOP":
            p.paragraph_format.space_after = gap
            p.paragraph_format.keep_with_next = True
        else:
            p.paragraph_format.space_before = gap
        pt = self._row(role)[3]
        self._text(p, label["open"], pt)
        if label["num"] is not None:  # 장별 번호: 장의 첫 캡션에서 번호를 1로 되돌림
            reset = " \\r 1" if label["num"] == 1 else ""
            instr = f" SEQ {seq_name(kind)} \\* ARABIC{reset} "
            fld = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="{escape(instr)}"><w:r><w:t>{label["num"]}</w:t></w:r></w:fldSimple>')
            p._p.append(fld)
            self._text(p, label["close"], pt)
        return p

    def _new_table(self, nrows: int, widths: list[int]):
        t = self.doc.add_table(rows=nrows, cols=len(widths))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        tblPr = t._tbl.tblPr
        _insert_ordered(tblPr, _w("tblW", w=tw(sum(widths)), type="dxa"), TBLPR_ORDER)
        mar = _w("tblCellMar")
        for side, v in zip(("top", "left", "bottom", "right"), (TABLE_IN_MARGIN[2], TABLE_IN_MARGIN[0], TABLE_IN_MARGIN[3], TABLE_IN_MARGIN[1])):
            mar.append(_w(side, w=tw(v), type="dxa"))
        _insert_ordered(tblPr, mar, TBLPR_ORDER)
        for ci, w in enumerate(widths):
            t.columns[ci].width = Twips(tw(w))
        for row in t.rows:
            for ci, cell in enumerate(row.cells):
                cell.width = Twips(tw(widths[ci]))
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        return t

    @staticmethod
    def _tc_style(tc, sides: list[str], fill: str | None, margin: tuple | None):
        tcPr = tc.get_or_add_tcPr()
        borders = _w("tcBorders")
        for side in ("top", "left", "bottom", "right"):
            if side in sides:
                borders.append(_w(side, val="single", sz=max(2, round(float(lib.BORDER_WIDTH.split()[0]) * EIGHTHS_PER_MM)),
                                  space=0, color=lib.BORDER_COLOR.lstrip("#")))
            else:
                borders.append(_w(side, val="nil"))
        _insert_ordered(tcPr, borders, TCPR_ORDER)
        if fill:
            _insert_ordered(tcPr, _w("shd", val="clear", color="auto", fill=fill.lstrip("#")), TCPR_ORDER)
        if margin:
            mar = _w("tcMar")
            for side, v in zip(("top", "left", "bottom", "right"), (margin[2], margin[0], margin[3], margin[1])):
                mar.append(_w(side, w=tw(v), type="dxa"))
            _insert_ordered(tcPr, mar, TCPR_ORDER)

    @staticmethod
    def _row_height(row, height: int, cant_split: bool = False, header: bool = False):
        row.height, row.height_rule = Twips(tw(height)), WD_ROW_HEIGHT_RULE.AT_LEAST
        trPr = row._tr.get_or_add_trPr()
        if cant_split:
            trPr.append(_w("cantSplit"))
        if header:
            trPr.append(_w("tblHeader"))

    def _cell_para(self, cell, text: str, role: str, align: str | None = None, bold: bool | None = None, first: bool = True):
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        p.style = self.doc.styles[self.style_names[self._row(role)[0]]]
        if align:
            p.paragraph_format.alignment = JC[align]
        if text:
            self._text(p, text, self._row(role)[3], bold=bold)
        return p

    def _transparent_row(self, t, ri: int, ncol: int, text: str, role: str, align: str):
        """표 안의 투명 병합 행 (단위, 자료). 표와 함께 이동"""
        cell = t.cell(ri, 0).merge(t.cell(ri, ncol - 1)) if ncol > 1 else t.cell(ri, 0)
        if ncol > 1:
            cell.width = Emu(sum(c.width for c in t.columns))
        self._cell_para(cell, text, role, align)
        self._tc_style(cell._tc, [], None, None)
        self._row_height(t.rows[ri], lib.ROW_BASE_H)

    def table(self, e: lib.Element):
        rows = [r + [""] * (max(len(x) for x in e.data["rows"]) - len(r)) for r in e.data["rows"]]
        nrow, ncol = len(rows), len(rows[0])
        widths = lib.calc_col_widths(rows, self.table_max_w)
        aligns = lib.col_aligns(rows, e.data["aligns"])
        hdr = lib.header_rows(rows)
        if nrow > lib.LONG_TABLE_ROWS and self.long_table_mode == "split":
            for sub in lib.split_table(e, rows):
                self.table(sub)
            return
        self.gap()
        unit, sources = e.data.get("unit"), e.data.get("sources", [])
        head = 1 if unit else 0
        long = nrow > lib.LONG_TABLE_ROWS and not e.data.get("split_done")
        self._spacer("table_pos", "before")
        if e.text:
            self._caption("TABLE", e.text, "TOP")
        t = self._new_table(nrow + head + len(sources), widths)
        vspans, hspans = lib.table_spans(rows, hdr, self.merge_mode)
        # 병합 영역: 앵커(행, 열) → (끝 행, 끝 열). 행은 머리글 포함 원고 행 기준
        anchors: dict[tuple[int, int], tuple[int, int]] = {}
        covered: set[tuple[int, int]] = set()
        for c, r0, r1 in vspans:
            a, b = hdr + r0, hdr + r1
            anchors[(a, c)] = (b, c)
            covered |= {(r, c) for r in range(a + 1, b + 1)}
        for ri, c0, c1 in hspans:
            anchors[(ri, c0)] = (ri, c1)
            covered |= {(ri, c) for c in range(c0 + 1, c1 + 1)}
        for (r0, c0), (r1, c1) in anchors.items():
            t.cell(head + r0, c0).merge(t.cell(head + r1, c1)).width = Twips(tw(sum(widths[c0:c1 + 1])))
        if unit:
            self._transparent_row(t, 0, ncol, unit, "unit", "RIGHT")
        for k, src in enumerate(sources):
            self._transparent_row(t, head + nrow + k, ncol, src, "source", "LEFT")
        heights = lib.row_heights(rows, widths, hdr, vspans, hspans)
        for ri, row in enumerate(rows):
            is_head = ri < hdr
            self._row_height(t.rows[head + ri], heights[ri], cant_split=not long,
                             header=long and ri < hdr)
            for ci, raw in enumerate(row):
                if (ri, ci) in covered:
                    continue
                cell = t.cell(head + ri, ci)
                fill = self.special_cell(raw)
                val, roles = lib.cell_lines("" if fill else raw, self._cell_line_role)
                role, align = lib.cell_role(aligns[ci], is_head)
                for li, line in enumerate(val.split("\n")):
                    br = roles[li] if li < len(roles) and not is_head else None
                    self._cell_para(cell, line, br or role, None if br else align, True if is_head else None, first=li == 0)
                r1, c1 = anchors.get((ri, ci), (ri, ci))
                sides = lib.edge_sides(ci, c1, ncol)  # 음영 셀도 표 좌우 바깥선 규칙을 따름 (hwpx와 같은 괘선)
                style_fill = fill or (lib.HEADER_FILL if is_head else None)
                for rr in range(ri, r1 + 1):  # 세로 병합의 이어진 칸에도 같은 괘선과 채움
                    tc = t.rows[head + rr]._tr.tc_lst[self._tc_index(t.rows[head + rr]._tr, ci)]
                    self._tc_style(tc, sides, style_fill, lib.CELL_MARGIN)
        if long and head:  # Word는 표 첫 행부터 이어진 행만 반복하므로 단위 행도 반복 대상에 넣음
            t.rows[0]._tr.get_or_add_trPr().append(_w("tblHeader"))
        if not long:  # 한/글 글자취급 표처럼 쪽을 넘기지 않고 통째로 이동
            for row in t.rows[:-1]:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        p.paragraph_format.keep_with_next = True
        self._spacer("table_pos", "after")
        self.after_block = True

    @staticmethod
    def _tc_index(tr, grid_col: int) -> int:
        """행 안에서 격자 열 위치에 해당하는 셀 순번 (가로 병합 칸 수 반영)"""
        col = 0
        for i, tc in enumerate(tr.tc_lst):
            if col >= grid_col:
                return i
            col += tc.grid_span
        return len(tr.tc_lst) - 1

    # ── 그림 ────────────────────────────────────────────────────────────
    def image(self, e: lib.Element):
        prepared = lib.prepare_image(e)
        if prepared is None:
            return
        data, _fmt, w_mm, h_mm = prepared
        self.gap()
        w_hu, h_hu = round(w_mm * 7200 / 25.4), round(h_mm * 7200 / 25.4)
        sources = e.data.get("sources", [])
        self._spacer("fig_pos", "before")
        if e.text and lib.FIG_CAPTION_SIDE == "TOP":
            self._caption("PICTURE", e.text, "TOP")
        if not lib.FIGURE_IN_TABLE:
            p = self._new_para("fig_pos", "CENTER")
            p.paragraph_format.line_spacing = 1.0  # 고정 줄 간격은 그림을 잘라 표시함
            p.add_run().add_picture(io.BytesIO(data), width=Emu(w_hu * EMU_PER_HU), height=Emu(h_hu * EMU_PER_HU))
            if e.text and lib.FIG_CAPTION_SIDE != "TOP":
                self._caption("PICTURE", e.text, "BOTTOM")
            for src in sources:
                self.para("source", src)
        else:  # 투명 1열 표에 그림과 자료 행을 담음 (hwpx와 같은 구성)
            t = self._new_table(1 + len(sources), [w_hu])
            cell = t.cell(0, 0)
            p = self._cell_para(cell, "", "fig_pos", "CENTER")
            p.paragraph_format.line_spacing = 1.0
            p.add_run().add_picture(io.BytesIO(data), width=Emu(w_hu * EMU_PER_HU), height=Emu(h_hu * EMU_PER_HU))
            self._tc_style(cell._tc, [], None, None)
            self._row_height(t.rows[0], h_hu + 2 * lib.ROW_LINE_H, cant_split=True)
            for k, src in enumerate(sources):
                self._transparent_row(t, 1 + k, 1, src, "source", "LEFT")
            cap_below = bool(e.text) and lib.FIG_CAPTION_SIDE != "TOP"
            # Why: 아래 캡션은 표 밖의 문단이므로 마지막 행까지 다음 문단과 묶어야 쪽 경계에서 캡션이 떨어지지 않음
            for row in (t.rows if cap_below else t.rows[:-1]):
                for p in row.cells[0].paragraphs:
                    p.paragraph_format.keep_with_next = True
            if cap_below:
                self._caption("PICTURE", e.text, "BOTTOM")
        self._spacer("fig_pos", "after")
        self.after_block = True

    # ── 밴드 ────────────────────────────────────────────────────────────
    def band_emit(self, spec: dict):
        """1행 3칸 표. 앞뒤 여백은 밴드 위치 스타일 값의 빈 문단, 다음 문단과 함께 두기는 스타일이 담당"""
        self.gap()
        first = self._spacer("band_pos", "before")
        t = self._new_table(1, spec["widths"])
        row = t.rows[0]
        self._row_height(row, spec["height"], cant_split=True)
        for ci, (text, pt, color, align, fill, border) in enumerate(spec["cells"]):
            cell = row.cells[ci]
            p = self._cell_para(cell, "", "band_cell", align)
            p.paragraph_format.keep_with_next = True
            self._run(p, text, "band_cell", pt=pt / HU_PER_PT, bold=True, color=color, font_key=self._row("chapter")[2])
            sides = ["top", "bottom", "left", "right"] if border else []
            self._tc_style(cell._tc, sides, fill, None)
            if border:  # 밴드 테두리는 표 괘선과 다른 색과 굵기
                for b in cell._tc.tcPr.find(qn("w:tcBorders")):
                    b.set(qn("w:color"), border.lstrip("#"))
                    b.set(qn("w:sz"), str(max(2, round(float(self.spec["band_border"].split()[0]) * EIGHTHS_PER_MM))))
        self._spacer("band_pos", "after")
        return first

    def heading_para(self, role: str, label: str, pt: int | None):
        p = self._new_para(role)
        numpr = p._p.get_or_add_pPr().get_or_add_numPr()
        numpr.get_or_add_numId().val = 0
        self._run(p, label, role, pt=pt / HU_PER_PT if pt else None, bold=True if pt else None)
        return p

    def chapter_started(self, p):
        """장의 캡션 번호 재시작은 캡션 필드(SEQ \\r 1)가 담당"""

    def after_table(self, e):
        """셀 채움과 번호 집계는 표 기록 중에 처리"""

    def after_image(self, e):
        """번호 집계는 캡션 기록 중에 처리"""

    # ── 앞부분(표지, 목차), 머리말, 쪽번호 ─────────────────────────────────
    def toc_lists(self) -> list[tuple[str, list[str]]]:
        return []

    def _front_matter(self, cover: dict | None):
        """표지와 목차를 앞 구역에 두고 본문 구역의 쪽번호를 1부터 시작. 문단 순서는 lib.front_matter_items"""
        if self.toc and not (self.p.get("toc_chapter") and self.p.get("toc_section")):
            print("WARN: 프로파일에 목차 스타일(toc_chapter, toc_section)이 없어 --toc 생략")
            self.toc = False
        body = self.doc.element.body
        anchor = body[0]
        made = [self._new_para("body")._p]   # 앞 구역 첫머리의 빈 바탕글 줄 (hwpx 구역 첫 문단과 같음)
        for it in lib.front_matter_items(cover, self.toc, self.headings, self.toc_depth, self.toc_lists()):
            kind = it["kind"]
            if kind in ("blank", "cover"):
                p = self._new_para("body", "CENTER")
                if kind == "cover":
                    self._run(p, it["text"], "body", pt=it["pt"] / HU_PER_PT, bold=it["bold"],
                              font_key=self._row("chapter")[2])
            elif kind == "toc_title":
                p = self._new_para("chapter", "CENTER")
                p.paragraph_format.page_break_before = it["page_break"]
                self._text(p, it["text"], self._row("chapter")[3])
            else:
                lv = it.get("level", 2)
                p = self._new_para("toc_chapter" if lv == 1 else "toc_section")
                if kind == "list_entry":
                    p.paragraph_format.left_indent = Twips(tw(lib.TOC_LIST_INDENT))
                elif lv >= 3:
                    p.paragraph_format.left_indent = Twips(tw(lib.TOC_SUB_INDENT))
                num = self.toc_numbers.get(it["text"])
                self._text(p, it["text"] + (f"\t{num}" if num else ""), self._row("toc_chapter" if lv == 1 else "toc_section")[3])
            made.append(p._p)
        for el in made:
            anchor.addprevious(el)
        body_sect = self.doc.sections[0]._sectPr   # 본문 첫 구역 (간지를 쓰면 본문이 여러 구역)
        front_sect = copy.deepcopy(self._sect_template())
        made[-1].get_or_add_pPr().append(front_sect)
        _insert_ordered(body_sect, _w("type", val="nextPage"), SECTPR_ORDER)
        _insert_ordered(body_sect, _w("pgNumType", start=1), SECTPR_ORDER)
        self._front_done = True

    def _sect_template(self):
        """머리말과 꼬리말 참조, 쪽번호 시작, 첫 쪽 구분을 뺀 구역 설정 사본 (쪽 크기, 여백, 단)"""
        sect = copy.deepcopy(self.doc.element.body.sectPr)
        for tag in ("headerReference", "footerReference", "pgNumType", "titlePg", "type"):
            for el in sect.findall(qn(f"w:{tag}")):
                sect.remove(el)
        return sect

    def _first_body_section(self):
        """머리말, 꼬리말을 정의할 본문 첫 구역. 이후 구역은 Word 규칙대로 앞 구역의 정의를 이어 씀"""
        return self.doc.sections[1 if getattr(self, "_front_done", False) else 0]

    def divider_emit(self, spec: dict, new_page: bool):
        """간지 쪽: 빈 바탕글 줄 뒤에 간지 밴드. 간지는 새 구역의 첫 쪽이며 첫 쪽 머리말과 꼬리말을 비워 쪽번호를 감춤
        Why 구역: Word는 쪽 단위 감추기 기능이 없어 구역의 '첫 쪽 다르게'로 한/글의 감추기와 같은 결과를 만듦"""
        body = self.doc.element.body
        if new_page:  # 앞 구역을 직전 문단에서 닫음. 새 구역은 다음 쪽에서 시작
            last = body[-2] if body[-1].tag == qn("w:sectPr") else body[-1]
            sect = self._sect_template()
            if getattr(self, "_section_divider", False):
                _insert_ordered(sect, _w("titlePg"), SECTPR_ORDER)
            _insert_ordered(sect, _w("type", val="nextPage"), SECTPR_ORDER)
            last.get_or_add_pPr().append(sect)
        self._section_divider = True
        final = body.sectPr
        _insert_ordered(final, _w("titlePg"), SECTPR_ORDER)   # 마지막 구역도 간지로 시작하면 첫 쪽 구분 유지
        self.pending_break = False
        for _ in range(self.spec["divider_top_lines"]):
            self._new_para("body")
        return self.band_emit(spec)

    def set_header(self, spec: str | None):
        if not spec:
            return
        parts = spec.split(";")
        text = parts[0].strip()
        align = next((kv.split("=", 1)[1].strip().upper() for kv in parts[1:] if kv.strip().lower().startswith("align=")), "RIGHT")
        header = self._first_body_section().header
        header.is_linked_to_previous = False
        p = header.paragraphs[0]
        p.style = self.doc.styles[self.style_names[self.base_names["header_text"]]]
        p.paragraph_format.alignment = JC[align]
        self._text(p, text, self.rows[self.base_names["header_text"]][3])

    def set_page_number(self, on: bool):
        """꼬리말 가운데 '- N -'. 앞 구역(표지, 목차)은 쪽번호 없음"""
        if not on:
            return
        footer = self._first_body_section().footer
        footer.is_linked_to_previous = False
        p = footer.paragraphs[0]
        p.style = self.doc.styles[self.style_names[self.base_names["page_number"]]]
        p.paragraph_format.alignment = JC["CENTER"]
        pt = self.rows[self.base_names["page_number"]][3]
        self._text(p, "- ", pt)
        for kind, text in (("begin", None), (None, " PAGE "), ("separate", None), (None, "1"), ("end", None)):
            r = p.add_run()
            if kind:
                r._r.append(_w("fldChar", fldCharType=kind))
            elif text == " PAGE ":
                it = OxmlElement("w:instrText")
                it.set(qn("xml:space"), "preserve")
                it.text = text
                r._r.append(it)
            else:
                r.text = text
        self._text(p, " -", pt)

    # ── 저장 ────────────────────────────────────────────────────────────
    def _write_footnotes(self):
        """각주 부품(footnotes.xml)과 설정의 구분선 참조. python-docx에 각주 API가 없어 부품을 직접 등록"""
        if not self._footnotes:
            return
        style_id = self.doc.styles[self.style_names[self.base_names["footnote"]]].style_id
        sep = ('<w:footnote w:type="separator" w:id="-1"><w:p><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr>'
               '<w:r><w:separator/></w:r></w:p></w:footnote>'
               '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr>'
               '<w:r><w:continuationSeparator/></w:r></w:p></w:footnote>')
        pt = self.rows[self.base_names["footnote"]][3]
        gap = self._space_twips(pt)

        def runs(text: str) -> str:  # 빈칸 조각에 한/글 빈칸 폭을 맞추는 자간 (본문 _text와 같은 규칙)
            out = []
            for part in SPACE_SPLIT_RE.split(text):
                if part:
                    rpr = f'<w:rPr><w:spacing w:val="{gap}"/></w:rPr>' if part[0] == " " and gap else ""
                    out.append(f'<w:r>{rpr}<w:t xml:space="preserve">{escape(part)}</w:t></w:r>')
            return "".join(out)
        notes = "".join(
            f'<w:footnote w:id="{fid}"><w:p><w:pPr><w:pStyle w:val="{style_id}"/></w:pPr>'
            f'<w:r><w:footnoteRef/></w:r>{runs(FOOTNOTE_SUFFIX + " " + text)}</w:p></w:footnote>'
            for fid, text in self._footnotes)
        blob = f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<w:footnotes {nsdecls("w")}>{sep}{notes}</w:footnotes>'
        part = Part(PackURI("/word/footnotes.xml"), CT.WML_FOOTNOTES, blob.encode("utf-8"), self.doc.part.package)
        self.doc.part.relate_to(part, RT.FOOTNOTES)
        settings = self.doc.settings.element
        fpr = parse_xml(f'<w:footnotePr {nsdecls("w")}><w:footnote w:id="-1"/><w:footnote w:id="0"/></w:footnotePr>')
        compat = settings.find(qn("w:compat"))
        compat.addprevious(fpr) if compat is not None else settings.append(fpr)

    def _blank_divider_pages(self):
        """간지 구역(첫 쪽 구분)의 첫 쪽 머리말과 꼬리말을 빈 정의로 명시
        Why: 첫 쪽 정의가 없으면 LibreOffice는 앞 구역의 기본 꼬리말(쪽번호)을 간지 쪽에 표시함"""
        for section in self.doc.sections:
            if section.different_first_page_header_footer:
                section.first_page_header.is_linked_to_previous = False
                section.first_page_footer.is_linked_to_previous = False

    def save(self, out: str):
        self._blank_divider_pages()
        self._write_footnotes()
        self.doc.save(os.path.abspath(out))
