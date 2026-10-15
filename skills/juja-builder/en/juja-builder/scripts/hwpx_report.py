#!/usr/bin/env python3
"""hwpx_report.py: a single CLI that typesets Markdown manuscripts on a skeleton extracted from a reference template.

Commands
  build     manuscript.md → document.hwpx      (python-hwpx 6.x, references skeleton styles)
  skeleton  reference.hwpx → skeleton.hwpx     (an empty template keeping only header.xml and the body section secPr)
  export    document.hwpx → manuscript.md      (export_rich_markdown)
  fill      document.hwpx + values.json → document   (placeholder substitution, filling table label paths)
  render    document.hwpx → PDF + PNG          (rhwp export-pdf + pdftoppm)
  check     document.hwpx schema validation + rhwp verify expected values

Verified environment: python-hwpx 6.3.0, lxml 6.0.2, pillow 12.1.1, python-docx 1.2.0, openpyxl 3.1.5,
Python 3.12.3, poppler 24.02, rhwp v0.8.6. Always install the latest versions first. The record above is for tracing regression causes.

Format notice: 본 제품은 한컴의 HWP 문서 파일(.hwp) 공개 문서를 참고하여 개발하였습니다.
Trademark notice: 한/글 (Hangul), Hancom Office, and HWPX are trademarks of Hancom Inc. (한글과컴퓨터)
Non-affiliation notice: this tool is an independent tool with no affiliation with Hancom Inc.

Why: one typesetter per input kind would spread norm changes to N places, so the typesetter is this single file and
     other inputs (DOCX, XLSX, HWPX) first converge into Markdown manuscripts.
     Differences by output format live only in the record step. The parser, inline notation, table layout calculation, figure size, and TOC order are format-agnostic functions,
     and docx_report.py (the docx writer) calls the same functions.
"""
from __future__ import annotations

import argparse
import json
import os
import pathlib
import re
import shutil
import subprocess
import sys
import tempfile
import warnings
import zipfile
from dataclasses import dataclass, field

warnings.filterwarnings("ignore", category=DeprecationWarning)

# ═══════════════════════════════════════════════════════════════════════
# Settings: typesetting values and the role-to-style mapping are adjusted only here
# ═══════════════════════════════════════════════════════════════════════
# Why: the terms of use of Hancom's published HWP and OWPML documents require this wording in the product's user interface, manual, help, and source.
# The wording is kept verbatim from the Hancom download center (https://www.hancom.com/support/downloadCenter/hwpOwpml). Printed at the bottom of CLI --help
HANCOM_FORMAT_NOTICE = "본 제품은 한컴의 HWP 문서 파일(.hwp) 공개 문서를 참고하여 개발하였습니다."
# Why: states trademark ownership and non-affiliation so the tool is not mistaken for an official Hancom tool
HANCOM_TRADEMARK_NOTICE = "한/글, 한컴오피스, HWPX는 한글과컴퓨터의 상표"
HANCOM_AFFILIATION_NOTICE = "이 도구는 한글과컴퓨터와 제휴 관계가 없는 독립 도구"
# For output at the bottom of CLI --help. One line per notice (line breaks kept with RawDescriptionHelpFormatter)
CLI_NOTICE = "\n".join((HANCOM_FORMAT_NOTICE, HANCOM_TRADEMARK_NOTICE, HANCOM_AFFILIATION_NOTICE))
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
ASSET_DIR = os.path.join(os.path.dirname(SCRIPT_DIR), "assets")

# role → skeleton style name. numbering: manual=numbers written in the body, auto=outline numbers attached
PROFILES = {
    "chapter": {  # chapter-structured documents (제N장 / N.N / N.N.N / 가. / □ / ❍ / - / ․ / ※)
        "numbering": "manual",
        "chapter_format": "제{n}장 {t}",
        "section_format": "{c}.{s} {t}",          # 1.1 title
        "subsection_format": "{c}.{s}.{ss} {t}",  # 1.1.1 title
        "caption_prefix": "{c}-",  # <표 1-3>, [그림 1-3]. Numbers restart per chapter
        "chapter": "장 제목", "section": "절 제목", "subsection": "항 제목", "ganada": "목 제목",
        "square": "논점", "circle": "주장", "dash": "근거", "dot": "상세", "note": "주석",
        "table_pos": "표 위치", "table_cell": "표 본문", "table_cell_left": "표 본문 왼쪽",
        "table_title": "표 캡션", "unit": "단위", "source": "출처",
        "fig_pos": "그림 위치", "fig_title": "그림 캡션", "body": "바탕글",
        "toc_chapter": "목차 장", "toc_section": "목차 절",  # optional: --toc static TOC
        "table_cell_bullet": "표 글머리", "table_cell_dash": "표 하이픈",  # optional: cell bullet lines
    },
    "flat": {  # documents without chapters (N.N / N.N.N / 가. / □ / ◦ / - / ․ / *)
        "numbering": "manual",           # standalone documents: bare 1. 2. numbers and 가나다 (written in the body instead of outline numbers)
        "chapter_format": "{t}",
        "section_format": "{s}. {t}",              # 1. title
        "subsection_format": "{s}.{ss} {t}",       # 1.1 title
        "caption_prefix": "",  # <표 3>, [그림 3]
        "chapter": "제목",  # documents without chapters: # is the document title (centered, 20pt bold)
        "section": "1.1", "subsection": "1.1.1", "ganada": "가. 제목",
        "square": "□ 내용", "circle": "◦ 내용", "dash": "- 내용", "dot": "․ 내용", "note": "* 내용",
        "table_pos": "표 위치", "table_cell": "표 내용", "table_cell_left": "표 내용",
        "table_title": "표 제목", "unit": "단위", "source": "출처",
        "fig_pos": "그림위치", "fig_title": "그림 제목", "body": "바탕글",
    },
}

# Table typesetting norms (HWPUNIT). Why: if a table's left edge sticks out past the text start line of the ○ level, it looks misaligned with the hierarchy
TABLE_INDENT_MARGIN = 1500      # level indent subtracted from the body width (table width cap = body width - this value)
TABLE_MIN_RATIO = 0.5           # table width floor (relative to the cap). Prevents too-narrow widths for 2-3 column label tables
COL_MIN_W = 3000                # minimum column width (about 10.6mm)
COL_UNIT_W = 500                # width of one weighted unit (one half-width character) at 10pt table text (one Korean character = 2 units)
COL_PAD_W = 1800                # left and right padding per column + glyph error margin
COL_ONE_LINE_UNITS = 20         # a column with a cell over this many units is a long-text column and left-aligned (10 Korean characters)
COL_SEARCH_STEPS = 40           # bisection steps for the column width baseline (converges below 1 HWPUNIT)
CELL_MARGIN = (510, 510, 141, 141)  # cell padding left, right, top, bottom (HWPUNIT)
CELL_UNIT_W = 500               # width of one weighted unit for row height estimation (10pt table text)
AMOUNT_CELL = re.compile(r"\d,\d{3}|\d\.\d")  # figures with thousands separators or decimal points. Only such columns are right-aligned
FIG_CAPTION_SIDE = "BOTTOM"    # figure caption position (TOP | BOTTOM)
MERGE_LABEL_COLS = 2            # maximum number of leading label columns subject to vertical merging (convention of merging consecutive equal values)
LONG_TABLE_ROWS = 25            # above this: treat-as-character off + engine splitting + header row repetition
HEADER_FILL = "#D9D9D9"         # table header background (conventional value in policy reports)
BORDER_WIDTH = "0.15 mm"        # cell rule width (conventional value in policy reports)
BORDER_COLOR = "#000000"
# Why: leaving the outer left and right table rules empty is the policy report convention. The leftmost cell has no left rule and the rightmost cell no right rule
EDGE_OPEN_SIDES = True
FIGURE_IN_TABLE = True          # put a figure in a transparent one-column table and set the table's numbering type to figure (hand-editing convention)
ROW_BASE_H = 2000               # one-line row height (한/글 expands to fit content, so this is a minimum)
ROW_LINE_H = 1400               # addition per folded or broken line
NUMERIC_CELL = re.compile(r"^[-+±(]?[\d,]+(?:\.\d+)?%?\)?$")
BULLET_PREFIXES = ("·", "•", "ㆍ", "∙", "○", "◦", "❍", "□", "■", "※", "- ", "▪")

# Figure size norms (mm, body width about 170mm). Why: a figure filling the full width looks misaligned with the body hierarchy,
# and portrait figures run past one page. Landscape is capped by width, square by width, and portrait by height
IMG_LANDSCAPE_W_MM = 140.0      # ratio 1.4 or more
IMG_SQUARE_W_MM = 90.0          # ratio over 0.8 and under 1.4
IMG_PORTRAIT_H_MM = 110.0       # ratio 0.8 or less
IMG_DEFAULT_DPI = 96.0
IMG_URL_TIMEOUT = 20
COVER_TOP_LINES = 8             # empty lines above the cover title
COVER_TITLE_PT = 2400           # cover title 24pt
COVER_SUB_PT = 1600             # subtitle 16pt
COVER_DATE_PT = 1400            # date 14pt
COVER_ORG_PT = 1600             # organization 16pt
COVER_GAP_LINES = 14            # empty lines between the date and the organization
TOC_TAB_POS = 48100             # right dotted tab position of TOC styles (skeleton tabPr value)
TOC_SUB_INDENT = 3000           # TOC subsection indent (twice the section's 1500, conventional policy report value)
TOC_UNIT_W = (700, 600)         # width of one weighted unit for estimating TOC tab width (chapter, section). 한/글 recalculates tabs, so this is for approximate renderer display
TOC_LIST_INDENT = 0             # left margin of list of tables and list of figures entries (overrides the TOC section style margin)
CAPTION_GAP = 850               # gap between caption and table (python-hwpx default, the value most common in the real corpus)
BLOCK_GAP_LINES = 0             # empty paragraphs before and after table and figure blocks. Spacing is handled by paragraph margins of the table position and figure position styles

# Labels: fixed strings printed into documents, manuscript notation the parser recognizes, and docx language tags. Defaults are the current Korean notation
# A label file (JSON, proposal.py `build --labels`) overrides some keys. Missing keys keep the default, unknown keys are an ERROR
# Why one table and one accessor (label): documents in other languages come from a label file with no per-language branch, and without overrides the document stays as before
LABEL_DEFAULTS = {
    # heading number formats. None keeps the active profile value (chapter_format, chapter_label, section_format, subsection_format, item_format)
    "formats": {"chapter": None, "chapter_label": None, "section": None, "subsection": None, "item": None},
    "letters": {"items": "가나다라마바사아자차카타파하", "format": "{x}."},   # letter sequence and notation of the 가나다 level. {g} in formats
    "toc": {"title": "목  차", "tables": "표 목차", "figures": "그림 목차"},
    # captions: word (label) and brackets (open, close) of numbered captions <표 1-1>, Word number field name (seq, None means label)
    "caption": {"table": {"label": "표", "open": "<", "close": ">", "seq": None},
                "figure": {"label": "그림", "open": "[", "close": "]", "seq": None}},
    # manuscript notation: source lines (출처:, (출처)), unit lines ((단위: )), table captions (<표>), figure captions ([그림]). Overrides only add tokens
    "notation": {"source": ["출처", "자료", "주"], "unit": ["단위"], "table": ["표"], "figure": ["그림"]},
    # docx default run language (w:lang). With Latin ko-KR, LibreOffice breaks English words and numbers mid-token (3,|456), so en-US
    "language": {"latin": "en-US", "east_asia": "ko-KR"},
    "output": {"proposal": "제안서", "report": "사전진단보고서", "plan": "기획보고서"},   # delivered file names per output type
}
FORMAT_PROFILE_KEYS = {"chapter": "chapter_format", "chapter_label": "chapter_label", "section": "section_format",
                       "subsection": "subsection_format", "item": "item_format"}
# placeholders per format. n chapter number, r Roman chapter number, c chapter, s section, ss subsection, g 가나다 letter (letters.format applied), i item number, t title
FORMAT_FIELDS = {"chapter": ("n", "r", "t"), "chapter_label": ("n", "r"), "section": ("c", "s", "t"),
                 "subsection": ("c", "s", "ss", "g", "t"), "item": ("i", "g", "t")}
SOURCE_PAREN_EXCLUDE = ("주",)   # tokens left out of the parenthesized source mark. (주) abbreviates 주식회사 (Inc.), so it is not a source line
LANG_TAG_RE = re.compile(r"[A-Za-z]{2,8}(-[A-Za-z0-9]{1,8})*")
_labels: dict = json.loads(json.dumps(LABEL_DEFAULTS))   # current labels (a copy of the defaults with load_labels overrides applied)

# rhwp renderer
RHWP_RELEASE_PAGE = "https://github.com/edwardkim/rhwp/releases/latest"
RENDER_DPI = 80

HP = "{http://www.hancom.co.kr/hwpml/2011/paragraph}"
MANUAL_SYMBOLS = {"square": "□ ", "circle": "○ ", "dash": "- ", "dot": "ㆍ ", "note": "※ "}


# ═══════════════════════════════════════════════════════════════════════
# 1. Markdown parser: normalizes manuscripts into a list of (kind, payload) elements
# ═══════════════════════════════════════════════════════════════════════
@dataclass
class Element:
    kind: str                # heading square circle dash dot note table image pagebreak
    text: str = ""
    level: int = 0           # heading: 1-4, bullet: indent rank
    data: dict = field(default_factory=dict)


HEADING_RE = re.compile(r"^(#{1,4})\s+(.*)$")
BULLET_RE = re.compile(r"^(\s*)[-*+]\s+(.*)$")
IMAGE_RE = re.compile(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$")
TABLE_SEP_RE = re.compile(r"^\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)*\|?$")
FOOTNOTE_DEF_RE = re.compile(r"^\[\^([^\]]+)\]:\s*(.*)$")
NUM_PREFIX_RE = re.compile(r"^(제\s*\d+\s*장|\d+(?:\.\d+)*\.?|[가-힣]\.)\s+")
SQUARE_MARK = ("□", "■", "☐")
CIRCLE_MARK = ("○", "❍", "◦", "ㅇ")
NOTE_MARK = ("※", "*")


def label(key: str):
    """Label value. key is a dotted path ("toc.title", "caption.table", "notation.source")"""
    v = _labels
    for k in key.split("."):
        v = v[k]
    return v


def letter_label(i: int) -> str:
    """Notation of the i-th letter of the 가나다 level (default "가.")"""
    items = label("letters.items")
    return label("letters.format").format(x=items[(i - 1) % len(items)])


def apply_format_labels(profile: dict) -> dict:
    """Apply formats.* overrides to a copy of the profile. Keys whose value is None keep the profile value"""
    out = dict(profile)
    for k, pk in FORMAT_PROFILE_KEYS.items():
        if _labels["formats"][k] is not None:
            out[pk] = _labels["formats"][k]
    return out


def _alt(tokens: list[str]) -> str:
    """Turn a token list into a regex alternation. Longer tokens come first to prevent prefix matches"""
    return "|".join(re.escape(t) for t in sorted(tokens, key=len, reverse=True))


def _compile_notation():
    """Rebuild the manuscript notation regexes from the current labels. With default labels the verdicts equal the previous regexes"""
    global CAPTION_RE, FIG_CAPTION_RE, UNIT_RE, SOURCE_RE, SOURCE_PAREN_RE
    n = _labels["notation"]
    paren = _alt([t for t in n["source"] if t not in SOURCE_PAREN_EXCLUDE])
    CAPTION_RE = re.compile(rf"^[<＜]\s*(?:{_alt(n['table'])})\s*[^>＞]*[>＞]\s*(.*)$")
    FIG_CAPTION_RE = re.compile(rf"^\[\s*(?:{_alt(n['figure'])})\s*[^\]]*\]\s*(.*)$")
    UNIT_RE = re.compile(rf"^\(\s*(?:{_alt(n['unit'])})\s*[:：].*\)$")
    SOURCE_RE = re.compile(rf"^(?:[(（]\s*(?:{paren})\s*[)）]|(?:{_alt(n['source'])})\s*[:：])")
    SOURCE_PAREN_RE = re.compile(rf"^[(（]\s*({paren})\s*[)）]\s*")   # turns the (출처) notation into "출처: "


_compile_notation()


def load_labels(path: str) -> list[str]:
    """Override the defaults with a label file (JSON). Missing keys keep the default. Unknown keys and format errors are collected and stop with ERROR
    notation.* only adds tokens (the default tokens stay recognized). Returns the list of overridden keys"""
    global _labels
    try:
        data = json.load(open(path, encoding="utf-8"))
    except (OSError, ValueError) as ex:
        sys.exit(f"ERROR: could not read label file {path} ({ex})")
    merged = json.loads(json.dumps(LABEL_DEFAULTS))
    errors: list[str] = []
    changed: list[str] = []

    def walk(src, dst: dict, base: dict, prefix: str):
        if not isinstance(src, dict):
            errors.append(f"{prefix or 'top level'}: JSON object required")
            return
        for k, v in src.items():
            key = f"{prefix}.{k}" if prefix else k
            if k not in base:
                errors.append(f"unknown key '{key}' (keys at this level: {', '.join(base)})")
            elif isinstance(base[k], dict):
                walk(v, dst[k], base[k], key)
            elif isinstance(base[k], list):  # manuscript notation: appended after the default tokens
                vals = [v] if isinstance(v, str) else v
                if not isinstance(vals, list) or not all(isinstance(x, str) and x.strip() for x in vals):
                    errors.append(f"{key}: string or list of strings required")
                    continue
                dst[k] = list(dict.fromkeys(base[k] + [x.strip() for x in vals]))
                changed.append(key)
            elif key == "letters.items" and isinstance(v, list) and v and all(isinstance(x, str) and x for x in v):
                dst[k] = v
                changed.append(key)
            elif not isinstance(v, str):
                errors.append(f"{key}: string required")
            else:
                dst[k] = v
                changed.append(key)
    walk(data, merged, LABEL_DEFAULTS, "")
    sample = {"n": 1, "r": "Ⅰ", "t": "T", "c": 1, "s": 1, "ss": 1, "g": "g", "i": 1}
    for k, fields in FORMAT_FIELDS.items():
        fmt = merged["formats"][k]
        if fmt is None:
            continue
        try:
            fmt.format(**{f: sample[f] for f in fields})
        except (KeyError, IndexError, ValueError):
            errors.append(f"formats.{k}: allowed placeholders are {', '.join('{' + f + '}' for f in fields)}")
    try:
        merged["letters"]["format"].format(x="g")
    except (KeyError, IndexError, ValueError):
        errors.append("letters.format: the allowed placeholder is {x}")
    if not merged["letters"]["items"]:
        errors.append("letters.items: at least one letter required")
    for kind, cap in merged["caption"].items():
        if not (cap["seq"] or cap["label"]).strip():
            errors.append(f"caption.{kind}: an empty label needs seq (the Word number field name)")
    for k, v in merged["language"].items():
        if not LANG_TAG_RE.fullmatch(v):
            errors.append(f"language.{k}: language tag format (for example ko-KR) required '{v}'")
    for k, v in merged["output"].items():
        if not v.strip() or re.search(r"[\\/]", v):
            errors.append(f"output.{k}: file name without path separators required '{v}'")
    if errors:
        sys.exit(f"ERROR: label file {path}: {len(errors)} errors\n" + "\n".join(f"   {e}" for e in errors))
    _labels = merged
    _compile_notation()
    return changed


def _strip_marker(text: str, marks: tuple[str, ...]) -> str | None:
    for m in marks:
        if text.startswith(m):
            return text[len(m):].lstrip()
    return None


def parse_markdown(path: str) -> tuple[list[Element], dict[str, str]]:
    lines = open(path, encoding="utf-8").read().splitlines()
    footnotes: dict[str, str] = {}
    indents = sorted({len(m.group(1).expandtabs(4)) for ln in lines
                      if (m := BULLET_RE.match(ln))})
    elems: list[Element] = []
    i = 0
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        stripped = line.strip()
        i += 1
        if not stripped:
            continue
        if m := FOOTNOTE_DEF_RE.match(stripped):
            footnotes[m.group(1)] = m.group(2)
            continue
        if stripped == "---":
            elems.append(Element("pagebreak"))
            continue
        if m := HEADING_RE.match(stripped):
            elems.append(Element("heading", m.group(2).strip(), len(m.group(1))))
            continue
        if stripped.startswith("|"):
            rows = []
            j = i - 1
            while j < len(lines) and lines[j].strip().startswith("|"):
                rows.append(lines[j].strip())
                j += 1
            i = j
            tbl = _parse_table(rows, elems)
            while i < len(lines) and not lines[i].strip():
                i += 1
            if i < len(lines) and UNIT_RE.match(lines[i].strip()) and not tbl.data.get("unit"):
                tbl.data["unit"] = lines[i].strip()  # a unit line written after the table also becomes the unit row above the table
                i += 1
            elems.append(tbl)
            continue
        if m := IMAGE_RE.match(stripped):
            elems.append(Element("image", m.group(1).strip(),
                                 data={"path": m.group(2).strip(), "base": os.path.dirname(path)}))
            continue
        if m := CAPTION_RE.match(stripped):
            elems.append(Element("table_caption", m.group(1).strip()))
            continue
        if m := FIG_CAPTION_RE.match(stripped):
            if elems and elems[-1].kind == "image":
                elems[-1].text = m.group(1).strip()
            continue
        if UNIT_RE.match(stripped):
            elems.append(Element("unit", stripped))
            continue
        if stripped.startswith(">"):  # a quote line gets its role from its content (source if it is a source notation, otherwise a note)
            body = stripped.lstrip("> ").strip()
            kind = "source" if SOURCE_RE.match(body) else "note"
            elems.append(Element(kind, SOURCE_PAREN_RE.sub(r"\1: ", body)))
            continue
        if SOURCE_RE.match(stripped):
            elems.append(Element("source", SOURCE_PAREN_RE.sub(r"\1: ", stripped)))
            continue
        if (t := _strip_marker(stripped, SQUARE_MARK)) is not None:
            elems.append(Element("square", t))
            continue
        if (t := _strip_marker(stripped, NOTE_MARK)) is not None and not stripped.startswith("**"):
            elems.append(Element("note", t))
            continue
        if m := BULLET_RE.match(raw):
            indent = len(m.group(1).expandtabs(4))
            rank = indents.index(indent) if indent in indents else len(indents)
            body = m.group(2).strip()
            if (t := _strip_marker(body, CIRCLE_MARK)) is not None:
                elems.append(Element("circle", t))
            elif (t := _strip_marker(body, SQUARE_MARK)) is not None:
                elems.append(Element("square", t))
            elif (t := _strip_marker(body, NOTE_MARK)) is not None:
                elems.append(Element("note", t))
            else:
                kind = ("circle", "dash", "dot")[min(rank, 2)]
                elems.append(Element(kind, body))
            continue
        if (t := _strip_marker(stripped, CIRCLE_MARK)) is not None:
            elems.append(Element("circle", t))
            continue
        elems.append(Element("circle", stripped))  # ordinary paragraphs are treated as ○ text
    return elems, footnotes


def _parse_table(rows: list[str], elems: list[Element]) -> Element:
    aligns: list[str | None] = []
    cells: list[list[str]] = []
    for r in rows:
        if TABLE_SEP_RE.match(r):
            for seg in r.strip("|").split("|"):
                s = seg.strip()
                aligns.append("LEFT" if s.startswith(":") and not s.endswith(":")
                              else "RIGHT" if s.endswith(":") and not s.startswith(":")
                              else "CENTER" if s.startswith(":") else None)
            continue
        cells.append([c.strip().replace("<br>", "\n") for c in r.strip().strip("|").split("|")])
    caption, unit = None, None
    for k in range(len(elems) - 1, max(len(elems) - 3, -1), -1):  # absorb the caption and (unit) lines within two lines before a table, in any order
        if elems[k].kind == "unit" and unit is None:
            unit = elems.pop(k).text
        elif elems[k].kind == "table_caption" and caption is None:
            caption = elems.pop(k).text
        else:
            break
    return Element("table", caption or "", data={"rows": cells, "aligns": aligns, "unit": unit})


# ═══════════════════════════════════════════════════════════════════════
# 2. Manuscript lint: reports manuscript defects before typesetting (conversion continues)
# ═══════════════════════════════════════════════════════════════════════
SHORT_SECTION_ELEMS = 8  # below this many elements per section (##), a page break may leave only blank space


def lint(elems: list[Element]) -> list[str]:
    issues = []
    prev_level = 0
    sec_sizes, cur = [], None
    for e in elems:
        if e.kind == "heading" and e.level <= 2:
            if cur is not None:
                sec_sizes.append(cur)
            cur = 0
        elif cur is not None:
            cur += 1
    if cur is not None:
        sec_sizes.append(cur)
    if sec_sizes:
        short = sum(1 for n in sec_sizes if n < SHORT_SECTION_ELEMS)
        issues.append(f"[page break HITL] of {len(sec_sizes)} chapters and sections, {short} are short (fewer than {SHORT_SECTION_ELEMS} elements): "
                      f"keep the default chapter. If the user asks to change page breaks, use --page-break section|none")
    for idx, e in enumerate(elems):
        if e.kind == "heading":
            if e.level > prev_level + 1 and prev_level:
                issues.append(f"[heading skipped] '{e.text[:30]}' (#{prev_level} followed by #{e.level})")
            prev_level = e.level
        if e.kind == "square":
            nxt = elems[idx + 1].kind if idx + 1 < len(elems) else None
            if nxt not in ("circle", "dash", "dot", "note"):
                issues.append(f"[lone □] no ○ or - text under '{e.text[:30]}'")
        if e.kind == "table" and len(e.data["rows"]) > LONG_TABLE_ROWS:
            issues.append(f"[long table HITL] table of {len(e.data['rows'])} rows: ask the user to choose --long-table engine|split")
        if e.kind == "table" and not e.text:
            issues.append("[table caption missing] a caption of the form '<표> title' on the line before the table is recommended")
    return issues


# ═══════════════════════════════════════════════════════════════════════
# 3. Typesetter: places elements on the skeleton with python-hwpx
# ═══════════════════════════════════════════════════════════════════════
def _weighted_len(text: str) -> int:
    return sum(2 if ("\uac00" <= ch <= "\ud7a3" or "\u4e00" <= ch <= "\u9fff") else 1 for ch in text)


def _line_units(cell: str) -> int:
    """Weighted length of the longest line of a cell (units needed to fit everything on one line)."""
    return max((_weighted_len(ln.strip()) for ln in cell.split("\n")), default=0)


def _word_units(cell: str) -> int:
    """Weighted length of the longest word of a cell (minimum units so no word is broken)."""
    return max((_weighted_len(w) for ln in cell.split("\n") for w in ln.split()), default=0)


def calc_col_widths(rows: list[list[str]], max_w: int) -> list[int]:
    """Column width allocation. The sum is the table's actual width.
    If every column fits on one line, each column gets its one-line width (scaled up proportionally if the table is narrower than the floor).
    If not, the table width is set to the cap, and bisection finds the baseline at which the sum of 'column width = min(one-line width, baseline)' equals the cap.
    Why: short columns keep their one-line width and only long columns fold at the baseline. Word width is guaranteed for every column to prevent line breaks in the middle of words."""
    ncol = max(len(r) for r in rows)
    cols = [[r[c] if c < len(r) else "" for r in rows] for c in range(ncol)]
    one_line = [max(COL_MIN_W, max(map(_line_units, col)) * COL_UNIT_W + COL_PAD_W) for col in cols]
    floor = [max(COL_MIN_W, max(map(_word_units, col)) * COL_UNIT_W + COL_PAD_W) for col in cols]
    if sum(one_line) <= max_w:
        target = max(sum(one_line), int(max_w * TABLE_MIN_RATIO))
        widths = [w * target / sum(one_line) for w in one_line]
    elif sum(floor) >= max_w:
        target = max_w
        widths = [w * target / sum(floor) for w in floor]
    else:
        target = max_w
        lo, hi = 0.0, float(max(one_line))
        for _ in range(COL_SEARCH_STEPS):
            cap = (lo + hi) / 2
            if sum(max(f, min(o, cap)) for o, f in zip(one_line, floor)) < target:
                lo = cap
            else:
                hi = cap
        widths = [max(f, min(o, hi)) for o, f in zip(one_line, floor)]
    out = [int(w) for w in widths]
    out[out.index(max(out))] += target - sum(out)  # the widest column absorbs the rounding remainder
    return out


def header_rows(rows: list[list[str]]) -> int:
    """Number of header rows. If rows merging upward with '^' follow the first row, the header is multi-level."""
    n = 1
    while n < len(rows) and any(v.strip() == "^" for v in rows[n]) and not any(
            NUMERIC_CELL.match(v.replace(" ", "")) for v in rows[n] if v.strip() and v.strip() not in ("^", "<")):
        n += 1
    return n


def merge_spans(rows: list[list[str]], mode: str, hdr: int = 1) -> list[tuple[int, int, int]]:
    """Vertical merge spans (col, start_row, end_row). Based on body rows (header excluded).
    mode auto: merge consecutive equal values in leading label columns (before the first numeric column, at most MERGE_LABEL_COLS).
    Lower columns merge only inside the merge spans of upper columns to keep the hierarchy consistent. A cell value '^' is an instruction to merge with the cell above."""
    body = rows[hdr:]
    ncol = len(rows[0])
    spans: list[tuple[int, int, int]] = []
    for c in range(ncol):  # vertical merges of multi-level headers (based on row 0, written with negative offsets)
        k = 0
        while k + 1 < hdr and rows[k + 1][c].strip() == "^":
            k += 1
        if k:
            spans.append((c, -hdr, -hdr + k))
    if not body:
        return spans
    label_cols = []
    if mode == "auto":
        for c in range(min(ncol, MERGE_LABEL_COLS)):
            vals = [r[c].strip() for r in body if r[c].strip() and r[c].strip() not in ("^", "<")]
            if vals and all(NUMERIC_CELL.match(v.replace(" ", "")) for v in vals):
                break
            label_cols.append(c)
    parent_bounds = [(0, len(body) - 1)]
    for c in range(ncol):
        explicit = any(r[c].strip() == "^" for r in body)
        if c not in label_cols and not explicit:
            continue
        new_bounds = []
        for lo, hi in parent_bounds:
            start = lo
            for i in range(lo + 1, hi + 2):
                same = i <= hi and (body[i][c].strip() == "^" or
                                    (c in label_cols and body[i][c].strip() == body[start][c].strip() and body[start][c].strip()))
                if not same:
                    if i - 1 > start:
                        spans.append((c, start, i - 1))
                    new_bounds.append((start, i - 1))
                    start = i
        parent_bounds = new_bounds
    return spans


def hmerge_spans(rows: list[list[str]]) -> list[tuple[int, int, int]]:
    """Horizontal merge spans (row, start_col, end_col). A cell value '<' is an instruction to merge with the left cell. Header rows included."""
    spans = []
    for ri, r in enumerate(rows):
        c = 0
        while c < len(r):
            if r[c].strip() == "<":
                c += 1
                continue
            end = c
            while end + 1 < len(r) and r[end + 1].strip() == "<":
                end += 1
            if end > c:
                spans.append((ri, c, end))
            c = end + 1
    return spans


def col_aligns(rows: list[list[str]], explicit: list[str | None]) -> list[str]:
    """Alignment is decided per column. Explicit marker > amount and decimal number columns RIGHT > other number columns (serial, year, count) CENTER > columns with bullets or long text LEFT > CENTER."""
    ncol = max(len(r) for r in rows)
    out = []
    for c in range(ncol):
        body = [r[c] for r in rows[1:] if c < len(r) and r[c].strip()]
        if c < len(explicit) and explicit[c]:
            out.append(explicit[c])
        elif body and all(NUMERIC_CELL.match(v.replace(" ", "")) for v in body):
            out.append("RIGHT" if any(AMOUNT_CELL.search(v) for v in body) else "CENTER")
        elif any(v.lstrip().startswith(BULLET_PREFIXES) or "\n" in v
                 or _weighted_len(v) > COL_ONE_LINE_UNITS for v in body):
            out.append("LEFT")
        else:
            out.append("CENTER")
    return out


# ── Format-agnostic typesetting judgments (shared by the hwpx and docx writers) ─────────────────────
INLINE_RE = re.compile(
    r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`|<sup>.*?</sup>|<sub>.*?</sub>|\^[^^\s]+\^"
    r"|\[\^[^\]]+\])")


def inline_segments(text: str) -> list[tuple[str, str]]:
    """Split inline notation into a list of (kind, text). Kinds are text, bold, italic, code, sup, sub, footnote (footnote key)"""
    out = []
    for seg in INLINE_RE.split(text):
        if not seg:
            continue
        if seg.startswith("**"):
            out.append(("bold", seg[2:-2]))
        elif seg.startswith("*"):
            out.append(("italic", seg[1:-1]))
        elif seg.startswith("`"):
            out.append(("code", seg[1:-1]))
        elif seg.startswith("<sup>") or seg.startswith("^"):
            out.append(("sup", seg[5:-6] if seg.startswith("<") else seg[1:-1]))
        elif seg.startswith("<sub>"):
            out.append(("sub", seg[5:-6]))
        elif seg.startswith("[^"):
            out.append(("footnote", seg[2:-1]))
        else:
            out.append(("text", seg))
    return out


def edge_sides(c0: int, c1: int, ncol: int) -> list[str]:
    """Sides to draw among the four cell rule sides. The outer left and right table rules stay empty (EDGE_OPEN_SIDES)"""
    sides = ["top", "bottom", "left", "right"]
    if EDGE_OPEN_SIDES and c0 == 0:
        sides.remove("left")
    if EDGE_OPEN_SIDES and c1 == ncol - 1:
        sides.remove("right")
    return sides


def cell_line_count(val: str, width: int) -> int:
    """Number of folded lines estimated from the cell text width and cell width"""
    inner = max(width - CELL_MARGIN[0] - CELL_MARGIN[1], 1)
    return max(1, val.count("\n") + 1,
               sum(-(-_weighted_len(ln) * CELL_UNIT_W // inner) for ln in val.split("\n")))


def lines_to_height(lines: int) -> int:
    return ROW_BASE_H + ROW_LINE_H * (lines - 1)


def row_heights(rows: list[list[str]], widths: list[int], hdr: int, vspans: list, hspans: list) -> list[int]:
    """Minimum height per row reflecting merges. Estimates folded lines from cell text width and column width (한/글 and Word expand to fit content).
    Row indexes count manuscript rows including the header.
    Why: counting the label of an automatic vertical merge span for every row makes every merged row as tall as the label's line count.
    Vertical merge labels are left out of row height calculation, and only if the sum of the span's row heights is smaller than the label height is the shortfall spread evenly over the span's rows."""
    anchors_v = {(hdr + r0, c): hdr + r1 for c, r0, r1 in vspans}
    covered = {(r, c) for (a, c), b in anchors_v.items() for r in range(a + 1, b + 1)}
    span_w = {}
    for ri, c0, c1 in hspans:
        span_w[(ri, c0)] = sum(widths[c0:c1 + 1])
        covered |= {(ri, c) for c in range(c0 + 1, c1 + 1)}
    heights = []
    for ri, row in enumerate(rows):
        lines = [cell_line_count(v, span_w.get((ri, ci), widths[ci])) for ci, v in enumerate(row)
                 if (ri, ci) not in covered and (ri, ci) not in anchors_v]
        heights.append(lines_to_height(max(lines, default=1)))
    for (a, c), b in anchors_v.items():
        need = lines_to_height(cell_line_count(rows[a][c], widths[c]))
        short = need - sum(heights[a:b + 1])
        if short > 0:
            n = b - a + 1
            for k in range(n):
                heights[a + k] += short // n + (1 if k < short % n else 0)
    return heights


def cell_lines(val: str, role_fn) -> tuple[str, list[str | None]]:
    """Bullet role per line of a cell value and the value with bullet symbols removed. Cells with the '^' or '<' merge instruction are empty"""
    lines_txt = val.split("\n")
    roles = [role_fn(ln) for ln in lines_txt]
    if any(roles):
        val = "\n".join(re.sub(r"^\s*[◦○ㅇ·•ㆍ-]\s*", "", ln) if br else ln for ln, br in zip(lines_txt, roles))
    return ("" if val.strip() in ("^", "<") else val), roles


def cell_role(align: str, is_head: bool) -> tuple[str, str | None]:
    """Role and alignment override of a cell paragraph. Long-text columns use the left style, amount columns right alignment, headers center"""
    role = "table_cell_left" if align == "LEFT" and not is_head else "table_cell"
    return role, ("RIGHT" if not is_head and align == "RIGHT" else None)


def table_spans(rows: list[list[str]], hdr: int, merge_mode: str) -> tuple[list, list]:
    """Vertical merges (col, r0, r1) and horizontal merges (row, c0, c1). Horizontal merges that overlap vertical merges are skipped with a notice"""
    if merge_mode == "none":
        return [], []
    vspans = merge_spans(rows, merge_mode, hdr)
    vcells = {(r, c) for c, r0, r1 in vspans for r in range(r0, r1 + 1)}
    hspans = []
    for ri, c0, c1 in hmerge_spans(rows):
        if any((ri - hdr, c) in vcells for c in range(c0, c1 + 1)):
            print(f"WARN: table row {ri} columns {c0}-{c1} horizontal merge overlaps a vertical merge, skipped")
            continue
        hspans.append((ri, c0, c1))
    return vspans, hspans


def split_table(e: Element, rows: list[list[str]]) -> list[Element]:
    """Physical split of a long table. Reattach the header row to each chunk and append (sequence) to the caption"""
    chunks = [rows[i:i + LONG_TABLE_ROWS - 1] for i in range(1, len(rows), LONG_TABLE_ROWS - 1)]
    return [Element("table", f"{e.text} ({k})" if e.text else "",
                    data={**e.data, "rows": [rows[0]] + chunk, "split_done": True})
            for k, chunk in enumerate(chunks, 1)]


def front_matter_items(cover: dict | None, toc: bool, headings: list[tuple[int, str]], toc_depth: int,
                       lists: list[tuple[str, list[str]]] | None = None) -> list[dict]:
    """Paragraph order of the cover and TOC (including the list of tables and list of figures). The writer only records this list in order
    kind: blank, cover(text, pt, bold), toc_title(text, page_break), toc_entry(level, text), list_entry(text)"""
    items: list[dict] = []
    if cover:
        items += [{"kind": "blank"}] * COVER_TOP_LINES
        items.append({"kind": "cover", "text": cover.get("title", ""), "pt": COVER_TITLE_PT, "bold": True})
        if cover.get("subtitle"):
            items += [{"kind": "blank"}, {"kind": "cover", "text": cover["subtitle"], "pt": COVER_SUB_PT, "bold": False}]
        items += [{"kind": "blank"}] * 2
        if cover.get("date"):
            items.append({"kind": "cover", "text": cover["date"], "pt": COVER_DATE_PT, "bold": False})
        items += [{"kind": "blank"}] * COVER_GAP_LINES
        if cover.get("org"):
            items.append({"kind": "cover", "text": cover["org"], "pt": COVER_ORG_PT, "bold": True})
    if toc:
        items.append({"kind": "toc_title", "text": label("toc.title"), "page_break": bool(cover)})
        items += [{"kind": "toc_entry", "level": lv, "text": text} for lv, text in headings if lv <= toc_depth]
        for title, entries in lists or []:
            if entries:
                items.append({"kind": "toc_title", "text": title, "page_break": True})
                items += [{"kind": "list_entry", "text": t} for t in entries]
    return items


def prepare_image(e: Element) -> tuple[bytes, str, float, float] | None:
    """(bytes, format, width mm, height mm) of a figure element. If the file is missing, report it and return None. webp, which 한/글 does not support, is converted to jpg"""
    path = e.data["path"]
    if path.startswith(("http://", "https://")):
        import urllib.request
        req = urllib.request.Request(path, headers={"User-Agent": "Mozilla/5.0"})
        with urllib.request.urlopen(req, timeout=IMG_URL_TIMEOUT) as resp:
            data = resp.read()
            ctype = resp.headers.get("Content-Type", "")
        fmt = {"image/jpeg": "jpg", "image/png": "png", "image/gif": "gif", "image/bmp": "bmp",
               "image/webp": "webp"}.get(ctype.split(";")[0].strip(), os.path.splitext(path)[1].lstrip(".").lower() or "png")
    else:
        if not os.path.isabs(path):
            path = os.path.join(e.data["base"], path)
        if not os.path.exists(path):
            print(f"WARN: image not found {path}")
            return None
        data = open(path, "rb").read()
        fmt = os.path.splitext(path)[1].lstrip(".").lower() or "png"
    if fmt == "jpeg":
        fmt = "jpg"
    w_mm, h_mm = image_mm(data)
    if fmt == "webp":
        from PIL import Image
        import io
        buf = io.BytesIO()
        Image.open(io.BytesIO(data)).convert("RGB").save(buf, "JPEG", quality=90)
        data, fmt = buf.getvalue(), "jpg"
    return data, fmt, w_mm, h_mm


def image_mm(data: bytes) -> tuple[float, float]:
    """Apply the figure size norms (width cap for landscape, width cap near square, height cap for portrait)"""
    try:
        from PIL import Image
        import io
        im = Image.open(io.BytesIO(data))
        px_w, px_h = im.size
        dpi = im.info.get("dpi", (IMG_DEFAULT_DPI, IMG_DEFAULT_DPI))[0] or IMG_DEFAULT_DPI
    except Exception:
        px_w, px_h, dpi = 800, 600, IMG_DEFAULT_DPI
    w_mm, h_mm = px_w / dpi * 25.4, px_h / dpi * 25.4
    ratio = w_mm / h_mm
    if ratio >= 1.4:
        scale = IMG_LANDSCAPE_W_MM / w_mm
    elif ratio > 0.8:
        scale = IMG_SQUARE_W_MM / w_mm
    else:
        scale = IMG_PORTRAIT_H_MM / h_mm
    return round(w_mm * scale, 1), round(h_mm * scale, 1)


def cell_line_role(line: str, profile: dict) -> str | None:
    """Map ◦ and - lines inside a cell to cell bullet roles (only when the profile has the roles)."""
    st = line.lstrip()
    if st.startswith(("◦", "○", "ㅇ", "·", "•", "ㆍ")) and profile.get("table_cell_bullet"):
        return "table_cell_bullet"
    if st.startswith("- ") and profile.get("table_cell_dash"):
        return "table_cell_dash"
    return None


class Builder:
    def __init__(self, skeleton: str, profile: dict, page_break: str = "chapter"):
        try:
            from hwpx import HwpxDocument
        except ImportError:
            sys.exit("ERROR: python-hwpx not installed. Run `pip install -U python-hwpx --break-system-packages` and retry")
        try:
            with zipfile.ZipFile(skeleton) as z:
                bad = z.testzip()
            if bad:
                sys.exit(f"ERROR: corrupted skeleton zip entry {bad}. Reinstall the skill package")
            self.doc = HwpxDocument.open(os.path.abspath(skeleton))
        except SystemExit:
            raise
        except Exception as ex:
            sys.exit(f"ERROR: failed to load the skeleton ({type(ex).__name__}: {ex}). "
                     "The skeleton file is intact, so a python-hwpx version problem is likely. Check `pip show python-hwpx` and reinstall the latest version. "
                     "Do not assemble a new document by calling the python-hwpx API directly. Update this script to the latest API")
        self.p = profile
        self.page_break = page_break  # chapter | section | none
        self.styles = {s.name: (sid, s) for sid, s in self.doc.styles.items()}
        missing = [v for k, v in profile.items()
                   if k not in ("numbering", "chapter_format", "section_format", "subsection_format", "caption_prefix") and v is not None and v not in self.styles]
        self.headings: list[tuple[int, str]] = []
        self.toc_numbers: dict[str, int] = {}
        if missing:
            sys.exit(f"ERROR: styles missing from the skeleton {missing}: fix the profile or check the list with the `skeleton` command")
        self.n = {"chapter": 0, "section": 0, "subsection": 0, "ganada": 0, "table": 0, "figure": 0,
                  "total_section": 0}
        self.header = self.doc.headers[0]
        self.body_w = self._body_width()
        self.table_max_w = self.body_w - TABLE_INDENT_MARGIN
        self._align_pp: dict[tuple[str, str], str] = {}
        self._bold_cp: dict[str, str] = {}
        self.pending_break = False
        self.manual_symbol = self._manual_symbols()
        self._warned: set[str] = set()
        self.numbered_square = False
        self.toc = False
        self.long_table_mode = "engine"
        self.after_block = False

    def _manual_symbols(self) -> dict[str, str]:
        """If the style's heading is BULLET, 한/글 draws the symbol. If NONE, the symbol must be written into the body."""
        out = {}
        for role, sym in MANUAL_SYMBOLS.items():
            if self.p.get(role) is None:
                continue
            pp_id = str(self.styles[self.p[role]][1].para_pr_id_ref)
            pp = self.header.paragraph_property(pp_id)
            htype = getattr(getattr(pp, "heading", None), "type", None) or getattr(pp, "heading_type", None)
            if htype is None:  # without the object attribute, read it directly from the XML
                el = self.header.element.find(f".//{{http://www.hancom.co.kr/hwpml/2011/head}}paraPr[@id='{pp_id}']")
                hd = el.find("{http://www.hancom.co.kr/hwpml/2011/head}heading") if el is not None else None
                htype = hd.get("type") if hd is not None else "NONE"
            if str(htype).upper() != "BULLET":
                out[role] = sym
        return out

    # ── Basic elements ───────────────────────────────────────────────────────
    def _body_width(self) -> int:
        sp = self.doc.sections[0].properties
        return sp.page_size.width - sp.page_margins.left - sp.page_margins.right

    def refs(self, role: str) -> dict:
        if self.p.get(role) is None and role not in self._warned:
            self._warned.add(role)
            print(f"WARN: no '{role}' style in the profile → output as 바탕글. The user must choose among deriving it, adding it to the skeleton, or absorbing it into an existing level")
        sid, s = self.styles[self.p[role] or self.p["body"]]
        return {"style_id_ref": sid, "para_pr_id_ref": s.para_pr_id_ref, "char_pr_id_ref": s.char_pr_id_ref}

    INLINE_RE = INLINE_RE

    def _run_style(self, base_cp: str, **fmt) -> str:
        """Clone the base charPr and add only bold, italic, and script elements.
        Why: the library's ensure_run returns an existing charPr that matches only bold (with a different size), which enlarges header text."""
        key = (base_cp,) + tuple(sorted(fmt.items()))
        if key in self._bold_cp:
            return self._bold_cp[key]
        HH = "{http://www.hancom.co.kr/hwpml/2011/head}"
        base = self.header.element.find(f".//{HH}charPr[@id='{base_cp}']")
        tags = {"bold": "bold", "italic": "italic", "sup": "supscript", "sub": "subscript"}
        wanted = [tags[k] for k, v in (("bold", fmt.get("bold")), ("italic", fmt.get("italic")),
                                       (fmt.get("script"), True)) if v and k in tags]

        def same(e):
            if e.get("height") != base.get("height") or e.get("textColor") != base.get("textColor"):
                return False
            fr, br = e.find(f"{HH}fontRef"), base.find(f"{HH}fontRef")
            if fr is None or br is None or fr.get("hangul") != br.get("hangul"):
                return False
            has = {c.tag.split("}")[1] for c in e}
            return all(t in has for t in wanted) and not ({"bold", "italic", "supscript", "subscript"} & has) - set(wanted)

        def add(e):
            import lxml.etree as ET
            anchor = e.find(f"{HH}offset")
            idx = list(e).index(anchor) + 1 if anchor is not None else len(e)
            for t in wanted:
                e.insert(idx, ET.Element(f"{HH}{t}"))
                idx += 1

        el = self.header.ensure_char_property(predicate=same, modifier=add, base_char_pr_id=base_cp)
        self._bold_cp[key] = el.get("id")
        return self._bold_cp[key]

    def para(self, role: str, text: str, footnotes: dict | None = None, no_heading: bool = False, **attrs):
        """Formatting applies only when the three style references are set together.
        Inline: **bold** *italic* ^superscript^ <sup></sup> <sub></sub> `verbatim` [^n] footnote (at that position)."""
        r = self.refs(role)
        if role == "square" and self.numbered_square and re.match(r"^\d+\.\s", text):
            r["para_pr_id_ref"] = self._no_bullet_pp(str(r["para_pr_id_ref"]))  # 1. 2. 3. numeric headings
        elif no_heading:  # the number is written in the body, so turn off the style's automatic outline number
            r["para_pr_id_ref"] = self._no_bullet_pp(str(r["para_pr_id_ref"]))
        else:
            text = self.manual_symbol.get(role, "") + text

        if self.pending_break:
            attrs["pageBreak"] = "1"
            self.pending_break = False
        p = self.doc.add_paragraph("", **r, **attrs)
        cp = str(r["char_pr_id_ref"])
        for kind, seg in inline_segments(text):
            if kind == "bold":
                p.add_run(seg, char_pr_id_ref=self._run_style(cp, bold=True))
            elif kind == "italic":
                p.add_run(seg, char_pr_id_ref=self._run_style(cp, italic=True))
            elif kind in ("sup", "sub"):
                p.add_run(seg, char_pr_id_ref=self._run_style(cp, script=kind))
            elif kind == "footnote":
                if footnotes is not None:
                    p.add_footnote(footnotes.get(seg, ""))
            else:  # text, code
                p.add_run(seg, char_pr_id_ref=cp)
        return p

    def _centered_refs(self, role: str) -> dict:
        """Table and figure paragraphs are centered regardless of style alignment. Even if a custom template's table position style is left-aligned, the table does not stick to the left."""
        r = self.refs(role)
        HH = "{http://www.hancom.co.kr/hwpml/2011/head}"
        el = self.header.element.find(f".//{HH}paraPr[@id='{r['para_pr_id_ref']}']")
        al = el.find(f"{HH}align") if el is not None else None
        if al is None or al.get("horizontal") != "CENTER":
            r["para_pr_id_ref"] = self._aligned_pp(str(r["para_pr_id_ref"]), "CENTER")
        return r

    def _aligned_pp(self, base_pp: str, align: str) -> str:
        key = (str(base_pp), align)
        if key not in self._align_pp:
            self._align_pp[key] = self.header.ensure_paragraph_format(base_para_pr_id=base_pp, alignment=align)
        return self._align_pp[key]

    def _no_bullet_pp(self, base_pp: str) -> str:
        key = (base_pp, "NOBULLET")
        if key not in self._align_pp:
            self._align_pp[key] = self.header.ensure_paragraph_format(
                base_para_pr_id=base_pp, heading={"type": "NONE", "idRef": "0", "level": "0"})
        return self._align_pp[key]


    def gap(self):
        for _ in range(BLOCK_GAP_LINES):
            self.doc.add_paragraph("", **self.refs("body"))


    def _cover_char(self, height: int, bold: bool) -> str:
        base = str(self.refs("chapter")["char_pr_id_ref"])
        HH = "{http://www.hancom.co.kr/hwpml/2011/head}"

        def mod(e):
            import lxml.etree as ET
            e.set("height", str(height))
            b = e.find(f"{HH}bold")
            if bold and b is None:
                anchor = e.find(f"{HH}offset")
                e.insert(list(e).index(anchor) + 1, ET.Element(f"{HH}bold"))
            if not bold and b is not None:
                e.remove(b)
        font = self.header.element.find(f".//{HH}charPr[@id='{base}']").find(f"{HH}fontRef").get("hangul")
        return self.header.ensure_char_property(
            predicate=lambda e: e.get("height") == str(height) and e.find(f"{HH}fontRef").get("hangul") == font
            and (e.find(f"{HH}bold") is not None) == bold, modifier=mod, base_char_pr_id=base).get("id")

    def _front_matter(self, cover: dict | None):
        """Put the cover and TOC in the front section (section0) and move the body into a new section. Body page numbers start from 1."""
        import copy
        import lxml.etree as ET
        s0 = self.doc.sections[0]
        s1 = self.doc.add_section(after=0)
        old = s1.paragraphs[0].element
        old.getparent().replace(old, copy.deepcopy(s0.paragraphs[0].element))
        for p in list(s0.element)[1:]:
            s0.element.remove(p)
            s1.element.append(p)
        self.body_section = s1
        sn = s1.paragraphs[0].element.find(f".//{HP}startNum")
        if sn is not None:
            sn.set("page", "1")
        body_ref = self.refs("body")
        center = self._aligned_pp(str(body_ref["para_pr_id_ref"]), "CENTER")

        def front(text="", cp=None, pp=center, **attrs):
            p = self.doc.add_paragraph("", **body_ref, **attrs)
            p.element.set("paraPrIDRef", pp)
            if text:
                p.add_run(text, char_pr_id_ref=cp or str(body_ref["char_pr_id_ref"]))
            s0.element.append(p.element)  # add_paragraph appends to the last section, so move it to the front section
            return p
        if self.toc and not (self.p.get("toc_chapter") and self.p.get("toc_section")):
            print("WARN: no TOC styles (toc_chapter, toc_section) in the profile, --toc skipped")
            self.toc = False
        sub_pp = list_pp = None
        for it in front_matter_items(cover, self.toc, self.headings, getattr(self, "toc_depth", 2), self.toc_lists()):
            kind = it["kind"]
            if kind == "blank":
                front()
            elif kind == "cover":
                front(it["text"], self._cover_char(it["pt"], it["bold"]))
            elif kind == "toc_title":
                ch = self.refs("chapter")
                front(it["text"], str(ch["char_pr_id_ref"]), self._aligned_pp(str(ch["para_pr_id_ref"]), "CENTER"),
                      pageBreak="1" if it["page_break"] else "0")
            else:  # toc_entry, list_entry
                lv = it.get("level", 2)
                text = it["text"]
                role = "toc_chapter" if lv == 1 else "toc_section"
                r = self.refs(role)
                pp = str(r["para_pr_id_ref"])
                left = TOC_SUB_INDENT if lv >= 3 else (0 if lv == 1 else 1500)
                if kind == "list_entry":
                    if list_pp is None:
                        list_pp = self.header.ensure_paragraph_format(base_para_pr_id=pp, margins={"left": TOC_LIST_INDENT})
                    pp, left = list_pp, TOC_LIST_INDENT
                elif lv >= 3:
                    if sub_pp is None:
                        sub_pp = self.header.ensure_paragraph_format(base_para_pr_id=pp, margins={"left": TOC_SUB_INDENT})
                    pp = sub_pp
                p = front(text, str(r["char_pr_id_ref"]), pp)
                p.element.set("styleIDRef", str(r["style_id_ref"]))
                num = self.toc_numbers.get(text)
                if num:
                    tnode = [t for t in p.element.iter(f"{HP}t") if t.text][-1]
                    unit = TOC_UNIT_W[0] if lv == 1 else TOC_UNIT_W[1]
                    width = max(TOC_TAB_POS - left - _weighted_len(text) * unit - len(str(num)) * unit, 1000)
                    ET.SubElement(tnode, f"{HP}tab", width=str(width), leader="3", type="2").tail = str(num)
        for sec in (s0, s1):
            sec.mark_dirty() if hasattr(sec, "mark_dirty") else None

    def toc_lists(self) -> list[tuple[str, list[str]]]:
        """List of tables and list of figures (title, entry list). The base typesetter has no lists"""
        return []

    def set_header(self, spec: str | None):
        """Header (OFF by default). "text;align=LEFT|CENTER|RIGHT". Applied only to body sections.
        python-hwpx writes the header in two places, inside secPr and in ctrl, so the copy inside secPr is removed and the paragraph alignment of the ctrl copy is set."""
        if not spec:
            return
        parts = spec.split(";")
        text = parts[0].strip()
        align = next((kv.split("=", 1)[1].strip().upper() for kv in parts[1:] if kv.strip().lower().startswith("align=")), "RIGHT")
        sec = self.doc.sections[-1]
        hf = self.doc.set_header_text(text, section=sec)
        body_pp = str(self.refs("body")["para_pr_id_ref"])
        pp = self._aligned_pp(body_pp, align)
        HPH = f"{HP}header"
        first = sec.paragraphs[0].element
        sec_pr = first.find(f".//{HP}secPr")
        if sec_pr is not None:
            for h in list(sec_pr.iter(HPH)):
                h.getparent().remove(h)
        for h in first.iter(HPH):
            for pe in h.iter(f"{HP}p"):
                pe.set("paraPrIDRef", pp)

    def set_page_number(self, on: bool):
        """Footer page numbers. OFF by default (prevents number conflicts when documents are combined). When ON, '- N -' at bottom center."""
        for sec in self.doc.sections:  # always remove from the front section, insert in body sections only when on
            fp = sec.paragraphs[0].element
            for ctrl in list(fp.iter(f"{HP}ctrl")):
                if ctrl.find(f"{HP}pageNum") is not None:
                    ctrl.getparent().remove(ctrl)
        first = self.doc.sections[-1].paragraphs[0].element
        for ctrl in list(first.iter(f"{HP}ctrl")):
            if ctrl.find(f"{HP}pageNum") is not None:
                ctrl.getparent().remove(ctrl)
        if on:
            import lxml.etree as ET
            run = first.find(f"{HP}run")
            ctrl = ET.Element(f"{HP}ctrl")
            ET.SubElement(ctrl, f"{HP}pageNum", pos="BOTTOM_CENTER", formatType="DIGIT", sideChar="-")
            # follows the 한/글 authoring order (section definition, column definition, page number). If the page number comes before the column definition, 한/글
            # reads the first section's margins as defaults (30mm left and right), shrinking the body width to 150mm so tables exceed the print area
            anchor = next((c for c in run.findall(f"{HP}ctrl") if c.find(f"{HP}colPr") is not None),
                          run.find(f"{HP}secPr"))
            run.insert(list(run).index(anchor) + 1 if anchor is not None else 0, ctrl)

    def _cell_line_role(self, line: str) -> str | None:
        return cell_line_role(line, self.p)

    def _border_fill(self, sides: list[str], fill: str | None = None) -> str:
        """Create the border definition and check and force the state of the four sides directly on the element.
        Why: depending on the library version, active_borders is interpreted differently, and there were cases where rules remained on transparent rows."""
        HH = "{http://www.hancom.co.kr/hwpml/2011/head}"
        bid = self.header.ensure_border_fill(border_color=BORDER_COLOR, border_width=BORDER_WIDTH,
                                             active_borders=sides, fill_color=fill)
        el = self.header.element.find(f".//{HH}borderFill[@id='{bid}']")
        if el is not None:
            for side in ("left", "right", "top", "bottom"):
                b = el.find(f"{HH}{side}Border")
                if b is not None:
                    b.set("type", "SOLID" if side in sides else "NONE")
        return str(bid)

    def _transparent_row(self, t, ri: int, ncol: int, text: str, role: str, align: str):
        """Transparent merged rows inside a table (unit, source). Kept inside the table so they move with it during hand editing."""
        no_border = self._border_fill([])
        u = self.refs(role)
        t.set_cell_text(ri, 0, text)
        for pe in t.cell(ri, 0).element.iter(f"{HP}p"):
            pe.set("paraPrIDRef", self._aligned_pp(str(u["para_pr_id_ref"]), align))
            pe.set("styleIDRef", str(u["style_id_ref"]))
            for run in pe.iter(f"{HP}run"):
                run.set("charPrIDRef", str(u["char_pr_id_ref"]))
        for ci in range(ncol):
            t.cell(ri, ci).element.set("borderFillIDRef", str(no_border))
            t.cell(ri, ci).set_size(height=ROW_BASE_H)
        if ncol > 1:
            t.merge_cells(ri, 0, ri, ncol - 1)

    def attach_caption(self, holder, kind: str, text: str, side: str, width: int):
        """Object-attached caption (the 한/글 caption feature). The number is an autoNum field and restarts at the start of a chapter."""
        role = "table_title" if kind == "TABLE" else "fig_title"
        r = self.refs(role)
        cap = holder.set_caption("", side=side, width=width, gap=CAPTION_GAP, char_pr_id_ref=r["char_pr_id_ref"])
        cp = cap.paragraphs[0]
        cp.element.set("paraPrIDRef", str(r["para_pr_id_ref"]))
        cp.element.set("styleIDRef", str(r["style_id_ref"]))
        for run in [c for c in cp.element if c.tag == f"{HP}run"]:
            cp.element.remove(run)
        cap = label("caption.table" if kind == "TABLE" else "caption.figure")
        open_b, close_b = f"{cap['open']}{cap['label']} ", f"{cap['close']} "
        prefix = self.p.get("caption_prefix", "").format(c=max(self.n["chapter"], 1))
        run = cp.add_run(open_b + prefix, char_pr_id_ref=r["char_pr_id_ref"])
        import lxml.etree as ET
        an = ET.SubElement(ET.SubElement(run.element, f"{HP}ctrl"), f"{HP}autoNum", num="1", numType=kind)
        ET.SubElement(an, f"{HP}autoNumFormat", type="DIGIT", userChar="", prefixChar="", suffixChar="", supscript="0")
        cp.add_run(close_b + text, char_pr_id_ref=r["char_pr_id_ref"])

    # ── Heading levels ───────────────────────────────────────────────────────
    def heading(self, level: int, text: str):
        text = NUM_PREFIX_RE.sub("", text).strip()
        role = ("chapter", "section", "subsection", "ganada")[level - 1]
        manual = self.p["numbering"] == "manual"
        if level == 1:
            self.n.update(chapter=self.n["chapter"] + 1, section=0, subsection=0, ganada=0)
            if self.page_break in ("chapter", "section") and self.n["chapter"] > 1:
                self.pending_break = True
            label = self.p["chapter_format"].format(n=self.n["chapter"], t=text)
        elif level == 2:
            self.n.update(section=self.n["section"] + 1, subsection=0, ganada=0)
            label = (self.p.get("section_format", "{c}.{s} {t}").format(c=max(self.n["chapter"], 1), s=self.n["section"], t=text)
                     if manual else text)
            # per-section breaks: the first section right after a chapter and the document's first section are already on a new page, so they are excluded
            if self.page_break == "section" and not (self.n["section"] == 1 and self.n["chapter"] > 0) \
                    and self.n["total_section"] > 0:
                self.pending_break = True
            self.n["total_section"] += 1
        elif level == 3:
            self.n.update(subsection=self.n["subsection"] + 1, ganada=0)
            label = (self.p.get("subsection_format", "{c}.{s}.{ss} {t}").format(
                c=max(self.n["chapter"], 1), s=max(self.n["section"], 1), ss=self.n["subsection"], t=text) if manual else text)
        else:
            self.n["ganada"] += 1
            label = f"{letter_label(self.n['ganada'])} {text}"
        p = self.para(role, label, no_heading=manual and level in (2, 3))
        if level <= 3:
            self.headings.append((level, label))
        if level == 1 and self.p.get("caption_prefix"):  # restart from <표 N-1> per chapter
            p.add_new_num(kind="TABLE", number=1)
            p.add_new_num(kind="PICTURE", number=1)

    # ── Tables ──────────────────────────────────────────────────────────────
    def table(self, e: Element):
        rows = [r + [""] * (max(len(x) for x in e.data["rows"]) - len(r)) for r in e.data["rows"]]
        nrow, ncol = len(rows), len(rows[0])
        widths = calc_col_widths(rows, self.table_max_w)
        aligns = col_aligns(rows, e.data["aligns"])
        hdr = header_rows(rows)
        table_w = sum(widths)
        if nrow > LONG_TABLE_ROWS and self.long_table_mode == "split":  # physical split: reattach the header row to each chunk
            for sub in split_table(e, rows):
                self.table(sub)
            return
        self.gap()
        r = self._centered_refs("table_pos")
        unit = e.data.get("unit")
        sources = e.data.get("sources", [])
        head = 1 if unit else 0  # unit row: a transparent merged cell above the header (hand-editing convention, aligned to the table's right end)
        t = self.doc.add_table(nrow + head + len(sources), ncol, width=table_w, **r)
        if e.text:
            self.attach_caption(t, "TABLE", e.text, "TOP", table_w)
        t.set_column_widths(widths)
        bf = {}
        if unit:
            self._transparent_row(t, 0, ncol, unit, "unit", "RIGHT")
        for k, src in enumerate(sources):  # source rows: transparent merged cells below the table, left-aligned
            self._transparent_row(t, head + nrow + k, ncol, src, "source", "LEFT")
        for ci in range(ncol):
            sides = edge_sides(ci, ci, ncol)
            bf[(ci, True)] = self._border_fill(sides, HEADER_FILL)
            bf[(ci, False)] = self._border_fill(sides)
        cell_cp = str(self.refs("table_cell")["char_pr_id_ref"])
        vspans, hspans = table_spans(rows, hdr, self.merge_mode)
        heights = row_heights(rows, widths, hdr, vspans, hspans)
        total_h = sum(heights)
        for ri, row in enumerate(rows):
            row_h = heights[ri]
            for ci, val in enumerate(row):
                val, bullet_roles = cell_lines(val, self._cell_line_role)
                t.set_cell_text(ri + head, ci, val, split_paragraphs=True)
                cell = t.cell(ri + head, ci)
                cell.set_size(height=row_h)
                is_head = ri < hdr
                role, align = cell_role(aligns[ci], is_head)
                ref = self.refs(role)
                pp = str(ref["para_pr_id_ref"])
                if align:
                    pp = self._aligned_pp(pp, align)
                cp = self._run_style(cell_cp, bold=True) if is_head else str(ref["char_pr_id_ref"])
                for li, pe in enumerate(cell.element.iter(f"{HP}p")):
                    br = bullet_roles[li] if li < len(bullet_roles) and not is_head else None
                    bref = self.refs(br) if br else None
                    pe.set("paraPrIDRef", str(bref["para_pr_id_ref"]) if bref else pp)
                    pe.set("styleIDRef", str(bref["style_id_ref"] if bref else ref["style_id_ref"]))
                    for run in pe.iter(f"{HP}run"):
                        run.set("charPrIDRef", str(bref["char_pr_id_ref"]) if bref else cp)
                cell.element.set("borderFillIDRef", str(bf[(ci, is_head)]))
                cell.element.set("hasMargin", "1")  # needed for cellMargin values to apply
                cm = cell.element.find(f"{HP}cellMargin")
                if cm is not None:
                    for side, mv in zip(("left", "right", "top", "bottom"), CELL_MARGIN):
                        cm.set(side, str(mv))
                if is_head:
                    cell.element.set("header", "1")
        sz = t.element.find(f"{HP}sz")
        if sz is not None:  # if the table height exceeds the sum of row heights, renderers stretch the last row
            sz.set("height", str(total_h + ROW_BASE_H * (head + len(sources))))
        pos = t.element.find(f"{HP}pos")
        if pos is not None:  # a treat-as-character table follows paragraph alignment, so pos alignment is also set to CENTER to remove interpretation conflicts
            pos.set("horzAlign", "CENTER")
        for c, r0, r1 in vspans:  # vertical merge of consecutive equal labels (convention). Negative values are header rows
            a, b_ = head + hdr + r0, head + hdr + r1
            for rr in range(a + 1, b_ + 1):
                t.set_cell_text(rr, c, "")
            t.merge_cells(a, c, b_, c)
        for ri, c0, c1 in hspans:  # horizontal merges from the '<' instruction
            for cc in range(c0 + 1, c1 + 1):
                t.set_cell_text(head + ri, cc, "")
            t.merge_cells(head + ri, c0, head + ri, c1)
        # after merging, reset the outer left and right rules by the actual spans (if a merged cell reaches the table edge, that side's rule is left empty)
        for tc in t.element.iter(f"{HP}tc"):
            addr = tc.find(f"{HP}cellAddr")
            span = tc.find(f"{HP}cellSpan")
            if addr is None or span is None:
                continue
            c0 = int(addr.get("colAddr")); r0 = int(addr.get("rowAddr"))
            c1 = c0 + int(span.get("colSpan")) - 1
            if r0 < head or r0 >= head + nrow:
                continue  # unit rows and source rows are transparent (nrow counts rows including the header)
            is_head = r0 < head + hdr
            tc.set("borderFillIDRef", self._border_fill(edge_sides(c0, c1, ncol), HEADER_FILL if is_head else None))
        if nrow > LONG_TABLE_ROWS and not e.data.get("split_done"):  # long table: treat-as-character off, 한/글 splits by rows at the page boundary
            t.element.set("repeatHeader", "1")
            t.element.set("pageBreak", "CELL")
            if pos is not None:
                pos.set("treatAsChar", "0")
                pos.set("flowWithText", "1")
        t.mark_dirty()
        self.after_block = True

    # ── Figures ────────────────────────────────────────────────────────────
    def image(self, e: Element):
        prepared = prepare_image(e)
        if prepared is None:
            return
        data, fmt, w_mm, h_mm = prepared
        self.gap()
        r = self._centered_refs("fig_pos")
        w_hu, h_hu = round(w_mm * 7200 / 25.4), round(h_mm * 7200 / 25.4)
        sources = e.data.get("sources", [])
        if not FIGURE_IN_TABLE:
            pic = self.doc.add_picture(data, fmt, width_mm=w_mm, height_mm=h_mm, **r)
            if e.text:
                self.attach_caption(pic, "PICTURE", e.text, "BOTTOM", w_hu)
            for src in sources:
                self.para("source", src)
            self.after_block = True
            return
        # put the figure in a transparent one-column table and set the table to the figure numbering type. The inner figure is excluded from the number count
        t = self.doc.add_table(1 + len(sources), 1, width=w_hu, **r)
        item = self.doc.add_image(data, fmt)
        item_id = getattr(item, "item_id", item)  # an object or a string depending on the library version
        cell = t.cell(0, 0)
        pic = cell.paragraphs[0].add_picture(str(item_id), width=w_hu, height=h_hu, char_pr_id_ref=r["char_pr_id_ref"])
        img = next((el for el in pic.element.iter() if el.tag.endswith("}img")), None)
        if img is None or not img.get("binaryItemIDRef"):
            sys.exit(f"ERROR: missing figure binary reference ({e.data['path']}). Check the return value of python-hwpx add_image")
        pic.element.set("numberingType", "NONE")
        for pe in cell.element.iter(f"{HP}p"):
            pe.set("paraPrIDRef", str(r["para_pr_id_ref"]))
            pe.set("styleIDRef", str(r["style_id_ref"]))
        cell.element.set("borderFillIDRef", self._border_fill([]))
        cell.set_size(height=h_hu + 2 * ROW_LINE_H)
        for k, src in enumerate(sources):
            self._transparent_row(t, 1 + k, 1, src, "source", "LEFT")
        t.element.set("numberingType", "PICTURE")
        pos = t.element.find(f"{HP}pos")
        if pos is not None:
            pos.set("horzAlign", "CENTER")
        sz = t.element.find(f"{HP}sz")
        if sz is not None:
            sz.set("height", str(h_hu + 2 * ROW_LINE_H + ROW_BASE_H * len(sources)))
        if e.text:
            self.attach_caption(t, "PICTURE", e.text, FIG_CAPTION_SIDE, w_hu)
        t.mark_dirty()
        self.after_block = True

    @staticmethod
    def _image_mm(data: bytes, fmt: str) -> tuple[float, float]:
        return image_mm(data)

    # ── Run ────────────────────────────────────────────────────────────
    def build(self, elems: list[Element], footnotes: dict, out: str):
        skip = 0
        for idx, e in enumerate(elems):
            if skip:
                skip -= 1
                continue
            if e.kind in ("table", "image"):  # ※, *, and 자료: lines right after are attached as transparent rows at the bottom of the block
                srcs, consumed = [], 0
                for nx in elems[idx + 1:]:
                    if nx.kind in ("note", "source"):
                        srcs.append(nx.text)
                    elif nx.kind == "unit" and e.kind == "table" and not e.data.get("unit"):
                        e.data["unit"] = nx.text  # a unit line written after the table is also pulled up to the unit row above the table
                    else:
                        break
                    consumed += 1
                e.data["sources"] = srcs
                skip = consumed
            if e.kind == "heading":
                self.heading(e.level, e.text)
            elif e.kind in ("square", "circle", "dash", "dot", "note"):
                self.last_para = self.para(e.kind, e.text, footnotes)
            elif e.kind == "unit":
                self.para("unit", e.text)
            elif e.kind == "source":  # a source line not attached to a table or figure moves into a footnote of the preceding paragraph
                if self.last_para is not None:
                    self.last_para.add_footnote(e.text)
                else:
                    self.para("note", e.text, footnotes)
            elif e.kind == "table":
                self.table(e)
            elif e.kind == "image":
                self.image(e)
            elif e.kind == "pagebreak":
                self.pending_break = True
        if self.toc or self.cover:
            self._front_matter(self.cover)
        self.set_page_number(self.page_number)
        self.set_header(self.header_spec)
        report = self.doc.validate()
        issues = getattr(report, "issues", [])
        if issues:
            print(f"WARN: {len(issues)} schema validation issues")
            for it in list(issues)[:5]:
                print("   ", it)
        self.doc.save_to_path(os.path.abspath(out))
        self.doc.close()


def repair_manifest(path: str) -> int:
    """Compare opf:item of content.hpf with binItem of header.xml and the BinData files, and register what is missing.
    Why: depending on the library version, only the first item gets registered in the manifest when inserting several images,
    and then 한/글 cannot open the other pictures. If everything matches, the file is left untouched."""
    MEDIA = {"jpg": "image/jpeg", "jpeg": "image/jpeg", "png": "image/png", "gif": "image/gif",
             "bmp": "image/bmp", "tif": "image/tiff", "tiff": "image/tiff", "svg": "image/svg+xml"}
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        data = {n: z.read(n) for n in names}
    hpf = data["Contents/content.hpf"].decode("utf-8")
    header = data.get("Contents/header.xml", b"").decode("utf-8")
    have = set(re.findall(r'<opf:item[^>]*href="BinData/([^"]+)"', hpf))
    wanted: list[str] = []
    for f in re.findall(r'<hh:binItem[^>]*BinData="([^"]+)"', header):  # priority 1: header.xml binItem
        if f not in wanted:
            wanted.append(f)
    for n in names:  # priority 2: actual BinData files (covers files not in binItem)
        if n.startswith("BinData/") and n.split("/", 1)[1] not in wanted:
            wanted.append(n.split("/", 1)[1])
    missing = [f for f in wanted if f not in have and f"BinData/{f}" in names]
    if not missing:
        return 0
    items = "".join(
        f'<opf:item id="{os.path.splitext(f)[0]}" href="BinData/{f}" '
        f'media-type="{MEDIA.get(os.path.splitext(f)[1].lstrip(".").lower(), "application/octet-stream")}" '
        f'isEmbeded="1"/>' for f in missing)
    hpf = hpf.replace("</opf:manifest>", items + "</opf:manifest>", 1)
    data["Contents/content.hpf"] = hpf.encode("utf-8")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr(zipfile.ZipInfo("mimetype"), data["mimetype"], compress_type=zipfile.ZIP_STORED)
        for n in names:
            if n != "mimetype":
                z.writestr(n, data[n])
    return len(missing)


def cmd_build(a):
    profile = PROFILES.get(a.profile) if a.profile in PROFILES else json.load(open(a.profile, encoding="utf-8"))
    skeleton = a.skeleton or os.path.join(ASSET_DIR, f"{a.profile}.hwpx")
    if not os.path.exists(skeleton):
        sys.exit(f"ERROR: skeleton not found {skeleton}")
    elems, footnotes = parse_markdown(a.input)
    for msg in lint(elems):
        print("LINT:", msg)
    out = a.output or os.path.splitext(a.input)[0] + ".hwpx"
    def make(numbers=None):
        b = Builder(skeleton, profile, page_break=a.page_break)
        b.numbered_square = a.numbered_square
        b.toc = a.toc
        b.page_number = a.page_number
        b.toc_depth = a.toc_depth
        b.cover = dict(kv.split("=", 1) for kv in a.cover.split(";") if "=" in kv) if a.cover else None
        b.header_spec = a.header
        b.merge_mode = a.merge
        b.long_table_mode = a.long_table
        b.toc_numbers = numbers or {}
        b.build(elems, footnotes, out)
        return b
    b = make()
    fixed = repair_manifest(out)
    if fixed:
        print(f"WARN: {fixed} images not registered in the manifest were corrected automatically (checking the python-hwpx version is recommended)")
    if a.toc and a.page_number:  # pass 2: find the page positions of titles by rendering and fill them into the TOC (approximate, finalized in 한/글)
        numbers = toc_page_numbers(out, [t for _, t in b.headings], a.rhwp)
        if numbers:
            make(numbers)
            repair_manifest(out)
            listed = [t for lv, t in b.headings if lv <= b.toc_depth]  # Why: titles are collected down to subsections, the TOC down to toc_depth
            print(f"TOC page numbers filled {sum(t in numbers for t in listed)}/{len(listed)} (approximate renderer)")
        else:
            print("WARN: TOC page numbers not filled because rhwp or pdftotext is missing")
    print(f"OK: {out} (elements {len(elems)}, footnotes {len(footnotes)})")


def toc_page_numbers(hwpx_path: str, headings: list[str], rhwp_opt: str | None) -> dict[str, int]:
    if not shutil.which("pdftotext") or not (rhwp_opt or find_rhwp() or find_soffice()):
        return {}
    pdf = os.path.splitext(hwpx_path)[0] + ".toc.pdf"
    engine = resolve_engine("auto", rhwp_opt)
    work = render_copy(hwpx_path, strip_newnum=(engine == "soffice"))
    try:
        to_pdf(work, pdf, engine, rhwp_opt, text_as_paths=False)
    except SystemExit:
        os.remove(work)
        return {}
    os.remove(work)
    txt = subprocess.run(["pdftotext", "-layout", pdf, "-"], capture_output=True, text=True).stdout
    os.remove(pdf)
    return match_toc_pages(txt.split("\f"), headings)


def _toc_norm(t: str) -> str:
    return re.sub(r"[\s.]+", "", t)  # bands put the number and title in different cells, so periods are ignored too


CHAPTER_TOKEN_RE = re.compile(r"제\d+장|[ⅠⅡⅢⅣⅤⅥⅦⅧⅨⅩ]")   # chapter numbers that intrude between titles on divider pages and bands
TOC_TOP_LINES = 6   # number of top-of-page lines searched first for chapter titles. Chapters start on a new page, so the title sits at the top
TOC_TITLE_WRAP_LINES = 3   # maximum lines one chapter title spans. A long divider title wraps to two lines with the chapter number line between them


def toc_titles() -> set[str]:
    """Title lines of TOC pages (full match after removing spaces). Follows the labels of the TOC, list of tables, and list of figures"""
    return {_toc_norm(label(k)) for k in ("toc.title", "toc.tables", "toc.figures")} - {""}


def chapter_token_re() -> re.Pattern:
    """Chapter number notation. Adds the number part (before {t}) of formats.chapter and formats.chapter_label overrides to the default notation"""
    pats = [CHAPTER_TOKEN_RE.pattern]
    for k in ("chapter", "chapter_label"):
        head = _toc_norm((label(f"formats.{k}") or "").split("{t}")[0])
        if "{n}" in head or "{r}" in head:
            pats.append(re.escape(head).replace(r"\{n\}", r"\d+").replace(r"\{r\}", "[ⅠⅡⅢⅣⅤⅥⅦⅧⅨⅩ]"))
    return re.compile("|".join(pats))


def match_toc_pages(pages: list[str], headings: list[str], extras: list[tuple[str, str]] | None = None,
                    top_titles: set[str] | None = None) -> dict[str, int]:
    """Find the body page number (first body page = 1) of TOC entries in the per-page text of the rendered PDF
    headings: title text. extras: list of (TOC entry text, text to search). Captions are searched without the number, by the closing bracket and title
    top_titles: titles searched first at the top of pages (chapter titles). Picks the chapter start page even if the same text appears earlier in a body table
    Why: approximate renderers draw automatic caption numbers blank. The title after the closing bracket is used as the search text so it is found without the number.
    Long chapter titles on divider pages fold into two lines and pdftotext puts the chapter number between them, so the chapter number is removed before searching"""
    norm = _toc_norm
    chapter_re, titles = chapter_token_re(), toc_titles()

    def clean(t: str) -> str:  # renderers insert spaces between characters as in "제 1장", so spaces are removed and then the chapter number
        return chapter_re.sub("", norm(t))

    def is_title_page(page: str) -> bool:  # the word "목차" inside body sentences is not a title line, so it is excluded
        return any(norm(ln) in titles for ln in page.splitlines())
    keys = [norm(t) for t in headings + [e for e, _ in extras or []]]

    def entry_index(line: str) -> int | None:
        return next((i for i, k in enumerate(keys) if k and line.startswith(k)), None)

    def page_indexes(page: str) -> list[int | None]:
        return [entry_index(ln) for ln in (norm(x) for x in page.splitlines()) if ln]
    # TOC span: from the first TOC title page through the consecutive title pages and continuation pages
    # a continuation page is one where every line is a TOC entry and the first line is an entry after the last entry of the previous page
    # Why the order condition: divider pages (one chapter title line) are also pages with only TOC entry lines, so order distinguishes them
    first = next((i for i, page in enumerate(pages) if is_title_page(page)), None)
    body_start = 0
    if first is not None:
        i, last = first, -1
        while i < len(pages):
            idx = page_indexes(pages[i])
            if is_title_page(pages[i]):
                last = max([x for x in idx if x is not None] + [last])
            elif idx and None not in idx and idx[0] > last:
                last = max(idx)
            else:
                break
            i += 1
        body_start = i
    top = top_titles or set()

    def top_lines(page: str) -> list[str]:
        return [x for x in page.splitlines() if x.strip()][:TOC_TOP_LINES]

    def title_alone_at_top(page: str, key: str) -> bool:
        # Why: band and divider chapter titles occupy their line alone. A table cell line is extracted together with the other
        # cells of the same row. A substring match alone mistakes a chapter-title phrase in a top-of-page table cell for the chapter start (reproduced in trial 6, R tracks)
        lines = top_lines(page)
        return any(clean("".join(lines[i:j])) == key
                   for i in range(len(lines)) for j in range(i + 1, min(i + TOC_TITLE_WRAP_LINES, len(lines)) + 1))
    found: dict[str, int] = {}
    for group in ([(h, h) for h in headings], extras or []):
        pos = body_start
        for entry, needle in group:
            key = clean(needle)
            if not key:
                continue
            hit = None
            if entry in top:
                hit = next((i for i in range(pos, len(pages)) if title_alone_at_top(pages[i], key)), None)
            if entry in top and hit is None:  # a page whose title line was extracted mixed with other text falls back to the previous substring match
                hit = next((i for i in range(pos, len(pages)) if key in clean("\n".join(top_lines(pages[i])))), None)
            if hit is None:
                hit = next((i for i in range(pos, len(pages)) if key in clean(pages[i])), None)
            if hit is not None:
                found[entry] = hit - body_start + 1
                pos = hit
    return found


# ═══════════════════════════════════════════════════════════════════════
# 4. skeleton: extract an empty template keeping only the style sheet and body section settings from a reference hwpx
# ═══════════════════════════════════════════════════════════════════════
def cmd_skeleton(a):
    src = zipfile.ZipFile(a.input)
    secs = sorted(n for n in src.namelist() if re.match(r"Contents/section\d+\.xml", n))
    sec_name = f"Contents/section{a.section}.xml" if a.section is not None else secs[-1]
    s = src.read(sec_name).decode("utf-8")
    first = s[: s.find("</hp:p>") + 7]
    first = re.sub(r"<hp:t\b[^>]*>.*?</hp:t>", "<hp:t/>", first, flags=re.S)
    first = re.sub(r"<hp:t\b[^>]*/>", "<hp:t/>", first)
    first = re.sub(r'(<hp:p\b[^>]*?)paraPrIDRef="\d+"', r'\1paraPrIDRef="0"', first, count=1)
    first = re.sub(r'(<hp:p\b[^>]*?)pageBreak="1"', r'\1pageBreak="0"', first, count=1)
    hpf = src.read("Contents/content.hpf").decode("utf-8")
    hpf = re.sub(r'<opf:item id="section[1-9]\d*"[^>]*/>', "", hpf)
    hpf = re.sub(r'<opf:itemref idref="section[1-9]\d*"[^>]*/>', "", hpf)
    hpf = re.sub(r'<opf:item [^>]*href="BinData/[^"]*"[^>]*/>', "", hpf)
    hpf = re.sub(r'(<opf:meta name="(creator|lastsaveby|date|CreatedDate|ModifiedDate)"[^>]*>)[^<]*(</opf:meta>)',
                 r"\1\3", hpf)  # remove author and date metadata
    with zipfile.ZipFile(a.output, "w", zipfile.ZIP_DEFLATED) as z:
        for it in src.infolist():
            n = it.filename
            if n.startswith("BinData/") or re.match(r"Contents/section[1-9]", n) or n == "Preview/PrvImage.png":
                continue
            data = src.read(n)
            if n == "Contents/section0.xml":
                data = (first + "</hs:sec>").encode("utf-8")
            elif n == "Preview/PrvText.txt":
                data = b""  # remove the original preview text
            elif n == "Contents/content.hpf":
                data = hpf.encode("utf-8")
            elif n == "Contents/header.xml":
                data = re.sub(rb'secCnt="\d+"', b'secCnt="1"', data)
            if n == "mimetype":
                z.writestr(zipfile.ZipInfo("mimetype"), data, compress_type=zipfile.ZIP_STORED)
            else:
                z.writestr(it, data)
    from hwpx import HwpxDocument
    doc = HwpxDocument.open(os.path.abspath(a.output))
    names = [s.name for s in doc.styles.values()]
    bullets = {bid: b.char for bid, b in doc.styles.bullets.items()} if hasattr(doc.styles, "bullets") else {}
    doc.close()
    profile = {"numbering": "manual", "chapter_format": "제{n}장 {t}"}
    guess = {"chapter": ("장 제목", "장"), "section": ("절 제목", "1.1", "절"), "subsection": ("1.1.1",), "ganada": ("가.", "가. 제목"),
             "square": ("□", "□ 내용"), "circle": ("❍ 내용1", "◦ 내용", "○"), "dash": ("- 내용2", "- 내용"),
             "dot": ("․ 내용3", "․ 내용"), "note": ("※ 내용4", "* 내용"), "table_pos": ("표위치", "표 위치"),
             "table_cell": ("표내용", "표 내용"), "table_cell_left": ("표내용(양쪽정렬)", "표내용", "표 내용"),
             "table_title": ("표 제목",), "unit": ("표,그림의 단위", "표 제목"), "source": ("표.그림의 자료및주", "출처", "* 내용"),
             "fig_pos": ("그림위치",), "fig_title": ("그림 제목",), "body": ("바탕글",)}
    for role, cands in guess.items():
        profile[role] = next((c for c in cands if c in names), None)
    prof_path = os.path.splitext(a.output)[0] + ".profile.json"
    json.dump(profile, open(prof_path, "w", encoding="utf-8"), ensure_ascii=False, indent=2)
    print(f"OK: {a.output} ({os.path.getsize(a.output)} bytes, section {sec_name})")
    print("Styles:", names)
    print("Bullets:", bullets)
    print(f"Draft profile: {prof_path}: set style names directly for None entries, then pass it with build --profile")


# ═══════════════════════════════════════════════════════════════════════
# 5. export / fill / render / check
# ═══════════════════════════════════════════════════════════════════════
ROLE_MD = {"chapter": "# ", "section": "## ", "subsection": "### ", "ganada": "#### ", "square": "□ ",
           "circle": "- ", "dash": "    + ", "dot": "        + ", "note": "※ ", "source": "> ",
           "table_title": "<표> ", "fig_title": "[그림] "}  # the source role uses the "> " prefix
LEADING_MARK_RE = re.compile(r"^(제\s*\d+\s*장|\d+(?:\.\d+)*\.?|[가-힣]\.|[□■○❍◦ㅇ\-ㆍ·※*]|<표[^>]*>|\[그림[^\]]*\])\s*")


def guess_role(text: str) -> str | None:
    """Role estimation for documents without style names (symbols typed by hand)."""
    for pat, role in ((r"^제\s*\d+\s*장\b", "chapter"), (r"^\d+\.\d+\.\d+\s", "subsection"), (r"^\d+\.\d+\s", "section"),
                      (r"^[가-힣]\.\s", "ganada"), (r"^[□■]", "square"), (r"^[○❍◦ㅇ]", "circle"),
                      (r"^-\s", "dash"), (r"^[ㆍ·•]", "dot"), (r"^[※*]", "note"), (r"^(자료|출처)\s*[:：]", "source"),
                      (r"^<표", "table_title"), (r"^\[그림", "fig_title")):
        if re.match(pat, text):
            return role
    return None


SYMBOL_ROLE = {"□": "square", "■": "square", "○": "circle", "❍": "circle", "◦": "circle", "ㅇ": "circle",
               "-": "dash", "\u00ad": "dash", "ㆍ": "dot", "·": "dot", "․": "dot", "※": "note", "*": "note"}
TEXT_ROLE_RE = [(re.compile(r"^제\s*\d+\s*장\s"), "chapter"), (re.compile(r"^\d+\.\d+\.\d+\s"), "subsection"),
                (re.compile(r"^\d+\.\d+\s"), "section"), (re.compile(r"^[가-힣]\.\s"), "ganada"),
                (re.compile(r"^[<＜]\s*표"), "table_title"), (re.compile(r"^\[\s*그림"), "fig_title"),
                (re.compile(r"^(자료|출처)\s*[:：]"), "source")]


def _guess_role(text: str, bullet_char: str | None) -> str | None:
    """For documents whose style names differ from the norms (reference reports): estimate the level by bullet character → leading symbol → number pattern, in that order."""
    if bullet_char and bullet_char in SYMBOL_ROLE:
        return SYMBOL_ROLE[bullet_char]
    if text and text[0] in SYMBOL_ROLE and (len(text) == 1 or text[1] == " "):
        return SYMBOL_ROLE[text[0]]
    for rx, role in TEXT_ROLE_RE:
        if rx.match(text):
            return role
    return None


def export_hierarchy(doc, profile: dict, image_dir: str | None) -> str:
    """After reverse-mapping style names, estimate the rest by symbol and number patterns and restore Markdown following the hierarchy norms (for round trips of re-edit and build)."""
    rev = {v: k for k, v in profile.items() if isinstance(v, str) and k in ROLE_MD}
    names = {sid: st.name for sid, st in doc.styles.items()}
    HH = "{http://www.hancom.co.kr/hwpml/2011/head}"
    hx = doc.headers[0].element
    bullet_char = {b.get("id"): b.get("char") for b in hx.iter(f"{HH}bullet")}
    pp_bullet = {}
    for pp in hx.iter(f"{HH}paraPr"):
        hd = pp.find(f"{HH}heading")
        if hd is not None and hd.get("type") == "BULLET":
            pp_bullet[pp.get("id")] = bullet_char.get(hd.get("idRef"))
    images = {it.item_id: it for it in doc.list_images()} if hasattr(doc, "list_images") else {}
    out, img_n = [], 0
    if image_dir:
        os.makedirs(image_dir, exist_ok=True)
    for sec in doc.sections:
        for p in sec.paragraphs:
            pics = list(p.element.iter(f"{HP}pic"))
            if pics and image_dir:
                for pic in pics:
                    ref = next((x for x in pic.iter() if x.tag.endswith("}img")), None)
                    item = images.get(ref.get("binaryItemIDRef")) if ref is not None else None
                    if item is None:
                        continue
                    img_n += 1
                    fn = f"img{img_n:03d}.{item.format}"
                    open(os.path.join(image_dir, fn), "wb").write(doc.package.read(item.href))
                    out.append(f"![]({os.path.basename(image_dir)}/{fn})")
            if p.tables:
                for t in p.tables:
                    grid, seen = [], set()
                    for row in t.get_cell_map():  # merged cells hold content only at the first position, the rest are empty
                        line = []
                        for pos in row:
                            key = id(pos.cell.element)
                            line.append("" if key in seen else pos.cell.text.replace("\n", "<br>"))
                            seen.add(key)
                        grid.append(line)
                    cap = t.caption.text.strip() if getattr(t, "caption", None) else ""
                    if cap:
                        out.append("<표> " + LEADING_MARK_RE.sub("", cap))
                    out.append("| " + " | ".join(grid[0]) + " |")
                    out.append("|" + "---|" * len(grid[0]))
                    out.extend("| " + " | ".join(r) + " |" for r in grid[1:])
                    out.append("")
                continue
            text = p.text.strip()
            if not text:
                continue
            role = rev.get(names.get(str(p.style_id_ref), ""), None) or guess_role(text)
            if role:
                out.append(ROLE_MD[role] + LEADING_MARK_RE.sub("", text))
            else:
                out.append(text)
    return "\n".join(out) + "\n"


def cmd_export(a):
    from hwpx import HwpxDocument
    doc = HwpxDocument.open(os.path.abspath(a.input))
    if a.profile:
        profile = PROFILES.get(a.profile) if a.profile in PROFILES else json.load(open(a.profile, encoding="utf-8"))
        md = export_hierarchy(doc, profile, a.image_dir)
        out = a.output or os.path.splitext(a.input)[0] + ".md"
        open(out, "w", encoding="utf-8").write(md)
        print(f"OK: {out} ({len(md)} characters, levels restored)")
        return
    kw = {"image_dir": a.image_dir, "image_ref_prefix": os.path.basename(a.image_dir) + "/"} if a.image_dir else {}
    md = doc.export_rich_markdown(detect_headings=True, **kw)
    out = a.output or os.path.splitext(a.input)[0] + ".md"
    open(out, "w", encoding="utf-8").write(md)
    print(f"OK: {out} ({len(md)} characters)")


def cmd_fill(a):
    from hwpx import HwpxDocument
    spec = json.load(open(a.mapping, encoding="utf-8"))
    doc = HwpxDocument.open(os.path.abspath(a.input))
    total = 0
    for old, new in spec.get("replace", {}).items():
        total += doc.replace_text_in_runs(old, new) or 0
    if spec.get("cells"):  # fill table cells by "label > right" paths
        res = doc.fill_by_path(spec["cells"])
        print(f"Table cells: applied {res['applied_count']} failed {res['failed_count']}")
    doc.save_to_path(os.path.abspath(a.output or a.input))
    print(f"OK: text substitutions {total} → {a.output or a.input}")


def find_rhwp() -> str | None:
    for cand in (os.path.join(SCRIPT_DIR, "rhwp", "rhwp"), os.path.join(SCRIPT_DIR, "rhwp", "rhwp.exe"),
                 shutil.which("rhwp"), os.path.expanduser("~/.local/bin/rhwp")):
        if cand and os.path.isfile(cand) and os.access(cand, os.X_OK):
            return cand
    return None


def rhwp_supports(rhwp: str, command: str, flag: str | None = None) -> bool:
    """Check that commands and flags exist from the help output. It is an immature project, so the CLI may change between versions."""
    res = subprocess.run([rhwp, command, "--help"], capture_output=True, text=True)
    text = res.stdout + res.stderr
    if res.returncode != 0 and "Usage" not in text and "사용법" not in text:
        return False
    return flag is None or flag in text


def writable_dir(path: str) -> str:
    """Directory for outputs. If next to the original is read-only (/mnt/skills, uploads), use the current working directory."""
    d = os.path.dirname(os.path.abspath(path))
    return d if os.access(d, os.W_OK) else os.getcwd()


def render_copy(src: str, strip_newnum: bool = False) -> str:
    """Review-only copy. Holds only substitutions that work around approximate renderer limits. The delivered file is unchanged.
    - soft hyphen bullets → hyphens (approximate renderers do not draw them)
    - strip_newnum: remove newNum controls in the middle of the body (H2Orestart only. Conversion fails on newNum outside the first paragraph)"""
    work = os.path.join(writable_dir(src), "_render_copy.hwpx")  # ASCII file name (avoids an H2Orestart locale problem)
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(work, "w", zipfile.ZIP_DEFLATED) as zout:
        for it in zin.infolist():
            d = zin.read(it.filename)
            if it.filename.startswith("Contents/section") and it.filename.endswith(".xml"):
                t = d.decode("utf-8").replace("\u00ad", "-")
                if strip_newnum:
                    t = re.sub(r'<hp:run charPrIDRef="\d+"><hp:ctrl><hp:newNum[^>]*/></hp:ctrl></hp:run>', "", t)
                d = t.encode("utf-8")
            if it.filename == "mimetype":
                zout.writestr(zipfile.ZipInfo("mimetype"), d, compress_type=zipfile.ZIP_STORED)
            else:
                zout.writestr(it.filename, d)
    return work


def find_soffice() -> str | None:
    return shutil.which("soffice") or shutil.which("libreoffice")


def resolve_engine(engine: str, rhwp_opt: str | None) -> str:
    rhwp = rhwp_opt or find_rhwp()
    if engine == "auto":
        return "rhwp" if rhwp else ("soffice" if find_soffice() else "")
    return engine


def to_pdf(src: str, pdf: str, engine: str, rhwp_opt: str | None, text_as_paths: bool = True) -> str:
    """hwpx → PDF. engine: rhwp | soffice (auto is settled beforehand by resolve_engine). Returns the name of the engine used.
    text_as_paths=False: characters are not frozen into paths, so the PDF keeps a text layer.
    Why: TOC page numbers are computed by reading the page positions of titles with pdftotext. A path conversion yields 0 characters."""
    rhwp = rhwp_opt or find_rhwp()
    soffice = find_soffice()
    engine = resolve_engine(engine, rhwp_opt)
    if engine == "rhwp":
        if not rhwp:
            sys.exit("ERROR: rhwp not available: obtain it with the procedure in the SKILL.md Preparation section, then pass --rhwp path or place it in scripts/rhwp/")
        if not rhwp_supports(rhwp, "export-pdf"):
            sys.exit("ERROR: this rhwp version has no export-pdf command: find the PDF output command in `rhwp --help` and run it manually")
        cmd = [rhwp, "export-pdf", src, "-o", pdf]
        if text_as_paths and rhwp_supports(rhwp, "export-pdf", "--text-as-paths"):
            cmd.append("--text-as-paths")
        res = subprocess.run(cmd, capture_output=True, text=True)
        if res.returncode != 0 or not os.path.exists(pdf):
            sys.exit(f"ERROR: rhwp conversion failed (may signal a document defect): {(res.stderr or res.stdout)[-400:]}")
        return "rhwp"
    if engine == "soffice":
        if not soffice:
            sys.exit("ERROR: LibreOffice not available: see the fallback installation procedure in the SKILL.md Preparation section")
        outdir = os.path.join(writable_dir(src), "_lo_out")
        shutil.rmtree(outdir, ignore_errors=True)
        env = {**os.environ, "LANG": os.environ.get("LANG") or "C.UTF-8"}
        # a separate profile per call: a fixed profile causes lock conflicts with other typesetting sessions running at the same time
        profile = tempfile.mkdtemp(prefix="lo-profile-")
        cmd = [soffice, f"-env:UserInstallation={pathlib.Path(profile).as_uri()}", "--headless",
               "--convert-to", "pdf:writer_pdf_Export", "--outdir", outdir, src]
        try:
            res = subprocess.run(cmd, capture_output=True, text=True, env=env, timeout=600)
        finally:
            shutil.rmtree(profile, ignore_errors=True)
        made = os.path.join(outdir, os.path.splitext(os.path.basename(src))[0] + ".pdf")
        is_hwpx = src.lower().endswith(".hwpx")
        if not os.path.exists(made):
            sys.exit(("ERROR: LibreOffice (H2Orestart) conversion failed: the H2Orestart extension is not installed, or it signals a document defect (element order, reference ID): "
                      if is_hwpx else "ERROR: LibreOffice conversion failed (may signal a document defect): ") + (res.stdout + res.stderr)[-300:])
        shutil.move(made, pdf)
        shutil.rmtree(outdir, ignore_errors=True)
        return "soffice(H2Orestart)" if is_hwpx else "soffice"
    sys.exit("ERROR: no renderer: obtain rhwp or LibreOffice+H2Orestart and rerun (SKILL.md Preparation section)")


PAGE_PNG_RE = re.compile(r"p-\d+\.png")   # pdftoppm page PNG name (prefix p). What render deletes and counts


def cmd_render(a):
    src = a.input
    is_docx = src.lower().endswith(".docx")
    engine = "soffice" if is_docx else resolve_engine(a.engine, a.rhwp)  # docx uses only LibreOffice
    work = src if a.keep_soft_hyphen or is_docx else render_copy(src, strip_newnum=(engine == "soffice"))
    stem = os.path.join(writable_dir(src), os.path.splitext(os.path.basename(src))[0])
    outdir = a.outdir or stem + "_pages"
    os.makedirs(outdir, exist_ok=True)
    # Why: delete only page PNGs from a previous render. Emptying the whole folder deletes the source or other files when --outdir points to their folder
    for f in os.listdir(outdir):
        if PAGE_PNG_RE.fullmatch(f):
            os.remove(os.path.join(outdir, f))
    # Why: with --outdir, the PDF also goes to that folder, so an engine cross-check render does not overwrite the existing PDF
    pdf = os.path.join(outdir, os.path.basename(stem) + ".pdf") if a.outdir else stem + ".pdf"
    used = to_pdf(work, pdf, engine, a.rhwp)
    if work != src:
        os.remove(work)
    if shutil.which("pdftoppm"):
        subprocess.run(["pdftoppm", "-png", "-r", str(a.dpi), pdf, os.path.join(outdir, "p")], check=True)
        n_png = len([f for f in os.listdir(outdir) if PAGE_PNG_RE.fullmatch(f)])
        print(f"OK[{used}]: {pdf} → {outdir}/ ({n_png} page(s))")
    else:
        print(f"OK[{used}]: {pdf} (PNG skipped because pdftoppm is missing)")


def cmd_check(a):
    from hwpx import HwpxDocument
    doc = HwpxDocument.open(os.path.abspath(a.input))
    report = doc.validate()
    issues = list(getattr(report, "issues", []))
    with zipfile.ZipFile(a.input) as z:
        hpf = z.read("Contents/content.hpf").decode("utf-8")
        have = set(re.findall(r'<opf:item[^>]*href="BinData/([^"]+)"', hpf))
        binitems = set(re.findall(r'<hh:binItem[^>]*BinData="([^"]+)"', z.read("Contents/header.xml").decode("utf-8")))
        files = {n.split("/", 1)[1] for n in z.namelist() if n.startswith("BinData/")}
        sec = "".join(z.read(n).decode("utf-8") for n in z.namelist() if n.startswith("Contents/section"))
    refs = set(re.findall(r'binaryItemIDRef="([^"]+)"', sec))
    stems = {os.path.splitext(f)[0] for f in files}
    if files - have:
        issues.append(f"images not registered in the manifest: {sorted(files - have)}")
    if binitems - files:
        issues.append(f"no BinData file for binItem: {sorted(binitems - files)}")
    if refs - stems:
        issues.append(f"no BinData for body references: {sorted(refs - stems)}")
    print(f"Images: BinData {len(files)}, binItem {len(binitems)}, manifest {len(have)}, body references {len(refs)}")
    print(f"Schema validation: {len(issues)} issues")
    for it in issues[:10]:
        print("   ", it)
    doc.close()
    rhwp = find_rhwp()
    if rhwp and a.expect and not rhwp_supports(rhwp, "verify"):
        print("rhwp verify skipped because this version lacks the command")
    elif rhwp and a.expect:
        import shlex
        cmd = [rhwp, "verify", a.input, "--json"] + shlex.split(a.expect)
        res = subprocess.run(cmd, capture_output=True, text=True)
        try:
            d = json.loads(res.stdout)
            print("rhwp verify:", d.get("verdict"), [(e["kind"], e.get("actual"), e["pass"]) for e in d["expectations"]])
        except json.JSONDecodeError:
            print("rhwp verify output could not be parsed:", res.stdout[:200], res.stderr[:200])
    sys.exit(1 if issues else 0)


# ═══════════════════════════════════════════════════════════════════════
# 6. Input adapters: DOCX and XLSX into Markdown manuscripts following the hierarchy norms (without retyping the body)
# ═══════════════════════════════════════════════════════════════════════
def _docx_para_key(p) -> str:
    """Level key of a paragraph: style name | left indent | automatic numbering | leading symbol."""
    pf = p.paragraph_format
    indent = int((pf.left_indent or 0) / 12700 * 10) if pf.left_indent is not None else 0
    numbered = p._p.pPr is not None and p._p.pPr.numPr is not None
    lead = p.text.lstrip()[:1] if p.text.strip() else ""
    lead = lead if lead in "□■○◦❍ㅇ-ㆍ·※*" else ""
    return f"{p.style.name}|{indent}|{'num' if numbered else ''}|{lead}"


def cmd_docx2md(a):
    import docx
    d = docx.Document(a.input)
    if not a.map:  # step 1: print the profile → settle the mapping with the user
        keys: dict[str, list[str]] = {}
        for p in d.paragraphs:
            if p.text.strip():
                keys.setdefault(_docx_para_key(p), []).append(p.text.strip()[:40])
        print("key = style name|indent(pt*10)|auto number|leading symbol  →  count, sample. Mapping JSON example: {\"본문|0||□\": \"square\"}")
        for k, v in sorted(keys.items(), key=lambda kv: -len(kv[1])):
            print(f"  {k!r}: {len(v)}  e.g.) {v[0]}")
        print("Roles: chapter section subsection ganada square circle dash dot note source skip")
        return
    mp = json.load(open(a.map, encoding="utf-8"))
    md_of = {"chapter": "# ", "section": "## ", "subsection": "### ", "ganada": "#### ", "square": "□ ",
             "circle": "- ", "dash": "    + ", "dot": "        + ", "note": "※ ", "source": "> "}
    out = []
    body = d.element.body
    pi, ti = 0, 0
    for child in body.iterchildren():
        tag = child.tag.split("}")[1]
        if tag == "p":
            p = d.paragraphs[pi]
            pi += 1
            if not p.text.strip():
                continue
            role = mp.get(_docx_para_key(p), "circle")
            if role == "skip":
                continue
            out.append(md_of[role] + LEADING_MARK_RE.sub("", p.text.strip()))
        elif tag == "tbl":
            t = d.tables[ti]
            ti += 1
            rows = [[c.text.strip().replace("\n", "<br>") for c in r.cells] for r in t.rows]
            out.append("| " + " | ".join(rows[0]) + " |")
            out.append("|" + "---|" * len(rows[0]))
            out.extend("| " + " | ".join(r) + " |" for r in rows[1:])
            out.append("")
    o = a.output or os.path.splitext(a.input)[0] + ".md"
    open(o, "w", encoding="utf-8").write("\n".join(out) + "\n")
    print(f"OK: {o} (paragraphs {pi}, tables {ti}): add table caption <표> lines and figures in the manuscript")


def cmd_xlsx2md(a):
    import openpyxl
    wb = openpyxl.load_workbook(a.input, data_only=True)
    out = []
    for ws in wb.worksheets:
        rows = [[("" if v is None else (f"{v:,}" if isinstance(v, (int, float)) and not isinstance(v, bool) else str(v)).strip())
                 for v in r] for r in ws.iter_rows(values_only=True)]
        rows = [r for r in rows if any(r)]
        if not rows:
            continue
        ncol = max(len(r) for r in rows)
        if ncol > 6 and not a.allow_wide:
            print(f"WARN: sheet '{ws.title}' has {ncol} columns: ask the user about transposing or splitting the table (--allow-wide proceeds as is)")
        out.append(f"<표> {ws.title}")
        out.append("| " + " | ".join(rows[0]) + " |")
        out.append("|" + "---|" * len(rows[0]))
        out.extend("| " + " | ".join(r + [""] * (ncol - len(r))) + " |" for r in rows[1:])
        out.append("")
    o = a.output or os.path.splitext(a.input)[0] + ".md"
    open(o, "w", encoding="utf-8").write("\n".join(out) + "\n")
    print(f"OK: {o} (sheets {len(wb.worksheets)}): add caption titles, (단위: ) lines, and 자료: lines in the manuscript")


# ═══════════════════════════════════════════════════════════════════════
def main():
    ap = argparse.ArgumentParser(description="hwpx_report CLI", epilog=CLI_NOTICE,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)

    b = sub.add_parser("build", help="manuscript.md → document.hwpx")
    b.add_argument("input")
    b.add_argument("output", nargs="?")
    b.add_argument("--profile", default="chapter", help="chapter | flat | profile JSON path")
    b.add_argument("--skeleton", help="skeleton hwpx (default assets/<profile>.hwpx)")
    b.add_argument("--page-break", default="chapter", choices=["chapter", "section", "none"],
                   help="new page unit: chapter=every # chapter (default), section=every # chapter and ## section, none=continuous (user's choice)")
    b.add_argument("--long-table", default="engine", choices=["engine", "split"],
                   help="tables over 25 rows: engine=한/글 engine split + header row repetition, split=physical split into chunk tables (user's choice)")
    b.add_argument("--toc", action="store_true", help="insert a static TOC of chapters and sections before the document (default OFF, on user request)")
    b.add_argument("--page-number", action="store_true", help="page number '- N -' centered in the footer (default OFF, prevents conflicts when combining)")
    b.add_argument("--toc-depth", type=int, default=2, choices=[1, 2, 3], help="TOC level depth: 1=chapter, 2=section (default), 3=subsection (user's choice)")
    b.add_argument("--cover", help='cover (default OFF): "title=title;subtitle=subtitle;date=YYYY. M.;org=organization"')
    b.add_argument("--header", help='header (default OFF): "text;align=LEFT|CENTER|RIGHT"')
    b.add_argument("--merge", default="auto", choices=["auto", "explicit", "none"],
                   help="cell merges: auto=consecutive equal values in leading label columns + '^' (vertical) and '<' (horizontal) instructions (default), explicit=instructions only, none=no merges")
    b.add_argument("--rhwp", help="rhwp path used for computing TOC page numbers")
    b.add_argument("--numbered-square", action="store_true",
                   help="output □ lines that start with a number, like '□ 1. title', as numeric headings without bullets (for documents without chapters)")
    b.set_defaults(func=cmd_build)

    s = sub.add_parser("skeleton", help="reference.hwpx → empty template + draft profile")
    s.add_argument("input")
    s.add_argument("output")
    s.add_argument("--section", type=int, help="body section number (default: the last section)")
    s.set_defaults(func=cmd_skeleton)

    e = sub.add_parser("export", help="document.hwpx → manuscript.md")
    e.add_argument("input")
    e.add_argument("output", nargs="?")
    e.add_argument("--image-dir")
    e.add_argument("--profile", help="if set, reverse-map style names and restore Markdown following the hierarchy norms (chapter | flat | JSON)")
    e.set_defaults(func=cmd_export)

    f = sub.add_parser("fill", help='substitution JSON {"replace": {...}, "cells": {"label > right": "value"}}')
    f.add_argument("input")
    f.add_argument("mapping")
    f.add_argument("-o", "--output")
    f.set_defaults(func=cmd_fill)

    r = sub.add_parser("render", help="rhwp export-pdf + pdftoppm")
    r.add_argument("input")
    r.add_argument("--outdir")
    r.add_argument("--dpi", type=int, default=RENDER_DPI)
    r.add_argument("--rhwp")
    r.add_argument("--keep-soft-hyphen", action="store_true", help="render the original as is without the review copy substitutions")
    r.add_argument("--engine", default="auto", choices=["auto", "rhwp", "soffice"],
                   help="auto=rhwp first, LibreOffice+H2Orestart fallback if missing")
    r.set_defaults(func=cmd_render)

    c = sub.add_parser("check", help="schema validation + rhwp verify")
    c.add_argument("input")
    c.add_argument("--expect", default="", help='rhwp verify argument string, e.g. "--expect-min-tables 2 --expect-contains 목차"')
    c.set_defaults(func=cmd_check)

    d = sub.add_parser("docx2md", help="DOCX → manuscript.md (step 1: print the profile, step 2: convert with --map JSON)")
    d.add_argument("input")
    d.add_argument("output", nargs="?")
    d.add_argument("--map", help="level mapping JSON path")
    d.set_defaults(func=cmd_docx2md)

    x = sub.add_parser("xlsx2md", help="XLSX → manuscript.md (a pipe table per sheet)")
    x.add_argument("input")
    x.add_argument("output", nargs="?")
    x.add_argument("--allow-wide", action="store_true")
    x.set_defaults(func=cmd_xlsx2md)


    a = ap.parse_args()
    a.func(a)


if __name__ == "__main__":
    main()
