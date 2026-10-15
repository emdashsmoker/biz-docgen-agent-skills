#!/usr/bin/env python3
"""docx_report.py: a writer that records the results of the common typesetting layer as docx.

Composition: manuscript parsing, inline notation, table layout calculation (column widths, row heights, merges, alignment, rule sides), figure size, and cover and TOC order
use the format-agnostic functions of hwpx_report.py as is. This file only moves those results into Word document elements.
Template values (font, size, alignment, margins, line spacing, level symbols, page setup) are read only from the spec the caller passes.
Why python-docx: package parts, relationships, pictures, headers, and footers are left to the library, and only formatting without an API
(bullet definitions, cell rules and fills, cell margins, fields, footnote parts) is added as OOXML elements to reduce Word compatibility defects.
Verified environment: python-docx 1.2.0, lxml 6.x, LibreOffice 24.x (render review)
"""
from __future__ import annotations

import copy
import io
import os
import sys
from datetime import datetime
from xml.sax.saxutils import escape

import hwpx_report as lib

try:
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
    from docx.opc.constants import CONTENT_TYPE as CT, RELATIONSHIP_TYPE as RT
    from docx.opc.packuri import PackURI
    from docx.opc.part import Part
    from docx.oxml import OxmlElement, parse_xml
    from docx.oxml.ns import nsdecls, qn
    from docx.shared import Emu, Pt, RGBColor, Twips
except ImportError:  # needed only at typesetting time. The hwpx-only path runs without python-docx
    Document = None

# ═══════════════════════════════════════════════════════════════════════
# Settings: unit conversion and Word-specific corrections
# ═══════════════════════════════════════════════════════════════════════
HU_PER_TWIP = 5            # HWPUNIT(1/7200 inch) → twip(1/1440 inch)
EMU_PER_HU = 127           # HWPUNIT → EMU(1/914400 inch)
HU_PER_PT = 100
TWIP_PER_PT = 20
EIGHTHS_PER_MM = 8 * 72 / 25.4   # rule width unit of w:sz (1/8 pt)
SPACER_MIN_TWIP = 20       # minimum height of a spacing paragraph (1pt). Word merges two tables with no paragraph between them
JC = {"LEFT": WD_ALIGN_PARAGRAPH.LEFT, "CENTER": WD_ALIGN_PARAGRAPH.CENTER, "RIGHT": WD_ALIGN_PARAGRAPH.RIGHT,
      "JUSTIFY": WD_ALIGN_PARAGRAPH.JUSTIFY, "DISTRIBUTE": WD_ALIGN_PARAGRAPH.DISTRIBUTE} if Document else {}
FOOTNOTE_SUFFIX = ")"      # character after the footnote number. Same as the hwpx footnote number format (number and closing parenthesis)
TABLE_IN_MARGIN = (510, 510, 141, 141)          # default table cell margins left, right, top, bottom (same as the python-hwpx table default)
# Schema order of formatting elements. Word judges elements out of order as corruption
TCPR_ORDER = ("cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar",
              "textDirection", "tcFitText", "vAlign", "hideMark")
TBLPR_ORDER = ("tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize", "tblW",
               "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook")
RPR_ORDER = ("rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike", "outline", "shadow",
             "emboss", "imprint", "noProof", "snapToGrid", "vanish", "webHidden", "color", "spacing", "w", "kern",
             "position", "sz", "szCs", "highlight", "u", "effect", "bdr", "shd", "fitText", "vertAlign", "rtl", "cs",
             "em", "lang", "eastAsianLayout", "specVanish", "oMath")
PPR_ORDER = ("pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
             "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct",
             "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind",
             "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection", "textAlignment",
             "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange")
SPACE_SPLIT_RE = __import__("re").compile(r"( +)")
SECTPR_ORDER = ("headerReference", "footerReference", "footnotePr", "endnotePr", "type", "pgSz", "pgMar", "paperSrc",
                "pgBorders", "lnNumType", "pgNumType", "cols", "formProt", "vAlign", "noEndnote", "titlePg",
                "textDirection", "bidi", "rtlGutter", "docGrid", "printerSettings")


def seq_name(kind: str) -> str:
    """Caption number field (SEQ) name. seq of the caption label, otherwise label. Word field names cannot hold spaces, so spaces become underscores"""
    cap = lib.label("caption.table" if kind == "TABLE" else "caption.figure")
    return "_".join((cap["seq"] or cap["label"]).split())


def tw(hu: float) -> int:
    return int(round(hu / HU_PER_TWIP))


def _insert_ordered(parent, child, order: tuple[str, ...]):
    """Insert child elements keeping schema order. An existing element of the same name is replaced"""
    name = child.tag.split("}")[1]
    old = parent.find(qn(f"w:{name}"))
    if old is not None:
        parent.replace(old, child)
        return child
    later = order[order.index(name) + 1:]
    for i, c in enumerate(parent):
        if c.tag.split("}")[-1] in later:
            parent.insert(i, child)
            return child
    parent.append(child)
    return child


def _w(tag: str, **attrs) -> "OxmlElement":
    el = OxmlElement(f"w:{tag}")
    for k, v in attrs.items():
        el.set(qn(f"w:{k}"), str(v))
    return el


class DocxBuilder:
    """Implements the same record functions as lib.Builder (hwpx) for docx. Judgment functions (table layout, inline, TOC order) call lib"""

    def __init__(self, spec: dict, profile: dict, page_break: str = "chapter"):
        if Document is None:
            sys.exit("ERROR: python-docx not installed. Run `pip install -U python-docx --break-system-packages` and retry")
        self.doc = Document()
        self.spec, self.p, self.page_break = spec, profile, page_break
        self.n = {"chapter": 0, "section": 0, "subsection": 0, "ganada": 0, "table": 0, "figure": 0, "total_section": 0}
        self.headings: list[tuple[int, str]] = []
        self.toc_numbers: dict[str, int] = {}
        self.pending_break = False
        self.after_block = False
        self.toc, self.cover, self.page_number, self.header_spec, self.toc_depth = False, None, False, None, 2
        self.merge_mode, self.long_table_mode = "auto", "engine"
        self.last_para = None
        self._warned: set[str] = set()
        self._footnotes: list[tuple[int, str]] = []
        page = spec["page"]
        self.body_w = page["width"] - page["left"] - page["right"]
        self.table_max_w = self.body_w - lib.TABLE_INDENT_MARGIN
        self.rows = {r[0]: r for r in spec["styles"]}          # style name → definition row
        self.base_names = {r[1]: r[0] for r in spec["styles"]}  # role → style name (header, page number, footnote)
        self._clean_template()
        self._setup_page()
        self._setup_defaults()
        self._num_ids = self._setup_numbering()
        self._setup_styles()
        self.manual_symbol = {role: sym for role, sym in lib.MANUAL_SYMBOLS.items()
                              if self.p.get(role) and self.p[role] in self.rows and not self.rows[self.p[role]][11]}
        # put an empty line at the start of the body, same as the first paragraph of an hwpx section (an empty 바탕글 line holding section settings)
        self._new_para("body")

    # ── Template: page, defaults, bullets, styles (read from spec in one place) ────────────────
    def _clean_template(self):
        """Remove the preview picture, custom XML, and author information of the python-docx default template"""
        pkg = self.doc.part.package
        for rid, rel in list(pkg.rels.items()):
            if rel.reltype == RT.THUMBNAIL:
                del pkg.rels[rid]
        for rid, rel in list(self.doc.part.rels.items()):
            if rel.reltype == RT.CUSTOM_XML:
                self.doc.part.drop_rel(rid)
        settings = self.doc.settings.element
        zoom = settings.find(qn("w:zoom"))
        if zoom is not None:  # zoom in the default template lacks the schema-required attribute (percent)
            zoom.set(qn("w:percent"), "100")
        for tag in ("savePreviewPicture", "rsids", "attachedTemplate"):
            for el in settings.findall(qn(f"w:{tag}")):
                settings.remove(el)
        cp = self.doc.core_properties
        cp.author = cp.last_modified_by = cp.comments = cp.title = cp.subject = cp.keywords = ""
        cp.revision = 1
        cp.created = cp.modified = datetime.now().replace(microsecond=0)

    def _setup_page(self):
        pg = self.spec["page"]
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Twips(tw(pg["width"])), Twips(tw(pg["height"]))
        sec.left_margin, sec.right_margin = Twips(tw(pg["left"])), Twips(tw(pg["right"]))
        # the 한/글 top margin runs from the page edge to the header area, and the header margin is that area's height. Word's top margin runs from the page edge to the body
        sec.top_margin, sec.header_distance = Twips(tw(pg["top"] + pg["header"])), Twips(tw(pg["top"]))
        sec.bottom_margin, sec.footer_distance = Twips(tw(pg["bottom"] + pg["footer"])), Twips(tw(pg["bottom"]))
        grid = sec._sectPr.find(qn("w:docGrid"))
        if grid is not None:  # turn off the character grid (so line spacing does not stretch to the grid)
            sec._sectPr.remove(grid)

    def _font(self, key: str) -> str:
        return self.spec["fonts"][key]

    def _setup_defaults(self):
        body = self.rows[self.p["body"]]
        f = self._font(body[2])
        styles = self.doc.styles.element
        old = styles.find(qn("w:docDefaults"))
        new = parse_xml(
            f'<w:docDefaults {nsdecls("w")}><w:rPrDefault><w:rPr>'
            f'<w:rFonts w:ascii="{f}" w:eastAsia="{f}" w:hAnsi="{f}" w:cs="{f}"/>'
            f'<w:sz w:val="{int(body[3] * 2)}"/><w:szCs w:val="{int(body[3] * 2)}"/>'
            f'<w:lang w:val="{lib.label("language.latin")}" w:eastAsia="{lib.label("language.east_asia")}" w:bidi="ar-SA"/>'
            f'</w:rPr></w:rPrDefault>'
            # 한/글 does not insert spacing between Korean and Latin text automatically. Turn off Word's automatic spacing
            f'<w:pPrDefault><w:pPr><w:autoSpaceDE w:val="0"/><w:autoSpaceDN w:val="0"/>'
            f'<w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr></w:pPrDefault></w:docDefaults>')
        styles.replace(old, new) if old is not None else styles.insert(0, new)

    def _bullet_geometry(self, row) -> tuple[int, int]:
        """(text start position, hanging indent) of a bullet paragraph in HWPUNIT. 한/글 puts the symbol in the left margin and starts the text after the symbol width and gap"""
        glyph = self.spec["glyphs"][row[11]]
        gw = self.spec["glyph_em"].get(glyph, 1.0) * row[3] * HU_PER_PT
        hang = gw + self.spec["bullet_offset"]
        return row[6] + hang, hang

    def _setup_numbering(self) -> dict[str, int]:
        """Create one level symbol definition (bullet) per style. Symbols are shown as Word bullets, not as body text"""
        numbering = self.doc.part.numbering_part.element
        ids: dict[str, int] = {}
        first_num = numbering.find(qn("w:num"))
        # assign numbers after the existing maximum so they do not collide with numbers used by list styles of the default template (List Bullet, etc.)
        base_a = max([int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))] or [0])
        base_n = max([int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))] or [0])
        for i, row in enumerate(r for r in self.spec["styles"] if r[11]):
            aid, nid = base_a + i + 1, base_n + i + 1
            glyph = self.spec["glyphs"][row[11]].replace("­", "-")
            font = self._font(row[2])
            text_at, hang = self._bullet_geometry(row)
            absn = parse_xml(
                f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{aid}"><w:multiLevelType w:val="singleLevel"/>'
                f'<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="bullet"/><w:lvlText w:val="{escape(glyph)}"/>'
                f'<w:lvlJc w:val="left"/><w:pPr><w:ind w:left="{tw(text_at)}" w:hanging="{tw(hang)}"/></w:pPr>'
                f'<w:rPr><w:rFonts w:ascii="{font}" w:eastAsia="{font}" w:hAnsi="{font}" w:hint="eastAsia"/></w:rPr>'
                f'</w:lvl></w:abstractNum>')
            num = parse_xml(f'<w:num {nsdecls("w")} w:numId="{nid}"><w:abstractNumId w:val="{aid}"/></w:num>')
            if first_num is not None:
                first_num.addprevious(absn)
            else:
                numbering.append(absn)
            numbering.append(num)
            ids[row[0]] = nid
        return ids

    def _style_line(self, row, pt: float | None = None) -> int:
        """Convert 한/글 percentage line spacing to fixed line spacing (font size × percentage). Removes renderer differences due to font height"""
        return int(round((pt or row[3]) * row[10] / 100 * TWIP_PER_PT))

    def _setup_styles(self):
        self.style_names: dict[str, str] = {}
        normal = self.doc.styles["Normal"]
        for row in self.spec["styles"]:
            name = row[0]
            if name == self.p["body"]:
                st = normal  # 바탕글 stays as Word's default style (Normal). Renaming it stops Word from recognizing it as the default style
            else:
                st = self.doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
                st.base_style = normal
                st.quick_style = True
            self.style_names[name] = st.name
            self._format_style(st, row)
            if st is normal:  # like 한/글, no automatic spacing between Korean and Latin text or numbers (LibreOffice reads the style value over the default)
                ppr = st.element.get_or_add_pPr()
                for tag in ("autoSpaceDE", "autoSpaceDN"):
                    _insert_ordered(ppr, _w(tag, val=0), PPR_ORDER)
        for name in ("Heading 1", "Heading 2", "Heading 3", "Title"):  # hide the default template's heading styles so their theme colors and fonts do not mix in
            try:
                self.doc.styles[name].hidden = True
            except KeyError:
                pass

    def _format_style(self, st, row):
        name, _role, fkey, pt, bold, align, left, indent, prev, nxt, ls, bullet, keep, widow = row
        font = self._font(fkey)
        st.font.name = font
        rfonts = st.element.get_or_add_rPr().get_or_add_rFonts()
        rfonts.set(qn("w:eastAsia"), font)
        rfonts.set(qn("w:cs"), font)
        st.font.size = Pt(pt)
        st.font.bold = bool(bold)
        st.font.color.rgb = RGBColor.from_string(self.spec["color"].lstrip("#"))
        pf = st.paragraph_format
        pf.alignment = JC[align]
        if bullet:
            text_at, hang = self._bullet_geometry(row)
            pf.left_indent, pf.first_line_indent = Twips(tw(text_at)), Twips(-tw(hang))
            numpr = st.element.get_or_add_pPr().get_or_add_numPr()
            numpr.get_or_add_ilvl().val = 0
            numpr.get_or_add_numId().val = self._num_ids[name]
        else:
            self._indent(pf, left, indent)
        pf.space_before, pf.space_after = Twips(tw(prev)), Twips(tw(nxt))
        pf.line_spacing = Twips(self._style_line(row))
        pf.keep_with_next = bool(keep)
        pf.widow_control = bool(widow)
        if name in (self.p.get("toc_chapter"), self.p.get("toc_section")):
            pf.tab_stops.add_tab_stop(Twips(tw(self.spec["toc_tab"])), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)

    @staticmethod
    def _indent(pf, left: int, indent: int):
        """Convert 한/글 margins to Word indents. A negative 한/글 hanging indent puts the first line in the margin and indents from the second line"""
        if indent < 0:
            pf.left_indent, pf.first_line_indent = Twips(tw(left - indent)), Twips(tw(indent))
        else:
            pf.left_indent, pf.first_line_indent = Twips(tw(left)), Twips(tw(indent))

    # ── Basic elements ───────────────────────────────────────────────────────
    def _row(self, role: str):
        name = self.p.get(role)
        if name is None and role not in self._warned:
            self._warned.add(role)
            print(f"WARN: no '{role}' style in the profile → output as 바탕글")
        return self.rows[name or self.p["body"]]

    def _new_para(self, role: str, align: str | None = None):
        row = self._row(role)
        p = self.doc.add_paragraph(style=self.style_names[row[0]])
        if align:
            p.paragraph_format.alignment = JC[align]
        if self.pending_break:
            p.paragraph_format.page_break_before = True
            self.pending_break = False
        return p

    def _run(self, p, text: str, role: str, pt: float | None = None, bold: bool | None = None,
             color: str | None = None, font_key: str | None = None):
        """Text with a changed size (heading levels, bands, cover). Paragraph line spacing is recalculated from the changed size"""
        if pt is not None:
            p.paragraph_format.line_spacing = Twips(self._style_line(self._row(role), pt))
        return self._text(p, text, pt or self._row(role)[3], bold=bold, color=color, font_key=font_key, size=pt)

    def _space_twips(self, pt: float) -> int:
        """Difference between the 한/글 fixed space width (half the font size) and the font's space width. Added to docx spaces as character spacing"""
        return int(round((self.spec["space_em"] - self.spec["glyph_em"][" "]) * pt * TWIP_PER_PT))

    def _text(self, p, text: str, pt: float, **fmt):
        """Record text split into space units. Space pieces get character spacing that matches the 한/글 space width
        fmt: bold, italic, script(sup|sub), color, font_key"""
        runs = []
        for part in SPACE_SPLIT_RE.split(text):
            if not part:
                continue
            r = p.add_run(part)
            if fmt.get("bold") is not None:
                r.font.bold = fmt["bold"]
            if fmt.get("italic"):
                r.font.italic = True
            if fmt.get("script") == "sup":
                r.font.superscript = True
            elif fmt.get("script") == "sub":
                r.font.subscript = True
            if fmt.get("color"):
                r.font.color.rgb = RGBColor.from_string(fmt["color"].lstrip("#"))
            if fmt.get("size"):
                r.font.size = Pt(fmt["size"])
            if fmt.get("font_key"):
                f = self._font(fmt["font_key"])
                r.font.name = f
                r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), f)
            if part[0] == " " and self._space_twips(pt):
                _insert_ordered(r._element.get_or_add_rPr(), _w("spacing", val=self._space_twips(pt)), RPR_ORDER)
            runs.append(r)
        return runs

    def para(self, role: str, text: str, footnotes: dict | None = None, no_heading: bool = False, **_):
        p = self._new_para(role)
        if no_heading:  # the number is written in the body, so turn off the style's bullet
            numpr = p._p.get_or_add_pPr().get_or_add_numPr()
            numpr.get_or_add_numId().val = 0
        else:
            text = self.manual_symbol.get(role, "") + text
        pt = self._row(role)[3]
        for kind, seg in lib.inline_segments(text):
            if kind == "footnote":
                if footnotes is not None:
                    self.add_footnote(p, footnotes.get(seg, ""))
                continue
            self._text(p, seg, pt, bold=True if kind == "bold" else None, italic=kind == "italic",
                       script=kind if kind in ("sup", "sub") else None)
        return p

    def add_footnote(self, p, text: str):
        """Footnote reference. The number is Word's automatic number, and the closing parenthesis after it is attached as text to match the hwpx footnote number format (1))"""
        fid = len(self._footnotes) + 1
        r = p.add_run()
        r.font.superscript = True
        r._r.append(_w("footnoteReference", id=fid))
        p.add_run(FOOTNOTE_SUFFIX).font.superscript = True
        self._footnotes.append((fid, text))

    def gap(self):
        for _ in range(lib.BLOCK_GAP_LINES):
            self._new_para("body")

    def _spacer(self, role: str, which: str):
        """Convert the spacing before and after paragraphs holding tables and figures into empty paragraph heights. 한/글 places a table like a character inside a paragraph, but a Word table sits outside paragraphs"""
        row = self._row(role)
        height = row[8] if which == "before" else row[9]
        if which == "before" and self.pending_break:  # no top spacing at the head of a page right after a page break, as in 한/글
            height = 0
        p = self._new_para(role)
        pf = p.paragraph_format
        pf.space_before = pf.space_after = Twips(0)
        pf.line_spacing = Twips(max(tw(height), SPACER_MIN_TWIP))
        return p

    # ── Tables ──────────────────────────────────────────────────────────────
    def _cell_line_role(self, line: str) -> str | None:
        return lib.cell_line_role(line, self.p)

    def special_cell(self, val: str) -> str | None:
        return None

    def caption_label(self, kind: str, text: str) -> dict:
        cap = lib.label("caption.table" if kind == "TABLE" else "caption.figure")
        label = f"{cap['open']}{text}{cap['close']}"   # tables use <title>, figures use [title] (brackets of the caption labels)
        return {"open": label, "num": None, "close": "", "text": label}

    def _caption(self, kind: str, text: str, side: str):
        role = "table_title" if kind == "TABLE" else "fig_title"
        label = self.caption_label(kind, text)
        p = self._new_para(role)
        gap = Twips(tw(lib.CAPTION_GAP))
        if side == "TOP":
            p.paragraph_format.space_after = gap
            p.paragraph_format.keep_with_next = True
        else:
            p.paragraph_format.space_before = gap
        pt = self._row(role)[3]
        self._text(p, label["open"], pt)
        if label["num"] is not None:  # per-chapter numbers: reset the number to 1 at the first caption of a chapter
            reset = " \\r 1" if label["num"] == 1 else ""
            instr = f" SEQ {seq_name(kind)} \\* ARABIC{reset} "
            fld = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="{escape(instr)}"><w:r><w:t>{label["num"]}</w:t></w:r></w:fldSimple>')
            p._p.append(fld)
            self._text(p, label["close"], pt)
        return p

    def _new_table(self, nrows: int, widths: list[int]):
        t = self.doc.add_table(rows=nrows, cols=len(widths))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        tblPr = t._tbl.tblPr
        _insert_ordered(tblPr, _w("tblW", w=tw(sum(widths)), type="dxa"), TBLPR_ORDER)
        mar = _w("tblCellMar")
        for side, v in zip(("top", "left", "bottom", "right"), (TABLE_IN_MARGIN[2], TABLE_IN_MARGIN[0], TABLE_IN_MARGIN[3], TABLE_IN_MARGIN[1])):
            mar.append(_w(side, w=tw(v), type="dxa"))
        _insert_ordered(tblPr, mar, TBLPR_ORDER)
        for ci, w in enumerate(widths):
            t.columns[ci].width = Twips(tw(w))
        for row in t.rows:
            for ci, cell in enumerate(row.cells):
                cell.width = Twips(tw(widths[ci]))
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        return t

    @staticmethod
    def _tc_style(tc, sides: list[str], fill: str | None, margin: tuple | None):
        tcPr = tc.get_or_add_tcPr()
        borders = _w("tcBorders")
        for side in ("top", "left", "bottom", "right"):
            if side in sides:
                borders.append(_w(side, val="single", sz=max(2, round(float(lib.BORDER_WIDTH.split()[0]) * EIGHTHS_PER_MM)),
                                  space=0, color=lib.BORDER_COLOR.lstrip("#")))
            else:
                borders.append(_w(side, val="nil"))
        _insert_ordered(tcPr, borders, TCPR_ORDER)
        if fill:
            _insert_ordered(tcPr, _w("shd", val="clear", color="auto", fill=fill.lstrip("#")), TCPR_ORDER)
        if margin:
            mar = _w("tcMar")
            for side, v in zip(("top", "left", "bottom", "right"), (margin[2], margin[0], margin[3], margin[1])):
                mar.append(_w(side, w=tw(v), type="dxa"))
            _insert_ordered(tcPr, mar, TCPR_ORDER)

    @staticmethod
    def _row_height(row, height: int, cant_split: bool = False, header: bool = False):
        row.height, row.height_rule = Twips(tw(height)), WD_ROW_HEIGHT_RULE.AT_LEAST
        trPr = row._tr.get_or_add_trPr()
        if cant_split:
            trPr.append(_w("cantSplit"))
        if header:
            trPr.append(_w("tblHeader"))

    def _cell_para(self, cell, text: str, role: str, align: str | None = None, bold: bool | None = None, first: bool = True):
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        p.style = self.doc.styles[self.style_names[self._row(role)[0]]]
        if align:
            p.paragraph_format.alignment = JC[align]
        if text:
            self._text(p, text, self._row(role)[3], bold=bold)
        return p

    def _transparent_row(self, t, ri: int, ncol: int, text: str, role: str, align: str):
        """Transparent merged rows inside a table (unit, source). They move with the table"""
        cell = t.cell(ri, 0).merge(t.cell(ri, ncol - 1)) if ncol > 1 else t.cell(ri, 0)
        if ncol > 1:
            cell.width = Emu(sum(c.width for c in t.columns))
        self._cell_para(cell, text, role, align)
        self._tc_style(cell._tc, [], None, None)
        self._row_height(t.rows[ri], lib.ROW_BASE_H)

    def table(self, e: lib.Element):
        rows = [r + [""] * (max(len(x) for x in e.data["rows"]) - len(r)) for r in e.data["rows"]]
        nrow, ncol = len(rows), len(rows[0])
        widths = lib.calc_col_widths(rows, self.table_max_w)
        aligns = lib.col_aligns(rows, e.data["aligns"])
        hdr = lib.header_rows(rows)
        if nrow > lib.LONG_TABLE_ROWS and self.long_table_mode == "split":
            for sub in lib.split_table(e, rows):
                self.table(sub)
            return
        self.gap()
        unit, sources = e.data.get("unit"), e.data.get("sources", [])
        head = 1 if unit else 0
        long = nrow > lib.LONG_TABLE_ROWS and not e.data.get("split_done")
        self._spacer("table_pos", "before")
        if e.text:
            self._caption("TABLE", e.text, "TOP")
        t = self._new_table(nrow + head + len(sources), widths)
        vspans, hspans = lib.table_spans(rows, hdr, self.merge_mode)
        # merge area: anchor (row, column) → (end row, end column). Rows count manuscript rows including the header
        anchors: dict[tuple[int, int], tuple[int, int]] = {}
        covered: set[tuple[int, int]] = set()
        for c, r0, r1 in vspans:
            a, b = hdr + r0, hdr + r1
            anchors[(a, c)] = (b, c)
            covered |= {(r, c) for r in range(a + 1, b + 1)}
        for ri, c0, c1 in hspans:
            anchors[(ri, c0)] = (ri, c1)
            covered |= {(ri, c) for c in range(c0 + 1, c1 + 1)}
        for (r0, c0), (r1, c1) in anchors.items():
            t.cell(head + r0, c0).merge(t.cell(head + r1, c1)).width = Twips(tw(sum(widths[c0:c1 + 1])))
        if unit:
            self._transparent_row(t, 0, ncol, unit, "unit", "RIGHT")
        for k, src in enumerate(sources):
            self._transparent_row(t, head + nrow + k, ncol, src, "source", "LEFT")
        heights = lib.row_heights(rows, widths, hdr, vspans, hspans)
        for ri, row in enumerate(rows):
            is_head = ri < hdr
            self._row_height(t.rows[head + ri], heights[ri], cant_split=not long,
                             header=long and ri < hdr)
            for ci, raw in enumerate(row):
                if (ri, ci) in covered:
                    continue
                cell = t.cell(head + ri, ci)
                fill = self.special_cell(raw)
                val, roles = lib.cell_lines("" if fill else raw, self._cell_line_role)
                role, align = lib.cell_role(aligns[ci], is_head)
                for li, line in enumerate(val.split("\n")):
                    br = roles[li] if li < len(roles) and not is_head else None
                    self._cell_para(cell, line, br or role, None if br else align, True if is_head else None, first=li == 0)
                r1, c1 = anchors.get((ri, ci), (ri, ci))
                sides = lib.edge_sides(ci, c1, ncol)  # shaded cells also follow the table's outer left and right rule (same rules as hwpx)
                style_fill = fill or (lib.HEADER_FILL if is_head else None)
                for rr in range(ri, r1 + 1):  # the same rules and fill for continued cells of a vertical merge
                    tc = t.rows[head + rr]._tr.tc_lst[self._tc_index(t.rows[head + rr]._tr, ci)]
                    self._tc_style(tc, sides, style_fill, lib.CELL_MARGIN)
        if long and head:  # Word repeats only rows continuous from the first table row, so the unit row is included in repetition
            t.rows[0]._tr.get_or_add_trPr().append(_w("tblHeader"))
        if not long:  # move whole without breaking across pages, like a 한/글 treat-as-character table
            for row in t.rows[:-1]:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        p.paragraph_format.keep_with_next = True
        self._spacer("table_pos", "after")
        self.after_block = True

    @staticmethod
    def _tc_index(tr, grid_col: int) -> int:
        """Cell index in the row for a grid column position (reflects the span of horizontal merges)"""
        col = 0
        for i, tc in enumerate(tr.tc_lst):
            if col >= grid_col:
                return i
            col += tc.grid_span
        return len(tr.tc_lst) - 1

    # ── Figures ────────────────────────────────────────────────────────────
    def image(self, e: lib.Element):
        prepared = lib.prepare_image(e)
        if prepared is None:
            return
        data, _fmt, w_mm, h_mm = prepared
        self.gap()
        w_hu, h_hu = round(w_mm * 7200 / 25.4), round(h_mm * 7200 / 25.4)
        sources = e.data.get("sources", [])
        self._spacer("fig_pos", "before")
        if e.text and lib.FIG_CAPTION_SIDE == "TOP":
            self._caption("PICTURE", e.text, "TOP")
        if not lib.FIGURE_IN_TABLE:
            p = self._new_para("fig_pos", "CENTER")
            p.paragraph_format.line_spacing = 1.0  # fixed line spacing clips pictures
            p.add_run().add_picture(io.BytesIO(data), width=Emu(w_hu * EMU_PER_HU), height=Emu(h_hu * EMU_PER_HU))
            if e.text and lib.FIG_CAPTION_SIDE != "TOP":
                self._caption("PICTURE", e.text, "BOTTOM")
            for src in sources:
                self.para("source", src)
        else:  # put the picture and the source row in a transparent one-column table (same composition as hwpx)
            t = self._new_table(1 + len(sources), [w_hu])
            cell = t.cell(0, 0)
            p = self._cell_para(cell, "", "fig_pos", "CENTER")
            p.paragraph_format.line_spacing = 1.0
            p.add_run().add_picture(io.BytesIO(data), width=Emu(w_hu * EMU_PER_HU), height=Emu(h_hu * EMU_PER_HU))
            self._tc_style(cell._tc, [], None, None)
            self._row_height(t.rows[0], h_hu + 2 * lib.ROW_LINE_H, cant_split=True)
            for k, src in enumerate(sources):
                self._transparent_row(t, 1 + k, 1, src, "source", "LEFT")
            cap_below = bool(e.text) and lib.FIG_CAPTION_SIDE != "TOP"
            # Why: a caption below is a paragraph outside the table, so the last row must also keep with next, or the caption can fall away at a page boundary
            for row in (t.rows if cap_below else t.rows[:-1]):
                for p in row.cells[0].paragraphs:
                    p.paragraph_format.keep_with_next = True
            if cap_below:
                self._caption("PICTURE", e.text, "BOTTOM")
        self._spacer("fig_pos", "after")
        self.after_block = True

    # ── Bands ────────────────────────────────────────────────────────────
    def band_emit(self, spec: dict):
        """A one-row, three-cell table. Spacing before and after is an empty paragraph of the band position style value, and keep-with-next is handled by the style"""
        self.gap()
        first = self._spacer("band_pos", "before")
        t = self._new_table(1, spec["widths"])
        row = t.rows[0]
        self._row_height(row, spec["height"], cant_split=True)
        for ci, (text, pt, color, align, fill, border) in enumerate(spec["cells"]):
            cell = row.cells[ci]
            p = self._cell_para(cell, "", "band_cell", align)
            p.paragraph_format.keep_with_next = True
            self._run(p, text, "band_cell", pt=pt / HU_PER_PT, bold=True, color=color, font_key=self._row("chapter")[2])
            sides = ["top", "bottom", "left", "right"] if border else []
            self._tc_style(cell._tc, sides, fill, None)
            if border:  # the band border has a different color and width from table rules
                for b in cell._tc.tcPr.find(qn("w:tcBorders")):
                    b.set(qn("w:color"), border.lstrip("#"))
                    b.set(qn("w:sz"), str(max(2, round(float(self.spec["band_border"].split()[0]) * EIGHTHS_PER_MM))))
        self._spacer("band_pos", "after")
        return first

    def heading_para(self, role: str, label: str, pt: int | None):
        p = self._new_para(role)
        numpr = p._p.get_or_add_pPr().get_or_add_numPr()
        numpr.get_or_add_numId().val = 0
        self._run(p, label, role, pt=pt / HU_PER_PT if pt else None, bold=True if pt else None)
        return p

    def chapter_started(self, p):
        """The caption field (SEQ \\r 1) handles restarting caption numbers per chapter"""

    def after_table(self, e):
        """Cell fill and number counting are handled while recording tables"""

    def after_image(self, e):
        """Number counting is handled while recording captions"""

    # ── Front matter (cover, TOC), headers, page numbers ─────────────────────────────────
    def toc_lists(self) -> list[tuple[str, list[str]]]:
        return []

    def _front_matter(self, cover: dict | None):
        """Put the cover and TOC in a front section and start page numbers of the body section from 1. Paragraph order is lib.front_matter_items"""
        if self.toc and not (self.p.get("toc_chapter") and self.p.get("toc_section")):
            print("WARN: no TOC styles (toc_chapter, toc_section) in the profile, --toc skipped")
            self.toc = False
        body = self.doc.element.body
        anchor = body[0]
        made = [self._new_para("body")._p]   # empty 바탕글 line at the start of the front section (same as the first paragraph of an hwpx section)
        for it in lib.front_matter_items(cover, self.toc, self.headings, self.toc_depth, self.toc_lists()):
            kind = it["kind"]
            if kind in ("blank", "cover"):
                p = self._new_para("body", "CENTER")
                if kind == "cover":
                    self._run(p, it["text"], "body", pt=it["pt"] / HU_PER_PT, bold=it["bold"],
                              font_key=self._row("chapter")[2])
            elif kind == "toc_title":
                p = self._new_para("chapter", "CENTER")
                p.paragraph_format.page_break_before = it["page_break"]
                self._text(p, it["text"], self._row("chapter")[3])
            else:
                lv = it.get("level", 2)
                p = self._new_para("toc_chapter" if lv == 1 else "toc_section")
                if kind == "list_entry":
                    p.paragraph_format.left_indent = Twips(tw(lib.TOC_LIST_INDENT))
                elif lv >= 3:
                    p.paragraph_format.left_indent = Twips(tw(lib.TOC_SUB_INDENT))
                num = self.toc_numbers.get(it["text"])
                self._text(p, it["text"] + (f"\t{num}" if num else ""), self._row("toc_chapter" if lv == 1 else "toc_section")[3])
            made.append(p._p)
        for el in made:
            anchor.addprevious(el)
        body_sect = self.doc.sections[0]._sectPr   # first body section (with divider pages the body has several sections)
        front_sect = copy.deepcopy(self._sect_template())
        made[-1].get_or_add_pPr().append(front_sect)
        _insert_ordered(body_sect, _w("type", val="nextPage"), SECTPR_ORDER)
        _insert_ordered(body_sect, _w("pgNumType", start=1), SECTPR_ORDER)
        self._front_done = True

    def _sect_template(self):
        """Copy of section settings without header and footer references, page number start, and first page distinction (page size, margins, columns)"""
        sect = copy.deepcopy(self.doc.element.body.sectPr)
        for tag in ("headerReference", "footerReference", "pgNumType", "titlePg", "type"):
            for el in sect.findall(qn(f"w:{tag}")):
                sect.remove(el)
        return sect

    def _first_body_section(self):
        """The first body section that defines headers and footers. Later sections inherit the previous section's definitions by Word's rules"""
        return self.doc.sections[1 if getattr(self, "_front_done", False) else 0]

    def divider_emit(self, spec: dict, new_page: bool):
        """Divider page: a divider band after an empty 바탕글 line. A divider page is the first page of a new section and hides the page number with an empty first-page header and footer
        Why a section: Word has no per-page hide feature, so a section's 'different first page' produces the same result as 한/글's hide feature"""
        body = self.doc.element.body
        if new_page:  # close the previous section at the preceding paragraph. The new section starts on the next page
            last = body[-2] if body[-1].tag == qn("w:sectPr") else body[-1]
            sect = self._sect_template()
            if getattr(self, "_section_divider", False):
                _insert_ordered(sect, _w("titlePg"), SECTPR_ORDER)
            _insert_ordered(sect, _w("type", val="nextPage"), SECTPR_ORDER)
            last.get_or_add_pPr().append(sect)
        self._section_divider = True
        final = body.sectPr
        _insert_ordered(final, _w("titlePg"), SECTPR_ORDER)   # keep the first page distinction if the last section also starts with a divider page
        self.pending_break = False
        for _ in range(self.spec["divider_top_lines"]):
            self._new_para("body")
        return self.band_emit(spec)

    def set_header(self, spec: str | None):
        if not spec:
            return
        parts = spec.split(";")
        text = parts[0].strip()
        align = next((kv.split("=", 1)[1].strip().upper() for kv in parts[1:] if kv.strip().lower().startswith("align=")), "RIGHT")
        header = self._first_body_section().header
        header.is_linked_to_previous = False
        p = header.paragraphs[0]
        p.style = self.doc.styles[self.style_names[self.base_names["header_text"]]]
        p.paragraph_format.alignment = JC[align]
        self._text(p, text, self.rows[self.base_names["header_text"]][3])

    def set_page_number(self, on: bool):
        """'- N -' centered in the footer. The front section (cover, TOC) has no page numbers"""
        if not on:
            return
        footer = self._first_body_section().footer
        footer.is_linked_to_previous = False
        p = footer.paragraphs[0]
        p.style = self.doc.styles[self.style_names[self.base_names["page_number"]]]
        p.paragraph_format.alignment = JC["CENTER"]
        pt = self.rows[self.base_names["page_number"]][3]
        self._text(p, "- ", pt)
        for kind, text in (("begin", None), (None, " PAGE "), ("separate", None), (None, "1"), ("end", None)):
            r = p.add_run()
            if kind:
                r._r.append(_w("fldChar", fldCharType=kind))
            elif text == " PAGE ":
                it = OxmlElement("w:instrText")
                it.set(qn("xml:space"), "preserve")
                it.text = text
                r._r.append(it)
            else:
                r.text = text
        self._text(p, " -", pt)

    # ── Save ────────────────────────────────────────────────────────────
    def _write_footnotes(self):
        """Footnote part (footnotes.xml) and the separator reference in settings. python-docx has no footnote API, so the part is registered directly"""
        if not self._footnotes:
            return
        style_id = self.doc.styles[self.style_names[self.base_names["footnote"]]].style_id
        sep = ('<w:footnote w:type="separator" w:id="-1"><w:p><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr>'
               '<w:r><w:separator/></w:r></w:p></w:footnote>'
               '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr>'
               '<w:r><w:continuationSeparator/></w:r></w:p></w:footnote>')
        pt = self.rows[self.base_names["footnote"]][3]
        gap = self._space_twips(pt)

        def runs(text: str) -> str:  # character spacing that matches the 한/글 space width in space pieces (same rule as the body _text)
            out = []
            for part in SPACE_SPLIT_RE.split(text):
                if part:
                    rpr = f'<w:rPr><w:spacing w:val="{gap}"/></w:rPr>' if part[0] == " " and gap else ""
                    out.append(f'<w:r>{rpr}<w:t xml:space="preserve">{escape(part)}</w:t></w:r>')
            return "".join(out)
        notes = "".join(
            f'<w:footnote w:id="{fid}"><w:p><w:pPr><w:pStyle w:val="{style_id}"/></w:pPr>'
            f'<w:r><w:footnoteRef/></w:r>{runs(FOOTNOTE_SUFFIX + " " + text)}</w:p></w:footnote>'
            for fid, text in self._footnotes)
        blob = f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<w:footnotes {nsdecls("w")}>{sep}{notes}</w:footnotes>'
        part = Part(PackURI("/word/footnotes.xml"), CT.WML_FOOTNOTES, blob.encode("utf-8"), self.doc.part.package)
        self.doc.part.relate_to(part, RT.FOOTNOTES)
        settings = self.doc.settings.element
        fpr = parse_xml(f'<w:footnotePr {nsdecls("w")}><w:footnote w:id="-1"/><w:footnote w:id="0"/></w:footnotePr>')
        compat = settings.find(qn("w:compat"))
        compat.addprevious(fpr) if compat is not None else settings.append(fpr)

    def _blank_divider_pages(self):
        """State empty first-page header and footer definitions for divider sections (first page distinction)
        Why: without a first-page definition, LibreOffice shows the previous section's default footer (page number) on the divider page"""
        for section in self.doc.sections:
            if section.different_first_page_header_footer:
                section.first_page_header.is_linked_to_previous = False
                section.first_page_footer.is_linked_to_previous = False

    def save(self, out: str):
        self._blank_divider_pages()
        self._write_footnotes()
        self.doc.save(os.path.abspath(out))
