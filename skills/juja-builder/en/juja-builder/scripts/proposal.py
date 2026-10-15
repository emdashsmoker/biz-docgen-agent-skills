#!/usr/bin/env python3
"""juja-builder (scripts/proposal.py): the single CLI of the research service proposal writing pipeline.

Commands
  rfp       RFP (hwpx|pdf|docx|md) → 01_RFP원문.md (converted text), empty 01_요구항목표.md and 01_배점표.md tables (the model reads the source and fills them)
  deck      proposal deck (pptx|pdf) → 발표자료/ (reference PDF, 원문.md, page PNGs, thumbnail sheets, embedded images)
  figure    capture a whole deck page or part of it as a diagram PNG for insertion (--grid makes a grid PNG for reading crop values)
  init      create the workspace (--kind proposal|report|plan. state.json, default skeleton _양식/proposal.hwpx. A proposal adds 00_제안전략서.md, 회사자료.md, and the common part templates 원고/41 61 62. A pre-diagnosis report adds 00_진단서.md, a planning report adds 00_기획설계.md)
  ledger    tally the deciding sections of 01_배점표.md and update attributed points and status in 02_섹션대장.md
  check     rfp | density | wording | tone | numbers | fill | volume | all (all is the check list for the output type. --skip leaves out named checks)
  build     merge 원고/*.md → 납품/{제안서|사전진단보고서|기획보고서}.hwpx, .docx (--format hwpx|docx|both. The skeleton is regenerated on every run.
            --labels is a document label file. The default labels are Korean)
  render    hwpx or docx → PDF + page PNGs (hwpx uses rhwp first, LibreOffice with H2Orestart if missing. docx uses LibreOffice)
  selftest  judge the render path by typesetting and rendering a sample manuscript (the start gate. Result in state.json checks.render)
  skeleton  client template.hwpx → skeleton.hwpx + draft profile

Composition: hwpx_report.py is the format-agnostic typesetting layer (manuscript parser, inline notation, table layout calculation, figure size, TOC order) and the hwpx writer,
docx_report.py is the docx writer. This file handles judgments per output type (heading numbers, bands, caption numbers, table width directives) and the verdict rules.
Why the split: one content source (MD manuscripts) and one format source (settings), with judgments computed once, so hwpx and docx record the same result.
Verified environment: Python 3.11-3.12, python-hwpx 6.x, python-docx 1.2.0, lxml, pillow, poppler 24.02, rhwp v0.8.6, LibreOffice (docx rendering)

Format notice: 본 제품은 한컴의 HWP 문서 파일(.hwp) 공개 문서를 참고하여 개발하였습니다.
Trademark notice: 한/글 (Hangul), Hancom Office, and HWPX are trademarks of Hancom Inc. (한글과컴퓨터)
Non-affiliation notice: this tool is an independent tool with no affiliation with Hancom Inc.
"""
from __future__ import annotations

import argparse
import glob
import json
import os
import re
import shutil
import subprocess
import sys
import copy
import warnings
import zipfile
from datetime import datetime

warnings.filterwarnings("ignore", category=DeprecationWarning)
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
ASSET_DIR = os.path.join(os.path.dirname(SCRIPT_DIR), "assets")
sys.path.insert(0, SCRIPT_DIR)
import hwpx_report as lib  # noqa: E402  format-agnostic layer such as the parser and table layout calculation, and the hwpx writer
import docx_report as docx_lib  # noqa: E402  docx writer
from lxml import etree  # noqa: E402  installed together as a python-hwpx dependency

HP = lib.HP
HH = "{http://www.hancom.co.kr/hwpml/2011/head}"
ROMAN = ["Ⅰ", "Ⅱ", "Ⅲ", "Ⅳ", "Ⅴ", "Ⅵ", "Ⅶ", "Ⅷ", "Ⅸ", "Ⅹ", "Ⅺ", "Ⅻ"]

# ═══════════════════════════════════════════════════════════════════════
# Settings 1: default skeleton and proposal profile
# ═══════════════════════════════════════════════════════════════════════
# init and build create the default skeleton in workspace/_양식/ by putting the definitions below on an empty new python-hwpx document.
# Why generated in code: no template file ships with the skill, and every style value stays reproducible from this settings part
SKELETON_FONT_BODY = "나눔바른고딕"   # fontface id 0: body, tables, notes
# fontface id 1: heading family. The bold font must be grouped under the same name for bold to show in both 한/글 and approximate renderers
# (NanumSquare splits names by weight, so bold displays differently per renderer)
SKELETON_FONT_HEAD = "나눔바른고딕"
SKELETON_TEXT_COLOR = "#000000"
SKELETON_BULLET_OFFSET_TYPE = "HWPUNIT"  # unit of the gap between symbol and text. PERCENT is proportional to symbol width, so the gap vanishes for narrow symbols
SKELETON_BULLET_OFFSET = "600"           # gap between symbol and text (6pt)
SKELETON_TOC_TAB_POS = 48100             # right tab position of TOC page numbers (margin of 90 from the body width 48190)
# Page setup (HWPUNIT, 1mm = 283.46)
SKELETON_PAGE_MARGIN = {"left": 5669, "right": 5669, "top": 4252, "bottom": 4252, "header": 2835, "footer": 2835}
# Level symbols (한/글 bullet feature). The special symbols □, ○, -, · of Article 2 of 「행정업무의 운영 및 혁신에 관한 규정 시행규칙」
# Level 4 (·) and notes (※) put the symbol in as body text (MANUAL_SYMBOLS). Narrow symbols then show the same gap in every renderer
SKELETON_GLYPHS = {"square": "□", "circle": "○", "dash": "-", "cell_bullet": "·", "cell_dash": "-"}
SKELETON_HANG_DOT = -800     # hanging indent of level 4 "· " (symbol and space width at 12pt)
SKELETON_HANG_NOTE = -1100   # hanging indent of note "※ " (symbol and space width at 10pt)
# Style definitions. This table is the single source of the skeleton styles and of the role → style name mapping of PROFILES
# font: body|head, pt: font size, align, left/indent/prev/next: HWPUNIT, ls: line spacing %,
# bullet: bullet key (SKELETON_GLYPHS), keep: keep with next, widow: widow and orphan control
SKELETON_STYLES = [
    # name           role key         font    pt    bold   align      left  indent prev  next  ls   bullet         keep widow
    ("장 제목",      "chapter",       "head", 20.0, True,  "LEFT",    0,    0,     0,    0,    140, None,          1, 1),
    ("절 제목",      "section",       "head", 15.0, True,  "LEFT",    0,    0,     2500, 300,  140, None,          1, 1),
    ("항 제목",      "subsection",    "head", 14.0, True,  "LEFT",    0,    0,     2000, 200,  140, None,          1, 1),
    ("목 제목",      "ganada",        "head", 13.0, True,  "LEFT",    300,  0,     1500, 100,  140, None,          1, 1),
    ("논점",         "square",        "head", 13.0, True,  "JUSTIFY", 500,  0,     1500, 0,    140, "square",      1, 1),
    ("주장",         "circle",        "body", 12.0, False, "JUSTIFY", 1000, 0,     1000, 0,    140, "circle",      0, 1),
    ("근거",         "dash",          "body", 12.0, False, "JUSTIFY", 2410, 0,     700,  0,    140, "dash",        0, 1),
    ("상세",         "dot",           "body", 12.0, False, "JUSTIFY", 3410, SKELETON_HANG_DOT, 700, 0, 140, None, 0, 1),
    ("주석",         "note",          "body", 10.0, False, "JUSTIFY", 3000, SKELETON_HANG_NOTE, 400, 0, 140, None, 0, 1),
    ("표 위치",      "table_pos",     "body", 10.0, False, "CENTER",  0,    0,     800,  1000, 160, None,          0, 0),
    ("표 본문",      "table_cell",    "body", 10.0, False, "CENTER",  0,    0,     0,    0,    130, None,          0, 0),
    ("표 본문 왼쪽", "table_cell_left", "body", 10.0, False, "LEFT",  0,    0,     0,    0,    130, None,          0, 0),
    ("표 글머리",    "table_cell_bullet", "body", 10.0, False, "LEFT", 0,   0,     0,    150,  130, "cell_bullet", 0, 0),
    ("표 하이픈",    "table_cell_dash", "body", 10.0, False, "LEFT",  1000, 0,     0,    150,  130, "cell_dash",   0, 0),
    ("표 캡션",      "table_title",   "head", 10.5, True,  "CENTER",  0,    0,     0,    0,    150, None,          1, 0),
    ("그림 캡션",    "fig_title",     "head", 10.5, True,  "CENTER",  0,    0,     0,    0,    150, None,          1, 0),
    ("그림 위치",    "fig_pos",       "body", 10.0, False, "CENTER",  0,    0,     1000, 1000, 160, None,          0, 0),
    ("단위",         "unit",          "body", 9.0,  False, "RIGHT",   0,    0,     300,  0,    150, None,          0, 0),
    ("출처",         "source",        "body", 9.0,  False, "LEFT",    0,    0,     300,  0,    140, None,          0, 0),
    ("밴드 위치",    "band_pos",      "body", 10.0, False, "CENTER",  0,    0,     1200, 600,  160, None,          1, 0),   # paragraph holding a band (spacing before and after, keep with next)
    ("밴드 칸",      "band_cell",     "head", 12.0, True,  "LEFT",    0,    0,     0,    0,    130, None,          0, 0),   # text style of the band number cell and title cell
    ("목차 장",      "toc_chapter",   "head", 14.0, True,  "LEFT",    0,    0,     1000, 0,    155, None,          0, 0),
    ("목차 절",      "toc_section",   "body", 12.0, False, "LEFT",    1500, 0,     500,  0,    155, None,          0, 0),
    ("문서 제목",    "doc_title",     "head", 20.0, True,  "CENTER",  0,    0,     0,    1500, 140, None,          1, 1),   # first line of a planning report (build --title)
]
# Base styles: styles already present in an empty new python-hwpx document. Overwritten with these values when the skeleton is created so hwpx and docx read the same values
SKELETON_BASE_STYLES = [
    ("바탕글",       "body",          "body", 10.0, False, "JUSTIFY", 0,    0,     0,    0,    160, None,          0, 0),
    ("쪽 번호",      "page_number",   "body", 10.0, False, "JUSTIFY", 0,    0,     0,    0,    160, None,          0, 0),
    ("머리말",       "header_text",   "body", 9.0,  False, "JUSTIFY", 0,    0,     0,    0,    150, None,          0, 0),
    ("각주",         "footnote",      "body", 9.0,  False, "LEFT",    0,    -1310, 0,    0,    130, None,          0, 0),   # prevents widened character spacing in source URLs and bibliographic notation
]
SKELETON_PAGE_SIZE = (59528, 84186)   # A4 width, height (HWPUNIT)
# NanumBarunGothic glyph widths (em). For computing symbol width of docx bullets. 한/글 measures directly from the font, so hwpx does not use this
SKELETON_GLYPH_EM = {"□": 0.892, "○": 0.892, "-": 0.583, "·": 0.446, "※": 0.892, " ": 0.224}
# Space width (em). The skeleton character shape turns off 'space suited to the font' (useFontSpace=0), so 한/글 fixes spaces at half the font size
# docx adds character spacing to spaces to match the same width
SKELETON_SPACE_EM = 0.5
# lib.Builder validates every profile value as a style name, so keys that are not style names are removed before passing to the typesetter (including keys ending in _pt)
PROFILE_EXTRA_KEYS = ("item_format", "chapter_label", "chapter_band", "toc_lists")
PROFILES = {
    "proposal": {
        **lib.PROFILES["chapter"],
        # role → style name comes from SKELETON_STYLES. Level symbols, fonts, and paragraph margins are set by the skeleton's styles
        **{role: name for name, role, *_ in SKELETON_STYLES},
        "body": "바탕글",                    # default style of an empty new python-hwpx document
        "chapter_format": "{r}. {t}",       # Ⅰ. 제안 개요
        "chapter_label": "{r}",             # chapter band number cell
        "section_format": "{s}. {t}",       # 1. 제안 배경 및 필요성
        "subsection_format": "{g} {t}",     # 가. subsection ({g} is the letter notation of the letters label, "가.")
        "item_format": "{i}) {t}",          # 1) item
        "caption_prefix": None,             # None=unnumbered caption <title> (proposal convention). "{c}-" gives automatic numbering <표 1-1>
        "chapter_pt": 2000, "section_pt": 1600, "subsection_pt": 1400,  # sizes for hierarchy. None means the skeleton size
    },
}
# Planning report: the same hierarchy as the proposal (Ⅰ., 1., 가., 1)) and unnumbered captions. The document title is build --title
PROFILES["plan"] = dict(PROFILES["proposal"])
# Pre-diagnosis report: final report convention (제1장, 1.1, 1.1.1, 가.) and caption numbers per chapter (<표 1-1>). The TOC includes the list of tables and list of figures
PROFILES["report"] = {
    **PROFILES["proposal"],
    "chapter_format": "제{n}장 {t}", "chapter_label": "제{n}장",
    "section_format": "{c}.{s} {t}", "subsection_format": "{c}.{s}.{ss} {t}", "item_format": "{g} {t}",
    "caption_prefix": "{c}-", "toc_lists": True,
    "chapter_band": {"label_w": 7600, "label_pt": 1800},   # one line up to "제10장". The width equals the Module band so title start lines align
}
# Output type: recorded in state.json by init --kind. Sets build's default profile. The delivered file name is the label output.<type> (hwpx_report.py LABEL_DEFAULTS)
KINDS = {"proposal": {"profile": "proposal"},
         "report": {"profile": "report"},
         "plan": {"profile": "plan"}}
# Color theme: applies only to bands. Chosen with build --theme. base fills the chapter number cell, tint fills the Module label cell,
# dark is the label text and the Task label border
THEMES = {
    "navy": {"base": "#6B8EC2", "tint": "#E3E9F3", "dark": "#35578F"},
    "mono": {"base": "#262626", "tint": "#D9D9D9", "dark": "#1A1A1A"},
}
DEFAULT_THEME = "navy"
# Table format keeps the typesetter defaults regardless of theme (black 0.15mm rules, gray #D9D9D9 header). Why: it must match the conventional format editors use when fixing tables in 한/글
lib.CELL_MARGIN = (510, 510, 280, 280)   # 1mm top and bottom cell margin. Keeps the first line of bullet cells off the rule
lib.FIG_CAPTION_SIDE = "BOTTOM"          # table captions above, figure captions below (convention). The caption attaches to the transparent table, so it comes after the source rows
lib.TOC_UNIT_W = (630, 540)              # approximate width of one half-width character at NanumBarunGothic 14pt and 12pt
# Bands: a one-row composition of number cell, empty cell, and title cell. Weight decreases from fill (chapter) to light fill (Module) to border (Task)
# Color values are theme keys (base, tint, dark) or "white". label_w: number cell width (HWPUNIT). Module and Task have the same width so title start lines align
# rows: multiple of the one-line row height
BAND_STYLE = {
    "chapter": {"label_w": 4200, "fill": "base", "border": None, "label_color": "white", "label_pt": 2000,
                "title_pt": 2000, "rows": 2.2},
    "module":  {"label_w": 7600, "fill": "tint", "border": None, "label_color": "dark", "label_pt": 1150,
                "title_pt": 1350, "rows": 1.5},
    "task":    {"label_w": 7600, "fill": None, "border": "dark", "label_color": "dark", "label_pt": 1100,
                "title_pt": 1200, "rows": 1.3},
    # divider page: the chapter band composition enlarged and placed near the middle of the page. One line up to "제10장"
    "divider": {"label_w": 11000, "fill": "base", "border": None, "label_color": "white", "label_pt": 2600,
                "title_pt": 2600, "rows": 3.5},
}
DIVIDER_TOP_LINES = 18   # empty 바탕글 lines above the divider band (288pt of about 700pt body height). The band sits slightly above the middle of the page
BAND_GAP_W = 1000        # width of the empty cell between the number cell and title cell (about 3.5mm). An empty cell instead of cell margins makes the gap look the same in every renderer
BAND_BORDER_W = "0.12 mm"
# addition per line for row height estimation. 10pt table text × 130% line spacing = 1300, plus 150 spacing after cell bullet paragraphs
lib.ROW_LINE_H = 1450
BAND_POS_ROLE = "band_pos"
NOTE_PREFIX = "※ "       # marker of note lines below tables and figures that are not source lines
# level 4 (detail) and notes have no bullet definition in the skeleton, so the typesetter writes the symbol as body text. The hanging indent is handled by the skeleton paragraph shape
lib.MANUAL_SYMBOLS = {**lib.MANUAL_SYMBOLS, "dot": "· ", "note": "※ "}
TASK_BAND_RE = re.compile(r"^(Module|Task|Step)\s*([\d\-.]+)[.\s]+(.*)$", re.I)
LONG_SPEC_RE = re.compile(r"^\(\s*장표\s*[:：]\s*(engine|split)(?:\s+(\d+))?\s*\)$")   # long-table directive line: (장표: engine) treat-as-character off and engine split, (장표: split 10) physical split every 10 rows
TABLE_PAGE_LINES = 30            # a table estimated at this many lines or more may run past one page and gets a [long table HITL] warning (approximation for body-width tables)
TABLE_LINE_UNITS = 80            # weighted character width of one table line (about 40 Korean characters). The basis of column width allocation
WIDTH_SPEC_RE = re.compile(r"^\(\s*(?:폭|열폭)\s*[:：]\s*(.+?)\s*\)$")   # table width directive line: (폭: 실적) group name, (폭: 3,2,2) ratios, (폭: 3,2,2, 80%) ratios and table width

# ═══════════════════════════════════════════════════════════════════════
# Settings 2: workspace file names and verdict criteria
# ═══════════════════════════════════════════════════════════════════════
F_STATE = "state.json"
F_STRATEGY = "00_제안전략서.md"
F_DIAGNOSIS = "00_진단서.md"   # design document of the pre-diagnosis report (corresponds to the proposal strategy file)
F_PLAN_DESIGN = "00_기획설계.md"   # design document of the planning report
# design document template per output type: type → (assets template, workspace file name). The proposal has its own branch in cmd_init
DESIGN_TEMPLATES = {"report": ("diagnosis.md", F_DIAGNOSIS), "plan": ("plan-design.md", F_PLAN_DESIGN)}
F_COMPANY = "회사자료.md"
F_RFP_TEXT = "01_RFP원문.md"
F_REQ = "01_요구항목표.md"
F_SCORE = "01_배점표.md"
F_LEDGER = "02_섹션대장.md"
F_STORYLINE = "02_스토리라인.md"
F_REPORT = "05_점검기록.md"
D_DECK = "발표자료"
D_MANUSCRIPT = "원고"
D_DELIVER = "납품"
D_SELFTEST = "_렌더시험"   # selftest output folder (sample manuscript, sample hwpx, PDF, PNG)
D_TEMPLATE = "_양식"       # default skeleton folder. Created by init and build
F_SKELETON = "proposal.hwpx"
SELFTEST_MIN_PAGES = 1     # render test PASS condition: one or more PNG pages
# Why the sample holds chapter, section, bands, □, ❍, evidence, and a table: to reveal at the start any renderer defect in drawing proposal-specific elements
SELFTEST_SAMPLE = """# 렌더 시험

## 표본 절

Module 1. 표본 모듈

Task 1-1. 표본 과업

□ 표본 논점

- (표본) 계위와 글머리 표시 확인
    + 근거 줄 표시 확인
        + 상세 줄 표시 확인

<표> 표본 표

| 항목 | 내용 |
|---|---|
| 셀 | 한글 표시 확인 |
"""
COMMON_TEMPLATES = {"41_제안사일반현황.md": "overview.md", "61_수행조직.md": "organization.md", "62_관리부문.md": "management.md"}
PREFIX_DENSITY = "5"          # check density target (execution part)
PREFIX_FILL = ("4", "6")      # check fill target (common part)
PREFIX_VOLUME = ("3", "4", "5", "6")    # check volume target. The lone □ verdict covers all manuscripts, composition unit lower bounds the original part (3, 5)
PREFIX_ORIGINAL = ("3", "5")
# Sequence prefix (optional). Attached as in `10-51_추진체계.md` when the chapter order of the RFP's specified TOC differs from the default order.
# Why: if the first digit of the file name served as both merge order and check type, the chapter order of a specified TOC could not be followed. The sequence prefix sets the order and the first digit after it sets the type
ORDER_PREFIX_RE = re.compile(r"^\d+-")

CONCRETE_MARK_RE = re.compile(r"\([^)]{0,6}(대상|자료|방법|기준|절차|범위|규모|구성|설계|산정|선정|편성|기록)[^)]{0,6}\)")
QUANTITY_MANY = 6   # warn of excessive listing if one Task has more figures than this
QUANTITY_RE = re.compile(r"\d[\d,]*\s*(건|명|회|개소|개|종|차|주|일|개월|기관|과제|표본|응답)")

# Volume units (usual composition of a qualitative technical proposal: 2-3 ❍ per □, 2 or more evidence hyphens or ※, one visual)
VOLUME_MIN_CIRCLE = 2       # lower bound of ❍ under □
VOLUME_MIN_EVIDENCE = 2     # lower bound of hyphens + ※ under □
VOLUME_MIN_BLOCK_PER_TASK = 1   # lower bound of tables or figures per Task
VOLUME_TASK_MIN_CHARS = 900     # lower bound of body characters per Task (about 1.5 pages). A usual Task runs 2-5 pages
VOLUME_SQUARE_MIN_CHARS = 350   # lower bound of characters per □ (about 0.5 page)
# Level count caps (SKILL.md Volume Units section). Overruns are check volume warnings and do not count toward the verdict
SQUARE_PER_SECTION_MAX = 4       # □ within a section (##) or subsection (###). Group into subsections at 5 or more
CIRCLE_PER_SQUARE_MAX = 4        # ❍ under □
DASH_PER_CIRCLE_MAX = 3          # hyphens under ❍. Convert to a table at 4 or more
SQUARE_TITLE_MAX = 40           # □ heading cap (13pt bold, one line at body width. Based on rhwp v0.8.6 rendering)
LINE_MAX = {"circle": 86, "dash": 84, "dot": 82, "note": 50}   # ❍, hyphen, and ․ are capped at 2 lines, ※ at 1 line (characters including spaces. Based on rhwp v0.8.6 rendering)
SCHEDULE_MARK = "■"              # period cell notation in the schedule chart. Converted to a shaded cell at typesetting
SCHEDULE_FILL = "#D9D9D9"      # fill of schedule period cells (the same gray as table headers)

# Verdict criteria (check wording/tone)
FORBIDDEN_DEFAULT = ["가용가능", "할 수 있다", "할 수 있음", "고려하고 있", "검토할 예정", "노력하겠", "가능할 것으로", "필요시 검토"]
SENTENCE_WARN = [(r"—", "em dash"), (r"가 아닌 ", "contrast construction 'X가 아닌 Y'"), (r"이 아니라 ", "contrast construction 'X가 아니라 Y'"),
                 (r"(으며|하며|이며)\s", "compound sentence with connective ending"), (r"(?<=[가-힣])(하고|되고)\s", "compound sentence with connective ending (-하고)"), (r"(그리고|그러나|그래서|하지만)\s", "conjunctive adverb"),
                 (r"(본 제안서|본 절에서|앞서 언급|다음과 같이|살펴보면|정리하면|요약하면)", "meta statement marker"),
                 (r"(MECE|피라미드 원칙|제1원칙|맥킨지식)", "thinking framework name exposed")]
FRAMEWORK_FAIL = [(r"\b(SCQA|SCR|SCQ|PREP)\b", "background framework name exposed"),
                  (r"\((S|C|Q|A|R)\s*(상황|전개|질문|답|해결)\)", "framework stage name exposed"),
                  (r"(상황|전개|질문|해결)\s*단계로\s*(구성|서술)", "framework stage name exposed")]
TONE_FAIL = [(r"(그간|그동안)[^.\n]{0,25}(소홀|미흡|부족|치중|그쳐|그친)", "implying client policy shortcomings"),
             (r"(형식적|관행적|소극적)[^.\n]{0,15}(운영|협의|추진|대응)", "criticizing client practice"),
             (r"(제안요청서|RFP)(에|의|가|는)[^.\n]{0,15}(더하여|보완하여|미비|흠결|불명확)", "correcting the RFP"),
             # 역산(?!업): excludes partial matches in compounds such as "지역산업연관표" and "지역산업진흥원"
             (r"(역산(?!업)|채점|평가위원|통제 ?배점|배점 귀속|완결조건|섹션 ?대장|적대적|채움본|미배치)", "internal logic vocabulary exposed (역산 is banned in schedule contexts too. Express as 체계적 일정 수립 or 효율적 일정 편성)"),
             (r"(?<![가-힣])(이면 ?니즈|표면 ?니즈|표면 ?목적|판단문|문형|구성 ?단위|변별력|결정 ?장소|지시 ?어휘|조판|치환|오리지널 ?파트|공통 ?파트)", "instruction vocabulary leak (terms of the skill's instructions. Describe the actual content that belongs there. references/lexicon.md replacement table)"),
             (r"(작업분해도|WBS|마일스톤|임계경로|워크 ?패키지|이터레이션|페인포인트|딜리버리|인사이트|온보딩)", "management loanwords"),
             (r"(?<![가-힣])(포인트|세팅|케이스|니즈|픽스|컨펌|어필|메리트)(?=[\s,.)을를은는이가의로와과에별도]|$)", "semi-colloquial vocabulary (replace with the formal words of the references/lexicon.md replacement table)"),
             (r"(입력값|출력값|인계|게이트|관문|입력(이|으로|되| ?자료)|출력(이|으로|물)|연결을 확보|매핑|인터페이스|파이프라인|연계 구조|산출물 연계)", "software design phrasing (in execution plans outside architecture descriptions, convert to '활용 자료', '~에 활용', '대응표')"),
             (r"(과업의 답|답이 되|답을 제시|전제 ?조건|누락 ?없이|빠짐 ?없이|대응 관계|과업의 핵심 ?답|과업 ?방향|수행 ?여건|일정 ?정합|근거 ?누적|Task 배치|도달 ?상태|도달 ?시점|분석형|조사형|도출형|지원형|수행내용 구성|열 번호|행 번호|(?<![가-힣])열은 |주차 단위|N주 단위)", "framework language and internal judgment exposed (the skill's design vocabulary and the author's checking language. Describe the resulting content directly)"),
             (r"(Task|Module|모듈)\s*\d+(-\d+)?(\s*(의|에서|을|를|과|와|은|는|에|로|으로)\s|\s+[가-힣])|(?<![가-힣])(앞|위|이전|전술|후술|다음|아래|해당|상기|하기)\s*(절|장|항)(?=[\s.,)에의을를은는이가과와로도만까부]|$)|(앞|이전|전술|후술|해당|상기)\s*(모듈|Task|Module)|(을|를)\s*참조|참조하", "cross reference (do not point to other sections, modules, or Tasks by number. Restate the needed content in place. Number references are allowed only in cells of the overview table, quick-reference table, schedule chart, and division of duties table)"),
             (r"\d+\s*개\s*(Module|Task|모듈|영역|축|전략|방안|과제|단계|원칙|방향|관점|부문|파트)", "count declaration (declaring the count of your own structure. Present the items themselves)"),
             (r"\d+\s*종(?!류|합|료|목|사)", "N-kinds declaration (list material and target names by their real names. Allowed only when a list table is in the same place)"),
             (r"(쓰임새|쓰임|짜임새|얼개|뼈대|밑그림|출발점|손질|채비|매만지|다듬|목표 ?값|(일정|보고|회의|과업)[을를의]? ?동기화|일정[을를]? ?정렬)", "unconventional vocabulary (use the replacements in references/lexicon.md)"),
             (r"\d+\s*(대|가지)\s*(전략|방안|과제|핵심|원칙|방향)", "count declaration (present the items themselves without declaring a count. If it quotes an external plan verbatim, record it as a misjudgment reason)"),
             (r"본 (과업|연구|제안|사업|용역|제안서)", "\"this engagement\" declaration (enter with content without self-reference. If a reference is needed, '과업' or the real project name)"),
             (r"(발주처|주관기관|주관부처)[^.\n]{0,20}(협조|제공|지원)[^.\n]{0,12}(전제|조건)|(을|를)\s*전제로|(확보|제공)(되어야|돼야)\s*(수행|추진|착수)", "precondition statement (a sentence that sets conditions for the client. Describe how to obtain materials as the proposer's procedure)"),
             (r"(범위|과업|연구)(에|에서)\s*(포함하지|포함되지|다루지)\s*않|제외\s*범위", "exclusion declaration (do not declare exclusions the client did not ask about. Reflect only exclusions the RFP states in the scope table)"),
             (r"(RFP|제안요청서|과업지시서)(가|는|이|에서)\s*(제시|요구|명시|규정|정)", "RFP as subject (do not make the RFP the subject of a sentence. Describe the presented content itself)"),
             (r"(아래|위|다음|상기|하기|해당)\s*(표|그림|도표)(는|은|에서|와 같이|과 같이)", "table-pointing meta statement (state what the table says as ❍ and place the table. Delete '아래 표는')"),
             (r"(첫째|둘째|셋째|넷째|다섯째|여섯째|첫\s*번째|두\s*번째|세\s*번째|네\s*번째|다섯\s*번째|하나는|다른 하나는|마지막으로)", "ordinal enumeration (do not chain with 첫째, 둘째. Give each item a ❍ or convert to a table)"),
             (r"(제안요청서|RFP|과업지시서|요구사항|요구항목|평가항목|평가기준)(의|에|에서)?\s*(요구|내용|취지|목적|기준)?(에|을|를|과|와)?\s*(대응|부합|충실|준수|맞추|맞게|따라|반영하여|고려하여)", "RFP compliance declaration (use the original requirement wording directly in Task names and descriptions)"),
             (r"것(임|이다|입니다|이었음)?[.)]?\s*$", "colloquial '것' ending ('남은 것', '해야 할 것', '과업의 포인트인 것'. Convert to nominal endings: '~ 잔존', '~ 필요', '~이 과업의 핵심')"),
             (r"(는데|은데|한데|인데|(?<!가)운데|더라|거든|잖아)[.)]?\s*$", "colloquial ending ('-는데' and '-인데' are explanatory endings. Break the sentence and end with a nominal form)"),
             (r"(는|은|을|를|할|될|된|인|한)\s*(게|거|걸|건)(?!\s*\d)(\s|$)", "colloquial contraction ('~하는 게' → '~하는 것이' or a nominal form, '~할 거' → '~ 예정', '~하는 건' → '~은')"),
             (r"\s(듯함|듯하다|듯|셈|편|만함|터)[.)]?\s*$|(로|고|다고)\s*(봄|보임|보여짐|여겨짐)[.)]?\s*$|(?<![가-힣])(식으로|쪽으로|쪽에서)\s", "colloquialism mistaken for formal style ('~듯', '~인 셈', '~하는 편', '~로 봄', '~식으로', '~쪽'. Judge with grounds or convert to '~로 판단', '~로 확인')"),
             (r"(^|\s)(일단|약간|좀|꽤|딱|그냥|어쨌든|아무튼|나름)\s", "colloquial adverb (delete or replace with figures or criteria)"),
             (r"^```", "code block fence (formulas go in two-column tables)")]
# Periods on hierarchy lines. One sentence per level item. Joining two sentences with a period or ending a line with a period is judged FAIL
# Targets are □, ❍(-), evidence (+), and ※ lines. Table cells, footnotes, source lines, and captions are excluded because dates and bibliographic notation use periods
HIER_LINE_RE = re.compile(r"^\s*(□|[-+*]\s|※)")
PERIOD_EXEMPT_RE = re.compile(r"^\s*(\[\^|※\s*출처|출처\s*[:：]|<표>|!\[)")
PERIOD_MID_RE = re.compile(r"[가-힣)\]\"”'’%]\.\s+[가-힣(A-Za-z\d]")   # the form "확정. 이후". Decimal points (3.5) and URLs are excluded
PERIOD_END_RE = re.compile(r"[가-힣)\]\"”'’%]\.\s*$")
# colloquial dependent nouns (듯, 셈, 편, 터, 만함) require a preceding space. Prevents conflicts with nominal endings such as "개편", "재편", "데이터", "센터"
TONE_WARN = [(r"(최대|최초|유일|최고|최상)의?\s", "superlative assertion. Check the comparison grounds"),
             (r"\bRFP\b|제안요청서|과업지시서", "RFP mention. Outside the header of the requirement quick-reference table, describe engagement content. If an organization must be named, use its real name"),
             (r"(획기적|혁신적|선도적|압도적|진정한)", "exaggerating modifier")]
# Agency naming verdict. Why: '발주처' is a working term the proposer side uses for the counterpart. Carried into manuscripts it makes supplier-view sentences that set the reader as a third party, so the real name is prompted
CLIENT_TERM = r"발주\s?(처|기관|부서)"
CLIENT_SUBJECT_RE = re.compile(CLIENT_TERM + r"(는|은|가|이|에서는|께서는)(?![가-힣])")
CLIENT_TERM_RE = re.compile(CLIENT_TERM)
# column names of track record tables are conventional notation for the contract counterpart of past engagements, so they are excluded from the verdict
CLIENT_HEADER_CELL_RE = re.compile(r"\|\s*" + CLIENT_TERM + r"\s*(?=\|)")
CLIENT_SLOT_RE = re.compile(r"\[기관명\]")
CLIENT_SUBJECT_LABEL = "client as subject (delete the subject and write target and state in the predicate. If a reference is needed, the real name from the Agency Naming section of 00_제안전략서.md)"
CLIENT_TERM_LABEL = "leftover '발주처' naming. Write the organization's real name (Agency Naming section of 00_제안전략서.md)"
CLIENT_SLOT_LABEL = "leftover agency naming slot (replace with the notation of the Agency Naming section of 00_제안전략서.md)"
SUM_LABEL_RE = re.compile(r"^\s*(합\s*계|계|총\s*계|소\s*계|합계)\s*$")
NUM_CELL_RE = re.compile(r"^[\d,]+(\.\d+)?%?$")
STAGES = ["strategy", "rfp", "ledger", "core", "common", "check", "build"]
CHECK_NAMES = ("rfp", "density", "volume", "wording", "tone", "numbers", "fill")   # names of the individual checks (targets of check all --skip)
# check all targets per type. Density (per Task) and common part substitution are proposal-only. Wording comes from the key wording dictionary of the proposal strategy file
CHECKS_BY_KIND = {"proposal": ("rfp", "density", "volume", "wording", "tone", "numbers", "fill"),
                  "report": ("rfp", "volume", "tone", "numbers", "fill"),
                  "plan": ("volume", "tone", "numbers", "fill")}   # fill for reports and plans only warns on leftover [확인]
# check tone whitelist sources. The key wording dictionary (00_제안전략서.md) and the Module and Task names of the module overview table (02_모듈총괄표.md).
# Original RFP terms and framework names the user confirmed are not tone verdict targets. Fixed in state.json at the first check tone to block later arbitrary additions
# The module overview table is a step 2 design output like 02_섹션대장 and 02_스토리라인, so it sits at the workspace root.
# Why: as a numbered file inside 원고/, the table's position would be tied to file name order and fall under an earlier chapter (proposer general status, and so on) in a specified TOC.
# The table's position is set by the insertion directive line `(삽입: 모듈총괄표)` in the manuscript, and build turns it into the table when merging. The table has one original
F_MODULE_TABLE = "02_모듈총괄표.md"
MODULE_TABLE_SUFFIX = "모듈총괄표.md"   # identifying suffix for compatibility with workspaces of the old method (원고/50_모듈총괄표.md)
INSERT_MODULE_RE = re.compile(r"^\(\s*삽입\s*[:：]\s*모듈총괄표\s*\)$")
WHITELIST_MASK = "◇"
WHITELIST_MIN_LEN = 4   # masking 2-3 character common nouns ("조사", "매핑") hides other verdicts such as "조사형", so they are excluded
TABLE_SEP_RE = lib.TABLE_SEP_RE
Element = lib.Element


# ═══════════════════════════════════════════════════════════════════════
# 1. Manuscript parsing: adds only Module/Task bands on top of the typesetter's parser
# ═══════════════════════════════════════════════════════════════════════
def parse_markdown(path: str) -> tuple[list[Element], dict[str, str]]:
    """Put proposal notation on top of lib.parse_markdown.
    - `(폭: ...)` table width directive lines are stripped before passing to the lib parser and recorded against the next table's sequence (so caption attachment does not break even when placed between caption and table)
    - `Module 1. title` and `Task 1-1. title` lines are turned into band elements"""
    raw = open(path, encoding="utf-8").read().split("\n")
    kept, specs, longs, table_idx, in_table = [], {}, {}, 0, False
    for ln in raw:
        st = ln.strip()
        if (m := WIDTH_SPEC_RE.match(st)):
            specs[table_idx] = m.group(1)
            continue
        if (m := LONG_SPEC_RE.match(st)):
            longs[table_idx] = (m.group(1), int(m.group(2)) if m.group(2) else None)
            continue
        is_row = st.startswith("|")
        if is_row and not in_table:
            table_idx += 1
        in_table = is_row
        kept.append(ln)
    src = path
    if specs or longs:  # image relative paths are based on the manuscript folder, so a temporary copy is placed in the same folder
        src = os.path.join(os.path.dirname(os.path.abspath(path)), f".{os.path.basename(path)}.tmp")
        open(src, "w", encoding="utf-8").write("\n".join(kept))
    try:
        elems, footnotes = lib.parse_markdown(src)
    finally:
        if src != path and os.path.exists(src):
            os.remove(src)
    ti = 0
    for e in elems:
        if e.kind == "table":
            if ti in specs:
                e.data["width_spec"] = specs[ti]
            if ti in longs:
                e.data["long_spec"] = longs[ti]
            ti += 1
        elif e.kind in ("circle", "square") and (m := TASK_BAND_RE.match(e.text)):
            e.kind = "band"
            e.data["kind"] = "module" if m.group(1).lower() == "module" else "task"
    return elems, footnotes


def estimate_table_lines(rows: list[list[str]]) -> int:
    """Approximate the table height in lines. Column widths are allocated in proportion to each column's maximum weighted length, and cell line counts sum the folded lines of each <br> segment."""
    ncol = max(len(r) for r in rows)
    need = [0] * ncol
    for r in rows:
        for c, v in enumerate(r):
            need[c] = max(need[c], max((lib._weighted_len(s.strip()) for s in v.split("<br>")), default=0), 4)
    total = sum(need) or 1
    caps = [max(6, TABLE_LINE_UNITS * n // total) for n in need]
    lines = 0
    for r in rows:
        row_lines = 1
        for c, v in enumerate(r):
            segs = [s.strip() for s in v.split("<br>")] or [""]
            row_lines = max(row_lines, sum(max(1, -(-lib._weighted_len(s) // caps[c])) for s in segs))
        lines += row_lines
    return lines


def lint(elems: list[Element]) -> list[str]:
    out = lib.lint(elems) + lib.lint_counts(elems) if hasattr(lib, "lint_counts") else lib.lint(elems)
    for e in elems:
        if e.kind != "table" or e.data.get("long_spec"):
            continue
        rows = e.data.get("rows") or []
        est = estimate_table_lines(rows)
        if est >= TABLE_PAGE_LINES and len(rows) <= lib.LONG_TABLE_ROWS:
            out.append(f"[long table HITL] '{e.text[:24]}' {len(rows)} rows estimated at {est} lines: may run past one page. Ask the user to choose (장표: engine) treat-as-character off or (장표: split N) physical split")
    return out


# ═══════════════════════════════════════════════════════════════════════
# 2. Typesetting: format-agnostic judgments (ProposalLayout) and writers per format (ProposalBuilder for hwpx, ProposalDocxBuilder for docx)
# ═══════════════════════════════════════════════════════════════════════
class ProposalLayout:
    """Typesetting judgments for proposals, pre-diagnosis reports, and planning reports. The hwpx and docx writers both inherit it and record the same results
    Judgment targets: heading numbers and page breaks, band composition, caption text and numbers, table width and long-table directives, notes below tables, element traversal
    Record functions the writers implement: para, heading_para, band_emit, chapter_started, add_footnote, after_table, after_image,
    _front_matter, set_page_number, set_header, save"""

    def _cell_line_role(self, line: str) -> str | None:
        """A cell line with only one symbol (○ in a comparison matrix, and so on) is not treated as a bullet.
        Why: lib judges lines starting with ○ as cell bullets, removes the symbol, and applies the bullet style. Size and alignment then mismatched ● and ◐."""
        if len(line.strip()) <= 1:
            return None
        return super()._cell_line_role(line)

    def heading(self, level: int, text: str):
        """Roman chapters (Ⅰ.), Arabic sections (1.), 가나다 subsections (가.), items (1)). Numbers are written in the body and outline numbers are turned off.
        The report profile uses 제1장, 1.1, 1.1.1, 가. The profile sets the format strings"""
        text = lib.NUM_PREFIX_RE.sub("", text).strip()
        role = ("chapter", "section", "subsection", "ganada")[level - 1]
        pt = None
        if level == 1:
            self.n.update(chapter=self.n["chapter"] + 1, section=0, subsection=0, ganada=0)
            # the document's first chapter gets no page break. The same for partial manuscripts whose chapter number is not 1 (--chapter-start)
            if self.page_break in ("chapter", "section") and self.n.get("chapters_done", 0) > 0:
                self.pending_break = True
            self.n["chapters_done"] = self.n.get("chapters_done", 0) + 1
            c = self.n["chapter"]
            label = self.p["chapter_format"].format(n=c, r=ROMAN[(c - 1) % len(ROMAN)], t=text)
            pt = self.p.get("chapter_pt")
        elif level == 2:
            self.n.update(section=self.n["section"] + 1, subsection=0, ganada=0)
            label = self.p["section_format"].format(c=max(self.n["chapter"], 1), s=self.n["section"], t=text)
            if self.page_break == "section" and not (self.n["section"] == 1 and self.n["chapter"] > 0) and self.n["total_section"] > 0:
                self.pending_break = True
            self.n["total_section"] += 1
            pt = self.p.get("section_pt")
        elif level == 3:
            self.n.update(subsection=self.n["subsection"] + 1, ganada=0)
            ss = self.n["subsection"]
            label = self.p["subsection_format"].format(c=max(self.n["chapter"], 1), s=max(self.n["section"], 1), ss=ss, g=lib.letter_label(ss), t=text)
            pt = self.p.get("subsection_pt")
        else:
            self.n["ganada"] += 1
            i = self.n["ganada"]
            label = self.p.get("item_format", "{g} {t}").format(i=i, g=lib.letter_label(i), t=text)
        if level == 1:  # chapter titles are bands like Module and Task. The label is recorded as is for TOC and page number matching
            self.cap_n = {"TABLE": 0, "PICTURE": 0}
            band_label = self.p.get("chapter_label", "{r}").format(n=c, r=ROMAN[(c - 1) % len(ROMAN)])
            if getattr(self, "divider", False):  # divider page: a page holding only the chapter title. The body starts from the first section on the next page
                self.pending_break = False
                p = self.divider_emit(self.band_spec(lib.Element("band", label, data={"kind": "divider", "label": band_label, "title": text})),
                                      new_page=self._body_started)
                self.after_block = True
                self.pending_break = True
            else:
                p = self.band(lib.Element("band", label, data={"kind": "chapter", "label": band_label, "title": text}))
        else:
            p = self.heading_para(role, label, pt)
        if level <= 3:
            self.headings.append((level, label))
        if level == 1:
            self.chapter_started(p)

    def band_spec(self, e: Element) -> dict:
        """Width and height of the one-row, three-cell band (number cell, empty cell, title cell), and per cell (text, size, text color, alignment, fill, border)"""
        kind = e.data["kind"]
        label, title = e.data.get("label"), e.data.get("title")
        if label is None:  # separate the number and title from a Module or Task manuscript line
            m = TASK_BAND_RE.match(e.text)
            label, title = (f"{m.group(1)} {m.group(2).rstrip('.')}", m.group(3).strip()) if m else ("", e.text)
        st = {**BAND_STYLE[kind], **(self.p.get("chapter_band", {}) if kind == "chapter" else {})}
        color = {**THEMES[getattr(self, "theme", DEFAULT_THEME)], "white": "#FFFFFF"}.get
        return {"widths": [st["label_w"], BAND_GAP_W, self.table_max_w - st["label_w"] - BAND_GAP_W],
                "height": int(lib.ROW_BASE_H * st["rows"]),
                "cells": [(label, st["label_pt"], color(st["label_color"]), "CENTER",
                           color(st["fill"]) if st["fill"] else None, color(st["border"]) if st["border"] else None),
                          ("", st["title_pt"], "#000000", "LEFT", None, None),
                          (title, st["title_pt"], "#000000", "LEFT", None, None)]}

    def band(self, e: Element):
        """Place a chapter, Module, or Task title as a one-row, three-cell table. Returns the paragraph holding the band (or the first paragraph) for restarting the chapter's caption numbers."""
        p = self.band_emit(self.band_spec(e))
        self.after_block = True
        return p

    def caption_label(self, kind: str, text: str) -> dict:
        """Caption text. Unnumbered profiles use <title> for tables and [title] for figures. Numbered profiles use <표 1-3> title and [그림 1-3] title
        Returns: open (before the number), num (the number, None if absent), close (after the number), text (the whole). Numbered captions are registered in the list of tables and list of figures
        The word and brackets are the caption labels (hwpx_report.py LABEL_DEFAULTS)"""
        prefix = self.p.get("caption_prefix")
        cap = lib.label("caption.table" if kind == "TABLE" else "caption.figure")
        if prefix is None:
            label = f"{cap['open']}{text}{cap['close']}"
            return {"open": label, "num": None, "close": "", "text": label}
        self.cap_n[kind] = self.cap_n.get(kind, 0) + 1
        num = self.cap_n[kind]
        open_b, close_b = f"{cap['open']}{cap['label']} ", f"{cap['close']} "
        head = open_b + prefix.format(c=max(self.n["chapter"], 1))
        full = f"{head}{num}{close_b}{text}"
        self.captions.append((kind, full, close_b.strip() + text))
        return {"open": head, "num": num, "close": close_b + text, "text": full}

    def toc_lists(self) -> list[tuple[str, list[str]]]:
        if not self.p.get("toc_lists"):
            return []
        return [(lib.label("toc.tables"), [t for k, t, _ in self.captions if k == "TABLE"]),
                (lib.label("toc.figures"), [t for k, t, _ in self.captions if k == "PICTURE"])]

    def unify_table_widths(self, elems: list[Element]):
        """Turn table width directives into column widths. A directive is a `(폭: ...)` line in the manuscript that the model sets qualitatively after checking whether a layout repeats.
        (폭: group name) tables with the same group name share a width computed once over all rows of the group (based on the largest needed width, so narrow widths do not propagate)
        (폭: 3,2,2) distributes the table width cap by ratio. (폭: 3,2,2, 80%) distributes 80% of the cap by ratio
        Tables without a directive are computed per table by the lib rules. Only with `--unify auto` are tables with the same header or first-column labels grouped automatically
        Why manual first: which tables share a layout is a document design judgment, and machine judgment creates mistakes."""
        tables = [e for e in elems if e.kind == "table" and e.data.get("rows")]
        groups: dict[tuple, list[Element]] = {}
        for e in tables:
            spec = e.data.get("width_spec")
            if not spec:
                continue
            parts = [x.strip() for x in spec.split(",")]
            if all(re.fullmatch(r"\d+(\.\d+)?%?", x) for x in parts):
                ratio = [float(x) for x in parts if not x.endswith("%")]
                pct = next((float(x[:-1]) / 100 for x in parts if x.endswith("%")), 1.0)
                ncol = len(e.data["rows"][0])
                if len(ratio) != ncol:
                    print(f"WARN: (폭: {spec}) ratio count {len(ratio)} differs from the column count {ncol}. Replaced with per-table calculation")
                    continue
                total = int(self.table_max_w * pct)
                widths = [int(total * r / sum(ratio)) for r in ratio]
                widths[-1] += total - sum(widths)
                e.data["col_widths"] = widths
            else:
                groups.setdefault(("G", spec), []).append(e)
        if getattr(self, "unify_mode", "manual") == "auto":
            def header_key(e): return ("H", tuple(c.strip() for c in e.data["rows"][0]))
            def label_key(e): return ("L", len(e.data["rows"][0]), tuple(r[0].strip() for r in e.data["rows"] if r))
            from collections import Counter
            free = [e for e in tables if not e.data.get("width_spec")]
            hc, lc = Counter(header_key(e) for e in free), Counter(label_key(e) for e in free)
            for e in free:
                if hc[header_key(e)] >= 2:
                    groups.setdefault(header_key(e), []).append(e)
                elif lc[label_key(e)] >= 2:
                    groups.setdefault(label_key(e), []).append(e)
        for key, group in groups.items():
            # tables with a different column count (module overview tables of modules with different Task counts, and so on) are unified with tables of the same column count
            by_ncol: dict[int, list] = {}
            for t in group:
                by_ncol.setdefault(len(t.data["rows"][0]), []).append(t)
            for sub in by_ncol.values():
                widths = lib.calc_col_widths([r for t in sub for r in t.data["rows"]], self.table_max_w)
                for t in sub:
                    t.data["col_widths"] = widths

    def table(self, e: Element):
        """Apply table width and long-table directives, then call the writer's table record"""
        forced = e.data.get("col_widths")
        original = lib.calc_col_widths
        if forced:
            lib.calc_col_widths = lambda rows, max_w: list(forced)  # use the group's common width instead of the lib calculation
        long_spec = e.data.get("long_spec")  # (장표: engine|split N) directive: applies long-table handling to this table only, regardless of row count
        saved_rows, saved_mode = lib.LONG_TABLE_ROWS, self.long_table_mode
        if long_spec:
            mode, chunk = long_spec
            self.long_table_mode = mode
            lib.LONG_TABLE_ROWS = 0 if mode == "engine" else (chunk or 12) + 1
        try:
            r = super().table(e)
        finally:
            lib.calc_col_widths = original
            lib.LONG_TABLE_ROWS, self.long_table_mode = saved_rows, saved_mode
        self.after_table(e)
        return r

    def image(self, e: Element):
        r = super().image(e)
        self.after_image(e)
        return r

    @staticmethod
    def special_cell(val: str) -> str | None:
        """A cell whose value is a single ■ is a period cell of the schedule chart. Conventionally the period is shown by fill instead of text. Returns the fill color"""
        return SCHEDULE_FILL if val.strip() == SCHEDULE_MARK else None

    def build(self, elems: list[Element], footnotes: dict, out: str):
        """Element traversal. Notes and source lines right after tables and figures are attached as bottom rows of the block, and source lines after body text move into a footnote of the preceding paragraph"""
        self.unify_table_widths(elems)
        self.last_para = None
        self._body_started = False   # whether the first body element was recorded. Used for page break decisions of divider pages
        if getattr(self, "doc_title", None):
            self.para("doc_title", self.doc_title)
            self._body_started = True
        skip = 0
        for idx, e in enumerate(elems):
            if skip:
                skip -= 1
                continue
            if e.kind in ("table", "image"):
                srcs, consumed = [], 0
                for nx in elems[idx + 1:]:
                    if nx.kind in ("note", "source"):  # source lines as is, other notes keep the ※ marker
                        text = inline_note_footnotes(nx.text, footnotes)
                        srcs.append(text if nx.kind == "source" or lib.SOURCE_RE.match(text) else NOTE_PREFIX + text)
                    elif nx.kind == "unit" and e.kind == "table" and not e.data.get("unit"):
                        e.data["unit"] = nx.text
                    else:
                        break
                    consumed += 1
                e.data["sources"] = srcs
                skip = consumed
            if e.kind == "heading":
                self.heading(e.level, e.text)
            elif e.kind == "band":
                self.band(e)
            elif e.kind in ("square", "circle", "dash", "dot", "note"):
                self.last_para = self.para(e.kind, e.text, footnotes)
            elif e.kind == "unit":
                self.para("unit", e.text)
            elif e.kind == "source":
                if self.last_para is not None:
                    self.add_footnote(self.last_para, e.text)
                else:
                    self.para("note", e.text, footnotes)
            elif e.kind == "table":
                self.table(e)
            elif e.kind == "image":
                self.image(e)
            elif e.kind == "pagebreak":
                self.pending_break = True
            if e.kind != "pagebreak":
                self._body_started = True
        if self.toc or self.cover:
            self._front_matter(self.cover)
        self.set_page_number(self.page_number)
        self.set_header(self.header_spec)
        self.save(out)


class ProposalBuilder(ProposalLayout, lib.Builder):
    """hwpx writer. Judgments by ProposalLayout, recording by python-hwpx"""

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.cap_n = {"TABLE": 0, "PICTURE": 0}
        self.captions: list[tuple[str, str, str]] = []

    def _sized_char(self, height: int, bold: bool) -> str:
        """charPr that changes only size and weight based on the chapter style font (heading hierarchy, band text)."""
        return self._cover_char(height, bold)

    def _border_fill(self, sides: list[str], fill: str | None = None) -> str:
        """Normalize hatchColor to 6 digits in lib's border definition.
        Why: python-hwpx writes hatchColor with 8 digits (#FF000000), but files authored in 한/글 use only 6 digits. Prevents wrong fill color display."""
        bid = super()._border_fill(sides, fill)
        el = self.header.element.find(f".//{HH}borderFill[@id='{bid}']")
        if el is not None:
            HC = "{http://www.hancom.co.kr/hwpml/2011/core}"
            for wb in el.iter(f"{HC}winBrush"):
                if len(wb.get("hatchColor", "")) == 9:
                    wb.set("hatchColor", "#" + wb.get("hatchColor")[3:])
        return bid

    def attach_caption(self, holder, kind: str, text: str, side: str, width: int):
        """If caption_prefix is None, attach an unnumbered caption <title> to the object (proposal convention). If set, lib's automatically numbered caption."""
        label = self.caption_label(kind, text)
        if label["num"] is not None:
            return super().attach_caption(holder, kind, text, side, width)
        role = "table_title" if kind == "TABLE" else "fig_title"
        r = self.refs(role)
        cap = holder.set_caption("", side=side, width=width, gap=lib.CAPTION_GAP, char_pr_id_ref=r["char_pr_id_ref"])
        cp = cap.paragraphs[0]
        cp.element.set("paraPrIDRef", str(r["para_pr_id_ref"]))
        cp.element.set("styleIDRef", str(r["style_id_ref"]))
        for run in [c for c in cp.element if c.tag == f"{HP}run"]:
            cp.element.remove(run)
        cp.add_run(label["text"], char_pr_id_ref=r["char_pr_id_ref"])

    def _band_char(self, height: int, color: str) -> str:
        """charPr that changes size, weight, and text color based on the chapter style font (bands only. White text for chapter bands)."""
        base = str(self.refs("chapter")["char_pr_id_ref"])

        def mod(e):
            e.set("height", str(height))
            e.set("textColor", color)
            if e.find(f"{HH}bold") is None:
                anchor = e.find(f"{HH}offset")
                e.insert(list(e).index(anchor) + 1, etree.Element(f"{HH}bold"))
        font = self.header.element.find(f".//{HH}charPr[@id='{base}']").find(f"{HH}fontRef").get("hangul")
        return self.header.ensure_char_property(
            predicate=lambda e: e.get("height") == str(height) and e.get("textColor") == color
            and e.find(f"{HH}fontRef").get("hangul") == font and e.find(f"{HH}bold") is not None,
            modifier=mod, base_char_pr_id=base).get("id")

    def heading_para(self, role: str, label: str, pt: int | None):
        p = self.para(role, label, no_heading=True)
        if pt:
            cp = self._sized_char(pt, True)
            for run in p.element.iter(f"{HP}run"):
                run.set("charPrIDRef", cp)
        return p

    def chapter_started(self, p):
        if self.p.get("caption_prefix"):  # restart from <표 N-1> per chapter
            p.add_new_num(kind="TABLE", number=1)
            p.add_new_num(kind="PICTURE", number=1)

    def band_emit(self, spec: dict):
        """Put a one-row, three-cell table in the band position paragraph. Why a table: table cell fill displays more consistently across renderers than paragraph shading"""
        self.gap()
        r = self._centered_refs(BAND_POS_ROLE)
        attrs = {}
        if self.pending_break:
            attrs["pageBreak"] = "1"
            self.pending_break = False
        p = self.doc.add_paragraph("", **r, **attrs)
        t = p.add_table(1, len(spec["cells"]), width=self.table_max_w)
        t.element.set("numberingType", "NONE")  # bands are excluded from the table number count (prevents contaminating automatic caption numbers)
        t.set_column_widths(spec["widths"])
        ref = self.refs("band_cell")
        height = spec["height"]
        empty = self._band_fill(None, None)
        fills = [self._band_fill(fill, border) if (fill or border) or ci == 0 else empty
                 for ci, (_, _, _, _, fill, border) in enumerate(spec["cells"])]
        for ci, (text, pt, color, align, _, _) in enumerate(spec["cells"]):
            t.set_cell_text(0, ci, text)
            cell = t.cell(0, ci)
            cp = self._band_char(pt, color)
            for pe in cell.element.iter(f"{HP}p"):
                pe.set("paraPrIDRef", self._aligned_pp(str(ref["para_pr_id_ref"]), align))
                pe.set("styleIDRef", str(ref["style_id_ref"]))
                for run in pe.iter(f"{HP}run"):
                    run.set("charPrIDRef", cp)
            cell.element.set("borderFillIDRef", fills[ci])
            cell.set_size(height=height)
            sub = cell.element.find(f"{HP}subList")
            if sub is not None:
                sub.set("vertAlign", "CENTER")
        sz = t.element.find(f"{HP}sz")
        if sz is not None:
            sz.set("height", str(height))
        pos = t.element.find(f"{HP}pos")
        if pos is not None:
            pos.set("horzAlign", "CENTER")
        t.mark_dirty()
        return p

    def divider_emit(self, spec: dict, new_page: bool):
        """Divider page: a divider band after an empty 바탕글 line. The header, footer, and page number of this page are hidden (한/글 hide)"""
        body = self.refs("body")
        first = None
        for i in range(DIVIDER_TOP_LINES):
            q = self.doc.add_paragraph("", **body, **({"pageBreak": "1"} if i == 0 and new_page else {}))
            first = first or q
        first.add_page_hiding(header=True, footer=True, page_num=True)
        return self.band_emit(spec)

    def _band_fill(self, fill: str | None, border: str | None) -> str:
        """Fill and border of band cells. The border color is set separately from table rules."""
        sides = ["top", "bottom", "left", "right"] if border else []
        bid = str(self.header.ensure_border_fill(border_color=border or lib.BORDER_COLOR, border_width=BAND_BORDER_W,
                                                 active_borders=sides, fill_color=fill))
        el = self.header.element.find(f".//{HH}borderFill[@id='{bid}']")
        if el is not None:
            for side in ("left", "right", "top", "bottom"):
                b = el.find(f"{HH}{side}Border")
                if b is not None:
                    b.set("type", "SOLID" if side in sides else "NONE")
            HC = "{http://www.hancom.co.kr/hwpml/2011/core}"
            for wb in el.iter(f"{HC}winBrush"):
                if len(wb.get("hatchColor", "")) == 9:
                    wb.set("hatchColor", "#" + wb.get("hatchColor")[3:])
        return bid

    def _last_tbl(self):
        tbls = list(self.doc.sections[-1].element.iter(f"{HP}tbl"))
        return tbls[-1] if tbls else None

    def after_table(self, e: Element):
        """Turn ■ cells into shaded cells, and in numbered caption profiles remove tables without captions from the number count (so list of tables numbers match 한/글 automatic numbers)"""
        tbl = self._last_tbl()
        if tbl is None:
            return
        if self.p.get("caption_prefix") and not e.text and tbl.get("numberingType") == "TABLE":
            tbl.set("numberingType", "NONE")
        rows = e.data.get("rows") or []
        if not any(self.special_cell(c) for row in rows for c in row):
            return
        ncol = int(tbl.get("colCnt"))
        for tc in tbl.iter(f"{HP}tc"):
            texts = [t for t in tc.iter(f"{HP}t")]
            fill = self.special_cell("".join((t.text or "") for t in texts))
            if not fill:
                continue
            for t in texts:
                t.text = ""
            # Why: shaded cells must also follow the table's outer left and right rule so no vertical lines appear on the first and last columns
            c0 = int(tc.find(f"{HP}cellAddr").get("colAddr"))
            c1 = c0 + int(tc.find(f"{HP}cellSpan").get("colSpan")) - 1
            tc.set("borderFillIDRef", self._border_fill(lib.edge_sides(c0, c1, ncol), fill))

    def after_image(self, e: Element):
        tbl = self._last_tbl()
        if tbl is not None and self.p.get("caption_prefix") and not e.text and tbl.get("numberingType") == "PICTURE":
            tbl.set("numberingType", "NONE")

    def add_footnote(self, p, text: str):
        p.add_footnote(text)

    def save(self, out: str):
        report = self.doc.validate()
        issues = getattr(report, "issues", [])
        if issues:
            print(f"WARN: {len(issues)} schema validation issues")
            for it in list(issues)[:5]:
                print("   ", it)
        self.doc.save_to_path(os.path.abspath(out))
        self.doc.close()


class ProposalDocxBuilder(ProposalLayout, docx_lib.DocxBuilder):
    """docx writer. Judgments by ProposalLayout, recording by docx_report.DocxBuilder"""

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.cap_n = {"TABLE": 0, "PICTURE": 0}
        self.captions: list[tuple[str, str, str]] = []


def load_state(work: str) -> dict:
    path = os.path.join(work, F_STATE)
    if os.path.exists(path):
        return json.load(open(path, encoding="utf-8"))
    return {"created": datetime.now().isoformat(timespec="seconds"), "stages": {}, "checks": {}, "files": {}}


def save_state(work: str, st: dict):
    st["updated"] = datetime.now().isoformat(timespec="seconds")
    json.dump(st, open(os.path.join(work, F_STATE), "w", encoding="utf-8"), ensure_ascii=False, indent=2)


def record(work: str, key: str, passed: bool, detail: dict):
    st = load_state(work)
    st["checks"][key] = {"pass": passed, "at": datetime.now().isoformat(timespec="seconds"), **detail}
    save_state(work, st)


LABELS_NONE = "none"   # --labels none: clear the label file recorded in the workspace state.json


def use_labels(work: str | None, opt: str | None) -> str | None:
    """Apply the label file. A --labels value is recorded in the workspace state.json. Without a value the recorded file is used. none clears the record
    Why record it: check (recognition of manuscript notation) and build (document labels) must read the manuscripts with the same labels, or verdicts and typesetting diverge"""
    path = None
    if opt and opt != LABELS_NONE:
        path = os.path.abspath(opt)
        if not os.path.exists(path):
            sys.exit(f"ERROR: label file not found {opt}")
    if work:
        st = load_state(work)
        if opt == LABELS_NONE and "labels" in st:
            st.pop("labels")
            save_state(work, st)
        elif path and st.get("labels") != path:
            st["labels"] = path
            save_state(work, st)
        elif not opt:
            path = st.get("labels")
            if path and not os.path.exists(path):
                sys.exit(f"ERROR: label file recorded in state.json not found {path}. Set it again with a --labels path or clear it with --labels none")
    if path:
        changed = lib.load_labels(path)
        print(f"labels: {path} ({len(changed)} keys overridden)")
    return path


def manuscript_kind(path: str) -> str:
    """File name without the sequence prefix. The first character is the check type (3, 5 original. 4, 6 common)"""
    return ORDER_PREFIX_RE.sub("", os.path.basename(path))


def manuscript_files(work: str, prefix: str = "") -> list[str]:
    """The merge order is the order of full file names. prefix applies to the name without the sequence prefix"""
    files = sorted(glob.glob(os.path.join(work, D_MANUSCRIPT, "*.md")))
    return [f for f in files if manuscript_kind(f).startswith(prefix)]


def legacy_module_table(work: str) -> str | None:
    found = [f for f in manuscript_files(work) if manuscript_kind(f).endswith(MODULE_TABLE_SUFFIX)]
    return found[0] if found else None


def module_table_path(work: str) -> str:
    root = os.path.join(work, F_MODULE_TABLE)
    return root if os.path.exists(root) else (legacy_module_table(work) or root)


def caption_start(line: str) -> bool:
    """Whether a line starts with the table caption notation (<표> title). The token is the notation.table label (default 표 plus tokens added by a label file)"""
    st = line.strip()
    return any(st.startswith(f"<{t}>") for t in lib.label("notation.table"))


def module_table_block(work: str) -> str:
    """Take from the caption (<표>) through the last table row in 02_모듈총괄표.md. Includes the (장표: ) and (폭: ) directive lines between caption and table."""
    lines = open(os.path.join(work, F_MODULE_TABLE), encoding="utf-8").read().split("\n")
    rows = [i for i, ln in enumerate(lines) if ln.strip().startswith("|")]
    if not rows:
        sys.exit(f"ERROR: {F_MODULE_TABLE} has no table")
    caps = [i for i, ln in enumerate(lines[:rows[0]]) if caption_start(ln)]
    return "\n".join(lines[(caps[-1] if caps else rows[0]):rows[-1] + 1])


def expand_module_table(work: str, text: str) -> tuple[str, int]:
    """Turn the insertion directive line of the manuscript into the table. Returns (changed body, number of replacements)"""
    out, count = [], 0
    for ln in text.split("\n"):
        if INSERT_MODULE_RE.match(ln.strip()):
            out += ["", module_table_block(work), ""]
            count += 1
        else:
            out.append(ln)
    return "\n".join(out), count


# ═══════════════════════════════════════════════════════════════════════
# Default skeleton creation (applies the SKELETON_* definitions of settings 1 to an empty new python-hwpx document)
# ═══════════════════════════════════════════════════════════════════════
SK_HC = "{http://www.hancom.co.kr/hwpml/2011/core}"


class SkeletonBaseError(Exception):
    """Raised when the empty new python-hwpx document is missing or its structure differs from expectations"""


def _skeleton_base_path() -> str:
    import hwpx
    return os.path.join(os.path.dirname(hwpx.__file__), "data", "Skeleton.hwpx")


def _sk_container(head, tag):
    el = head.find(f".//{HH}{tag}")
    if el is None:
        raise SkeletonBaseError(f"header.xml has no {tag} element")
    return el


def _sk_renumber(container):
    container.set("itemCnt", str(len(container)))


def _sk_set_fonts(head):
    """Replace the font lists of every language with two fonts, body (0) and heading (1). Base styles (바탕글, 각주, and so on) use the body font"""
    for ff in head.iter(f"{HH}fontface"):
        for f, face in zip(ff.findall(f"{HH}font"), (SKELETON_FONT_BODY, SKELETON_FONT_HEAD)):
            f.set("face", face)
            ti = f.find(f"{HH}typeInfo")
            if ti is not None:
                ti.set("familyType", "FCAT_GOTHIC")
    for cp in head.iter(f"{HH}charPr"):
        ref = cp.find(f"{HH}fontRef")
        for k in list(ref.attrib):
            ref.set(k, "0")


def _sk_add_char(head, template, font_id, pt, bold) -> str:
    cps = _sk_container(head, "charProperties")
    el = copy.deepcopy(template)
    el.set("id", str(len(cps)))
    el.set("height", str(int(round(pt * 100))))
    el.set("textColor", SKELETON_TEXT_COLOR)
    ref = el.find(f"{HH}fontRef")
    for k in list(ref.attrib):
        ref.set(k, str(font_id))
    for b in el.findall(f"{HH}bold"):
        el.remove(b)
    if bold:
        el.insert(list(el).index(el.find(f"{HH}offset")) + 1, etree.Element(f"{HH}bold"))
    cps.append(el)
    _sk_renumber(cps)
    return el.get("id")


def _sk_add_bullet(head, char) -> str:
    bullets = head.find(f".//{HH}bullets")
    if bullets is None:  # an empty new document has no bullet list, so create it after the numbering list
        bullets = etree.Element(f"{HH}bullets", itemCnt="0")
        _sk_container(head, "numberings").addnext(bullets)
    el = etree.SubElement(bullets, f"{HH}bullet", id=str(len(bullets) + 1), char=char, useImage="0")
    etree.SubElement(el, f"{HH}paraHead", level="0", align="LEFT", useInstWidth="0", autoIndent="1",
                     widthAdjust="0", textOffsetType=SKELETON_BULLET_OFFSET_TYPE, textOffset=SKELETON_BULLET_OFFSET,
                     numFormat="DIGIT", charPrIDRef="4294967295", checkable="0")
    _sk_renumber(bullets)
    return el.get("id")


def _sk_add_toc_tab(head) -> str:
    """Right-aligned dotted tab of TOC paragraphs. The basis on which 한/글 aligns page numbers to the right end"""
    tabs = _sk_container(head, "tabProperties")
    tp = etree.SubElement(tabs, f"{HH}tabPr", id=str(len(tabs)), autoTabLeft="0", autoTabRight="0")
    sw = etree.SubElement(tp, f"{HP}switch")
    case = etree.SubElement(sw, f"{HP}case")
    case.set(f"{HP}required-namespace", "http://www.hancom.co.kr/hwpml/2016/HwpUnitChar")
    etree.SubElement(case, f"{HH}tabItem", pos=str(SKELETON_TOC_TAB_POS), type="RIGHT", leader="DOT", unit="HWPUNIT")
    default = etree.SubElement(sw, f"{HP}default")
    etree.SubElement(default, f"{HH}tabItem", pos=str(SKELETON_TOC_TAB_POS * 2), type="RIGHT", leader="DOT")
    _sk_renumber(tabs)
    return tp.get("id")


def _sk_add_para(head, template, align, left, indent, prev, nxt, ls, bullet_id, keep, widow, tab_id=None) -> str:
    pps = _sk_container(head, "paraProperties")
    el = copy.deepcopy(template)
    el.set("id", str(len(pps)))
    if tab_id is not None:
        el.set("tabPrIDRef", tab_id)
    el.find(f"{HH}align").set("horizontal", align)
    hd = el.find(f"{HH}heading")
    hd.set("type", "BULLET" if bullet_id else "NONE")
    hd.set("idRef", bullet_id or "0")
    hd.set("level", "0")
    bs = el.find(f"{HH}breakSetting")
    bs.set("breakNonLatinWord", "KEEP_WORD")  # line breaks by word
    bs.set("keepWithNext", str(keep))
    bs.set("widowOrphan", str(widow))
    for m in el.iter(f"{HH}margin"):
        for tag, val in (("intent", indent), ("left", left), ("prev", prev), ("next", nxt)):
            m.find(f"{SK_HC}{tag}").set("value", str(val))
    for s in el.iter(f"{HH}lineSpacing"):
        s.set("value", str(ls))
    pps.append(el)
    _sk_renumber(pps)
    return el.get("id")


def _sk_add_style(head, name, pp, cp):
    styles = _sk_container(head, "styles")
    sid = str(len(styles))
    etree.SubElement(styles, f"{HH}style", id=sid, type="PARA", name=name, engName="", paraPrIDRef=pp,
                     charPrIDRef=cp, nextStyleIDRef=sid, langID="1042", lockForm="0")
    _sk_renumber(styles)


def _sk_build_header(xml: bytes) -> bytes:
    head = etree.fromstring(xml)
    _sk_set_fonts(head)
    char_tpl = head.find(f".//{HH}charPr[@id='0']")
    para_tpl = head.find(f".//{HH}paraPr[@id='0']")
    if char_tpl is None or para_tpl is None:
        raise SkeletonBaseError("header.xml has no character shape or paragraph shape with id 0")
    char_tpl, para_tpl = copy.deepcopy(char_tpl), copy.deepcopy(para_tpl)
    toc_tab = _sk_add_toc_tab(head)
    for name, role, font, pt, bold, align, left, indent, prev, nxt, ls, bkey, keep, widow in SKELETON_STYLES:
        cp = _sk_add_char(head, char_tpl, 1 if font == "head" else 0, pt, bold)
        bullet_id = _sk_add_bullet(head, SKELETON_GLYPHS[bkey]) if bkey else None
        pp = _sk_add_para(head, para_tpl, align, left, indent, prev, nxt, ls, bullet_id, keep, widow,
                          toc_tab if role.startswith("toc_") else None)
        _sk_add_style(head, name, pp, cp)
    _sk_apply_base_styles(head)
    return etree.tostring(head, xml_declaration=True, encoding="UTF-8", standalone=True)


def _sk_apply_base_styles(head):
    """Set the base styles of the empty new document (바탕글, 쪽 번호, 머리말, 각주) to the settings values. If values are equal, the XML stays as is"""
    for name, _role, _font, pt, bold, align, left, indent, prev, nxt, ls, _b, keep, widow in SKELETON_BASE_STYLES:
        st = head.find(f".//{HH}style[@name='{name}']")
        if st is None:
            raise SkeletonBaseError(f"base style {name} missing from header.xml")
        cp = head.find(f".//{HH}charPr[@id='{st.get('charPrIDRef')}']")
        pp = head.find(f".//{HH}paraPr[@id='{st.get('paraPrIDRef')}']")
        cp.set("height", str(int(round(pt * 100))))
        cp.set("textColor", SKELETON_TEXT_COLOR)
        for b in cp.findall(f"{HH}bold"):
            cp.remove(b)
        if bold:
            cp.insert(list(cp).index(cp.find(f"{HH}offset")) + 1, etree.Element(f"{HH}bold"))
        pp.find(f"{HH}align").set("horizontal", align)
        bs = pp.find(f"{HH}breakSetting")
        bs.set("keepWithNext", str(keep))
        bs.set("widowOrphan", str(widow))
        for m in pp.iter(f"{HH}margin"):
            # the default branch of hp:switch records old units (twice HWPUNIT). The same way as files authored in 한/글
            scale = 2 if m.getparent().tag == f"{HP}default" else 1
            for tag, val in (("intent", indent), ("left", left), ("prev", prev), ("next", nxt)):
                m.find(f"{SK_HC}{tag}").set("value", str(val * scale))
        for sp in pp.iter(f"{HH}lineSpacing"):
            sp.set("value", str(ls))


def _sk_build_section(xml: bytes) -> bytes:
    s = xml.decode("utf-8")
    m = re.search(r"<hp:margin [^>]*/>", s)
    if m is None:
        raise SkeletonBaseError("section0.xml has no page margin element")
    new = ('<hp:margin header="{header}" footer="{footer}" gutter="0" left="{left}" right="{right}" '
           'top="{top}" bottom="{bottom}"/>').format(**SKELETON_PAGE_MARGIN)
    s = s.replace(m.group(0), new, 1)
    pm = re.search(r'<hp:pagePr [^>]*>', s)
    if pm is None:
        raise SkeletonBaseError("section0.xml has no page size element")
    sized = re.sub(r'width="\d+"', f'width="{SKELETON_PAGE_SIZE[0]}"', pm.group(0))
    sized = re.sub(r'height="\d+"', f'height="{SKELETON_PAGE_SIZE[1]}"', sized)
    s = s.replace(pm.group(0), sized, 1)
    # the first paragraph's line layout cache (horzsize) remains at the value for the empty new document's default margins (30mm left and right)
    # as in files authored in 한/글, the cache is matched to the actual body width so margin settings and cache do not diverge
    text_width = SKELETON_PAGE_SIZE[0] - SKELETON_PAGE_MARGIN["left"] - SKELETON_PAGE_MARGIN["right"]
    s = re.sub(r'(<hp:lineseg [^>]*?horzsize=")\d+(")', rf'\g<1>{text_width}\g<2>', s)
    return s.encode("utf-8")


def default_skeleton_path(base_dir: str) -> str:
    return os.path.join(base_dir, D_TEMPLATE, F_SKELETON)


def write_default_skeleton(base_dir: str) -> str:
    """Create the default skeleton in base_dir/_양식/ and return its path. Rewritten with the same content on every call
    Why regenerate every time: keeps the workspace template from diverging from the settings part even when the skill is updated. Change the template by editing the settings part or with build --skeleton"""
    base = _skeleton_base_path()
    if not os.path.exists(base):
        sys.exit(f"ERROR: empty new python-hwpx document not found {base}. Check the python-hwpx installation and version (API call site fixes in the Preparation section)")
    out = default_skeleton_path(base_dir)
    os.makedirs(os.path.dirname(out), exist_ok=True)
    tmp = out + ".tmp"
    try:
        with zipfile.ZipFile(base) as zin, zipfile.ZipFile(tmp, "w") as zout:
            for n in zin.namelist():
                d = zin.read(n)
                if n == "Contents/header.xml":
                    d = _sk_build_header(d)
                elif n == "Contents/section0.xml":
                    d = _sk_build_section(d)
                elif n == "Preview/PrvText.txt":
                    d = b""
                method = zipfile.ZIP_STORED if n == "mimetype" else zipfile.ZIP_DEFLATED
                zout.writestr(n, d, compress_type=method)
    # kinds of structural change: missing required elements (SkeletonBaseError, AttributeError), missing parts (KeyError), order mismatch (ValueError), XML corruption
    except (SkeletonBaseError, AttributeError, KeyError, ValueError, etree.XMLSyntaxError) as e:
        if os.path.exists(tmp):
            os.remove(tmp)
        sys.exit(f"ERROR: failed to create the default skeleton ({e}). The structure of the empty new python-hwpx document presumably changed (API call site fixes in the Preparation section)")
    os.replace(tmp, out)
    return out


def cmd_init(a):
    work = a.work
    kind = a.kind
    st = load_state(work)
    if st.get("kind") and st["kind"] != kind:
        sys.exit(f"ERROR: the output type of this workspace is {st['kind']}. To change the type, run init in a new workspace")
    os.makedirs(os.path.join(work, D_MANUSCRIPT), exist_ok=True)
    os.makedirs(os.path.join(work, D_DELIVER), exist_ok=True)
    made = []
    if kind == "proposal":  # the proposal strategy file, company data, and common part templates are proposal-only
        for src, dst in (("strategy.md", F_STRATEGY), ("company.md", F_COMPANY)):
            target = os.path.join(work, dst)
            if not os.path.exists(target):
                shutil.copy(os.path.join(ASSET_DIR, src), target)
                made.append(dst)
        for dst, src in COMMON_TEMPLATES.items():
            target = os.path.join(work, D_MANUSCRIPT, dst)
            if not os.path.exists(target):
                shutil.copy(os.path.join(ASSET_DIR, "common", src), target)
                made.append(f"{D_MANUSCRIPT}/{dst}")
    elif kind in DESIGN_TEMPLATES:  # the diagnosis file of reports, the planning design of plans
        src, dst = DESIGN_TEMPLATES[kind]
        target = os.path.join(work, dst)
        if not os.path.exists(target):
            shutil.copy(os.path.join(ASSET_DIR, src), target)
            made.append(dst)
    made.append(os.path.relpath(write_default_skeleton(work), work))
    st = load_state(work)
    st["kind"] = kind
    st["stages"]["init"] = datetime.now().isoformat(timespec="seconds")
    save_state(work, st)
    print(f"OK: {work} ({kind}) created files {made}")
    if kind == "proposal":
        print("Next: confirm render path PASS with `selftest`, then run `rfp` and the kickoff questions. Fill the [확인] items of 00_제안전략서.md and 회사자료.md from user answers")
    else:
        guide = "references/report.md" if kind == "report" else "references/planning.md"
        print(f"Next: confirm render path PASS with `selftest`, then follow {guide} for the procedure")


# ═══════════════════════════════════════════════════════════════════════
# 5. RFP conversion: source text and empty table templates (the model does the extraction)
# ═══════════════════════════════════════════════════════════════════════
MERGE_UP = "^"      # merge marker same as the manuscript notation: vertical merge with the cell above
MERGE_LEFT = "<"    # horizontal merge with the left cell


def _hwpx_paragraph_lines(p) -> list[str]:
    """Convert one paragraph into text lines and table lines. Tables sit in runs inside the paragraph, so output is split in character order"""
    out, buf = [], []
    for run in p.findall(f"{HP}run"):
        for child in run:
            tag = etree.QName(child).localname
            if tag == "t":
                buf.append("".join(child.itertext()))
            elif tag == "tbl":
                if "".join(buf).strip():
                    out.append("".join(buf).strip())
                buf = []
                out += [""] + _hwpx_table_lines(child) + [""]
            else:  # objects with sub-paragraphs such as text boxes output those paragraphs in sequence
                for sub in child.iter(f"{HP}subList"):
                    for sp in sub.findall(f"{HP}p"):
                        out += _hwpx_paragraph_lines(sp)
    if "".join(buf).strip():
        out.append("".join(buf).strip())
    return out


def _hwpx_cell_text(tc) -> str:
    """Join the paragraphs of a cell with <br>. Tables inside a cell have only their text joined (a table inside a table cannot be expressed as a markdown table)"""
    parts = []
    for sub in tc.findall(f"{HP}subList"):
        for sp in sub.findall(f"{HP}p"):
            text = " ".join(line for line in _hwpx_paragraph_lines(sp) if line and not line.startswith("|"))
            nested = [" ".join(t.itertext()).strip() for t in sp.iter(f"{HP}tbl")]
            parts.append(" ".join([text] + [n for n in nested if n]).strip())
    return "<br>".join(x for x in parts if x).replace("|", "/")


def _hwpx_table_lines(tbl) -> list[str]:
    """Fill the grid by cell addresses and merge spans. python-hwpx export_markdown pulls cells after a vertical merge forward, which misaligns points"""
    rows, cols = int(tbl.get("rowCnt", "0")), int(tbl.get("colCnt", "0"))
    if not rows or not cols:
        return []
    grid = [["" for _ in range(cols)] for _ in range(rows)]
    for tr in tbl.findall(f"{HP}tr"):
        for tc in tr.findall(f"{HP}tc"):
            addr, span = tc.find(f"{HP}cellAddr"), tc.find(f"{HP}cellSpan")
            r, c = int(addr.get("rowAddr")), int(addr.get("colAddr"))
            rs = int(span.get("rowSpan", "1")) if span is not None else 1
            cs = int(span.get("colSpan", "1")) if span is not None else 1
            if r >= rows or c >= cols:
                continue
            grid[r][c] = _hwpx_cell_text(tc)
            for rr in range(r, min(r + rs, rows)):
                for cc in range(c, min(c + cs, cols)):
                    if (rr, cc) != (r, c):
                        grid[rr][cc] = MERGE_LEFT if rr == r else MERGE_UP
    if all(cell in ("", MERGE_UP, MERGE_LEFT) for row in grid for cell in row):
        return []  # empty tables used for page layout carry no source information and are skipped
    lines = ["| " + " | ".join(grid[0]) + " |", "|" + "---|" * cols]
    return lines + ["| " + " | ".join(row) + " |" for row in grid[1:]]


def hwpx_to_markdown(path: str) -> str:
    with zipfile.ZipFile(path) as z:
        names = sorted((n for n in z.namelist() if re.match(r"Contents/section\d+\.xml$", n)),
                       key=lambda n: int(re.search(r"\d+", os.path.basename(n)).group()))
        lines = []
        for n in names:
            root = etree.fromstring(z.read(n))
            for p in root.findall(f"{HP}p"):
                lines += _hwpx_paragraph_lines(p)
    return "\n".join(lines)


def read_document(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    if ext == ".hwpx":
        return hwpx_to_markdown(path)
    if ext == ".pdf":
        if not shutil.which("pdftotext"):
            sys.exit("ERROR: pdftotext missing. Run `apt-get install -y poppler-utils` and retry")
        return subprocess.run(["pdftotext", "-layout", path, "-"], capture_output=True, text=True).stdout
    if ext == ".docx":
        import docx
        d = docx.Document(path)
        out = [p.text for p in d.paragraphs]
        for t in d.tables:
            rows = [[c.text.strip().replace("\n", " ") for c in r.cells] for r in t.rows]
            if rows:
                out.append("| " + " | ".join(rows[0]) + " |")
                out.append("|" + "---|" * len(rows[0]))
                out.extend("| " + " | ".join(r) + " |" for r in rows[1:])
        return "\n".join(out)
    if ext in (".md", ".txt"):
        return open(path, encoding="utf-8").read()
    sys.exit(f"ERROR: unsupported input format {ext} (hwpx, pdf, docx, md)")


REQ_TABLE_HEADER = ["번호", "상위 항목", "요구항목 원문", "배치 섹션", "비고"]
SCORE_TABLE_HEADER = ["평가항목", "세부평가항목", "평가기준 원문", "배점", "결정 섹션"]


def _md_table(header: list[str], rows: list[list[str]]) -> list[str]:
    out = ["| " + " | ".join(header) + " |", "|" + "---|" * len(header)]
    out += ["| " + " | ".join(r) + " |" for r in rows]
    return out


def cmd_rfp(a):
    """Convert the RFP into text and place empty table templates.
    Why: RFP section structure, bullets, and table use differ by client, so regex extraction yields 0 items on table-form RFPs and broken lines on PDFs.
    The model reads 01_RFP원문.md and copies requirement items and points into the tables verbatim. Later commands (ledger, check rfp) read only the tables."""
    work = a.work
    if not os.path.isfile(a.input):
        sys.exit(f"ERROR: RFP file not found: {a.input}")
    os.makedirs(work, exist_ok=True)
    text = read_document(a.input)

    head = ["# RFP 원문", f"- 입력 파일 {os.path.basename(a.input)}", f"- 변환 문자 수 {len(text)}",
            f"- In tables, `{MERGE_UP}` marks a cell merged with the cell above and `{MERGE_LEFT}` with the left cell", "", "## 원문 전체", "", text]
    open(os.path.join(work, F_RFP_TEXT), "w", encoding="utf-8").write("\n".join(head))

    kind = load_state(work).get("kind") or "proposal"
    place = ("section IDs of 02_섹션대장.md" if kind == "proposal" else "section numbers of the TOC design table in 02_스토리라인.md (for example 2.1)")
    req = ["# 요구항목 전수 추출표",
           "- Read the task content sections (detailed execution content, engagement scope) of 01_RFP원문.md and transfer the lowest-level items row by row verbatim. No summarizing or restating",
           "- In the parent item column, enter the parent headings the item belongs to (e.g. 1. 현황 분석 > 가. 정책 동향). The grounds for placement decisions",
           f"- Enter {place} in the placement section column (commas for multiple). Run `check rfp` after filling every row",
           "- Transfer items inside source tables as rows too. Replace `|` inside cells with `/`. In source tables, `^` marks a cell merged with the cell above and `<` with the left cell", "",
           *_md_table(REQ_TABLE_HEADER, [["R1", "", "", "", ""]])]
    open(os.path.join(work, F_REQ), "w", encoding="utf-8").write("\n".join(req) + "\n")

    sc = ["# 배점표",
          "- Read the proposal evaluation section of 01_RFP원문.md (evaluation items, points, evaluation criteria) and fill the breakdown table by detailed evaluation item. The points column holds numbers only",
          "- The deciding section is the ID of the section read to assign that score (commas for multiple, split evenly). Quantitative items (credit rating and other appended forms) get `정량`",
          "- Confirm basic project conditions (project name, period, budget, eligibility, deadline) and writing rules (ban on vague expressions, quantification, volume cap, reference tables, abbreviations) in the source and enter them in 00_제안전략서.md", "",
          "## 전개표", "", *_md_table(SCORE_TABLE_HEADER, [["", "", "", "", ""]])]
    if kind == "proposal":  # the scoring table and ledger are proposal-only. Report placement sections are section numbers
        open(os.path.join(work, F_SCORE), "w", encoding="utf-8").write("\n".join(sc) + "\n")

    st = load_state(work)
    st["stages"]["rfp"] = datetime.now().isoformat(timespec="seconds")
    st["files"]["rfp"] = os.path.abspath(a.input)
    save_state(work, st)
    if kind == "proposal":
        print(f"OK: {F_RFP_TEXT} ({len(text)} characters), {F_REQ} and {F_SCORE} placed as empty tables")
        print("Next: read 01_RFP원문.md, fill 01_요구항목표.md and the breakdown table of 01_배점표.md, then run `ledger`")
    else:
        print(f"OK: {F_RFP_TEXT} ({len(text)} characters), {F_REQ} placed as an empty table")
        print("Next: read 01_RFP원문.md and fill 01_요구항목표.md. The placement section column holds section numbers of the step 2 TOC")


# ═══════════════════════════════════════════════════════════════════════
# 5-1. Deck input: normalize pptx and pdf into a reference PDF and convert the source text (the main body is deck_source.py)
# ═══════════════════════════════════════════════════════════════════════
def cmd_deck(a):
    """Treat the deck the same way as the RFP. The body text is converted to a file. The model that read the full source does the extraction and judgment"""
    import deck_source
    work = a.work
    os.makedirs(work, exist_ok=True)
    deck_source.parse(a.input, os.path.join(work, D_DECK), allow_garbled=a.allow_garbled)
    st = load_state(work)
    st["stages"]["deck"] = datetime.now().isoformat(timespec="seconds")
    st["files"]["deck"] = os.path.abspath(a.input)
    save_state(work, st)
    deck_source.say(f"Next: read the full text of {D_DECK}/원문.md (the same way as 01_RFP원문.md). Check diagrams in the thumbnail sheets of {D_DECK}/sheets")


def cmd_figure(a):
    import deck_source
    out_dir = os.path.join(a.work, D_DECK)
    if a.grid:
        deck_source.grid(out_dir, a.page)
    else:
        deck_source.figure(out_dir, a.page, a.crop, a.name, not a.no_trim, a.pdf)


# ═══════════════════════════════════════════════════════════════════════
# 6. Section ledger
# ═══════════════════════════════════════════════════════════════════════
def parse_table_rows(md: str, header_key: str) -> list[list[str]]:
    """Return the data rows of the pipe table that contains header_key (separator excluded)."""
    lines = md.splitlines()
    for i, ln in enumerate(lines):
        if ln.strip().startswith("|") and header_key in ln:
            rows = []
            for ln2 in lines[i + 1:]:
                if not ln2.strip().startswith("|"):
                    break
                if TABLE_SEP_RE.match(ln2.strip()):
                    continue
                rows.append([c.strip() for c in ln2.strip().strip("|").split("|")])
            return rows
    return []


def score_by_section(work: str) -> tuple[dict[str, float], int]:
    rows = parse_table_rows(open(os.path.join(work, F_SCORE), encoding="utf-8").read(), "결정 섹션")
    totals: dict[str, float] = {}
    n = 0
    for r in rows:
        if len(r) < 5 or not r[3].strip():
            continue
        try:
            pts = float(re.sub(r"[^\d.]", "", r[3]))
        except ValueError:
            continue
        ids = [x.strip() for x in re.split(r"[,，/]", r[4]) if x.strip()]
        if not ids:
            continue
        n += 1
        for sid in ids:
            totals[sid] = totals.get(sid, 0) + pts / len(ids)
    return totals, n


def section_sort_key(sid: str) -> tuple:
    """Sort section IDs such as Ⅱ-1 and Ⅲ-2 by chapter number and section number."""
    m = re.match(r"([ⅠⅡⅢⅣⅤⅥⅦⅧⅨⅩ]|[IVX]+|\d+)[-.]?(\d+)?", sid)
    if not m:
        return (99, 0, sid)
    ch = m.group(1)
    ci = ROMAN.index(ch) + 1 if ch in ROMAN else (["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X"].index(ch) + 1
                                                  if ch in ("I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X") else int(ch))
    return (ci, int(m.group(2) or 0), sid)


def cmd_ledger(a):
    work = a.work
    totals, n = score_by_section(work)
    path = os.path.join(work, F_LEDGER)
    st = load_state(work)
    existing = parse_table_rows(open(path, encoding="utf-8").read(), "섹션ID") if os.path.exists(path) else []
    by_id = {r[0]: r for r in existing if r and r[0]}
    ids = sorted(dict.fromkeys(list(by_id) + [k for k in totals if k != "정량"]), key=section_sort_key)
    out = ["# 섹션대장", f"- Tallied {n} breakdown rows of the scoring table. Attributed points are computed automatically from the deciding-section column, and the type is 오리지널 or 공통",
           "- The manuscript file is the file name under 원고/. Status is updated from check results", "",
           "| 섹션ID | 섹션명 | 유형 | 통제배점 | 원고파일 | 상태 |", "|---|---|---|---|---|---|"]
    for sid in ids:
        r = by_id.get(sid, [sid, "", "", "", "", ""]) + [""] * 6
        pts = totals.get(sid)
        pts_s = f"{pts:g}" if pts is not None else r[3]
        status = []
        for key, val in st.get("checks", {}).items():
            if key.startswith("density:") and r[4] and r[4] in key:
                status.append("density " + ("PASS" if val["pass"] else "FAIL"))
        out.append(f"| {sid} | {r[1]} | {r[2]} | {pts_s} | {r[4]} | {' '.join(status) or r[5]} |")
    unassigned = "quantitative " + f"{totals['정량']:g} points" if "정량" in totals else ""
    open(path, "w", encoding="utf-8").write("\n".join(out) + "\n")
    st["stages"]["ledger"] = datetime.now().isoformat(timespec="seconds")
    save_state(work, st)
    print(f"OK: {F_LEDGER} sections {len(ids)}, point total {sum(totals.values()):g} points {unassigned}")
    if not existing:
        print("Next: fill the section name, type, and manuscript file columns, then run `ledger` again (attributed points are kept automatically)")


# ═══════════════════════════════════════════════════════════════════════
# 7. Checks (check)
# ═══════════════════════════════════════════════════════════════════════
def _print_result(name: str, passed: bool, fails: list[str], warns: list[str] | None = None):
    print(f"[{name}] {'PASS' if passed else 'FAIL'}")
    for f in fails[:40]:
        print("   FAIL", f)
    if len(fails) > 40:
        print(f"   ... and {len(fails) - 40} more")
    for w in (warns or [])[:20]:
        print("   WARN", w)


OUT_OF_SCOPE = "범위 외"   # notation for requirement items outside the output range of a pre-diagnosis report (report.md Procedure section)
SECTION_NO_RE = re.compile(r"\d+\.\d+")


def rfp_place_ids(work: str) -> tuple[set[str], str, str]:
    """Comparison source for the placement section column. For proposals, section IDs of the section ledger; for pre-diagnosis reports, section numbers of the storyline TOC design table
    Returns: (set of valid values, source name, guidance text when the source is missing)"""
    kind = load_state(work).get("kind") or "proposal"
    if kind == "report":
        path = os.path.join(work, F_STORYLINE)
        rows = parse_table_rows(open(path, encoding="utf-8").read(), "작성 방식") if os.path.exists(path) else []
        ids = {m.group() for r in rows if len(r) > 1 for m in SECTION_NO_RE.finditer(r[1])}
        return ids, "the storyline TOC design table", f"{F_STORYLINE} has no section numbers in its TOC design table."
    path = os.path.join(work, F_LEDGER)
    ledger = parse_table_rows(open(path, encoding="utf-8").read(), "섹션ID") if os.path.exists(path) else []
    return {r[0] for r in ledger if r and r[0]}, "the section ledger", f"{F_LEDGER} missing (before the step 2 ledger)."


def check_rfp(work: str) -> tuple[bool, list[str], list[str]]:
    if not os.path.exists(os.path.join(work, F_REQ)):
        return False, [f"{F_REQ} missing. The `rfp` command has not run"], []
    rows = parse_table_rows(open(os.path.join(work, F_REQ), encoding="utf-8").read(), "요구항목 원문")
    ids, source, missing = rfp_place_ids(work)
    fails, warns = [], []
    if not ids:
        warns.append(f"{missing} Placement section values not compared. Rerun after writing that file")
    for r in rows:
        if len(r) < 4 or not r[0]:
            continue
        if not r[3].strip():
            fails.append(f"{r[0]} placement section empty: {r[2][:40]}")
            continue
        for sid in re.split(r"[,，/]", r[3]):
            sid = sid.strip()
            if sid and ids and sid not in ids and sid != OUT_OF_SCOPE:
                fails.append(f"{r[0]} placement section '{sid}' not in {source} (check the ID)")
    if not rows:
        fails.append("the requirement table has no rows")
    return not fails, fails, warns


def _segments(lines: list[str]) -> list[tuple[str, list[str]]]:
    """Split manuscripts by Task or Step band. A Task span ends before the next Task, a Module band, or a chapter or section heading (#).
    Why: cutting spans only at the next Task adds the subsection after the last Task or the next Module introduction to that Task and causes misjudgments"""
    segs: list[tuple[str, list[str]]] = []
    cur_name, cur = None, []
    for ln in lines:
        s = ln.strip()
        m = TASK_BAND_RE.match(s)
        if m or re.match(r"#{1,6}\s", s):
            if cur_name is not None:
                segs.append((cur_name, cur))
            is_task = m is not None and m.group(1).lower() != "module"
            cur_name, cur = (s[:40] if is_task else None), []
        else:
            cur.append(ln)
    if cur_name is not None:
        segs.append((cur_name, cur))
    return segs


def check_density(path: str) -> tuple[bool, list[str], list[str]]:
    lines = open(path, encoding="utf-8").read().splitlines()
    fails, warns = [], []
    elems, _ = parse_markdown(path)
    for msg in lint(elems):
        if "lone □" in msg:
            fails.append(msg)
        elif "HITL" not in msg:
            warns.append(msg)
    segs = _segments(lines)
    for name, body in segs:
        text = "\n".join(body)
        miss = []
        if "(목적)" not in text:
            miss.append("(목적) description")
        # Why: a figure adopted from a source also counts as a Task's evidence block (same criterion as check volume)
        if not any(ln.strip().startswith(("|", "![")) for ln in body):
            miss.append("table or figure (a filled table fitting the analysis method for analysis-centered Tasks, a criteria table or draft for derivation-centered Tasks)")
        if "산출물" not in text:
            miss.append("(산출물) description")
        if not QUANTITY_RE.search(text) and not CONCRETE_MARK_RE.search(text):  # figures, or specification of target, material, method, or criterion
            miss.append("grounds of concreteness (scale figures or a description of target, material, method, criterion, or procedure)")
        n_num = len(QUANTITY_RE.findall(text))
        if n_num > QUANTITY_MANY:
            warns.append(f"{name}: {n_num} figures. Check whether figures were attached to uncountable work")
        if miss:
            fails.append(f"{name}: {', '.join(miss)} missing")
    if not segs and not any(ln.strip().startswith("|") for ln in lines):
        warns.append("manuscript without tables. One or more evidence tables recommended")
    return not fails, fails, warns


def _strategy_wording(work: str) -> tuple[list[list[str]], list[str]]:
    md = open(os.path.join(work, F_STRATEGY), encoding="utf-8").read()
    terms = parse_table_rows(md, "RFP 용어")
    forbidden = []
    m = re.search(r"^#+\s*금지 표현.*?$(.*?)(?=^#+\s|\Z)", md, re.S | re.M)
    if m:
        forbidden = [re.sub(r"^\s*[-+*]\s*", "", ln).strip() for ln in m.group(1).splitlines() if re.match(r"^\s*[-+*]\s+\S", ln)]
    return terms, forbidden


def check_wording(work: str) -> tuple[bool, list[str], list[str]]:
    terms, forbidden = _strategy_wording(work)
    # section IDs in the responding section column are turned into manuscript files of the section ledger to set the scope. File name fragments are compared as written
    ledger = os.path.join(work, F_LEDGER)
    sec_files = {r[0]: os.path.basename(r[4]) for r in parse_table_rows(open(ledger, encoding="utf-8").read(), "섹션ID")
                 if len(r) > 4 and r[0] and r[4]} if os.path.exists(ledger) else {}
    files = manuscript_files(work)
    texts = {os.path.basename(f): open(f, encoding="utf-8").read() for f in files}
    fails, warns = [], []
    if not files:
        return False, ["no files in 원고/"], []
    for r in terms:
        if len(r) < 2 or not r[0] or r[0].startswith("["):
            continue
        rfp_term, use_term = r[0], (r[1] or r[0])
        sec = r[2] if len(r) > 2 else ""
        wanted = {sec_files.get(x.strip(), x.strip()) for x in re.split(r"[,，/]", sec) if x.strip()}
        scope = {k: v for k, v in texts.items() if not wanted or any(w in k for w in wanted)} or texts
        hits = sum(v.count(use_term) for v in scope.values())
        if hits == 0:
            fails.append(f"wording '{use_term}' (RFP '{rfp_term}') appears 0 times (scope {sec or 'all'})")
        variants = {k for k, v in texts.items() if rfp_term != use_term and rfp_term in v}
        if variants:
            warns.append(f"RFP notation '{rfp_term}' mixed with proposal notation '{use_term}': {sorted(variants)}")
    for term in dict.fromkeys(FORBIDDEN_DEFAULT + forbidden):
        if not term:
            continue
        for k, v in texts.items():
            for i, ln in enumerate(v.splitlines(), 1):
                if term in ln:
                    fails.append(f"prohibited expression '{term}' {k}:{i}")
    for pat, label in SENTENCE_WARN:
        n = 0
        for k, v in texts.items():
            for i, ln in enumerate(v.splitlines(), 1):
                if ln.strip().startswith("|") or ln.strip().startswith(">"):
                    continue
                if re.search(pat, ln):
                    n += 1
                    if n <= 3:
                        warns.append(f"{label} {k}:{i} {ln.strip()[:50]}")
        if n > 3:
            warns.append(f"{label} and {n - 3} more")
    return not fails, fails, warns


def check_residual(work: str) -> tuple[bool, list[str], list[str]]:
    """Locations of leftover [확인] in pre-diagnosis and planning reports. Not counted in the verdict and used as the list of user questions at step 5"""
    warns = []
    for f in manuscript_files(work):
        for i, ln in enumerate(open(f, encoding="utf-8").read().splitlines(), 1):
            if "[확인" in ln:
                warns.append(f"leftover [확인] {os.path.basename(f)}:{i} {ln.strip()[:60]} (resolve by asking the user before delivery)")
    return True, [], warns


MANUAL_HEADING = "## 수동 기록"   # section of 05_점검기록.md preserved when check all reruns


def write_check_report(work: str, auto_lines: list[str]):
    """Rewrite the automatic verdict results while preserving the manual records section and below. Keeps records of self-review, render review, and discretionary adjustments from being erased by reruns"""
    path = os.path.join(work, F_REPORT)
    old = open(path, encoding="utf-8").read() if os.path.exists(path) else ""
    i = old.find(MANUAL_HEADING)
    manual = old[i:] if i >= 0 else f"{MANUAL_HEADING}\n\n- Records of self-review, render review, and discretionary adjustments. This section and below are preserved even when check all reruns\n"
    open(path, "w", encoding="utf-8").write("\n".join(auto_lines).rstrip() + "\n\n" + manual)


def manuscript_hash(work: str) -> str:
    """Fingerprint of the manuscripts check all judged. The basis on which build detects manuscript changes after the verdict. Includes the module overview table because build inserts it into the manuscripts"""
    import hashlib
    h = hashlib.sha256()
    for f in manuscript_files(work) + [p for p in [module_table_path(work)] if os.path.exists(p)]:
        h.update(os.path.basename(f).encode("utf-8"))
        h.update(open(f, "rb").read())
    labels = load_state(work).get("labels")
    if labels and os.path.exists(labels):  # the label file changes how manuscript notation is recognized, so it is part of the fingerprint. Without a record the fingerprint is as before
        h.update(b"labels")
        h.update(open(labels, "rb").read())
    return h.hexdigest()


def check_fill(work: str) -> tuple[bool, list[str], list[str]]:
    fails, warns = [], []
    for f in [x for pre in PREFIX_FILL for x in manuscript_files(work, pre)]:
        name = os.path.basename(f)
        for i, ln in enumerate(open(f, encoding="utf-8").read().splitlines(), 1):
            if "[확인" in ln:
                fails.append(f"leftover substitution {name}:{i} {ln.strip()[:60]}")
        rows = parse_table_rows(open(f, encoding="utf-8").read(), "참여율")
        if rows:
            hdr_line = next(ln for ln in open(f, encoding="utf-8") if "참여율" in ln and ln.strip().startswith("|"))
            cols = [c.strip() for c in hdr_line.strip().strip("|").split("|")]
            ci = next(i for i, c in enumerate(cols) if "참여율" in c)
            total = 0.0
            for r in rows:
                if len(r) > ci and re.search(r"\d", r[ci]):
                    v = float(re.sub(r"[^\d.]", "", r[ci]) or 0)
                    total += v
                    if v > 100:
                        fails.append(f"participation rate over 100% {name}: {r[0]} {v:g}%")
            warns.append(f"{name} participation rate total {total:g}% (compare with the person-month calculation)")
    # [확인] in the original part is a temporary marker while writing. The verdict is unchanged, and only leftover locations are reported for user questions before delivery
    for f in [x for pre in PREFIX_ORIGINAL for x in manuscript_files(work, pre)]:
        name = os.path.basename(f)
        for i, ln in enumerate(open(f, encoding="utf-8").read().splitlines(), 1):
            if "[확인" in ln:
                warns.append(f"leftover [확인] in the original part {name}:{i} {ln.strip()[:60]} (resolve by asking the user before delivery)")
    return not fails, fails, warns


ACRONYM_RE = re.compile(r"[A-Z][A-Z0-9&]+")   # uppercase English abbreviations (RFP, ISP) go into the whitelist even when short. Short Korean nouns are excluded because they hide other verdicts


def _whitelist_terms(work: str) -> list[str]:
    """Collect terms of the key wording dictionary (proposal) or the RFP terms table of the diagnosis file (pre-diagnosis report), and Module and Task names of the module overview table."""
    terms: list[str] = []
    if os.path.exists(os.path.join(work, F_STRATEGY)):
        for r in _strategy_wording(work)[0]:
            terms += [c.strip() for c in r[:2] if c.strip() and "[확인]" not in c]
    diag = os.path.join(work, F_DIAGNOSIS)
    if os.path.exists(diag):  # pre-diagnosis report: original RFP terms and manuscript notation in the RFP terms table of the diagnosis file (domain terms such as '평가위원' in R&D evaluation engagements)
        for r in parse_table_rows(open(diag, encoding="utf-8").read(), "RFP 용어"):
            terms += [c.strip() for c in r[:2] if c.strip() and "[확인]" not in c]
    mod = module_table_path(work)
    if os.path.exists(mod):
        for r in parse_table_rows(open(mod, encoding="utf-8").read(), "Module"):
            for c in r[:2]:
                name = TASK_BAND_RE.sub(lambda m: m.group(3), c.strip())
                if name:
                    terms.append(name)
    terms = [t for t in dict.fromkeys(terms)
             if (len(t) >= WHITELIST_MIN_LEN or ACRONYM_RE.fullmatch(t)) and "[확인]" not in t]
    return sorted(terms, key=len, reverse=True)   # mask longer terms first to prevent partial matches


def _whitelist_ready(work: str) -> bool:
    """The fixing point is after step 2 is complete: the module overview table exists and no [확인] remains in the key wording dictionary.
    Why: fixing before the framework is settled misjudges normal dictionary filling and module name finalization as a source change FAIL."""
    if not os.path.exists(module_table_path(work)):
        return False
    rows = _strategy_wording(work)[0] if os.path.exists(os.path.join(work, F_STRATEGY)) else []
    return bool(rows) and not any("[확인]" in c for r in rows for c in r[:2])


def _lock_whitelist(work: str, terms: list[str], reset: bool) -> tuple[str | None, str | None]:
    """When the fixing conditions are met, fix the whitelist in state.json. If the sources change afterward, return (FAIL reason, None); if not fixed, return (None, WARN reason).
    Why: blocks appending terms to the dictionary or module names to avoid FAIL. Release is allowed only with --reset-whitelist on the user's instruction."""
    st = load_state(work)
    locked = st.get("whitelist")
    if locked is None and not reset and not _whitelist_ready(work):
        return None, "whitelist not fixed (module overview table missing or [확인] left in the key wording dictionary). Fixed at the first check tone after step 2 is complete"
    if locked is None or reset:
        st["whitelist"] = {"terms": terms, "locked_at": datetime.now().isoformat(timespec="seconds")}
        save_state(work, st)
        return None, None
    added = sorted(set(terms) - set(locked["terms"]))
    removed = sorted(set(locked["terms"]) - set(terms))
    if added or removed:
        return f"whitelist source changed (fixed {locked['locked_at']}). Added {added[:5]} removed {removed[:5]}. With a user instruction, `check tone --reset-whitelist`", None
    return None, None


def _mask_whitelist(ln: str, terms: list[str]) -> str:
    for t in terms:
        if t in ln:
            ln = ln.replace(t, WHITELIST_MASK * len(t))
    return ln


def _client_term_findings(ln: str) -> tuple[list[str], list[str]]:
    """Agency naming verdict. '발주처' in subject position and unfilled naming slots are FAIL, other '발주처' are WARN.
    Why: some uses such as '발주처 협의' conventionally remain, so a blanket FAIL gives many false positives. Only repeated subjects are a certain defect."""
    body = CLIENT_HEADER_CELL_RE.sub("|", ln)
    fails, warns = [], []
    if CLIENT_SLOT_RE.search(body):
        fails.append(CLIENT_SLOT_LABEL)
    if CLIENT_SUBJECT_RE.search(body):
        fails.append(CLIENT_SUBJECT_LABEL)
    elif CLIENT_TERM_RE.search(body):
        warns.append(CLIENT_TERM_LABEL)
    return fails, warns


def check_tone(work: str, reset_whitelist: bool = False, kind: str = "proposal") -> tuple[bool, list[str], list[str]]:
    fails, warns = [], []
    terms = _whitelist_terms(work)
    # whitelist fixing is proposal-only (the key wording dictionary is in the proposal strategy file). Reports and plans mask only module overview table names
    lock_fail, lock_warn = _lock_whitelist(work, terms, reset_whitelist) if kind == "proposal" else (None, None)
    if lock_fail:
        fails.append(lock_fail)
    if lock_warn:
        warns.append(lock_warn)
    for f in manuscript_files(work):
        name = os.path.basename(f)
        for i, raw in enumerate(open(f, encoding="utf-8").read().splitlines(), 1):
            ln = _mask_whitelist(raw, terms)
            st = ln.strip()
            in_cell = st.startswith("|")
            for pat, label in TONE_FAIL + FRAMEWORK_FAIL:
                if in_cell and "cross reference" in label:  # Task and Module numbers in table cells are normal use in overview, quick-reference, and schedule tables
                    continue
                if (TASK_BAND_RE.match(st) or st.startswith("<표>") or WIDTH_SPEC_RE.match(st)) and "cross reference" in label:
                    continue
                if re.search(pat, ln):
                    fails.append(f"{label} {name}:{i} {raw.strip()[:50]}")
            for pat, label in TONE_WARN:
                if re.search(pat, ln):
                    warns.append(f"{label} {name}:{i} {raw.strip()[:50]}")
            client_fail, client_warn = _client_term_findings(ln)
            fails += [f"{label} {name}:{i} {raw.strip()[:50]}" for label in client_fail]
            warns += [f"{label} {name}:{i} {raw.strip()[:50]}" for label in client_warn]
            if HIER_LINE_RE.match(raw) and not PERIOD_EXEMPT_RE.match(raw):
                if PERIOD_MID_RE.search(raw):
                    fails.append(f"two sentences on one level item (sentences joined by a period. Split the second sentence into a separate level) {name}:{i} {raw.strip()[:50]}")
                elif PERIOD_END_RE.search(raw):
                    fails.append(f"period at line end (no period after a nominal ending) {name}:{i} {raw.strip()[:50]}")
    return not fails, fails, warns


def _table_blocks(lines: list[str]) -> list[tuple[int, list[list[str]]]]:
    blocks, i = [], 0
    while i < len(lines):
        if lines[i].strip().startswith("|"):
            j = i
            rows = []
            while j < len(lines) and lines[j].strip().startswith("|"):
                if not TABLE_SEP_RE.match(lines[j].strip()):
                    rows.append([c.strip() for c in lines[j].strip().strip("|").split("|")])
                j += 1
            blocks.append((i + 1, rows))
            i = j
            continue
        i += 1
    return blocks


def check_numbers(work: str) -> tuple[bool, list[str], list[str]]:
    """Recalculate the total rows and composition ratio columns of tables. Replaces the model's mental arithmetic."""
    fails, warns = [], []
    for f in manuscript_files(work):
        name = os.path.basename(f)
        lines = open(f, encoding="utf-8").read().splitlines()
        for start, rows in _table_blocks(lines):
            if len(rows) < 3:
                continue
            body = [r for r in rows[1:] if not SUM_LABEL_RE.match(r[0])]
            total_rows = [r for r in rows[1:] if SUM_LABEL_RE.match(r[0])]
            for tr in total_rows:
                for c in range(1, min(len(tr), len(rows[0]))):
                    if not NUM_CELL_RE.match(tr[c].replace(" ", "")):
                        continue
                    vals = []
                    for r in body:
                        if c < len(r) and NUM_CELL_RE.match(r[c].replace(" ", "")):
                            vals.append(float(r[c].replace(",", "").rstrip("%")))
                    if len(vals) < 2:
                        continue
                    shown = float(tr[c].replace(",", "").replace(" ", "").rstrip("%"))
                    calc = sum(vals)
                    if abs(shown - calc) > max(0.05, abs(calc) * 0.005):
                        fails.append(f"total mismatch {name}:{start} column '{rows[0][c]}' shown {shown:g} calculated {calc:g}")
            for c in range(1, len(rows[0])):
                if not re.search(r"(비율|비중|구성비|점유율|%)", rows[0][c]):
                    continue
                vals = [float(r[c].replace(",", "").rstrip("%")) for r in body
                        if c < len(r) and NUM_CELL_RE.match(r[c].replace(" ", ""))]
                if len(vals) >= 3 and 90 < sum(vals) < 110 and abs(sum(vals) - 100) > 0.5:
                    warns.append(f"composition ratio total {sum(vals):.1f}% {name}:{start} column '{rows[0][c]}'")
    return not fails, fails, warns




# ═══════════════════════════════════════════════════════════════════════
# 7-2. Volume unit verdict (check volume)
# ═══════════════════════════════════════════════════════════════════════
def count_warnings(elems: list[Element]) -> list[str]:
    """Return as warnings the points where the number of sibling items exceeds the cap for the level. Reported once per group, at the first overrun.
    Why: sibling items over the cap blur the point, so this signals splitting into subsections or converting to a table"""
    out: list[str] = []
    squares = circles = dashes = 0
    sec_name = sq_name = ci_name = ""
    for e in elems:
        if e.kind in ("heading", "band"):
            squares = circles = dashes = 0
            sec_name, sq_name, ci_name = e.text[:24], "", ""
        elif e.kind == "square":
            squares += 1
            if squares == SQUARE_PER_SECTION_MAX + 1:
                out.append(f"[level count] '{sec_name}' has {squares} or more □ (cap {SQUARE_PER_SECTION_MAX}). Group into subsections (###)")
            circles = dashes = 0
            sq_name, ci_name = e.text[:24], ""
        elif e.kind == "circle":
            circles += 1
            if sq_name and circles == CIRCLE_PER_SQUARE_MAX + 1:
                out.append(f"[level count] □ '{sq_name}' has {circles} or more ❍ (cap {CIRCLE_PER_SQUARE_MAX}). Split the □")
            dashes, ci_name = 0, e.text[:24]
        elif e.kind == "dash":
            dashes += 1
            if ci_name and dashes == DASH_PER_CIRCLE_MAX + 1:
                out.append(f"[level count] ❍ '{ci_name}' has {dashes} or more hyphens (cap {DASH_PER_CIRCLE_MAX}). Convert to a table")
    return out


def check_volume(path: str, original: bool | None = None, style: bool | None = None) -> tuple[bool, list[str], list[str]]:
    """Two or more ❍ and two or more evidence lines (hyphen, ※) per □, one or more tables or figures per Task, and lower bounds on body characters.
    Why: well-written proposals build volume by filling composition units without stretching sentences. It is the only path that keeps both the short-sentence discipline and the volume."""
    elems, _ = parse_markdown(path)
    fails, warns = [], []
    squares: list[dict] = []
    tasks: list[dict] = []
    cur_sq = cur_task = None
    if original is None:
        original = manuscript_kind(path)[:1] in PREFIX_ORIGINAL
    if style is None:  # □ heading line cap and level count warnings. Planning reports get only these warnings, without lower bounds
        style = original
    for e in elems:
        if e.kind == "band" and e.data.get("kind") == "task":
            cur_task = {"name": e.text[:30], "blocks": 0, "chars": 0}
            tasks.append(cur_task)
            cur_sq = None
        elif e.kind == "band" or e.kind == "heading":
            cur_sq = None
        elif e.kind == "square":
            cur_sq = {"name": e.text[:30], "title": e.text, "circle": 0, "evidence": 0, "chars": len(e.text), "block_first": False}
            squares.append(cur_sq)
        elif e.kind in ("circle", "dash", "dot", "note"):
            if len(e.text) > LINE_MAX[e.kind]:
                warns.append(f"{e.kind} item {len(e.text)} characters ({'1-line' if e.kind == 'note' else '2-line'} cap {LINE_MAX[e.kind]} characters). Split the sentence and add a level, or convert to a table: {e.text[:24]}")
            if cur_sq is not None:
                cur_sq["chars"] += len(e.text)
                if e.kind == "circle":
                    cur_sq["circle"] += 1
                else:
                    cur_sq["evidence"] += 1
        if cur_sq is not None and e.kind in ("table", "image") and cur_sq["circle"] == 0:
            cur_sq["block_first"] = True   # a table or figure comes right under □ without ❍
        if cur_task is not None:
            if e.kind in ("table", "image"):
                cur_task["blocks"] += 1
                cur_task["chars"] += sum(len(c) for row in e.data.get("rows", []) for c in row)
            elif e.kind in ("square", "circle", "dash", "dot", "note"):
                cur_task["chars"] += len(e.text)
    for s in squares:
        if s["circle"] == 0:
            fails.append(f"□ '{s['name']}' alone. □ is a point and needs at least one ❍ below it. If only a table follows, first state what the table says as ❍")
        elif s["block_first"]:
            fails.append(f"□ '{s['name']}' has a table or figure right below it without ❍. Place a ❍ stating what the table says before it")
        elif original and s["circle"] < VOLUME_MIN_CIRCLE:
            fails.append(f"□ '{s['name']}' has {s['circle']} ❍ (lower bound {VOLUME_MIN_CIRCLE}). Split claims or merge □")
        if style and len(s["title"]) > SQUARE_TITLE_MAX:
            warns.append(f"□ '{s['name']}' heading {len(s['title'])} characters (one-line cap {SQUARE_TITLE_MAX} characters). Keep only the key words of the judgment and move elaboration to ❍")
        if not original:
            continue
        if s["evidence"] < VOLUME_MIN_EVIDENCE:
            warns.append(f"□ '{s['name']}' evidence lines (hyphen, ※) {s['evidence']} (recommended {VOLUME_MIN_EVIDENCE} or more). Add figures, material names, and cases on the evidence level")
        if s["chars"] < VOLUME_SQUARE_MIN_CHARS:
            warns.append(f"□ '{s['name']}' character count {s['chars']} (recommended {VOLUME_SQUARE_MIN_CHARS} or more)")
    if style:
        warns += count_warnings(elems)
    for t in tasks:
        if t["blocks"] < VOLUME_MIN_BLOCK_PER_TASK:
            fails.append(f"{t['name']} tables or figures {t['blocks']}. Place a filled table fitting the analysis method or a criteria table")
        if t["chars"] < VOLUME_TASK_MIN_CHARS:
            warns.append(f"{t['name']} body character count {t['chars']} (recommended {VOLUME_TASK_MIN_CHARS} or more). Add a named analysis target list, a filled preliminary survey table, a diagnosis item table, or a comparison matrix")
    return not fails, fails, warns


def resolve_manuscript(work: str, name: str) -> str:
    """Resolve the --file path. Searched relative to the current folder, the workspace, and the manuscript folder, in that order"""
    for cand in (name, os.path.join(work, name), os.path.join(work, D_MANUSCRIPT, os.path.basename(name))):
        if os.path.exists(cand):
            return cand
    sys.exit(f"ERROR: manuscript file not found {name} (not found relative to the current folder, the workspace, or {D_MANUSCRIPT}/)")


def cmd_check(a):
    work = a.work
    kind = work_kind(work)
    results = {}
    labels = use_labels(work, a.labels)
    # check all --skip: run without the named checks. The skip and its reason go into 05_점검기록.md and state.json, and build accepts this PASS record
    skips = [x.strip() for x in (a.skip or "").split(",") if x.strip()]
    if skips and a.what != "all":
        sys.exit("ERROR: --skip is for check all only")
    if (bad := [x for x in skips if x not in CHECK_NAMES]):
        sys.exit(f"ERROR: unknown check names {bad} (check names: {', '.join(CHECK_NAMES)})")
    if (extra := [x for x in skips if x not in CHECKS_BY_KIND[kind]]):
        print(f"WARN: --skip {extra} not in the check all list of the {kind} type. Ignored")
    skips = [x for x in CHECKS_BY_KIND[kind] if x in skips]
    reason = a.skip_reason or "no reason given (state one with --skip-reason)"
    if skips:
        print(f"[skipped] {', '.join(skips)} (reason: {reason})")

    def run(name: str) -> bool:
        """Individual checks run regardless of type. all runs the check list for the type (CHECKS_BY_KIND) minus the --skip targets"""
        return a.what == name or (a.what == "all" and name in CHECKS_BY_KIND[kind] and name not in skips)
    if run("rfp"):
        ok, fails, warns = check_rfp(work)
        _print_result("requirement placement", ok, fails, warns)
        record(work, "rfp", ok, {"fails": fails})
        results["requirement placement"] = (ok, fails, warns)
    if run("density"):
        targets = [resolve_manuscript(work, a.file)] if a.file else manuscript_files(work, PREFIX_DENSITY)
        for f in targets:
            ok, fails, warns = check_density(f)
            _print_result(f"density {os.path.basename(f)}", ok, fails, warns)
            record(work, f"density:{os.path.basename(f)}", ok, {"fails": fails})
            results[f"density {os.path.basename(f)}"] = (ok, fails, warns)
    if run("volume"):
        if kind == "proposal":
            targets = [resolve_manuscript(work, a.file)] if a.file else [f for p in PREFIX_VOLUME for f in manuscript_files(work, p)]
        else:
            targets = [resolve_manuscript(work, a.file)] if a.file else manuscript_files(work)
        original = None if kind == "proposal" else kind == "report"   # reports apply composition unit lower bounds to all manuscripts
        for f in targets:
            ok, fails, warns = check_volume(f, original, None if kind == "proposal" else True)
            _print_result(f"volume composition {os.path.basename(f)}", ok, fails, warns)
            record(work, f"volume:{os.path.basename(f)}", ok, {"fails": fails})
            results[f"volume composition {os.path.basename(f)}"] = (ok, fails, warns)
    if run("wording"):
        ok, fails, warns = check_wording(work)
        _print_result("wording", ok, fails, warns)
        record(work, "wording", ok, {"fails": fails})
        results["wording"] = (ok, fails, warns)
    if run("tone"):
        ok, fails, warns = check_tone(work, a.reset_whitelist, kind)
        _print_result("tone and vocabulary", ok, fails, warns)
        record(work, "tone", ok, {"fails": fails})
        results["tone and vocabulary"] = (ok, fails, warns)
    if run("numbers"):
        ok, fails, warns = check_numbers(work)
        _print_result("figure check", ok, fails, warns)
        record(work, "numbers", ok, {"fails": fails})
        results["figure check"] = (ok, fails, warns)
    if run("fill"):
        name = "common part substitution" if kind == "proposal" else "leftover confirmation markers"
        ok, fails, warns = check_fill(work) if kind == "proposal" else check_residual(work)
        _print_result(name, ok, fails, warns)
        record(work, "fill", ok, {"fails": fails})
        results[name] = (ok, fails, warns)
    if a.what == "all":
        out = ["# 점검기록", f"- Run {datetime.now().isoformat(timespec='minutes')}"] + ([f"- Labels {labels}"] if labels else []) + [""]
        if skips:
            out += ["## Skipped checks", f"- (skipped) {', '.join(skips)} (check all --skip)", f"- (reason) {reason}", ""]
        for name, (ok, fails, warns) in results.items():
            out.append(f"## {name} {'PASS' if ok else 'FAIL'}")
            out += [f"- (defect) {f}" for f in fails] + [f"- (warning) {w}" for w in warns] + [""]
        write_check_report(work, out)
        allok = all(v[0] for v in results.values())
        st = load_state(work)
        st["stages"]["check"] = datetime.now().isoformat(timespec="seconds")
        st["checks"]["all"] = {"pass": allok, "hash": manuscript_hash(work),
                               **({"skipped": skips, "skip_reason": reason} if skips else {})}
        save_state(work, st)
        print(f"{'OK' if allok else 'FAIL'}: {F_REPORT} recorded. {'build may proceed' if allok else 'fix defects and rerun'}")
        sys.exit(0 if allok else 1)


# ═══════════════════════════════════════════════════════════════════════
# 8. Typesetting, render, and skeleton commands (the typesetter body is hwpx_report.py)
# ═══════════════════════════════════════════════════════════════════════
def _profile_for_builder(profile: dict) -> dict:
    """lib.Builder validates every profile value as a style name, so pass a copy without keys that are not style names (sizes, format strings, band settings)."""
    return {k: v for k, v in profile.items() if not k.endswith("_pt") and k not in PROFILE_EXTRA_KEYS}


def docx_spec() -> dict:
    """Template definition passed to the docx writer. Uses only the settings values (SKELETON_*) shared with the hwpx skeleton"""
    return {"fonts": {"body": SKELETON_FONT_BODY, "head": SKELETON_FONT_HEAD}, "color": SKELETON_TEXT_COLOR,
            "styles": SKELETON_BASE_STYLES + SKELETON_STYLES, "glyphs": SKELETON_GLYPHS, "glyph_em": SKELETON_GLYPH_EM,
            "bullet_offset": int(SKELETON_BULLET_OFFSET), "toc_tab": SKELETON_TOC_TAB_POS, "band_border": BAND_BORDER_W,
            "space_em": SKELETON_SPACE_EM, "divider_top_lines": DIVIDER_TOP_LINES,
            "page": {"width": SKELETON_PAGE_SIZE[0], "height": SKELETON_PAGE_SIZE[1], **SKELETON_PAGE_MARGIN}}


def work_kind(work: str | None) -> str:
    """Output type of the workspace (init --kind). Proposal if there is no record"""
    return (load_state(work).get("kind") if work else None) or "proposal"


def module_table_parent(files: list[str]) -> str | None:
    """Return the file whose heading (# or ##) the module overview table sits under. If the overview table file has a heading before the table, the file itself.
    Why: if the overview table file has no chapter or section heading, the table is typeset under the last section of the previous file (proposer general status in the default composition)."""
    parent = None
    for f in files:
        is_table_file = manuscript_kind(f).endswith(MODULE_TABLE_SUFFIX)
        for ln in open(f, encoding="utf-8").read().split("\n"):
            st = ln.strip()
            if is_table_file and (st.startswith("|") or caption_start(st)):
                return parent
            if re.match(r"^#{1,2} ", st):
                parent = f
    return None


FOOTNOTE_REF_RE = re.compile(r"\[\^([^\]]+)\]")


def inline_note_footnotes(text: str, footnotes: dict) -> str:
    """Note rows below tables and figures are transparent table rows that cannot hold footnotes, so footnote notation prints as literal text.
    Write the footnote content out in parentheses and warn to fix the manuscript"""
    def repl(m):
        body = footnotes.get(m.group(1), "")
        label = re.sub(r"^f\d+-", "", m.group(1))   # the file scope prefix added when merging is not in the manuscript, so it is removed in the notice
        print(f"WARN: footnote [^{label}] on a note line below a table typeset as a parenthesized source. Move it to a `※ 출처:` line in the manuscript")
        return f" ({body})" if body else ""
    return FOOTNOTE_REF_RE.sub(repl, text)
FOOTNOTE_DEF_LINE_RE = re.compile(r"^\[\^([^\]]+)\]:", re.M)


def scope_footnotes(text: str, scope: str) -> tuple[str, list[str]]:
    """Attach a file scope prefix to footnote labels. The merge parser collects footnote definitions into one dictionary keyed by label,
    so using [^1] in every file lets a later file's definition overwrite an earlier file's footnote. Returns: (converted manuscript, list of warnings)"""
    defs = FOOTNOTE_DEF_LINE_RE.findall(text)
    warns = [f"duplicate footnote definition [^{x}]" for x in sorted({d for d in defs if defs.count(d) > 1})]
    undefined = sorted(set(FOOTNOTE_REF_RE.findall(text)) - set(defs))
    warns += [f"footnote without definition [^{x}] (typeset as an empty footnote)" for x in undefined]
    return FOOTNOTE_REF_RE.sub(lambda m: f"[^{scope}-{m.group(1)}]", text), warns


def cmd_build(a):
    work = a.input if os.path.isdir(a.input) else None
    kind = work_kind(work)
    use_labels(work, a.labels)
    profile_key = a.profile or KINDS[kind]["profile"]
    profile = PROFILES.get(profile_key) if profile_key in PROFILES else json.load(open(profile_key, encoding="utf-8"))
    profile = lib.apply_format_labels(profile)   # label formats.* override the heading number formats when present
    formats = ("hwpx", "docx") if a.format == "both" else (a.format,)
    if "docx" in formats and (a.skeleton or profile_key not in PROFILES):
        sys.exit("ERROR: docx supports only the default template (settings definitions). Typeset client templates (--skeleton, profile JSON) with --format hwpx")
    skeleton = None
    if "hwpx" in formats:
        if a.skeleton:
            skeleton = a.skeleton
        elif profile_key in PROFILES:  # default profile: created in _양식/ of the workspace (for a single manuscript, the folder holding it)
            skeleton = write_default_skeleton(work or os.path.dirname(os.path.abspath(a.input)))
        else:
            sys.exit("ERROR: with a profile JSON, also set the skeleton hwpx with --skeleton")
        if not os.path.exists(skeleton):
            sys.exit(f"ERROR: skeleton not found {skeleton}")
    if work:  # workspace: merge 원고/*.md in file name order
        files = manuscript_files(work)
        if not files:
            sys.exit(f"ERROR: no manuscripts in {work}/{D_MANUSCRIPT}/")
        merged = os.path.join(work, D_DELIVER, "_합본.md")
        os.makedirs(os.path.dirname(merged), exist_ok=True)
        has_design = os.path.exists(os.path.join(work, F_MODULE_TABLE))
        inserted = 0
        with open(merged, "w", encoding="utf-8") as fo:
            for i, f in enumerate(files, 1):
                text, fn_warns = scope_footnotes(open(f, encoding="utf-8").read(), f"f{i}")
                for w in fn_warns:
                    print(f"WARN: {os.path.basename(f)} {w}")
                if has_design:
                    text, n = expand_module_table(work, text)
                    inserted += n
                fo.write(text.rstrip() + "\n\n")
        # a proposal requires the module overview table once at the start of the execution measures. Reports and plans are optional (0 or 1 time)
        need_one = kind == "proposal"
        if has_design and (inserted != 1 if need_one else inserted > 1) and not legacy_module_table(work):
            msg = (f"module overview table insertion directive lines: {inserted}. After the ❍ text where the execution measures start, "
                   "write one `(삽입: 모듈총괄표)` line and then build")
            if not a.force:
                sys.exit("ERROR: " + msg)
            print("WARN: " + msg)
        src = merged
        out_base = os.path.splitext(a.output)[0] if a.output else os.path.join(work, D_DELIVER, lib.label(f"output.{kind}"))
        parent = module_table_parent(files)
        if kind == "proposal" and parent and manuscript_kind(parent)[:1] in PREFIX_FILL:  # --force is for trial typesetting, so it warns instead of stopping
            msg = (f"the module overview table sits under a common part heading ({os.path.basename(parent)}). "
                   f"Move the chapter heading (#) or section heading (##) where the execution measures start to the first line of the module overview table file, then build")
            if not a.force:
                sys.exit("ERROR: " + msg)
            print("WARN: " + msg)
        st = load_state(work)
        rec = st.get("checks", {}).get("all", {})
        if not rec.get("pass") and not a.force:
            sys.exit("ERROR: no `check all` PASS record. Fix the defects, pass check all, then build. Use --force only for trial typesetting")
        if rec.get("hash") and rec["hash"] != manuscript_hash(work) and not a.force:
            sys.exit("ERROR: manuscripts, the module overview table, or the label file changed after check all passed. Rerun check all, then build")
        if rec.get("pass") and rec.get("skipped"):
            print(f"NOTE: the check all PASS record skipped {', '.join(rec['skipped'])} (reason: {rec.get('skip_reason', '')})")
        print(f"merged {len(files)} files → {merged}")
    else:
        src = a.input
        out_base = os.path.splitext(a.output or a.input)[0]
    elems, footnotes = parse_markdown(src)
    msgs = lint(elems)
    for msg in msgs:
        print("LINT:", msg)
    if a.long_table is None and any("[long table HITL]" in m for m in msgs):
        sys.exit("ERROR: a table needs long-table handling but --long-table is not set. Set the kickoff answer (engine or split) with --long-table, then build")

    def make(fmt: str, out: str, numbers: dict[str, int] | None = None) -> ProposalLayout:
        """One typesetting run. With numbers, rewrite the typeset copy with page numbers filled in the TOC. Manuscript elements change during typesetting, so a copy is passed"""
        if fmt == "hwpx":
            b = ProposalBuilder(skeleton, _profile_for_builder(profile), page_break=a.page_break)
            b.p = profile
        else:
            b = ProposalDocxBuilder(docx_spec(), profile, page_break=a.page_break)
        b.numbered_square = False
        b.toc = a.toc
        b.page_number = a.page_number
        b.toc_depth = a.toc_depth
        b.cover = dict(kv.split("=", 1) for kv in a.cover.split(";") if "=" in kv) if a.cover else None
        b.header_spec = a.header
        b.merge_mode = a.merge
        b.long_table_mode = a.long_table or "split"
        b.toc_numbers = numbers or {}
        b.unify_mode = a.unify
        b.theme = a.theme
        b.doc_title = a.title
        b.divider = (a.divider or ("on" if "report" in (kind, profile_key) else "off")) == "on"   # divider default: ON for reports, OFF otherwise
        b.n["chapter"] = a.chapter_start - 1   # align chapter numbers of a partial manuscript with the full TOC
        b.build(copy.deepcopy(elems), footnotes, out)
        return b

    made = {}
    for fmt in formats:
        out = f"{out_base}.{fmt}"
        b = make(fmt, out)
        if fmt == "hwpx":
            fixed = lib.repair_manifest(out)
            if fixed:
                print(f"WARN: {fixed} images not registered in the manifest were corrected automatically")
        if a.toc and a.page_number:  # pass 2: find the page positions of titles and captions by rendering and fill them into the TOC (approximate, finalized in 한/글 and Word)
            numbers = toc_page_numbers(out, fmt, [t for _, t in b.headings], [(t, needle) for _, t, needle in b.captions],
                                       {t for lv, t in b.headings if lv == 1})
            if numbers:
                make(fmt, out, numbers)
                if fmt == "hwpx":
                    lib.repair_manifest(out)
                # count only entries in the TOC. Why: titles down to subsections (###) and all captions go into detection, but the TOC holds only up to toc_depth and the list of tables setting
                listed = [t for lv, t in b.headings if lv <= b.toc_depth] + [t for _, items in b.toc_lists() for t in items]
                print(f"TOC page numbers filled {sum(t in numbers for t in listed)}/{len(listed)} ({fmt}, approximate renderer)")
            else:
                print(f"WARN: renderer or pdftotext missing, so {fmt} TOC page numbers not filled")
        made[fmt] = out
        print(f"OK: {out} (elements {len(elems)}, footnotes {len(footnotes)})")
    if len(made) == 2:  # both formats produced together: compare the character order of paragraphs
        diffs = check_parity(made["hwpx"], made["docx"])
        if diffs:
            print(f"WARN: hwpx and docx paragraph text mismatch in {len(diffs)} places (signals a typesetting code defect)")
            for d in diffs[:5]:
                print("   ", d)
        else:
            print(f"OK: hwpx and docx paragraph text match ({len(hwpx_paragraph_texts(made['hwpx'])[0])} paragraphs, TOC page numbers excluded)")
    if work:
        st = load_state(work)
        st["stages"]["build"] = datetime.now().isoformat(timespec="seconds")
        for fmt, out in made.items():
            st["files"][fmt] = os.path.abspath(out)
        save_state(work, st)


W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
TOC_NUM_RE = re.compile(r"\t\d+$")   # TOC page numbers depend on each format's rendering, so they are excluded from the comparison


def hwpx_paragraph_texts(path: str) -> tuple[list[str], list[str]]:
    """Paragraph text of the hwpx body (including table cells and captions, excluding headers) and footnote text. In document order"""
    body, notes = [], []
    with zipfile.ZipFile(path) as z:
        for n in sorted(x for x in z.namelist() if re.match(r"Contents/section\d+\.xml$", x)):
            for p in etree.fromstring(z.read(n)).iter(f"{HP}p"):
                anc = {a.tag for a in p.iterancestors()}
                if f"{HP}header" in anc or f"{HP}footer" in anc:
                    continue
                buf = []
                for t in p.iter(f"{HP}t"):
                    owner = next(a for a in t.iterancestors() if a.tag == f"{HP}p")
                    if owner is not p:
                        continue
                    buf.append(t.text or "")
                    buf += ["\t" + (c.tail or "") for c in t if c.tag == f"{HP}tab"]
                text = "".join(buf)
                if text.strip():
                    (notes if f"{HP}footNote" in anc else body).append(text)
    return body, notes


def docx_paragraph_texts(path: str) -> tuple[list[str], list[str]]:
    """Paragraph text of the docx body and footnote text. Caption number fields and the parenthesis after footnote numbers are removed to match hwpx automatic numbers"""
    with zipfile.ZipFile(path) as z:
        root = etree.fromstring(z.read("word/document.xml"))
        fn = etree.fromstring(z.read("word/footnotes.xml")) if "word/footnotes.xml" in z.namelist() else None
    body = []
    for p in root.iter(f"{W_NS}p"):
        buf, after_ref = [], False
        for el in p.iter(f"{W_NS}t", f"{W_NS}tab", f"{W_NS}footnoteReference"):
            if el.tag == f"{W_NS}footnoteReference":
                after_ref = True
            elif any(a.tag == f"{W_NS}fldSimple" for a in el.iterancestors()):
                continue
            elif el.tag == f"{W_NS}tab":
                buf.append("\t")
            elif after_ref and el.text == docx_lib.FOOTNOTE_SUFFIX:
                after_ref = False
            else:
                buf.append(el.text or "")
        if "".join(buf).strip():
            body.append("".join(buf))
    notes = []
    for note in (fn.findall(f"{W_NS}footnote") if fn is not None else []):
        if not note.get(f"{W_NS}type"):
            text = "".join(t.text or "" for t in note.iter(f"{W_NS}t"))
            notes.append(text[len(docx_lib.FOOTNOTE_SUFFIX):].strip())
    return body, notes


def check_parity(hwpx_path: str, docx_path: str) -> list[str]:
    """Compare the character order of paragraphs of hwpx and docx made from the same manuscript. Returns a list of differences (empty if they match)"""
    hb, hn = hwpx_paragraph_texts(hwpx_path)
    db, dn = docx_paragraph_texts(docx_path)
    hb, db = [TOC_NUM_RE.sub("", t) for t in hb], [TOC_NUM_RE.sub("", t) for t in db]
    diffs = [f"{i + 1}: paragraph mismatch, hwpx '{a[:30]}' docx '{b[:30]}'" for i, (a, b) in enumerate(zip(hb, db)) if a != b]
    if len(hb) != len(db):
        diffs.append(f"paragraph count hwpx {len(hb)} docx {len(db)}")
    if [t.strip() for t in hn] != dn:
        diffs.append(f"footnotes hwpx {len(hn)} docx {len(dn)} mismatch")
    return diffs


def toc_page_numbers(path: str, fmt: str, headings: list[str], captions: list[tuple[str, str]],
                     chapters: set[str] | None = None) -> dict[str, int]:
    """Render the typeset copy into a PDF that keeps a text layer and find the body page numbers of TOC entries. hwpx uses rhwp first, docx uses LibreOffice"""
    if not shutil.which("pdftotext"):
        return {}
    pdf = os.path.splitext(path)[0] + ".toc.pdf"
    if fmt == "hwpx":
        if not (lib.find_rhwp() or lib.find_soffice()):
            return {}
        engine = lib.resolve_engine("auto", None)
        src = lib.render_copy(path, strip_newnum=(engine == "soffice"))
    else:
        if not lib.find_soffice():
            return {}
        engine, src = "soffice", path
    try:
        lib.to_pdf(src, pdf, engine, None, text_as_paths=False)
    except SystemExit:
        return {}
    finally:
        if src != path and os.path.exists(src):
            os.remove(src)
    pages = subprocess.run(["pdftotext", "-layout", pdf, "-"], capture_output=True, text=True).stdout.split("\f")
    os.remove(pdf)
    return lib.match_toc_pages(pages, headings, captions, chapters)


_ORIGINAL_COL_ALIGNS = lib.col_aligns


def _col_aligns_symbol_neutral(rows: list[list[str]], explicit: list[str | None]) -> list[str]:
    """Cells with only one symbol (○ in a comparison matrix, ■ in a schedule chart) do not count as bullets in the column alignment judgment.
    Why: lib sets a column to LEFT if it has a cell starting with ○. Symbol columns leaned left together with ● and ◐. Explicit markers take priority as before."""
    masked = [rows[0]] + [["" if v.strip() in lib.BULLET_PREFIXES else v for v in r] for r in rows[1:]]
    return _ORIGINAL_COL_ALIGNS(masked, explicit)


lib.col_aligns = _col_aligns_symbol_neutral

_ORIGINAL_RENDER_COPY = lib.render_copy


def _cell_bullet_paras(header_xml: str) -> dict[str, str]:
    """paraPr id of table cell bullet styles (style name containing '표', heading BULLET) → bullet character."""
    bullets = dict(re.findall(r'<hh:bullet id="(\d+)"[^>]*char="([^"]*)"', header_xml))
    out = {}
    for m in re.finditer(r'<hh:paraPr id="(\d+)".*?</hh:paraPr>', header_xml, re.S):
        pp_id, body = m.group(1), m.group(0)
        h = re.search(r'<hh:heading type="BULLET" idRef="(\d+)"', body)
        if h and re.search(rf'<hh:style [^>]*name="[^"]*표[^"]*"[^>]*paraPrIDRef="{pp_id}"', header_xml):
            out[pp_id] = bullets.get(h.group(1), "◦").replace("\u00ad", "-")
    return out


def _render_copy_with_bullets(src: str, strip_newnum: bool = False) -> str:
    """In the review copy, replace soft hyphens in bullet definitions (header.xml) with hyphens and write table cell bullets as body text.
    Why: approximate renderers do not draw soft hyphen bullets or bullets inside table cells, so levels and cell items look empty. The original hwpx is unchanged."""
    work = _ORIGINAL_RENDER_COPY(src, strip_newnum)
    tmp = work + ".tmp"
    with zipfile.ZipFile(work) as zin:
        cell_bullets = _cell_bullet_paras(zin.read("Contents/header.xml").decode("utf-8"))
    with zipfile.ZipFile(work) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for it in zin.infolist():
            d = zin.read(it.filename)
            if it.filename == "Contents/header.xml":
                d = d.decode("utf-8").replace('char="\u00ad"', 'char="-"').encode("utf-8")
            # rhwp (as of v0.8.6) draws cell bullets directly. Character insertion applies only to the soffice path (strip_newnum) to prevent double display
            if it.filename.startswith("Contents/section") and it.filename.endswith(".xml") and cell_bullets and strip_newnum:
                t = d.decode("utf-8")
                for pp_id, ch in cell_bullets.items():
                    t = re.sub(rf'(<hp:p [^>]*paraPrIDRef="{pp_id}"[^>]*>.*?<hp:t>)', rf'\g<1>{ch} ', t, flags=re.S)
                d = t.encode("utf-8")
            if it.filename == "mimetype":
                zout.writestr(zipfile.ZipInfo("mimetype"), d, compress_type=zipfile.ZIP_STORED)
            else:
                zout.writestr(it.filename, d)
    os.replace(tmp, work)
    return work


def cmd_render(a):
    lib.render_copy = _render_copy_with_bullets
    return lib.cmd_render(a)


def _run_self(args: list[str]) -> subprocess.CompletedProcess:
    """Run this CLI as a subprocess. Why: catches sys.exit errors of build and render to record them as verdict results"""
    return subprocess.run([sys.executable, os.path.abspath(__file__), *args], capture_output=True, text=True)


def cmd_selftest(a):
    """Typeset and render a sample manuscript to judge whether the render path works. The result is recorded in state.json checks.render.
    Why: finds renderer defects at the start, when the context has room, to prevent failures to secure a renderer at step 6"""
    out_dir = os.path.abspath(os.path.join(a.work, D_SELFTEST))  # Why: keeps the state.json record path independent of the working location
    os.makedirs(out_dir, exist_ok=True)
    sample_md = os.path.join(out_dir, "표본.md")
    sample_hwpx = os.path.join(out_dir, "표본.hwpx")
    open(sample_md, "w", encoding="utf-8").write(SELFTEST_SAMPLE)

    built = _run_self(["build", sample_md, sample_hwpx, "--skeleton", write_default_skeleton(a.work)])
    if built.returncode != 0 or not os.path.exists(sample_hwpx):
        record(a.work, "render", False, {"stage": "build", "error": (built.stderr or built.stdout)[-400:]})
        sys.exit(f"FAIL[build]: sample typesetting failed. Check dependency installation (Preparation section)\n{(built.stderr or built.stdout)[-400:]}")

    render_args = ["render", sample_hwpx, "--engine", a.engine] + (["--rhwp", a.rhwp] if a.rhwp else [])
    rendered = _run_self(render_args)
    pages_dir = os.path.splitext(sample_hwpx)[0] + "_pages"
    pages = len(glob.glob(os.path.join(pages_dir, "*.png")))
    engine = m.group(1) if (m := re.search(r"OK\[([^\]]+)\]", rendered.stdout)) else ""
    passed = rendered.returncode == 0 and pages >= SELFTEST_MIN_PAGES
    detail = {"stage": "render", "engine": engine, "pages": pages, "png_dir": pages_dir}
    if not passed:
        detail["error"] = (rendered.stderr or rendered.stdout)[-400:]
    record(a.work, "render", passed, detail)
    if not passed:
        sys.exit(f"FAIL[render]: engine {engine or 'none'}, PNG {pages} page(s). Move to the next-priority renderer in the Preparation section\n{detail['error']}")
    print(f"PASS[{engine}]: sample rendered {pages} page(s) → {pages_dir}/ (open the PNG and check bullets, bands, and tables)")
    # docx path: typeset the same sample as docx and render with LibreOffice. Result in checks.render_docx (a gate only for docx delivery)
    sample_docx = os.path.join(out_dir, "표본_docx.docx")   # a separate name so the PNG folder does not overlap with the hwpx sample
    built = _run_self(["build", sample_md, sample_docx, "--format", "docx"])
    if built.returncode != 0 or not os.path.exists(sample_docx):
        record(a.work, "render_docx", False, {"stage": "build", "error": (built.stderr or built.stdout)[-400:]})
        print(f"WARN[docx]: sample docx typesetting failed. Check the python-docx installation (Preparation section)\n{(built.stderr or built.stdout)[-400:]}")
        return
    if not lib.find_soffice():
        record(a.work, "render_docx", False, {"stage": "render", "error": "LibreOffice missing"})
        print("WARN[docx]: docx render review impossible because LibreOffice is missing. For docx delivery, install LibreOffice per the Preparation section")
        return
    rendered = _run_self(["render", sample_docx])
    docx_pages_dir = os.path.splitext(sample_docx)[0] + "_pages"
    n_docx = len(glob.glob(os.path.join(docx_pages_dir, "*.png")))
    ok = rendered.returncode == 0 and n_docx >= SELFTEST_MIN_PAGES
    record(a.work, "render_docx", ok, {"stage": "render", "pages": n_docx, "png_dir": docx_pages_dir})
    print(f"{'PASS' if ok else 'WARN'}[docx]: sample rendered {n_docx} page(s) → {docx_pages_dir}/")


def main():
    ap = argparse.ArgumentParser(description="juja-builder CLI", epilog=lib.CLI_NOTICE,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)

    p = sub.add_parser("init", help="create the workspace")
    p.add_argument("work")
    p.add_argument("--kind", default="proposal", choices=sorted(KINDS), help="output type: proposal (default), report (pre-diagnosis report), plan (planning report)")
    p.set_defaults(func=cmd_init)

    p = sub.add_parser("rfp", help="RFP (hwpx|pdf|docx|md) → 01_RFP원문.md and empty requirement and scoring tables")
    p.add_argument("input")
    p.add_argument("work")
    p.set_defaults(func=cmd_rfp)

    p = sub.add_parser("ledger", help="create and update the section ledger by tallying the scoring table")
    p.add_argument("work")
    p.set_defaults(func=cmd_ledger)

    p = sub.add_parser("deck", help="proposal deck (pptx|pdf) → 발표자료/ 원문.md and diagram capture material")
    p.add_argument("input")
    p.add_argument("work")
    p.add_argument("--allow-garbled", action="store_true", help="proceed with an incomplete text layer PDF on the premise of reading page PNGs (only on the user's decision)")
    p.set_defaults(func=cmd_deck)

    p = sub.add_parser("figure", help="deck page → diagram PNG for insertion (발표자료/figures)")
    p.add_argument("work")
    p.add_argument("page", type=int)
    p.add_argument("--grid", action="store_true", help="create only a 10% grid PNG for reading crop values")
    p.add_argument("--crop", help="left,top,right,bottom percentages (e.g. 4,28,97,96)")
    p.add_argument("--name", help="file name suffix")
    p.add_argument("--no-trim", action="store_true", help="turn off automatic removal of outer margins")
    p.add_argument("--pdf", help="replacement capture source PDF. Capture from the PDF saved by PowerPoint a diagram whose text overflowed in the pptx conversion")
    p.set_defaults(func=cmd_figure)

    p = sub.add_parser("check", help="rfp | density | volume | wording | tone | numbers | fill | all")
    p.add_argument("what", choices=[*CHECK_NAMES, "all"])
    p.add_argument("work")
    p.add_argument("--file", help="one manuscript for density/volume")
    p.add_argument("--reset-whitelist", action="store_true", help="refix the tone whitelist (only on the user's instruction)")
    p.add_argument("--skip", help="checks to leave out of check all (comma separated, for example density,wording,tone,numbers). Only for the Korean-specific checks of manuscripts not written in Korean")
    p.add_argument("--skip-reason", help="reason for --skip. Recorded in 05_점검기록.md and state.json")
    p.add_argument("--labels", help="label file (JSON). Applied to the recognition of manuscript notation and recorded in state.json (none clears the record)")
    p.set_defaults(func=cmd_check)

    p = sub.add_parser("build", help="workspace (merge of 원고/*.md) or manuscript.md → hwpx, docx")
    p.add_argument("input")
    p.add_argument("output", nargs="?", help="output path (--format sets the extension)")
    p.add_argument("--profile", default=None, help="proposal | report | plan | profile JSON path (default: the workspace output type, proposal for a single manuscript)")
    p.add_argument("--format", default="hwpx", choices=["hwpx", "docx", "both"], help="output format (kickoff answer). Both formats are typeset from the same manuscripts and template definitions")
    p.add_argument("--title", help="document title on the first line of the body (planning report)")
    p.add_argument("--chapter-start", type=int, default=1, help="first chapter number (the number in the full TOC when producing only some chapters)")
    p.add_argument("--divider", choices=["on", "off"], default=None, help="chapter divider pages (default: on for pre-diagnosis reports, off for proposals and planning reports)")
    p.add_argument("--skeleton", help="skeleton hwpx (default: workspace/_양식/proposal.hwpx generated on every run)")
    p.add_argument("--page-break", default="chapter", choices=["chapter", "section", "none"])
    p.add_argument("--long-table", default=None, choices=["engine", "split"], help="default handling of long tables (kickoff answer). split=split into chunk tables with the header row repeated, engine=한/글 engine split. build stops if a table needs it and this is not set")
    p.add_argument("--merge", default="auto", choices=["auto", "explicit", "none"])
    p.add_argument("--toc", action="store_true", help="insert a static TOC (default OFF. The delivered TOC uses the 한/글 table of contents feature)")
    p.add_argument("--toc-depth", type=int, default=2, choices=[1, 2, 3])
    p.add_argument("--page-number", action="store_true", help="footer page numbers (default OFF)")
    p.add_argument("--cover", help='cover: "title=title;subtitle=subtitle;date=YYYY. M.;org=organization"')
    p.add_argument("--header", help='header: "text;align=RIGHT"')
    p.add_argument("--unify", default="manual", choices=["manual", "auto"], help="table width unification: manual=apply only (폭: ...) directives in manuscripts (default), auto=also group tables of the same layout automatically")
    p.add_argument("--theme", default=DEFAULT_THEME, choices=sorted(THEMES), help="band color theme: navy (default), mono (black and gray). Table format does not depend on the theme")
    p.add_argument("--labels", help="label file (JSON). Overrides the document's fixed strings, manuscript notation, language tags, and delivered file names. For a workspace it is recorded in state.json (none clears the record)")
    p.add_argument("--force", action="store_true")
    p.set_defaults(func=cmd_build)

    p = sub.add_parser("render", help="hwpx or docx → PDF + PNG (delegated to hwpx_report.py render. docx uses LibreOffice)")
    p.add_argument("input")
    p.add_argument("--outdir")
    p.add_argument("--dpi", type=int, default=lib.RENDER_DPI)
    p.add_argument("--rhwp")
    p.add_argument("--keep-soft-hyphen", action="store_true")
    p.add_argument("--engine", default="auto", choices=["auto", "rhwp", "soffice"])
    p.set_defaults(func=cmd_render)

    p = sub.add_parser("selftest", help="judge the render path by typesetting and rendering a sample manuscript (start gate, recorded in state.json checks.render)")
    p.add_argument("work")
    p.add_argument("--rhwp")
    p.add_argument("--engine", default="auto", choices=["auto", "rhwp", "soffice"])
    p.set_defaults(func=cmd_selftest)

    p = sub.add_parser("skeleton", help="client template.hwpx → skeleton + draft profile (delegated to hwpx_report.py skeleton)")
    p.add_argument("input")
    p.add_argument("output")
    p.add_argument("--section", type=int)
    p.set_defaults(func=lib.cmd_skeleton)

    a = ap.parse_args()
    a.func(a)


if __name__ == "__main__":
    main()
