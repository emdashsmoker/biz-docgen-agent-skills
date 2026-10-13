#!/usr/bin/env python3
"""validate.py: structural validation of the finished PPTX package (run after finalize.py)

WHY: PowerPoint refuses to open files whose package structure is off. Rendering review runs on LibreOffice,
     so defects that LibreOffice opens leniently must be caught separately.

Checks (exit code 1 if there is at least 1 defect)
1. ZIP opens and [Content_Types].xml exists
2. Every part has a content type declaration (extension Default or Override)
3. Every XML part is well-formed (lxml parsing)
4. Internal target parts of relationships (.rels) exist
5. Slide list ids in presentation.xml are unique and their relationship ids exist
6. Shape ids (cNvPr) are unique per slide
7. python-pptx opens the file and iterates the slides
8. (Optional) If the OOXML_PML_XSD environment variable points to a PresentationML schema file, slide XML is schema-validated

Usage: python3 validate.py <file.pptx>
"""
from __future__ import annotations

import os
import posixpath
import sys
import zipfile
from collections import Counter

from lxml import etree

CONTENT_TYPES = "[Content_Types].xml"
NS = {
    "ct": "http://schemas.openxmlformats.org/package/2006/content-types",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}
XML_SUFFIXES = (".xml", ".rels")
SLIDE_ID_MIN = 256  # Lower bound of sldId in ECMA-376


def rels_source_dir(rels_name: str) -> str:
    """Folder of the source part a relationship file describes. The source of _rels/x.xml.rels is x.xml"""
    folder = posixpath.dirname(posixpath.dirname(rels_name))
    return folder


def check_content_types(zf: zipfile.ZipFile, parts: set[str], errors: list[str]) -> None:
    if CONTENT_TYPES not in parts:
        errors.append("[Content_Types].xml missing")
        return
    root = etree.fromstring(zf.read(CONTENT_TYPES))
    defaults = {el.get("Extension", "").lower() for el in root.findall("ct:Default", NS)}
    overrides = {el.get("PartName", "").lstrip("/") for el in root.findall("ct:Override", NS)}
    for name in sorted(parts):
        if name == CONTENT_TYPES:
            continue
        ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
        if name not in overrides and ext not in defaults:
            errors.append(f"Part without content type declaration: {name}")
    for name in sorted(overrides - parts):
        errors.append(f"Override declared for a missing part: {name}")


def check_xml_parts(zf: zipfile.ZipFile, parts: set[str], errors: list[str]) -> dict[str, etree._Element]:
    trees = {}
    for name in sorted(parts):
        if not name.endswith(XML_SUFFIXES):
            continue
        try:
            trees[name] = etree.fromstring(zf.read(name))
        except etree.XMLSyntaxError as exc:
            errors.append(f"XML syntax error {name}: {exc}")
    return trees


def check_relationships(trees: dict, parts: set[str], errors: list[str]) -> dict[str, dict[str, str]]:
    """Whether relationship targets exist. Returns: {rId: target part} per source part"""
    rel_map: dict[str, dict[str, str]] = {}
    for name, root in trees.items():
        if not name.endswith(".rels"):
            continue
        src_dir = rels_source_dir(name)
        src_part = posixpath.join(src_dir, posixpath.basename(name)[: -len(".rels")])
        targets = {}
        for rel in root.findall("rel:Relationship", NS):
            if rel.get("TargetMode") == "External":
                continue
            target = rel.get("Target", "")
            resolved = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(src_dir, target))
            targets[rel.get("Id")] = resolved
            if resolved not in parts:
                errors.append(f"Relationship target missing {name} {rel.get('Id')} → {resolved}")
        rel_map[src_part.lstrip("/")] = targets
    return rel_map


def check_presentation(trees: dict, rel_map: dict, errors: list[str]) -> None:
    pres = trees.get("ppt/presentation.xml")
    if pres is None:
        errors.append("ppt/presentation.xml missing")
        return
    ids = [el.get("id") for el in pres.findall("p:sldIdLst/p:sldId", NS)]
    for sid, n in Counter(ids).items():
        if n > 1:
            errors.append(f"Duplicate slide id: {sid}")
    for sid in ids:
        if sid is None or not sid.isdigit() or int(sid) < SLIDE_ID_MIN:
            errors.append(f"Slide id out of range: {sid}")
    rid_attr = f"{{{NS['r']}}}id"
    known = rel_map.get("ppt/presentation.xml", {})
    for el in pres.findall("p:sldIdLst/p:sldId", NS):
        if el.get(rid_attr) not in known:
            errors.append(f"Slide relationship id missing: {el.get(rid_attr)}")


def check_shape_ids(trees: dict, errors: list[str]) -> None:
    for name, root in trees.items():
        if not (name.startswith("ppt/slides/slide") and name.endswith(".xml")):
            continue
        ids = [el.get("id") for el in root.iter(f"{{{NS['p']}}}cNvPr")]
        dup = [i for i, n in Counter(ids).items() if n > 1]
        if dup:
            errors.append(f"Duplicate shape id {name}: {dup[:5]}")


def check_python_pptx(path: str, errors: list[str]) -> int:
    try:
        from pptx import Presentation
    except ImportError:
        errors.append("python-pptx not installed (pip install python-pptx)")
        return 0
    try:
        prs = Presentation(path)
        count = 0
        for slide in prs.slides:
            for _shape in slide.shapes:
                pass
            count += 1
        return count
    except Exception as exc:  # Report the reason for the open failure as is
        errors.append(f"python-pptx failed to open: {exc}")
        return 0


def check_schema(trees: dict, errors: list[str]) -> str:
    xsd_path = os.environ.get("OOXML_PML_XSD")
    if not xsd_path:
        return "skipped (OOXML_PML_XSD not set)"
    try:
        schema = etree.XMLSchema(etree.parse(xsd_path))
    except Exception as exc:
        errors.append(f"Schema load failed {xsd_path}: {exc}")
        return "failed"
    for name, root in trees.items():
        if name.startswith("ppt/slides/slide") and name.endswith(".xml") and not schema.validate(root):
            errors.append(f"Schema mismatch {name}: {schema.error_log.last_error}")
    return "done"


def validate(path: str) -> list[str]:
    errors: list[str] = []
    try:
        zf = zipfile.ZipFile(path)
    except zipfile.BadZipFile as exc:
        return [f"ZIP open failed: {exc}"]
    with zf:
        parts = {i.filename for i in zf.infolist() if not i.is_dir()}
        check_content_types(zf, parts, errors)
        trees = check_xml_parts(zf, parts, errors)
        rel_map = check_relationships(trees, parts, errors)
        check_presentation(trees, rel_map, errors)
        check_shape_ids(trees, errors)
        schema_state = check_schema(trees, errors)
    slides = check_python_pptx(path, errors)
    print(f"validate: {len(parts)} parts, {slides} slides, schema validation {schema_state}")
    return errors


def main() -> int:
    if len(sys.argv) != 2:
        print(__doc__)
        return 2
    errors = validate(sys.argv[1])
    for e in errors:
        print("FAIL", e)
    print("validate: PASS" if not errors else f"validate: FAIL {len(errors)} issues")
    return 0 if not errors else 1


if __name__ == "__main__":
    sys.exit(main())
