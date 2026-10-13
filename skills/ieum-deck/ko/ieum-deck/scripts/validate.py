#!/usr/bin/env python3
"""validate.py: 완성 PPTX 패키지의 구조 검증 (finalize.py 이후 실행)

WHY: PowerPoint 는 패키지 구조가 어긋난 파일의 개방을 거부한다. 렌더링 검수는 LibreOffice 로
     수행되므로 LibreOffice 가 관대하게 여는 결함을 별도로 잡아야 한다.

검사 항목 (결함 1건 이상이면 종료 코드 1)
1. ZIP 개방과 [Content_Types].xml 존재
2. 모든 부품의 콘텐츠 형식 선언 (확장자 Default 또는 Override)
3. 모든 XML 부품의 문법 정합 (lxml 파싱)
4. 관계(.rels)의 내부 대상 부품 실재
5. presentation.xml 의 슬라이드 목록 id 고유성과 관계 id 실재
6. 슬라이드별 도형 id(cNvPr) 고유성
7. python-pptx 개방과 슬라이드 순회
8. (선택) 환경변수 OOXML_PML_XSD 가 PresentationML 스키마 파일을 가리키면 슬라이드 XML 스키마 검증

사용: python3 validate.py <file.pptx>
"""
from __future__ import annotations

import os
import posixpath
import sys
import zipfile
from collections import Counter

from lxml import etree

CONTENT_TYPES = "[Content_Types].xml"
NS = {
    "ct": "http://schemas.openxmlformats.org/package/2006/content-types",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}
XML_SUFFIXES = (".xml", ".rels")
SLIDE_ID_MIN = 256  # ECMA-376 의 sldId 하한


def rels_source_dir(rels_name: str) -> str:
    """관계 파일이 설명하는 원본 부품의 폴더. _rels/x.xml.rels 의 원본은 x.xml"""
    folder = posixpath.dirname(posixpath.dirname(rels_name))
    return folder


def check_content_types(zf: zipfile.ZipFile, parts: set[str], errors: list[str]) -> None:
    if CONTENT_TYPES not in parts:
        errors.append("[Content_Types].xml 없음")
        return
    root = etree.fromstring(zf.read(CONTENT_TYPES))
    defaults = {el.get("Extension", "").lower() for el in root.findall("ct:Default", NS)}
    overrides = {el.get("PartName", "").lstrip("/") for el in root.findall("ct:Override", NS)}
    for name in sorted(parts):
        if name == CONTENT_TYPES:
            continue
        ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
        if name not in overrides and ext not in defaults:
            errors.append(f"콘텐츠 형식 미선언 부품: {name}")
    for name in sorted(overrides - parts):
        errors.append(f"실재하지 않는 부품의 Override 선언: {name}")


def check_xml_parts(zf: zipfile.ZipFile, parts: set[str], errors: list[str]) -> dict[str, etree._Element]:
    trees = {}
    for name in sorted(parts):
        if not name.endswith(XML_SUFFIXES):
            continue
        try:
            trees[name] = etree.fromstring(zf.read(name))
        except etree.XMLSyntaxError as exc:
            errors.append(f"XML 문법 오류 {name}: {exc}")
    return trees


def check_relationships(trees: dict, parts: set[str], errors: list[str]) -> dict[str, dict[str, str]]:
    """관계 대상 실재 여부. 반환: 원본 부품별 {rId: 대상 부품}"""
    rel_map: dict[str, dict[str, str]] = {}
    for name, root in trees.items():
        if not name.endswith(".rels"):
            continue
        src_dir = rels_source_dir(name)
        src_part = posixpath.join(src_dir, posixpath.basename(name)[: -len(".rels")])
        targets = {}
        for rel in root.findall("rel:Relationship", NS):
            if rel.get("TargetMode") == "External":
                continue
            target = rel.get("Target", "")
            resolved = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(src_dir, target))
            targets[rel.get("Id")] = resolved
            if resolved not in parts:
                errors.append(f"관계 대상 없음 {name} {rel.get('Id')} → {resolved}")
        rel_map[src_part.lstrip("/")] = targets
    return rel_map


def check_presentation(trees: dict, rel_map: dict, errors: list[str]) -> None:
    pres = trees.get("ppt/presentation.xml")
    if pres is None:
        errors.append("ppt/presentation.xml 없음")
        return
    ids = [el.get("id") for el in pres.findall("p:sldIdLst/p:sldId", NS)]
    for sid, n in Counter(ids).items():
        if n > 1:
            errors.append(f"슬라이드 id 중복: {sid}")
    for sid in ids:
        if sid is None or not sid.isdigit() or int(sid) < SLIDE_ID_MIN:
            errors.append(f"슬라이드 id 범위 오류: {sid}")
    rid_attr = f"{{{NS['r']}}}id"
    known = rel_map.get("ppt/presentation.xml", {})
    for el in pres.findall("p:sldIdLst/p:sldId", NS):
        if el.get(rid_attr) not in known:
            errors.append(f"슬라이드 관계 id 없음: {el.get(rid_attr)}")


def check_shape_ids(trees: dict, errors: list[str]) -> None:
    for name, root in trees.items():
        if not (name.startswith("ppt/slides/slide") and name.endswith(".xml")):
            continue
        ids = [el.get("id") for el in root.iter(f"{{{NS['p']}}}cNvPr")]
        dup = [i for i, n in Counter(ids).items() if n > 1]
        if dup:
            errors.append(f"도형 id 중복 {name}: {dup[:5]}")


def check_python_pptx(path: str, errors: list[str]) -> int:
    try:
        from pptx import Presentation
    except ImportError:
        errors.append("python-pptx 미설치 (pip install python-pptx)")
        return 0
    try:
        prs = Presentation(path)
        count = 0
        for slide in prs.slides:
            for _shape in slide.shapes:
                pass
            count += 1
        return count
    except Exception as exc:  # 개방 실패 사유를 그대로 보고
        errors.append(f"python-pptx 개방 실패: {exc}")
        return 0


def check_schema(trees: dict, errors: list[str]) -> str:
    xsd_path = os.environ.get("OOXML_PML_XSD")
    if not xsd_path:
        return "생략 (OOXML_PML_XSD 미지정)"
    try:
        schema = etree.XMLSchema(etree.parse(xsd_path))
    except Exception as exc:
        errors.append(f"스키마 로드 실패 {xsd_path}: {exc}")
        return "실패"
    for name, root in trees.items():
        if name.startswith("ppt/slides/slide") and name.endswith(".xml") and not schema.validate(root):
            errors.append(f"스키마 불일치 {name}: {schema.error_log.last_error}")
    return "수행"


def validate(path: str) -> list[str]:
    errors: list[str] = []
    try:
        zf = zipfile.ZipFile(path)
    except zipfile.BadZipFile as exc:
        return [f"ZIP 개방 실패: {exc}"]
    with zf:
        parts = {i.filename for i in zf.infolist() if not i.is_dir()}
        check_content_types(zf, parts, errors)
        trees = check_xml_parts(zf, parts, errors)
        rel_map = check_relationships(trees, parts, errors)
        check_presentation(trees, rel_map, errors)
        check_shape_ids(trees, errors)
        schema_state = check_schema(trees, errors)
    slides = check_python_pptx(path, errors)
    print(f"validate: 부품 {len(parts)}개, 슬라이드 {slides}장, 스키마 검증 {schema_state}")
    return errors


def main() -> int:
    if len(sys.argv) != 2:
        print(__doc__)
        return 2
    errors = validate(sys.argv[1])
    for e in errors:
        print("FAIL", e)
    print("validate: PASS" if not errors else f"validate: FAIL {len(errors)}건")
    return 0 if not errors else 1


if __name__ == "__main__":
    sys.exit(main())
