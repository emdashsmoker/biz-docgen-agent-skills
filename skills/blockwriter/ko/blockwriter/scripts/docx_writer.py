"""DOCX 조판기: HWPX와 같은 원고 요소와 같은 양식 정의(form.py)로 Word 문서를 만든다.

양식 대응 원칙
- 스타일 수치(글꼴, 크기, 굵기, 장평, 자간, 정렬, 여백, 줄 간격)는 form.py 값을 단위만 바꿔 옮긴다
- 장·절·소절 제목은 Word 기본 제목 스타일에 입혀 목차 필드와 탐색 창이 인식하게 한다
- 글머리 기호는 Word 글머리 목록으로 만들어 줄바꿈 시 둘째 줄이 본문 시작선에 맞게 한다
- 표 번호는 SEQ 필드, 목차는 TOC 필드로 넣고 필드 결과값을 미리 채워 필드를 갱신하지 않는 뷰어에서도 보이게 한다
- 목차는 Word가 만드는 목차와 같은 구조(제목 책갈피, 항목 하이퍼링크, PAGEREF 쪽번호 필드, 목차 콘텐츠 컨트롤)로 넣는다
"""
from __future__ import annotations

import datetime as dt
import io
import re
import sys

import config as C
import manuscript as M
from config import msg
from form import FORMS, Char, Para
from hwpx_writer import MANUAL_SYMBOLS, cell_sides, load_profile, parse_cover, parse_header_spec

try:
    from docx import Document
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.opc.packuri import PackURI
    from docx.opc.part import XmlPart
    from docx.oxml import OxmlElement, parse_xml
    from docx.oxml.ns import nsdecls, qn
    from docx.shared import Emu
    from lxml import etree
except ImportError:   # pragma: no cover
    sys.exit(msg("err_no_docx"))

ALIGN = {"JUSTIFY": "both", "LEFT": "left", "CENTER": "center", "RIGHT": "right", "DISTRIBUTE": "distribute"}
FOOTNOTES_CT = "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"
EMU_PER_HWPUNIT = 127             # 914400 EMU / 7200 HWPUNIT


def tw(hwpunit: float) -> int:
    """HWPUNIT(1/7200 in) → twip(1/1440 in)"""
    return round(hwpunit / 5)


def el(tag: str, **attrs) -> "OxmlElement":
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k.replace("__", ":")), str(v))
    return e


def sub(parent, tag: str, **attrs):
    e = el(tag, **attrs)
    parent.append(e)
    return e


def bullet_hanging_pt(head: tuple, size: float) -> float:
    """글머리 기호 폭 + 기호 뒤 간격 + 폭 조정. 한/글 글머리 자동 내어쓰기와 같은 위치에 둘째 줄을 맞춘다."""
    _, char, offset, adjust = head
    return M.bullet_head_hu(char, size * 100, offset, adjust) / 100


def rpr_xml(char: Char, bold: bool | None = None) -> "OxmlElement":
    """글자 모양 → w:rPr. Word는 동아시아 글꼴 칸이 하나이므로 한글 슬롯 값을 쓴다."""
    faces, ratio, spacing = char.faces, char.ratios[0], char.spacings[0]
    r = el("w:rPr")
    sub(r, "w:rFonts", w__ascii=faces[1], w__hAnsi=faces[1], w__eastAsia=faces[0], w__cs=faces[0])
    b = char.bold if bold is None else bold
    sub(r, "w:b", w__val=int(b))
    sub(r, "w:bCs", w__val=int(b))
    sub(r, "w:color", w__val="000000")
    if spacing:
        sub(r, "w:spacing", w__val=round(char.size * spacing / 100 * 20))
    if ratio != 100:
        sub(r, "w:w", w__val=ratio)
    sub(r, "w:sz", w__val=round(char.size * 2))
    sub(r, "w:szCs", w__val=round(char.size * 2))
    return r


def ppr_xml(para: Para, char: Char, num_id: int | None = None, align: str | None = None,
            before_extra: int = 0, after_extra: int = 0, exact_line: int | None = None) -> "OxmlElement":
    """문단 모양 → w:pPr. 줄 간격은 글자 크기 대비 %이므로 최소 줄 높이로 옮긴다."""
    intent, left, right, prev, nxt = para.margin
    p = el("w:pPr")
    if para.keep_next:
        sub(p, "w:keepNext")
    if num_id is not None:
        numpr = sub(p, "w:numPr")
        sub(numpr, "w:ilvl", w__val=0)
        sub(numpr, "w:numId", w__val=num_id)
    sub(p, "w:snapToGrid", w__val=0)
    if exact_line is not None:
        sub(p, "w:spacing", w__before=tw(prev + before_extra), w__after=tw(nxt + after_extra),
            w__line=exact_line, w__lineRule="exact")
    else:
        sub(p, "w:spacing", w__before=tw(prev + before_extra), w__after=tw(nxt + after_extra),
            w__line=round(char.size * para.line / 100 * 20), w__lineRule="atLeast")
    head = para.head if isinstance(para.head, tuple) and para.head[0] == "BULLET" else None
    if head and num_id is not None:
        hang = round(bullet_hanging_pt(head, char.size) * 20)
        sub(p, "w:ind", w__left=tw(left) + hang, w__right=tw(right), w__hanging=hang)
    elif intent < 0:
        sub(p, "w:ind", w__left=tw(left - intent), w__right=tw(right), w__hanging=tw(-intent))
    else:
        sub(p, "w:ind", w__left=tw(left), w__right=tw(right), w__firstLine=tw(intent))
    sub(p, "w:jc", w__val=ALIGN[align or para.align])
    return p


class DocxBuilder:
    def __init__(self, form_name: str, page_break: str):
        self.form = FORMS[form_name]
        self.form_name = form_name
        self.p = load_profile(form_name)
        self.doc = Document()
        self.numbering = M.Numbering(page_break)
        self.style_ids: dict[str, str] = {}
        self.num_ids: dict[str, int] = {}
        self.pending_break = False
        self.headings: list[tuple[int, str]] = []
        self.heading_marks: list[str] = []      # self.headings와 같은 순서의 제목 책갈피 이름
        self.toc_numbers: dict[int, int] = {}   # 목차 항목 순번 → 쪽번호
        self.caption_n = {"TABLE": 0, "PICTURE": 0}
        self.fn_part = None
        self.fn_count = 0
        self.last_para = None
        self.body_w = self.form.page.width - self.form.page.left - self.form.page.right
        self.table_max_w = self.body_w - C.TABLE_INDENT_MARGIN
        self.numbered_square = False
        self.toc = False
        self.toc_depth = 2
        self.cover = None
        self.page_number = False
        self.header_spec = None
        self.merge_mode = "auto"
        self.long_table_mode = "engine"
        self._setup_styles()
        self._font_table()

    # ── 글꼴 표 ────────────────────────────────────────────────────────
    def _font_table(self):
        """양식 글꼴마다 영문 이름, 문자 집합, 계열을 글꼴 표에 기록한다. 글꼴 이름 자체는 바꾸지 않는다."""
        part = next(r.target_part for r in self.doc.part.rels.values() if r.reltype == RT.FONT_TABLE)
        root = parse_xml(part.blob)
        known = {f.get(qn("w:name")) for f in root.findall(qn("w:font"))}
        faces = {f for st in self.form.styles for f in st.char.faces} | set(self.form.default_char.faces)
        for face in sorted(faces - known):
            font = sub(root, "w:font", w__name=face)
            if face in C.DOCX_FONT_ALT_NAMES:
                sub(font, "w:altName", w__val=C.DOCX_FONT_ALT_NAMES[face])
            sub(font, "w:charset", w__val="81")
            family = "roman" if "Myeongjo" in C.DOCX_FONT_ALT_NAMES.get(face, "") else "swiss"
            sub(font, "w:family", w__val=family)
            sub(font, "w:pitch", w__val="variable")
        part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    # ── 스타일 ─────────────────────────────────────────────────────────
    def _style_el(self, style_id: str, name: str, kind: str = "paragraph"):
        styles = self.doc.styles.element
        st = next((s for s in styles.findall(qn("w:style")) if s.get(qn("w:styleId")) == style_id), None)
        if st is None:
            st = el("w:style", w__type=kind, w__styleId=style_id)
            sub(st, "w:name", w__val=name)
            sub(st, "w:qFormat")
            styles.append(st)
        for tag in ("w:pPr", "w:rPr", "w:link", "w:basedOn", "w:next"):
            for old in st.findall(qn(tag)):
                st.remove(old)
        return st

    def _setup_styles(self):
        form = self.form
        body = form.style("body")
        # 문서 기본값: 양식 바탕글 글꼴, 한국어 언어 설정, 문단 뒤 간격 없음
        defaults = self.doc.styles.element.find(qn("w:docDefaults"))
        rpr_def = defaults.find(qn("w:rPrDefault")).find(qn("w:rPr"))
        rpr_def.clear()
        for child in rpr_xml(body.char):
            rpr_def.append(child)
        # 라틴 언어를 ko-KR로 두면 LibreOffice가 숫자 중간(3,|456)에서 줄을 바꾼다. 기본값은 Word 한국어판 기본 문서와 같은 조합
        sub(rpr_def, "w:lang", w__val=C.label("language.latin"), w__eastAsia=C.label("language.east_asia"), w__bidi="ar-SA")
        ppr_def = defaults.find(qn("w:pPrDefault")).find(qn("w:pPr"))
        ppr_def.clear()
        # 한글과 영문, 한글과 숫자 사이 자동 간격 끔 (HWPX 문단 모양의 autoSpacing과 같게)
        sub(ppr_def, "w:autoSpaceDE", w__val=0)
        sub(ppr_def, "w:autoSpaceDN", w__val=0)
        sub(ppr_def, "w:spacing", w__after=0, w__line=240, w__lineRule="auto")
        self._define_bullets()
        heading_map = C.DOCX_HEADING_STYLES[self.form_name]
        for st in form.styles:
            if st.kind != "PARA":
                continue
            if st.role == "body":
                sid = "Normal"
            elif st.role in heading_map:
                sid = heading_map[st.role].replace(" ", "")
            elif st.role == "footnote":
                sid = "FootnoteText"
            elif st.role == "toc_chapter":
                sid = "TOC1"
            elif st.role == "toc_section":
                sid = "TOC2"
            else:
                sid = "RB" + "".join(w.capitalize() for w in st.role.split("_"))
            self.style_ids[st.role] = sid
            builtin_name = {"Normal": "Normal", "FootnoteText": "footnote text", "TOC1": "toc 1", "TOC2": "toc 2"}.get(sid)
            if sid.startswith("Heading"):
                builtin_name = "heading " + sid[-1]
            elif sid == "Title":
                builtin_name = "Title"
            el_ = self._style_el(sid, builtin_name or st.name)
            if builtin_name is None:
                el_.find(qn("w:name")).set(qn("w:val"), st.name)
            if sid != "Normal":
                el_.insert(1, el("w:basedOn", w__val="Normal"))
                el_.insert(2, el("w:next", w__val=sid if st.next_self else "Normal"))
            outline = None
            if sid.startswith("Heading"):
                outline = int(sid[-1]) - 1
            ppr = ppr_xml(st.para, st.char, self.num_ids.get(st.role))
            if st.role in ("toc_chapter", "toc_section"):
                tabs = el("w:tabs")
                sub(tabs, "w:tab", w__val="right", w__leader="dot", w__pos=tw(C.TOC_TAB_POS))
                ppr.insert(2, tabs)
            if outline is not None:
                ppr.append(el("w:outlineLvl", w__val=outline))
            el_.append(ppr)
            el_.append(rpr_xml(st.char))
        toc2 = form.style("toc_section")
        if toc2:   # 목차 소절은 절 모양에 들여쓰기만 더한다
            st3 = self._style_el("TOC3", "toc 3")
            st3.insert(1, el("w:basedOn", w__val="Normal"))
            para3 = Para(**{**toc2.para.__dict__, "margin": (0, C.TOC_SUB_INDENT) + toc2.para.margin[2:]})
            ppr = ppr_xml(para3, toc2.char)
            tabs = el("w:tabs")
            sub(tabs, "w:tab", w__val="right", w__leader="dot", w__pos=tw(C.TOC_TAB_POS))
            ppr.insert(1, tabs)
            st3.append(ppr)
            st3.append(rpr_xml(toc2.char))
            self.style_ids["toc_sub"] = "TOC3"
        ref = self._style_el("FootnoteReference", "footnote reference", "character")
        rpr = el("w:rPr")
        sub(rpr, "w:vertAlign", w__val="superscript")
        ref.append(rpr)
        self._fix_style_order()

    def _fix_style_order(self):
        """pPr 안 자식 순서를 스키마 순서로 정렬 (Word는 순서가 어긋난 문서를 손상으로 판정)"""
        order = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "widowControl", "numPr", "pBdr", "shd", "tabs",
                 "suppressAutoHyphens", "autoSpaceDE", "autoSpaceDN", "snapToGrid", "spacing", "ind", "jc", "outlineLvl",
                 "sectPr"]
        for ppr in self.doc.styles.element.iter(qn("w:pPr")):
            sort_children(ppr, order)

    def _define_bullets(self):
        """글머리 역할마다 Word 글머리 목록 정의를 만든다 (단일 수준)"""
        numbering = self.doc.part.numbering_part.element
        next_abs = max([int(a.get(qn("w:abstractNumId"))) for a in numbering.findall(qn("w:abstractNum"))] + [-1]) + 1
        next_num = max([int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))] + [0]) + 1
        for st in self.form.styles:
            head = st.para.head
            if not (isinstance(head, tuple) and head[0] == "BULLET"):
                continue
            abstract = el("w:abstractNum", w__abstractNumId=next_abs)
            sub(abstract, "w:multiLevelType", w__val="singleLevel")
            lvl = sub(abstract, "w:lvl", w__ilvl=0)
            sub(lvl, "w:start", w__val=1)
            sub(lvl, "w:numFmt", w__val="bullet")
            sub(lvl, "w:lvlText", w__val=head[1])
            sub(lvl, "w:lvlJc", w__val="left")
            hang = round(bullet_hanging_pt(head, st.char.size) * 20)
            lppr = sub(lvl, "w:pPr")
            sub(lppr, "w:ind", w__left=tw(st.para.margin[1]) + hang, w__hanging=hang)
            lrpr = sub(lvl, "w:rPr")
            face = st.char.faces[0]
            sub(lrpr, "w:rFonts", w__ascii=face, w__hAnsi=face, w__eastAsia=face, w__cs=face, w__hint="eastAsia")
            sub(lrpr, "w:b", w__val=int(st.char.bold))
            sub(lrpr, "w:sz", w__val=round(st.char.size * 2))
            # abstractNum은 num보다 앞에 와야 한다
            first_num = numbering.find(qn("w:num"))
            if first_num is not None:
                first_num.addprevious(abstract)
            else:
                numbering.append(abstract)
            num = sub(numbering, "w:num", w__numId=next_num)
            sub(num, "w:abstractNumId", w__val=next_abs)
            self.num_ids[st.role] = next_num
            next_abs += 1
            next_num += 1

    # ── 문단 ───────────────────────────────────────────────────────────
    def _body_append(self, element):
        """본문 끝(문서 구역 설정 앞)에 요소를 붙인다"""
        body = self.doc.element.body
        body.insert(len(body) - 1, element)
        return element

    def _new_p(self, role: str | None, ppr_extra=None):
        p = el("w:p")
        ppr = sub(p, "w:pPr")
        if role:
            sub(ppr, "w:pStyle", w__val=self.style_ids.get(role, "Normal"))
        if ppr_extra is not None:
            for child in list(ppr_extra):
                ppr.append(child)
        if self.pending_break and role is not None:
            ppr.append(el("w:pageBreakBefore"))
            self.pending_break = False
        sort_children(ppr, ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "widowControl", "numPr", "pBdr",
                            "shd", "tabs", "suppressAutoHyphens", "snapToGrid", "spacing", "ind", "jc", "outlineLvl"])
        return p

    def _run(self, text: str, rpr=None):
        r = el("w:r")
        if rpr is not None:
            r.append(rpr)
        parts = text.split("\t")
        for i, part in enumerate(parts):
            if i:
                sub(r, "w:tab")
            if part:
                t = sub(r, "w:t")
                t.text = part
                t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        return r

    def _fmt_rpr(self, kind: str, bold: bool = False):
        rpr = el("w:rPr")
        if bold or kind == "bold":   # 머리글 셀은 모든 run에 굵게를 더함
            sub(rpr, "w:b")
            sub(rpr, "w:bCs")
        if kind == "italic":
            sub(rpr, "w:i")
            sub(rpr, "w:iCs")
        elif kind in ("sup", "sub"):
            sub(rpr, "w:vertAlign", w__val="superscript" if kind == "sup" else "subscript")
        return rpr

    def para(self, role: str, text: str, footnotes: dict | None = None, no_bullet: bool = False):
        st = self.form.style(role) or self.form.style("body")
        head = st.para.head
        is_bullet = isinstance(head, tuple) and head[0] == "BULLET"
        extra = None
        numbered = role == "square" and self.numbered_square and M.NUMBERED_SQUARE_RE.match(text)
        if no_bullet or numbered:
            # 숫자 헤더: 글머리를 끄고 왼쪽 여백 유지. 번호 폭만큼 내어써서 둘째 줄을 번호 뒤 글자에 맞춘다
            hang = round(st.char.size * 100 * M.number_prefix_em(text, C.SPACE_EM_DOCX)) if numbered else 0
            extra = el("w:pPr")
            numpr = sub(extra, "w:numPr")
            sub(numpr, "w:ilvl", w__val=0)
            sub(numpr, "w:numId", w__val=0)
            sub(extra, "w:ind", w__left=tw(st.para.margin[1] + hang), w__hanging=tw(hang))
            if numbered:   # 양쪽 정렬은 접힌 첫 줄의 빈칸을 늘려 번호 뒤 글자 위치를 옮기므로 왼쪽 정렬
                sub(extra, "w:jc", w__val="left")
        elif not is_bullet and role in MANUAL_SYMBOLS:
            text = MANUAL_SYMBOLS[role] + text
        p = self._new_p(role, extra)
        self._append_inline(p, text, footnotes)
        return self._body_append(p)

    def _append_inline(self, p, text: str, footnotes: dict | None = None):
        """인라인 서식(굵게, 기울임, 첨자)을 run으로 나누어 붙인다. 각주 참조는 footnotes가 있을 때만 처리"""
        for kind, seg in M.inline_segments(text):
            if kind == "footnote":
                if footnotes is not None:
                    self._footnote_ref(p, footnotes.get(seg, ""))
            elif kind in ("bold", "italic", "sup", "sub"):
                p.append(self._run(seg, self._fmt_rpr(kind)))
            else:
                p.append(self._run(seg))

    # ── 각주 ───────────────────────────────────────────────────────────
    def _footnotes(self):
        if self.fn_part is None:
            xml = (f'<w:footnotes {nsdecls("w", "r")}>'
                   '<w:footnote w:type="separator" w:id="-1"><w:p><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr>'
                   '<w:r><w:separator/></w:r></w:p></w:footnote>'
                   '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr>'
                   '<w:r><w:continuationSeparator/></w:r></w:p></w:footnote></w:footnotes>')
            self.fn_part = XmlPart(PackURI("/word/footnotes.xml"), FOOTNOTES_CT, parse_xml(xml), self.doc.part.package)
            self.doc.part.relate_to(self.fn_part, RT.FOOTNOTES)
            settings = self.doc.settings.element
            fpr = el("w:footnotePr")
            sub(fpr, "w:footnote", w__id=-1)
            sub(fpr, "w:footnote", w__id=0)
            insert_settings_child(settings, fpr)
        return self.fn_part.element

    def _footnote_ref(self, p, text: str):
        """본문 각주 번호는 위 첨자, 번호 뒤 문자(한/글 각주 모양의 ')')를 같은 첨자로 붙인다"""
        notes = self._footnotes()
        self.fn_count += 1
        nid = self.fn_count
        r = el("w:r")
        rpr = sub(r, "w:rPr")
        sub(rpr, "w:rStyle", w__val="FootnoteReference")
        sub(r, "w:footnoteReference", w__id=nid)
        p.append(r)
        sup = el("w:rPr")
        sub(sup, "w:vertAlign", w__val="superscript")
        p.append(self._run(self.form.page.note_suffix, sup))
        note = sub(notes, "w:footnote", w__id=nid)
        np_ = sub(note, "w:p")
        ppr = sub(np_, "w:pPr")
        sub(ppr, "w:pStyle", w__val="FootnoteText")
        fr = sub(np_, "w:r")
        frpr = sub(fr, "w:rPr")
        sub(frpr, "w:vertAlign", w__val="baseline")   # 각주 영역의 번호는 본문 크기 (한/글 각주 모양과 같게)
        sub(fr, "w:footnoteRef")
        np_.append(self._run(self.form.page.note_suffix + "\t"))
        self._append_inline(np_, text)   # 각주 본문도 본문과 같은 인라인 서식 적용

    # ── 헤더 계위 ───────────────────────────────────────────────────────
    def heading(self, level: int, text: str):
        text = M.NUM_PREFIX_RE.sub("", text).strip()
        role = ("chapter", "section", "subsection", "ganada")[level - 1]
        if self.numbering.advance(level):
            self.pending_break = True
        label = M.heading_label(self.p, self.numbering.n, level, text)
        p = self.para(role, label)
        if level <= 3:
            self.headings.append((level, label))
            self.heading_marks.append(self._bookmark(p))
        if level == 1 and self.p.get("caption_prefix"):
            self.caption_n = {"TABLE": 0, "PICTURE": 0}

    # ── 캡션 ───────────────────────────────────────────────────────────
    def _caption(self, kind: str, text: str, before: int, after: int, obj_w: int):
        """<표 1-3> 제목. 번호는 SEQ 필드이며 장 체계 양식은 제목 1 스타일마다 번호를 새로 시작(\\s 1).
        obj_w는 캡션이 딸린 표나 그림 표의 폭. HWPX 캡션 개체처럼 이 폭 안에서 줄을 바꾸도록 양쪽을 들여씀"""
        self.caption_n[kind] += 1
        role = "table_title" if kind == "TABLE" else "fig_title"
        st = self.form.style(role)
        extra = el("w:pPr")
        if kind == "TABLE":   # 표 캡션은 표 위에 붙어 함께 이동
            sub(extra, "w:keepNext")
        intent, left, right, prev, nxt = st.para.margin
        sub(extra, "w:spacing", w__before=tw(prev + before), w__after=tw(nxt + after),
            w__line=round(st.char.size * st.para.line / 100 * 20), w__lineRule="atLeast")
        side = max((self.body_w - obj_w) // 2, 0)   # 표와 그림은 본문 폭 가운데에 놓임
        if intent < 0:
            sub(extra, "w:ind", w__left=tw(left + side - intent), w__right=tw(right + side), w__hanging=tw(-intent))
        else:
            sub(extra, "w:ind", w__left=tw(left + side), w__right=tw(right + side), w__firstLine=tw(intent))
        p = self._new_p(role, extra)
        open_b, close_b = C.caption_marks(kind)
        prefix = self.p.get("caption_prefix", "").format(c=max(self.numbering.n["chapter"], 1))
        p.append(self._run(open_b + prefix))
        seq = C.caption_seq(kind)
        instr = f" SEQ {seq} \\* ARABIC" + (" \\s 1 " if self.p.get("caption_prefix") else " ")
        for r in field_runs(instr, str(self.caption_n[kind])):
            p.append(r)
        p.append(self._run(close_b + text))
        return self._body_append(p)

    # ── 표 ──────────────────────────────────────────────────────────────
    def table(self, e: M.Element):
        if len(e.data["rows"]) > C.LONG_TABLE_ROWS and self.long_table_mode == "split" and not e.data.get("split_done"):
            for sub_e in M.split_long_table(e):
                self.table(sub_e)
            return
        plan = M.table_plan(e, self.merge_mode)
        rows, hdr, aligns = plan["rows"], plan["hdr"], plan["aligns"]
        nrow, ncol = len(rows), len(rows[0])
        font = self._cell_font()
        widths = M.table_col_widths(rows, e.data.get("colwidth"), self.table_max_w, font, e.text)
        unit = e.data.get("unit")
        sources = e.data.get("sources", [])
        pos = self.form.style("table_pos")
        long_table = nrow > C.LONG_TABLE_ROWS and not e.data.get("split_done")
        if e.text:
            self._caption("TABLE", e.text, before=pos.para.margin[3], after=C.CAPTION_GAP, obj_w=sum(widths))
        for ri, c0, c1 in plan["skipped"]:
            print(msg("warn_hmerge_overlap", row=ri, c0=c0, c1=c1))
        # 격자: (행, 열) → 셀 정보. 병합으로 가려지는 칸은 None
        grid = [[{"text": v, "span": 1, "vmerge": None} for v in row] for row in rows]
        for c, r0, r1 in plan["vspans"]:
            a, b = hdr + r0, hdr + r1
            grid[a][c]["vmerge"] = "restart"
            for rr in range(a + 1, b + 1):
                grid[rr][c] = {"text": "", "span": 1, "vmerge": "continue"}
        for ri, c0, c1 in plan["hspans"]:
            grid[ri][c0]["span"] = c1 - c0 + 1
            for cc in range(c0 + 1, c1 + 1):
                grid[ri][cc] = None
        tbl = self._tbl(widths)
        if self.pending_break:   # 캡션 없는 표 앞의 쪽 나눔은 빈 문단으로 옮긴다
            self._page_break_para()
        cell_style = self.form.style("table_cell")
        header_repeat = long_table
        if unit:
            self._transparent_row(tbl, widths, unit, "unit", "right", header=header_repeat)
        line_counts = M.row_line_counts(rows, widths, font, plan["vspans"], hdr)
        for ri, row in enumerate(rows):
            is_head = ri < hdr
            row_h = M.row_height(line_counts[ri])
            tr = sub(tbl, "w:tr")
            trpr = sub(tr, "w:trPr")
            sub(trpr, "w:cantSplit")
            sub(trpr, "w:trHeight", w__val=tw(row_h), w__hRule="atLeast")
            if is_head and header_repeat:
                sub(trpr, "w:tblHeader")
            c = 0
            while c < ncol:
                cell = grid[ri][c]
                if cell is None:
                    c += 1
                    continue
                span = cell["span"]
                tc = sub(tr, "w:tc")
                tcpr = sub(tc, "w:tcPr")
                sub(tcpr, "w:tcW", w__w=tw(sum(widths[c:c + span])), w__type="dxa")
                if span > 1:
                    sub(tcpr, "w:gridSpan", w__val=span)
                if cell["vmerge"]:
                    sub(tcpr, "w:vMerge", w__val=cell["vmerge"])
                self._cell_borders(tcpr, cell_sides(c, c + span - 1, ncol))
                if is_head:
                    sub(tcpr, "w:shd", w__val="clear", w__color="auto", w__fill=C.HEADER_FILL.lstrip("#"))
                sub(tcpr, "w:vAlign", w__val="center")
                text = "" if cell["vmerge"] == "continue" or cell["text"].strip() in ("^", "<") else cell["text"]
                self._cell_paragraphs(tc, text, aligns[c], is_head, cell_style, keep=not long_table)
                c += span
        for k, src in enumerate(sources):
            self._transparent_row(tbl, widths, src, "source", "left")
        self._body_append(tbl)
        self._spacer(pos.para.margin[4])

    def _tbl(self, widths: list[int]):
        tbl = el("w:tbl")
        tblpr = sub(tbl, "w:tblPr")
        sub(tblpr, "w:tblW", w__w=tw(sum(widths)), w__type="dxa")
        sub(tblpr, "w:jc", w__val="center")
        sub(tblpr, "w:tblLayout", w__type="fixed")
        mar = sub(tblpr, "w:tblCellMar")
        sub(mar, "w:top", w__w=tw(C.CELL_MARGIN_TB), w__type="dxa")
        sub(mar, "w:left", w__w=tw(C.CELL_MARGIN_LR), w__type="dxa")
        sub(mar, "w:bottom", w__w=tw(C.CELL_MARGIN_TB), w__type="dxa")
        sub(mar, "w:right", w__w=tw(C.CELL_MARGIN_LR), w__type="dxa")
        sub(tblpr, "w:tblLook", w__val="0000", w__firstRow=0, w__lastRow=0, w__firstColumn=0, w__lastColumn=0,
            w__noHBand=1, w__noVBand=1)
        grid = sub(tbl, "w:tblGrid")
        for w in widths:
            sub(grid, "w:gridCol", w__w=tw(w))
        return tbl

    @staticmethod
    def _cell_borders(tcpr, sides: list[str]):
        b = sub(tcpr, "w:tcBorders")
        size = round(C.BORDER_WIDTH_MM / 25.4 * 72 * 8)   # 1/8pt 단위
        for side in ("top", "left", "bottom", "right"):
            if side in sides:
                sub(b, f"w:{side}", w__val="single", w__sz=size, w__space=0, w__color=C.BORDER_COLOR.lstrip("#"))
            else:
                sub(b, f"w:{side}", w__val="nil")

    def _cell_font(self) -> M.CellFont:
        """표 셀 줄 스타일 (열 폭과 행 줄 수 산정용)"""
        def line_style(role: str) -> M.LineStyle | None:   # 셀 글머리 스타일이 없는 양식은 None
            st = self.form.style(role)
            if st is None:
                return None
            head = st.para.head if isinstance(st.para.head, tuple) and st.para.head[0] == "BULLET" else None
            lead = st.para.margin[1] + st.para.margin[2] + (bullet_hanging_pt(head, st.char.size) * 100 if head else 0)
            return M.LineStyle(st.char.size * 100, st.char.ratios[0], st.char.spacings[0], lead)
        plain = tuple(s for s in (line_style("table_cell"), line_style("table_cell_left")) if s)
        return M.CellFont(C.SPACE_EM_DOCX, plain, line_style("table_cell_bullet"), line_style("table_cell_dash"))

    def _cell_paragraphs(self, tc, text: str, align: str, is_head: bool, cell_style, keep: bool):
        lines = text.split("\n")
        for ln in lines:
            role = "table_cell"
            body = ln
            if not is_head:
                st = ln.lstrip()
                if st.startswith(M.CELL_BULLET_MARKS) and self.form.style("table_cell_bullet"):
                    role, body = "table_cell_bullet", M.CELL_MARK_RE.sub("", ln)
                elif st.startswith("- ") and self.form.style("table_cell_dash"):
                    role, body = "table_cell_dash", M.CELL_MARK_RE.sub("", ln)
                elif align == "LEFT":
                    role = "table_cell_left" if self.form.style("table_cell_left") else "table_cell"
            extra = el("w:pPr")
            if keep:
                sub(extra, "w:keepNext")
            if role in ("table_cell", "table_cell_left"):
                st_ = self.form.style(role)
                if align == "LEFT" and not is_head and st_.para.align not in ("LEFT", "JUSTIFY"):
                    sub(extra, "w:jc", w__val="left")
                elif align == "RIGHT" and not is_head:
                    sub(extra, "w:jc", w__val="right")
                elif is_head:
                    sub(extra, "w:jc", w__val=ALIGN[cell_style.para.align])
            p = el("w:p")
            ppr = sub(p, "w:pPr")
            sub(ppr, "w:pStyle", w__val=self.style_ids[role])
            for child in list(extra):
                ppr.append(child)
            sort_children(ppr, ["pStyle", "keepNext", "numPr", "spacing", "ind", "jc"])
            for kind, seg in M.inline_segments(body):
                if kind == "footnote":   # 셀 안의 각주 참조는 지원하지 않아 표기 그대로 둠
                    kind, seg = "plain", f"[^{seg}]"
                rpr = self._fmt_rpr(kind, bold=is_head)
                p.append(self._run(seg, rpr if len(rpr) else None))
            tc.append(p)

    def _transparent_row(self, tbl, widths, text: str, role: str, align: str, header: bool = False):
        """표 안의 투명 병합 행 (단위, 출처). 표와 함께 이동한다."""
        tr = sub(tbl, "w:tr")
        trpr = sub(tr, "w:trPr")
        sub(trpr, "w:cantSplit")
        sub(trpr, "w:trHeight", w__val=tw(C.ROW_BASE_H), w__hRule="atLeast")
        if header:
            sub(trpr, "w:tblHeader")
        tc = sub(tr, "w:tc")
        tcpr = sub(tc, "w:tcPr")
        sub(tcpr, "w:tcW", w__w=tw(sum(widths)), w__type="dxa")
        if len(widths) > 1:
            sub(tcpr, "w:gridSpan", w__val=len(widths))
        self._cell_borders(tcpr, [])
        sub(tcpr, "w:vAlign", w__val="center")
        p = sub(tc, "w:p")
        ppr = sub(p, "w:pPr")
        sub(ppr, "w:pStyle", w__val=self.style_ids[role])
        sub(ppr, "w:keepNext")
        sub(ppr, "w:jc", w__val=align)
        p.append(self._run(text))

    def _page_break_para(self):
        """표·그림 앞 쪽 나눔. 새 쪽 첫 줄을 차지하지 않도록 높이가 최소인 빈 문단에 쪽 나눔을 건다."""
        p = el("w:p")
        ppr = sub(p, "w:pPr")
        sub(ppr, "w:pageBreakBefore")
        sub(ppr, "w:snapToGrid", w__val=0)
        sub(ppr, "w:spacing", w__before=0, w__after=0, w__line=C.DOCX_SPACER_PT * 20, w__lineRule="exact")
        rpr = sub(ppr, "w:rPr")
        sub(rpr, "w:sz", w__val=C.DOCX_SPACER_PT * 2)
        self._body_append(p)
        self.pending_break = False

    def _spacer(self, after_hwpunit: int):
        """표 아래 간격. 한/글에서는 표를 담은 문단의 아래 여백이 맡는 간격을 빈 문단으로 옮긴다."""
        p = el("w:p")
        ppr = sub(p, "w:pPr")
        sub(ppr, "w:snapToGrid", w__val=0)
        sub(ppr, "w:spacing", w__before=0, w__after=tw(after_hwpunit), w__line=C.DOCX_SPACER_PT * 20, w__lineRule="exact")
        rpr = sub(ppr, "w:rPr")
        sub(rpr, "w:sz", w__val=C.DOCX_SPACER_PT * 2)
        self._body_append(p)

    # ── 그림 ────────────────────────────────────────────────────────────
    def image(self, e: M.Element):
        loaded = M.load_image(e)
        if loaded is None:
            return
        data, fmt, (w_mm, h_mm) = loaded
        w_hu, h_hu = round(w_mm * 7200 / 25.4), round(h_mm * 7200 / 25.4)
        pos = self.form.style("fig_pos")
        sources = e.data.get("sources", [])
        # 그림 문단: python-docx로 그림을 넣은 뒤 표 셀로 옮긴다 (관계는 문서 파트에 남는다)
        holder = self.doc.add_paragraph()
        holder.add_run().add_picture(io.BytesIO(data), width=Emu(w_hu * EMU_PER_HWPUNIT), height=Emu(h_hu * EMU_PER_HWPUNIT))
        pic_p = holder._p
        pic_p.getparent().remove(pic_p)
        ppr = pic_p.get_or_add_pPr()
        sub(ppr, "w:pStyle", w__val=self.style_ids["fig_pos"])
        sub(ppr, "w:keepNext")
        sub(ppr, "w:spacing", w__before=tw(pos.para.margin[3]), w__after=0, w__line=240, w__lineRule="auto")
        sub(ppr, "w:jc", w__val="center")
        sort_children(ppr, ["pStyle", "keepNext", "spacing", "jc"])
        if self.pending_break:
            self._page_break_para()
        if not C.FIGURE_IN_TABLE:
            self._body_append(pic_p)
            for src in sources:
                self.para("source", src)
        else:
            # 셀 좌우 여백을 더해야 그림이 셀 안 폭에 들어간다 (고정 폭 표라 넘치는 부분이 잘림)
            cell_w = w_hu + 2 * C.CELL_MARGIN_LR
            tbl = self._tbl([cell_w])
            tr = sub(tbl, "w:tr")
            trpr = sub(tr, "w:trPr")
            sub(trpr, "w:cantSplit")
            tc = sub(tr, "w:tc")
            tcpr = sub(tc, "w:tcPr")
            sub(tcpr, "w:tcW", w__w=tw(cell_w), w__type="dxa")
            self._cell_borders(tcpr, [])
            tc.append(pic_p)
            for src in sources:
                self._transparent_row(tbl, [cell_w], src, "source", "left")
            self._body_append(tbl)
        if e.text:
            fig_w = w_hu + 2 * C.CELL_MARGIN_LR if C.FIGURE_IN_TABLE else w_hu   # HWPX 캡션 개체 폭과 같게
            self._caption("PICTURE", e.text, before=C.CAPTION_GAP, after=pos.para.margin[4], obj_w=fig_w)
        else:
            self._spacer(pos.para.margin[4])

    # ── 앞부분 구역, 머리말, 쪽번호 ─────────────────────────────────────
    def _first_line(self):
        """한/글 구역 설정 문단(문서 첫 빈 문단)과 같은 높이의 빈 문단. 높이는 줄 배치 저장값의 줄 높이 + 줄 간격"""
        _, height, _, _, gap, _, _ = self.form.first_line
        p = el("w:p")
        ppr = sub(p, "w:pPr")
        sub(ppr, "w:snapToGrid", w__val=0)
        sub(ppr, "w:spacing", w__before=0, w__after=0, w__line=tw(height + gap), w__lineRule="exact")
        return p

    def _page_setup(self, sect, body: bool):
        pg = self.form.page
        for tag in ("w:pgSz", "w:pgMar", "w:pgNumType", "w:titlePg"):
            for old in sect.findall(qn(tag)):
                sect.remove(old)
        pgsz = el("w:pgSz", w__w=tw(pg.width), w__h=tw(pg.height))
        # 한/글은 위쪽 여백 안쪽에 머리말 영역을 더 두므로 Word 본문 시작 = 위쪽 여백 + 머리말
        pgmar = el("w:pgMar", w__top=tw(pg.top + pg.header), w__right=tw(pg.right), w__bottom=tw(pg.bottom + pg.footer),
                   w__left=tw(pg.left), w__header=tw(pg.top), w__footer=tw(pg.bottom), w__gutter=0)
        anchor = sect.find(qn("w:cols"))
        for node in (pgsz, pgmar):
            if anchor is not None:
                anchor.addprevious(node)
            else:
                sect.append(node)
        if body and (self.cover or self.toc):
            pgnum = el("w:pgNumType", w__start=1)
            pgmar.addnext(pgnum)
        sort_children(sect, ["headerReference", "footerReference", "footnotePr", "endnotePr", "type", "pgSz", "pgMar",
                             "paperSrc", "pgBorders", "lnNumType", "pgNumType", "cols", "formProt", "vAlign",
                             "noEndnote", "titlePg", "textDirection", "bidi", "rtlGutter", "docGrid", "printerSettings"])

    def _header_footer(self):
        section = self.doc.sections[0]
        if self.header_spec:
            text, align = parse_header_spec(self.header_spec)
            hdr = section.header
            hdr.is_linked_to_previous = False
            p = hdr.paragraphs[0]._p
            for child in list(p):
                p.remove(child)
            body = self.form.style("body")
            ppr = ppr_xml(body.para, self.form.default_char, align=align)
            p.append(ppr)
            p.append(self._run(text, rpr_xml(self.form.default_char)))
        if self.page_number:
            ftr = section.footer
            ftr.is_linked_to_previous = False
            p = ftr.paragraphs[0]._p
            for child in list(p):
                p.remove(child)
            body = self.form.style("body")
            p.append(ppr_xml(body.para, self.form.default_char, align="CENTER"))
            num_style = self.form.style("page_number")
            rpr = lambda: rpr_xml(num_style.char if num_style else self.form.default_char)
            p.append(self._run("- ", rpr()))
            for r in field_runs(" PAGE ", "1", rpr):
                p.append(r)
            p.append(self._run(" -", rpr()))

    def _front_matter(self):
        """표지와 목차를 앞부분 구역에 둔다. 쪽번호와 머리말은 본문 구역에만 적용한다."""
        body_el = self.doc.element.body
        front: list = [self._first_line()]
        body = self.form.style("body")
        chapter = self.form.style("chapter")

        def line(text="", char=None, para=None, page_break=False):
            p = el("w:p")
            ppr = ppr_xml(para or body.para, char or body.char, align="CENTER")
            if page_break:
                ppr.insert(0, el("w:pageBreakBefore"))
            p.append(ppr)
            if text:
                p.append(self._run(text, rpr_xml(char)))
            front.append(p)
        if self.cover:
            cover_char = lambda size, bold: Char(chapter.char.face, size, bold=bold)
            for _ in range(C.COVER_TOP_LINES):
                line()
            line(self.cover.get("title", ""), cover_char(C.COVER_TITLE_PT, True))
            if self.cover.get("subtitle"):
                line()
                line(self.cover["subtitle"], cover_char(C.COVER_SUB_PT, False))
            line()
            line()
            if self.cover.get("date"):
                line(self.cover["date"], cover_char(C.COVER_DATE_PT, False))
            for _ in range(C.COVER_GAP_LINES):
                line()
            if self.cover.get("org"):
                line(self.cover["org"], cover_char(C.COVER_ORG_PT, True))
        if self.toc:
            if not (self.form.style("toc_chapter") and self.form.style("toc_section")):
                print(msg("warn_no_toc_style"))
            else:
                line(C.label("toc.title"), chapter.char, chapter.para, page_break=bool(self.cover))
                front.extend(self._toc_paragraphs())
        # 앞부분 구역 끝: 쪽 설정만 담은 구역 나눔 (머리말·꼬리말·쪽번호 없음)
        if front[-1].tag != qn("w:p"):   # 목차 콘텐츠 컨트롤 뒤에는 구역 나눔을 담을 최소 높이 문단을 둔다
            front.append(self._tiny_para())
        last = front[-1]
        ppr = last.find(qn("w:pPr"))
        if ppr is None:
            ppr = el("w:pPr")
            last.insert(0, ppr)
        sect = el("w:sectPr")
        sub(sect, "w:type", w__val="nextPage")
        sub(sect, "w:cols", w__space=720)
        sub(sect, "w:docGrid", w__linePitch=360)
        self._page_setup(sect, body=False)
        ppr.append(sect)
        for i, p in enumerate(front):
            body_el.insert(i, p)

    def _bookmark(self, p) -> str:
        """제목 문단 전체를 책갈피로 감싸고 이름을 돌려준다. 목차 항목의 하이퍼링크와 PAGEREF 필드가 이 책갈피를 가리킨다."""
        bid = len(self.heading_marks) + 1
        name = C.DOCX_TOC_BOOKMARK.format(n=bid)
        start = el("w:bookmarkStart", w__id=bid, w__name=name)
        ppr = p.find(qn("w:pPr"))
        if ppr is None:
            p.insert(0, start)
        else:
            ppr.addnext(start)
        p.append(el("w:bookmarkEnd", w__id=bid))
        return name

    def _tiny_para(self):
        """내용 없는 최소 높이 문단"""
        p = el("w:p")
        ppr = sub(p, "w:pPr")
        sub(ppr, "w:spacing", w__before=0, w__after=0, w__line=C.DOCX_SPACER_PT * 20, w__lineRule="exact")
        return p

    def _toc_paragraphs(self) -> list:
        """목차 필드. Word가 만드는 목차와 같은 구조로 넣고 결과값(항목과 쪽번호)을 미리 채운다.
        항목마다 제목 책갈피로 가는 하이퍼링크와 PAGEREF 쪽번호 필드를 둔다. Word에서 '쪽 번호만 업데이트'를 고르면
        PAGEREF가 책갈피 위치로 쪽번호를 다시 계산하고, '목차 전체 업데이트'를 고르면 Word가 항목을 새로 만든다."""
        show_numbers = self.page_number
        instr = f' TOC \\o "1-{self.toc_depth}" \\h \\z' + ("" if show_numbers else " \\n") + " "
        entries = [(lv, text, mark) for (lv, text), mark in zip(self.headings, self.heading_marks)
                   if lv <= self.toc_depth]

        def web_hidden():   # \z 스위치: 웹 모양 보기에서 탭과 쪽번호를 숨긴다 (Word 목차와 같은 표시)
            rpr = el("w:rPr")
            sub(rpr, "w:webHidden")
            return rpr
        sdt = el("w:sdt")
        part = sub(sub(sdt, "w:sdtPr"), "w:docPartObj")
        sub(part, "w:docPartGallery", w__val="Table of Contents")
        sub(part, "w:docPartUnique")
        content = sub(sdt, "w:sdtContent")
        last = None
        for i, (lv, text, mark) in enumerate(entries or [(1, "", None)]):
            p = sub(content, "w:p")
            ppr = sub(p, "w:pPr")
            sub(ppr, "w:pStyle", w__val={1: "TOC1", 2: "TOC2"}.get(lv, "TOC3"))
            if i == 0:
                p.extend(field_runs(instr, None))
            holder = sub(p, "w:hyperlink", w__anchor=mark, w__history=1) if mark else p
            holder.append(self._run(text))
            if show_numbers and mark:
                tab = sub(holder, "w:r")
                tab.append(web_hidden())
                sub(tab, "w:tab")
                num = self.toc_numbers.get(i)
                holder.extend(field_runs(f" PAGEREF {mark} \\h ", str(num) if num else "", web_hidden))
            last = p
        end = sub(last, "w:r")
        sub(end, "w:fldChar", w__fldCharType="end")
        return [sdt]

    def _finish(self, out: str):
        sect = self.doc.element.body.find(qn("w:sectPr"))
        self._page_setup(sect, body=True)
        self._header_footer()
        if self.cover or self.toc:
            self._front_matter()
        if self.toc and self.page_number:   # Word가 열 때 목차 쪽번호를 다시 계산
            insert_settings_child(self.doc.settings.element, el("w:updateFields", w__val="true"))
        cp = self.doc.core_properties
        cp.author = cp.last_modified_by = cp.comments = ""
        cp.title = (self.cover or {}).get("title", "")
        now = dt.datetime.now(dt.timezone.utc).replace(microsecond=0, tzinfo=None)
        cp.created = cp.modified = now
        cp.revision = 1
        drop_thumbnail(self.doc)
        self.doc.save(out)

    def build(self, elems: list[M.Element], footnotes: dict, out: str):
        self.doc.element.body.insert(0, self._first_line())
        for e in M.attach_block_notes(elems):
            if e.kind == "heading":
                self.heading(e.level, e.text)
                self.last_para = None   # 제목 뒤 출처 줄이 앞 절 문단의 각주로 붙지 않게 연결을 끊음
            elif e.kind in ("square", "circle", "dash", "dot", "note"):
                self.last_para = self.para(e.kind, e.text, footnotes)
            elif e.kind == "unit":
                self.para("unit", e.text)
            elif e.kind == "source":   # 표·그림에 붙지 않은 출처 줄은 직전 문단의 각주로 이동
                if self.last_para is not None:
                    self._footnote_ref(self.last_para, e.text)
                else:
                    self.para("note", e.text, footnotes)
            elif e.kind == "table":
                self.table(e)
            elif e.kind == "image":
                self.image(e)
            elif e.kind == "pagebreak":
                self.pending_break = True
        self._finish(out)


# ── 보조 함수 ───────────────────────────────────────────────────────────────
def field_runs(instr: str, result: str | None, rpr_factory=None) -> list:
    """복합 필드(begin, 명령, separate, 결과, end) 실행 요소. result가 None이면 end 없이 앞 3개만."""
    def run():
        r = el("w:r")
        if rpr_factory:
            r.append(rpr_factory())
        return r
    begin = run()
    sub(begin, "w:fldChar", w__fldCharType="begin")
    ins = run()
    t = sub(ins, "w:instrText")
    t.text = instr
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    sep = run()
    sub(sep, "w:fldChar", w__fldCharType="separate")
    if result is None:
        return [begin, ins, sep]
    res = run()
    rt = sub(res, "w:t")
    rt.text = result
    end = run()
    sub(end, "w:fldChar", w__fldCharType="end")
    return [begin, ins, sep, res, end]


def sort_children(parent, order: list[str]):
    rank = {name: i for i, name in enumerate(order)}
    kids = list(parent)
    kids.sort(key=lambda k: rank.get(k.tag.split("}")[1], len(order)))
    for k in kids:
        parent.remove(k)
        parent.append(k)


SETTINGS_ORDER = ["writeProtection", "view", "zoom", "removePersonalInformation", "removeDateAndTime", "doNotDisplayPageBoundaries",
                  "displayBackgroundShape", "printPostScriptOverText", "printFractionalCharacterWidth", "printFormsData",
                  "embedTrueTypeFonts", "embedSystemFonts", "saveSubsetFonts", "saveFormsData", "mirrorMargins",
                  "alignBordersAndEdges", "bordersDoNotSurroundHeader", "bordersDoNotSurroundFooter", "gutterAtTop",
                  "hideSpellingErrors", "hideGrammaticalErrors", "activeWritingStyle", "proofState", "formsDesign",
                  "attachedTemplate", "linkStyles", "stylePaneFormatFilter", "stylePaneSortMethod", "documentType",
                  "mailMerge", "revisionView", "trackRevisions", "doNotTrackMoves", "doNotTrackFormatting",
                  "documentProtection", "autoFormatOverride", "styleLockTheme", "styleLockQFSet", "defaultTabStop",
                  "autoHyphenation", "consecutiveHyphenLimit", "hyphenationZone", "doNotHyphenateCaps", "showEnvelope",
                  "summaryLength", "clickAndTypeStyle", "defaultTableStyle", "evenAndOddHeaders", "bookFoldRevPrinting",
                  "bookFoldPrinting", "bookFoldPrintingSheets", "drawingGridHorizontalSpacing", "drawingGridVerticalSpacing",
                  "displayHorizontalDrawingGridEvery", "displayVerticalDrawingGridEvery", "doNotUseMarginsForDrawingGridOrigin",
                  "drawingGridHorizontalOrigin", "drawingGridVerticalOrigin", "doNotShadeFormData", "noPunctuationKerning",
                  "characterSpacingControl", "printTwoOnOne", "strictFirstAndLastChars", "noLineBreaksAfter",
                  "noLineBreaksBefore", "savePreviewPicture", "doNotValidateAgainstSchema", "saveInvalidXml",
                  "ignoreMixedContent", "alwaysShowPlaceholderText", "doNotDemarcateInvalidXml", "saveXmlDataOnly",
                  "useXSLTWhenSaving", "saveThroughXslt", "showXMLTags", "alwaysMergeEmptyNamespace", "updateFields",
                  "hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids", "mathPr",
                  "attachedSchema", "themeFontLang", "clrSchemeMapping", "doNotIncludeSubdocsInStats",
                  "doNotAutoCompressPictures", "forceUpgrade", "captions", "readModeInkLockDown", "smartTagType",
                  "schemaLibrary", "shapeDefaults", "doNotEmbedSmartTags", "decimalSymbol", "listSeparator"]


def insert_settings_child(settings, child):
    """w:settings 자식은 스키마 순서를 지켜야 Word가 문서를 연다"""
    for old in settings.findall(child.tag):
        settings.remove(old)
    settings.append(child)
    sort_children(settings, SETTINGS_ORDER)


def drop_thumbnail(doc):
    """기본 서식 파일에 딸린 미리보기 그림은 문서 내용과 무관하므로 제거"""
    rels = doc.part.package.rels
    for rid, rel in list(rels.items()):
        if rel.reltype.endswith("/thumbnail"):
            rels.pop(rid)
            getattr(rels, "_target_parts_by_rId", {}).pop(rid, None)


def build_docx(a, elems, footnotes, out: str):
    """build 명령의 DOCX 경로. 목차와 쪽번호를 함께 켜면 LibreOffice 렌더링으로 제목 쪽을 찾아 목차 결과값을 채운다."""
    import render

    def make(numbers=None):
        b = DocxBuilder(a.profile, a.page_break)
        b.numbered_square = a.numbered_square
        b.toc, b.toc_depth = a.toc, a.toc_depth
        b.page_number = a.page_number
        b.cover = parse_cover(a.cover)
        b.header_spec = a.header
        b.merge_mode = a.merge
        b.long_table_mode = a.long_table
        b.toc_numbers = numbers or {}
        b.build(elems, footnotes, out)
        return b
    b = make()
    has_toc = bool(b.form.style("toc_chapter") and b.form.style("toc_section"))
    if a.toc and a.page_number and b.headings and has_toc:   # 2회 조판: LibreOffice 렌더링으로 제목 쪽을 찾아 목차 결과값에 기입
        render.fill_toc_numbers(make, out, [t for _, t in M.toc_entries(b.headings, a.toc_depth)], None)


def check_docx(path: str) -> int:
    """DOCX 구성 요약. 문서를 열 수 있는지와 각 요소 수를 확인한다."""
    import zipfile
    doc = Document(path)
    body = doc.element.body
    with zipfile.ZipFile(path) as z:
        notes = z.read("word/footnotes.xml").decode("utf-8").count("<w:footnote ") - 2 if "word/footnotes.xml" in z.namelist() else 0
        images = sum(1 for n in z.namelist() if n.startswith("word/media/"))
    toc = any("TOC" in (t.text or "") for t in body.iter(qn("w:instrText")))
    print(msg("docx_check_parts", paras=len(body.findall(qn("w:p"))), tables=len(body.findall(qn("w:tbl"))),
              images=images, notes=max(notes, 0), toc=toc))
    return 0
